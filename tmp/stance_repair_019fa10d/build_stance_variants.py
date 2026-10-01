import json
import os
import sys

import openpyxl

root = sys.argv[1]
source = sys.argv[2]
out_dir = sys.argv[3]
sys.path.insert(0, os.path.join(root, "fec"))
from moneyball_core import overall_display_tier
from tier_util import legislative_press_tiers
from stance.workbook_stance_views import sync_category_stance_sheet, sync_coalition_profiles_sheet

category = json.load(open(os.path.join(root, "category_stance.json"), encoding="utf-8"))
matrix = category["matrix"]
caps = (category.get("gate") or {}).get("overall_capped") or {}
tier, _ = legislative_press_tiers(list(matrix))

base = openpyxl.load_workbook(source)
ra = base["Rank & Allocation"]
members = [ra.cell(row, 1).value for row in range(3, 536) if ra.cell(row, 1).value]
overall = {name: overall_display_tier(matrix[name], tier.get(name, 0), name, caps) for name in members}
base.close()

category_book = openpyxl.load_workbook(source)
for row in range(3, 536):
    name = category_book["Rank & Allocation"].cell(row, 1).value
    if name:
        category_book["Rank & Allocation"].cell(row, 7).value = overall[name]
sync_category_stance_sheet(category_book, category, members, overall)
category_book.save(os.path.join(out_dir, "variant_category.xlsx"))
category_book.close()

coalition_book = openpyxl.load_workbook(source)
sync_coalition_profiles_sheet(coalition_book)
coalition_book.save(os.path.join(out_dir, "variant_coalition.xlsx"))
coalition_book.close()
