import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

source = r"C:\tmp\mb_original.xlsx"

book = openpyxl.load_workbook(source)
cell = book["Coalition Profiles"]["Z30"]
cell.value = "Style test"
cell.fill = PatternFill("solid", fgColor="253A5A")
cell.font = Font(color="FFFFFF", bold=True)
cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
book.save(r"C:\tmp\variant_style.xlsx")
book.close()

book = openpyxl.load_workbook(source)
book["Coalition Profiles"]["B18"] = "=COUNTIF('Rank & Allocation'!$G$3:$G$535,$A18)"
book.save(r"C:\tmp\variant_formula.xlsx")
book.close()
