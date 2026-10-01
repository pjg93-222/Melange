# -*- coding: utf-8 -*-
"""Build the verified stance repair with the repository's Excel-compatible writer."""
from __future__ import annotations

import argparse
import collections
import json
import os
import sys

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    root = os.path.abspath(args.root)
    sys.path.insert(0, os.path.join(root, "fec"))
    from moneyball_core import overall_display_tier
    from tier_util import legislative_press_tiers
    from stance.workbook_stance_views import (
        stance_view_contract,
        sync_category_stance_sheet,
        sync_coalition_profiles_sheet,
    )

    with open(os.path.join(root, "category_stance.json"), encoding="utf-8") as handle:
        category = json.load(handle)
    matrix = category["matrix"]
    caps = (category.get("gate") or {}).get("overall_capped") or {}
    tier, tier_rank = legislative_press_tiers(list(matrix))

    wb = openpyxl.load_workbook(os.path.abspath(args.input))
    ra = wb["Rank & Allocation"]
    members = []
    overall = {}
    distribution = collections.Counter()
    for row in range(3, 536):
        member = ra.cell(row, 1).value
        if not member:
            continue
        stance = overall_display_tier(matrix.get(member, {}), tier.get(member, 0), member, caps)
        ra.cell(row, 7).value = stance
        members.append(member)
        overall[member] = stance
        distribution[stance] += 1

    receipt = stance_view_contract(category, members, overall, tier)
    sync_category_stance_sheet(wb, category, members, overall)
    sync_coalition_profiles_sheet(wb)
    cp = wb["Coalition Profiles"]
    cp["A16"] = "BY CURRENT CU-STANCE TIER"
    for col in range(2, 7):
        cp.cell(16, col).value = None
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output = os.path.abspath(args.output)
    os.makedirs(os.path.dirname(output), exist_ok=True)
    wb.save(output)
    wb.close()

    print("MEMBERS", len(members))
    print("CATEGORIES", len(category["cats"]))
    print("DISTRIBUTION", dict(distribution))
    print("CHAMPIONS", " | ".join(receipt["champions"]))
    print("SAVED", output)


if __name__ == "__main__":
    main()
