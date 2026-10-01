from __future__ import annotations

import collections
import json
import os
import sys
import zipfile

import openpyxl


root = os.path.abspath(sys.argv[1])
path = os.path.abspath(sys.argv[2])
sys.path.insert(0, os.path.join(root, "fec"))
from tier_util import legislative_press_tiers
from stance.workbook_stance_views import GENERAL_STANCE_TIERS

with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None

category = json.load(open(os.path.join(root, "category_stance.json"), encoding="utf-8"))
matrix = category["matrix"]
tier, _ = legislative_press_tiers(list(matrix))

formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
values = openpyxl.load_workbook(path, data_only=True, read_only=False)
assert len(formulas.sheetnames) == 30

ra = values["Rank & Allocation"]
members = [ra.cell(row, 1).value for row in range(3, 536) if ra.cell(row, 1).value]
assert len(members) == 533 and len(set(members)) == 533
overall = {ra.cell(row, 1).value: ra.cell(row, 7).value for row in range(3, 536) if ra.cell(row, 1).value}
distribution = collections.Counter(overall.values())
champions = {member for member, stance in overall.items() if stance == "Champion"}
tier_one = {member for member, value in tier.items() if value == 1}
assert champions == tier_one and len(champions) == 12

cs = formulas["Category Stance"]
assert [cs.cell(1, col).value for col in range(3, 21)] == category["cats"]
cs_rows = {cs.cell(row, 1).value: row for row in range(2, 535) if cs.cell(row, 1).value}
assert set(cs_rows) == set(matrix)
general_bad = sum(cs.cell(cs_rows[member], 2).value != overall[member] for member in members)
category_bad = sum(
    cs.cell(cs_rows[member], col).value != matrix[member].get(cat, "—")
    for member in members
    for col, cat in enumerate(category["cats"], start=3)
)
assert general_bad == 0 and category_bad == 0

cp_f = formulas["Coalition Profiles"]
cp_v = values["Coalition Profiles"]
assert tuple(cp_f.cell(row, 1).value for row in range(18, 29)) == GENERAL_STANCE_TIERS
assert tuple(cp_f.cell(3, col).value for col in range(10, 21)) == GENERAL_STANCE_TIERS
assert cp_v["B29"].value == 533
assert abs(cp_v["D29"].value - 2_000_000) < 1
assert cp_v["F29"].value == 133
dominant = [cp_v.cell(row, 21).value for row in range(4, 13)]
assert dominant == [
    "Ally", "Persuadable", "Ally", "Ally", "Highly Persuadable",
    "Highly Persuadable", "Highly Persuadable", "Persuadable", "Ally",
]

error_tokens = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}
formula_errors = []
for sheet in values.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value in error_tokens:
                formula_errors.append((sheet.title, cell.coordinate, cell.value))
assert not formula_errors, formula_errors[:20]

print("ZIP_OK")
print("SHEETS", len(formulas.sheetnames))
print("MEMBERS", len(members))
print("DISTRIBUTION", dict(distribution))
print("CHAMPIONS", " | ".join(sorted(champions)))
print("CATEGORY_GENERAL_MISMATCHES", general_bad)
print("CATEGORY_CELL_MISMATCHES", category_bad)
print("COALITION_TOTALS", cp_v["B29"].value, cp_v["D29"].value, cp_v["F29"].value)
print("DOMINANT_STANCES", dominant)
print("FORMULA_ERRORS", len(formula_errors))
