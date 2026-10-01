from __future__ import annotations

import csv
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems")
BASELINE = (
    ROOT
    / "audits"
    / "legfix5"
    / "production_baseline_full_downstream"
    / "files"
)
INTEGRATION = (
    ROOT
    / "audits"
    / "legfix5"
    / "staged_integration"
    / "greenlight-20260728a"
)
AMENDMENTS = (
    ROOT
    / "audits"
    / "legfix5"
    / "opus_adjudication"
    / "amendment_culling_opus5_20260728a"
    / "amendment_partition_3703.ndjson"
)


def load_ndjson(path: Path) -> list[dict[str, object]]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def normalized(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.casefold().replace("’", "'")
    value = re.sub(r"\b(jr|sr|ii|iii|iv)\b", " ", value)
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def sponsor_name(label: str) -> str:
    visible = label.split("[", 1)[0].strip()
    if "," not in visible:
        return visible
    last, first = visible.split(",", 1)
    return f"{first.strip()} {last.strip()}"


edges = [
    row
    for row in load_ndjson(INTEGRATION / "integrated_membership_edges.ndjson")
    if row.get("scoring_eligible") is True
]
keys = {str(row["canonical_key"]) for row in edges}
amendments = {
    str(row["canonical_key"]): row
    for row in load_ndjson(AMENDMENTS)
    if str(row.get("canonical_key")) in keys
}

book = load_workbook(
    BASELINE / "Moneyball Main Version - Blended.xlsx",
    read_only=True,
    data_only=True,
)
roster = [
    row[0]
    for row in book["Rank & Allocation"].iter_rows(
        min_row=3, max_row=535, min_col=1, max_col=1, values_only=True
    )
    if row[0]
]
book.close()
roster_by_normalized = {normalized(name): name for name in roster}

legislators = []
with (ROOT / "fec" / "legislators.csv").open(
    encoding="utf-8-sig", newline=""
) as handle:
    legislators = list(csv.DictReader(handle))
bioguide_by_roster = {}
for member in roster:
    key = normalized(member)
    matches = [
        row
        for row in legislators
        if normalized(row.get("full_name") or "") == key
        or normalized(
            f"{row.get('first_name') or ''} {row.get('middle_name') or ''} "
            f"{row.get('last_name') or ''}"
        )
        == key
        or normalized(
            f"{row.get('nickname') or ''} {row.get('last_name') or ''}"
        )
        == key
    ]
    if len(matches) == 1:
        bioguide_by_roster[member] = matches[0].get("bioguide_id")

mapped = {}
unmapped = {}
for key, row in sorted(amendments.items()):
    label = str(row.get("sponsor") or "")
    parsed = sponsor_name(label)
    normalized_parsed = normalized(parsed)
    direct = roster_by_normalized.get(normalized_parsed)
    if direct:
        mapped[key] = {
            "member": direct,
            "bioguide_id": bioguide_by_roster.get(direct),
            "sponsor": label,
        }
        continue
    parts = normalized_parsed.split()
    candidates = [
        member
        for member in roster
        if parts
        and normalized(member).split()[-1:] == parts[-1:]
        and normalized(member).split()[:1] == parts[:1]
    ]
    if len(candidates) == 1:
        member = candidates[0]
        mapped[key] = {
            "member": member,
            "bioguide_id": bioguide_by_roster.get(member),
            "sponsor": label,
        }
    else:
        unmapped[key] = {
            "parsed": parsed,
            "sponsor": label,
            "candidates": candidates,
        }

print(
    json.dumps(
        {
            "scoring_edges": len(edges),
            "scoring_measure_count": len(keys),
            "amendment_records": len(amendments),
            "mapped": len(mapped),
            "unmapped": len(unmapped),
            "mapped_records": mapped,
            "unmapped_records": unmapped,
        },
        indent=2,
        ensure_ascii=False,
        sort_keys=True,
    )
)
