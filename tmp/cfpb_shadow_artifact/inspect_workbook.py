from pathlib import Path

from openpyxl import load_workbook


path = Path(
    r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems\audits\legfix5"
    r"\production_baseline_full_downstream\files"
    r"\Moneyball Main Version - Blended.xlsx"
)
columns = [
    "CU",
    "CV",
    "CW",
    "CX",
    "CY",
    "CZ",
    "DA",
    "DB",
    "EO",
    "EP",
    "EQ",
    "ER",
    "ES",
    "ET",
    "EU",
    "EV",
]

formula_book = load_workbook(path, read_only=True, data_only=False)
value_book = load_workbook(path, read_only=True, data_only=True)
raw_formula = formula_book["Legislative Score - Raw"]
raw_value = value_book["Legislative Score - Raw"]
normalized_formula = formula_book["Legislative Score - Normalized"]
normalized_value = value_book["Legislative Score - Normalized"]

print("RAW")
raw_matrix = list(
    raw_value.iter_rows(
        min_row=3,
        max_row=535,
        min_col=99,
        max_col=152,
        values_only=True,
    )
)
normalized_matrix = list(
    normalized_value.iter_rows(
        min_row=3,
        max_row=535,
        min_col=99,
        max_col=152,
        values_only=True,
    )
)


def column_number(label: str) -> int:
    result = 0
    for character in label:
        result = result * 26 + ord(character) - 64
    return result


for column in columns:
    offset = column_number(column) - 99
    values = [row[offset] for row in raw_matrix]
    numeric = [float(value) for value in values if isinstance(value, (int, float))]
    print(
        column,
        {
            "header": raw_formula[f"{column}1"].value,
            "weight": raw_formula[f"{column}2"].value,
            "nonzero": sum(value != 0 for value in numeric),
            "sum": sum(numeric),
            "min": min(numeric) if numeric else None,
            "max": max(numeric) if numeric else None,
        },
    )

print("NORMALIZED")
for column in columns:
    offset = column_number(column) - 99
    values = [row[offset] for row in normalized_matrix]
    numeric = [float(value) for value in values if isinstance(value, (int, float))]
    print(
        column,
        {
            "header": normalized_formula[f"{column}1"].value,
            "weight": normalized_formula[f"{column}2"].value,
            "formula3": normalized_formula[f"{column}3"].value,
            "nonzero": sum(value != 0 for value in numeric),
            "sum": sum(numeric),
            "min": min(numeric) if numeric else None,
            "max": max(numeric) if numeric else None,
        },
    )

for sheet, cells in {
    "Rank & Allocation": ["D3", "E3"],
    "Blended Model": ["D3", "G3", "J3", "K3", "O3", "P3", "S3", "T3"],
}.items():
    print(sheet)
    for cell in cells:
        print(
            cell,
            {
                "formula": formula_book[sheet][cell].value,
                "value": value_book[sheet][cell].value,
            },
        )

formula_book.close()
value_book.close()
