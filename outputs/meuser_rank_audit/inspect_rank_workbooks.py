import json
import os
import sys

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "fec"))
import stance.media_rules as media_rules


def inspect(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    bm = wb["Blended Model"]
    ra = wb["Rank & Allocation"]
    members = [ra.cell(r, 1).value for r in range(3, 536) if ra.cell(r, 1).value]
    rows = {member: index + 3 for index, member in enumerate(members)}
    row = rows.get("Daniel Meuser")
    weights = {r: bm.cell(r, 28).value for r in (2, 3, 4, 5)}
    out = {
        "file": os.path.basename(path),
        "mtime": os.path.getmtime(path),
        "weights": weights,
        "row": row,
        "influence": bm.cell(row, 5).value if row else None,
        "demo_raw": bm.cell(row, 6).value if row else None,
        "leg_raw": bm.cell(row, 15).value if row else None,
        "score_cache": bm.cell(row, 19).value if row else None,
        "rank_cache": bm.cell(row, 20).value if row else None,
    }
    rw = wb["Press Release - Raw"]
    rrows = {rw.cell(r, 1).value: r for r in range(1, rw.max_row + 1) if rw.cell(r, 1).value}
    rr = rrows.get("Daniel Meuser")
    cols = list(range(5, 24))
    all_cols = cols + [24]
    headers = {col: rw.cell(1, col).value for col in all_cols}
    raw_all = {
        col: {member: (rw.cell(rrows[member], col).value or 0.0) for member in rrows}
        for col in all_cols
    }
    by_header = {headers[col]: raw_all[col] for col in all_cols if headers[col]}
    by_header, _ = media_rules.disperse_general(by_header, list(rrows))
    normalized = {
        col: {
            member: max(-1.0, min(1.0, float(value)))
            for member, value in by_header.get(headers[col], raw_all[col]).items()
        }
        for col in cols
    }
    pn = wb["Press Release - Normalized"]
    category_weights = {col: (pn.cell(535, col).value or 0.0) for col in cols}
    pillars = {
        member: sum(normalized[col].get(member, 0.0) * category_weights[col] for col in cols)
        for member in rrows
    }
    maximum = max(pillars.values()) or 1.0
    ordered = sorted(pillars, key=lambda member: -pillars[member])
    out["press_pillar"] = pillars.get("Daniel Meuser")
    out["press_pct"] = pillars.get("Daniel Meuser", 0.0) / maximum
    out["press_rank"] = ordered.index("Daniel Meuser") + 1
    out["press_contributions"] = [
        [headers[col], normalized[col].get("Daniel Meuser", 0.0) * category_weights[col]]
        for col in cols
        if normalized[col].get("Daniel Meuser", 0.0) * category_weights[col]
    ]
    if rr:
        values = []
        for col in range(5, 25):
            val = rw.cell(rr, col).value
            if val not in (None, 0, 0.0):
                values.append([rw.cell(1, col).value, val])
        out["media_raw_nonzero"] = values
    print(json.dumps(out, ensure_ascii=False, separators=(",", ":")))
    wb.close()


for arg in sys.argv[1:]:
    inspect(os.path.abspath(arg))
