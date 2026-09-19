from pathlib import Path
import collections, hashlib, json, re, sys
from decimal import Decimal
import openpyxl
sys.path.insert(0, r'C:\Users\pgoff\Mentat\mentat\tools')
from report_directory import report_directory
OUT=Path(__file__).parent

def money(v):
    return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
def table(book,name,header):
    sheet=book[name]
    names=[c.value for c in next(sheet.iter_rows(min_row=header,max_row=header))]
    return [dict(zip(names,row),_row=rn) for rn,row in enumerate(sheet.iter_rows(min_row=header+1,values_only=True),header+1) if any(v is not None for v in row)]
result={}
for industry in ['Crypto','Banking']:
    path=report_directory()/f'{industry} Industry Spending Overview.xlsx'
    with path.open('rb') as f: digest=hashlib.file_digest(f,'sha256').hexdigest()
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    data={'file':str(path),'sha256':digest}
    overview=table(book,'Crypto Spending Overview' if industry=='Crypto' else 'Spending Overview',5)
    data['overview']={}
    for cycle in ['2021-2022','2023-2024','2025-2026']:
        rows=[r for r in overview if r.get('Cycle')==cycle]
        data['overview'][cycle]={'total_giving_all_rows':str(sum((money(r.get('Total Giving')) for r in rows),Decimal(0))), 'by_scope':{str(r.get('Scope')):str(money(r.get('Total Giving'))) for r in rows}}
    if industry=='Crypto':
        data['cycles']={}
        for cycle in ['2021-2022','2023-2024','2025-2026']:
            rows=table(book,cycle,4)
            flags={}
            for flag in ['Superseded Amendment','Duplicate Filing Version','Duplicate Memo','Memo Leg','Conduit Leg']:
                selected=[r for r in rows if r.get(flag)==1]
                included=[r for r in selected if r.get('Include in Analysis')=='Yes']
                flags[flag]={'rows':len(selected),'included_rows':len(included),'included_analysis_dollars':str(sum((money(r.get('Analysis Amount')) for r in included),Decimal(0)))}
            data['cycles'][cycle]={'row_count':len(rows),'analysis_amount':str(sum((money(r.get('Analysis Amount')) for r in rows),Decimal(0))),'flags':flags,'flagged_large_rows':[{k:r.get(k) for k in ['_row','Source ID','Company','Contributor','Amount','Analysis Amount','Recipient','Committee ID','Date','Include in Analysis','Superseded Amendment','Duplicate Filing Version','Duplicate Memo']} for r in rows if (r.get('Superseded Amendment')==1 or r.get('Duplicate Filing Version')==1) and abs(money(r.get('Amount')))>=1000000]}
        data['coinbase_23m_rows']=[r for r in table(book,'2023-2024',4) if r.get('Company')=='Coinbase' and money(r.get('Amount'))==23000000]
        recipient=table(book,'Recipient List',5)
        data['recipient_list_sums']={}
        for cycle in ['2021-2022','2023-2024','2025-2026']:
            data['recipient_list_sums'][cycle]={view:str(sum((money(r.get('Total Giving')) for r in recipient if r.get('Cycle')==cycle and r.get('View')==view),Decimal(0))) for view in ['Industry','Company']}
    if industry=='Banking':
        employees=table(book,'Employee List',5)
        slim={'source_file':str(path),'sha256':digest,'sheet':'Employee List','scope':'Published Employee List rows across cycles 2021-2022, 2023-2024, and 2025-2026', 'columns':[{'name':'ZIPs','type':'category'},{'name':'Total Giving','type':'number'},{'name':'Cycle','type':'category'}], 'rows':[[r.get('ZIPs'),r.get('Total Giving'),r.get('Cycle')] for r in employees], 'row_numbers':[r['_row'] for r in employees]}
        (OUT/'banking_employee_chart_data.json').write_text(json.dumps(slim),encoding='utf-8')
        slim_overview={'source_file':str(path),'sha256':digest,'sheet':'Spending Overview','columns':[{'name':'Cycle','type':'category'},{'name':'Total Giving','type':'number'},{'name':'Scope','type':'category'}], 'rows':[[r.get('Cycle'),r.get('Total Giving'),r.get('Scope')] for r in overview], 'row_numbers':[r['_row'] for r in overview]}
        (OUT/'banking_overview_chart_data.json').write_text(json.dumps(slim_overview),encoding='utf-8')
        data['employee_headers']=list(employees[0])
        data['employee_examples']=employees[:5]
        zip_header=next((k for k in employees[0] if 'zip' in str(k).lower()),None)
        amount_header=next((k for k in employees[0] if k=='Total Giving'),None)
        data['employee_split_zip']={'zip_header':zip_header,'amount_header':amount_header}
        if zip_header and amount_header:
            multiple=[r for r in employees if len([p for p in re.split(r'\s*[;|]\s*',str(r.get(zip_header) or '')) if p])>1]
            base=sum((money(r.get(amount_header)) for r in employees),Decimal(0))
            exploded=sum((money(r.get(amount_header))*max(1,len([p for p in re.split(r'\s*[;|]\s*',str(r.get(zip_header) or '')) if p])) for r in employees),Decimal(0))
            data['employee_split_zip'].update({'rows':len(employees),'multi_zip_rows':len(multiple),'base_sum':str(base),'naive_split_sum':str(exploded),'inflation':str(exploded-base),'examples':sorted(multiple,key=lambda r:money(r.get(amount_header)),reverse=True)[:5]})
    result[industry]=data
    book.close()
(OUT/'published_quantification.json').write_text(json.dumps(result,indent=2,default=str),encoding='utf-8')
print(str(OUT/'published_quantification.json'))
