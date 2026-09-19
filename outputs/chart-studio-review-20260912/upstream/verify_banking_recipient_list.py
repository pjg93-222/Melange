from pathlib import Path
import sys,json,hashlib
from decimal import Decimal
import openpyxl
sys.path.insert(0,r'C:\Users\pgoff\Mentat\mentat\tools')
from report_directory import report_directory
out=Path(__file__).parent
path=report_directory()/'Banking Industry Spending Overview.xlsx'
with path.open('rb') as f: digest=hashlib.file_digest(f,'sha256').hexdigest()
book=openpyxl.load_workbook(path,read_only=True,data_only=True)
sheet=book['Recipient List']
header=[c.value for c in next(sheet.iter_rows(min_row=5,max_row=5))]
rows=[dict(zip(header,r),_row=n) for n,r in enumerate(sheet.iter_rows(min_row=6,values_only=True),6) if any(v is not None for v in r)]
def amount(r):
    v=r.get('Total Giving')
    return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
totals={}
for cycle in ['All cycles','2021-2022','2023-2024','2025-2026']:
    selected=[r for r in rows if cycle=='All cycles' or r.get('Cycle')==cycle]
    totals[cycle]={view:{'rows':sum(r.get('View')==view for r in selected),'amount':str(sum((amount(r) for r in selected if r.get('View')==view),Decimal(0)))} for view in ['Industry','Company']}
    totals[cycle]['all_rows']={'rows':len(selected),'amount':str(sum((amount(r) for r in selected),Decimal(0)))}
data={'source_file':str(path),'sha256':digest,'sheet':'Recipient List','columns':[{'name':'Cycle','type':'category'},{'name':'Total Giving','type':'number'},{'name':'View','type':'category'},{'name':'Company','type':'category'}], 'rows':[[r.get('Cycle'),r.get('Total Giving'),r.get('View'),r.get('Company')] for r in rows], 'row_numbers':[r['_row'] for r in rows]}
(out/'banking_recipient_chart_data.json').write_text(json.dumps(data),encoding='utf-8')
result={'source_file':str(path),'sha256':digest,'sheet':'Recipient List','header_row':5,'total_giving_column':openpyxl.utils.get_column_letter(header.index('Total Giving')+1),'first_data_row':6,'last_data_row':max(r['_row'] for r in rows),'totals':totals}
(out/'banking_recipient_reconciliation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
book.close()
