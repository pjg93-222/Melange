import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.abspath(__file__)));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""FULL-BODY AUDIT — every module in one report.
Consolidates the integrity (22) + frontier (11) audits and adds correctness checks for the
new modules: money-response, electoral plays, co-sponsor routes, LT-priming, stance, determinism."""
import json,sys,os,subprocess,re,collections,importlib.util,csv
import config_overlay as CO
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
B=MB_ROOT; FE=os.path.dirname(os.path.abspath(__file__)); DATA_FE=B+r'\fec'
def _pid_alive(pid):
    try:
        os.kill(int(pid),0); return True
    except (OSError,ValueError,TypeError):
        return False
def assert_no_foreign_pipeline_writer(lock_path=None,parent_pid=None):
    """A standalone audit must never read a mixed-generation artifact set."""
    lock_path=lock_path or os.path.join(B,'.pipeline_run.lock')
    parent_pid=os.getppid() if parent_pid is None else int(parent_pid)
    try:
        with open(lock_path,encoding='utf-8') as handle: owner=json.load(handle).get('pid')
    except (FileNotFoundError,OSError,ValueError,TypeError):
        return
    if owner not in (os.getpid(),parent_pid) and _pid_alive(owner):
        raise RuntimeError('refusing audit while a foreign Moneyball pipeline owns the artifact set (pid %s)'%owner)
assert_no_foreign_pipeline_writer()
P,F=[],[]
def ck(name,cond,detail=''):
    (P if cond else F).append(name); print(('  OK   ' if cond else 'FAIL  ')+name+((' :: '+detail) if (detail and not cond) else ''))
def run(script):
    return subprocess.run([sys.executable,FE+'\\'+script],capture_output=True,text=True,encoding='utf-8',errors='replace').stdout


print('========== FULL-BODY AUDIT ==========\n[1] Existing audits')
for s,nm,exp in [('audit.py','integrity (22)',22),('audit_frontier.py','frontier (11)',11)]:   # S3-03 audit run-2: assert COUNT too — a silently-skipped check must fail here, not vanish
    out=run(s); m=re.search(r'(\d+) passed, (\d+) FAILED',out)
    ck('%s sub-audit clean'%nm, bool(m) and m.group(2)=='0' and int(m.group(1))>=exp, out.strip().splitlines()[-1] if out else 'no output')
print('\n[2] Money-response (diminishing returns + capacity)')
spec=importlib.util.spec_from_file_location('mr',FE+r'\money_response.py'); mr=importlib.util.module_from_spec(spec); spec.loader.exec_module(mr)
g='Jared Golden'
s1,s2,s4=mr.shift_lean(g,'cu',100000),mr.shift_lean(g,'cu',200000),mr.shift_lean(g,'cu',400000)
ck('money curve concave (doubling $ < doubling effect)', s2<2*s1 and s4<2*s2, '%.1f %.1f %.1f'%(s1,s2,s4))
big=mr.shift_lean('Jared Golden','cu',100000); small=mr.shift_lean('Steve Daines','cu',100000)
ck('capacity-aware (bigger-receipts member moves less)', abs(big)<abs(small))
# 2026-07-02: PT loss-frame asymmetry (Linde-Vis / K&T) is BY DESIGN: comp = -LOSSFRAME x cu, not the plain mirror.
ck('competitor money = -LOSSFRAME x cu (loss-frame asymmetry, PT)', abs(mr.shift_lean(g,'comp',100000)+mr.LOSSFRAME*s1)<1e-9)
_valid_ov=set(mr.ORDER); _deft={mr.defection(m,-30) for m in mr.C}-{None}
ck('defection yields ONLY valid overall tiers (no taxonomy leak)', _deft<=_valid_ov, 'invalid: %s'%(_deft-_valid_ov))
_probe=next(iter(mr.C)); _probe_old=mr.OV.get(_probe); mr.OV[_probe]='Opponent'
_opp_immune=mr.defection(_probe,-99) is None
if _probe_old is None: mr.OV.pop(_probe,None)
else: mr.OV[_probe]=_probe_old
ck('committed Opponents immune to money defection (structural probe)',_opp_immune)
# 2026-07-02: test the RUNG, not a named member — Durbin (the old exemplar) legitimately graduated to
# full Opponent once the tiered media weights priced his interchange opposition.
_hsk=[m for m in mr.C if mr.PROJ.get(mr.OV.get(m),mr.OV.get(m))=='Highly Skeptical']
ck('Highly Skeptical deep-but-movable (sub-threshold absorbed, big shock -> Opponent)',
   bool(_hsk) and all(mr.defection(m,-5) is None and mr.defection(m,-12)=='Opponent' for m in _hsk),
   '%d members at rung'%len(_hsk))
import moneyball_core as _band_core
_ob=dict((tier,cut) for cut,tier in _band_core.OVERALL_BANDS)
ck('overall bends are the strict fixed ladder, not roster-population quotas',
   _ob=={'Ally':3.0,'Highly Persuadable':2.0,'Persuadable':1.0,
         'Lean Persuadable':.4,'True Neutral':0.0,'Lean Skeptic':-.3,
         'Skeptic':-.75,'Highly Skeptical':-1.5} and
   _band_core.overall_tier_from_value(0)=='True Neutral' and
   _band_core.overall_tier_from_value(.01) not in ('Lean Skeptic','Skeptic','Highly Skeptical','Opponent') and
   _band_core.overall_tier_from_value(-.01) not in ('Highly Persuadable','Persuadable','Lean Persuadable','Ally'),
   str(_ob))
print('\n[3] Electoral plays (non-partisan, gated)')
EP=json.load(open(B+r'\electoral_plays.json'))
contest=[m for m,d in EP.items() if d.get('play')=='contest']; defend=[m for m,d in EP.items() if d.get('play')=='defend']
ck('contest plays gated to Opponents only', all(EP[m]['overall']=='Opponent' for m in contest))
ck('defend plays gated to allies only', all(EP[m]['overall'] in ('Champion','Defender','Ally') for m in defend))
pb=collections.Counter(EP[m]['party'] for m in contest)
_contest_n=sum(pb.values())
ck('contest plays not partisan-skewed (<70%% one party; enforced at n>=12)',
   _contest_n<12 or max(pb['D'],pb['R'])<=0.70*_contest_n,
   'D%d R%d (n=%d)'%(pb['D'],pb['R'],_contest_n))
print('\n[4] Stance matrix sanity')
CS=json.load(open(B+r'\category_stance.json')); MAT=CS['matrix']
_overall_bands=((CS.get('thresholds') or {}).get('overall') or {}).get('bands') or []
ck('overall band table matches kernel (all cutoffs + tiers, provenance)',
   len(_overall_bands)==len(mr.MC.OVERALL_BANDS) and
   all(abs((_overall_bands[_i].get('minimum') or 0)-_cut)<1e-12 and _overall_bands[_i].get('tier')==_tier
       for _i,(_cut,_tier) in enumerate(mr.MC.OVERALL_BANDS)))
ck('Durbin is Interchange Opponent', MAT['Richard Durbin']['Interchange']=='Opponent')
ck('Durbin anti on Credit Card (Opponent/Skeptic)', MAT['Richard Durbin']['Credit Card Mandates'] in ('Opponent','Skeptic'))
cnt=collections.Counter(MAT[m][c] for m in MAT for c in CS['cats'])
ck('champion rate disciplined (<2000 of ~9000 cells)', cnt['Champion']<2000, str(cnt['Champion']))
print('\n[5] Co-sponsor route (Pareto + MC reinforcement + determinism)')
def front(env):
    e=dict(os.environ,TOPIC='Mortgage',**env)
    out=subprocess.run([sys.executable,FE+r'\cosponsor_route.py'],capture_output=True,text=True,encoding='utf-8',errors='replace',env=e).stdout
    d={}
    for ln in out.splitlines():
        m=re.search(r'([0-9,]+) +-> +\+([0-9.]+) +co-sponsors +risk +([0-9.]+)',ln)
        if m: d[int(m.group(1).replace(',',''))]=float(m.group(2))
    return d
tr=front({}); mc=front({'MCMODE':'1','MC':'30'}); mc2=front({'MCMODE':'1','MC':'30'})
shared=sorted(set(tr)&set(mc))
ck('route front non-empty', len(tr)>3, '%d pts'%len(tr))
ck('MC-exact >= tree-reach at shared costs (reinforcement)', all(mc[c]>=tr[c]-0.05 for c in shared) if shared else False)
ck('MC route deterministic (same seed -> identical)', mc==mc2)
print('\n[6] LT-priming validity')
pr=json.load(open(B+r'\lt_priming.json'))
ck('priming sorted cheapest-first', all(pr[i]['cost']<=pr[i+1]['cost'] for i in range(len(pr)-1)))
ck('every recipe closes the gap (pressure+money)', all(p['cost']>0 and p['gap']>0 for p in pr))
ck('no committed Opponents/Champions in primeable set', all(p['tier'] not in ('Champion','Defender','Opponent') for p in pr))
print('\n[7] Scenario runner (end-to-end)')
def scen():
    return subprocess.run([sys.executable,FE+r'\scenario.py','--money','Jared Golden:comp:500000'],capture_output=True,text=True,encoding='utf-8',errors='replace').stdout
o1=scen(); o2=scen()
net=re.search(r'([+-][0-9.]+) co-sponsors\)',o1)   # the coalition/mobilization NET line
ck('scenario runs end-to-end (reports a net influence change)', bool(net), o1.strip().splitlines()[-1] if o1 else 'no output')
ck('scenario deterministic (same scenario -> identical)', o1==o2)
ck('competitor shock reduces CU influence (sign sane)', bool(net) and float(net.group(1))<=0, net.group(1) if net else '?')
# H2/H3 regression: competitor shock must reduce the FLIP channel too (caught the imputed-vs-raw ind footing bug)
flnet=re.search(r'winnability [0-9.]+ -> [0-9.]+  \(([+-][0-9.]+) expected flips\)',o1)
ck('competitor shock reduces winnability too (flip channel, consistent footing)', bool(flnet) and float(flnet.group(1))<=0, flnet.group(1) if flnet else '?')
# scenario CU spend now registers via flip channel (regression on the bug we just fixed)
cu_o=subprocess.run([sys.executable,FE+r'\scenario.py','--money','Tim Scott:cu:300000'],capture_output=True,text=True,encoding='utf-8',errors='replace').stdout
fl=re.search(r'([+-][0-9.]+) expected flips',cu_o)
ck('CU spend registers gradual flip-value gain (not absorbed)', bool(fl) and float(fl.group(1))>0, fl.group(1) if fl else '?')
print('\n[8] Joint budget optimizer (fixed-state lazy greedy / modeled replanning)')
def opt(env): return subprocess.run([sys.executable,FE+r'\optimize.py'],capture_output=True,text=True,encoding='utf-8',errors='replace',env=dict(os.environ,**env)).stdout
# OPT_OUT -> temp json so the test run doesn't clobber the production optimize_result.json (mirrors FRONTIER_OUT below).
# INC=5000 (divides PMAX=10000): optimize.py's cap check requires room for a FULL chunk, so INC>PMAX allocates nothing.
_OE={'BUDGET':'300000','INC':'5000','OPT_OUT':B+r'\_audit_tmp_optimize.json'}
oo1=opt(_OE); oo2=opt(_OE)
gn=re.search(r'\(([+-][0-9.]+) expected flips\)',oo1)
ck('optimizer runs + reports a winnability frontier', bool(gn), oo1.strip().splitlines()[1] if oo1 else 'no output')
ck('optimizer improves winnability (monotone, sign sane)', bool(gn) and float(gn.group(1))>0, gn.group(1) if gn else '?')
ck('optimizer deterministic (same budget -> identical)', oo1==oo2)
# H1 regression: winning a member (conversion) must be a GAIN -> the spend->winnability path is monotone non-decreasing
_opt=json.load(open(B+r'\_audit_tmp_optimize.json')); _pp=[v for _,v in _opt['path']]
_drops=sum(1 for i in range(1,len(_pp)) if _pp[i]<_pp[i-1]-1e-9)
ck('optimizer path monotone (conversion is a gain, not the old persuadability-loss)', _drops==0, '%d drops; base %.2f final %.2f'%(_drops,_pp[0],_pp[-1]))
ck('optimizer reports mode, observation status, and guarantee scope honestly',
   _opt.get('mode') in ('fixed_state_lazy_greedy','modeled_threshold_replanning') and
   isinstance(_opt.get('observations'),list) and isinstance(_opt.get('guarantee'),dict) and
   ((_opt.get('certificate_valid') and _opt['guarantee'].get('kind')=='exact_discrete_separable_concave')
    or (not _opt.get('certificate_valid') and _opt['guarantee'].get('kind')=='none')),
   str({k:_opt.get(k) for k in ('mode','observations','guarantee','certificate_valid')}))
ck('optimizer exact fixed-state certificate is internally consistent',
   (not _opt.get('certificate_valid')) or
   (_opt.get('fixed_state_concave') and not _opt.get('concavity_violations') and
    abs(_opt.get('opt_bound',0)-_opt.get('final',1))<=0.001 and
    _opt['guarantee'].get('certified_gain_ratio')==1.0),
   str(_opt.get('guarantee')))
# H3 regression: members with NO FEC data must NOT rank as top flip targets (ind=0 -> pflip~1 was misreading "missing" as "no opposition")
_s=importlib.util.spec_from_file_location('fmA',FE+r'\flip_model.py'); _fm=importlib.util.module_from_spec(_s); _s.loader.exec_module(_fm)
_wn=collections.Counter()
for c in _fm.cats:
    cs,ct,nA,nB,mAt,mBt,tgt=_fm.topic_setup(c,_fm.CS0)
    for n in _fm.members: _wn[n]+=tgt[n]*_fm.pflip(_fm.cu0[n],_fm.ind0[n],_fm.R[n],nA[n],nB[n],mAt,mBt)
_nd=sum(1 for n,_ in _wn.most_common(20) if n in _fm.nodata)
ck('no zero-FEC-data member ranks top-20 winnability (missing != no-opposition)', _nd==0, '%d of top-20 are data-less'%_nd)
# HIGH-3 (corrected): impute ONLY no-FEC-record (receipts=0) blind members; trust FEC-verified-real ind=0 zeros
_blind_ok=all(_fm.ind0[n]>0 for n in _fm.ind_imputed)
_realzero=[n for n in _fm.members if (_fm.C[n].get('receipts') or 0)>0 and ((_fm.C[n]['banks'] or 0)+(_fm.C[n]['finserv'] or 0))==0]
_real_kept=all(_fm.ind0[n]==0 for n in _realzero)
ck('imputation limited to no-FEC-record members; verified-real ind=0 zeros trusted', _blind_ok and _real_kept, 'blind %d imputed, %d real-zeros trusted'%(len(_fm.ind_imputed),len(_realzero)))
# HIGH-1 regression: frontier cascade now uses edge-keyed CRN -> a single-topic run is deterministic across repeats
def bf(): return subprocess.run([sys.executable,FE+r'\budget_frontier.py'],capture_output=True,text=True,encoding='utf-8',errors='replace',env=dict(os.environ,TOPIC='Interchange',FRONTIER_OUT=B+r'\_audit_tmp_frontier.json')).stdout   # temp out -> don't clobber the all-topics production json
bf1=bf(); bf2=bf()
ck('budget_frontier deterministic w/ edge-keyed CRN (single topic)', bf1==bf2 and 'Interchange' in bf1)

print('\n[NEW] Salience shock + press SNA layers (Noble-Turner / HCoI / Fink / Zhou-Elejalde / Joo)')
import json as _j
if os.path.exists(B+r'\salience_advisory.json'):
    _adv=_j.load(open(B+r'\salience_advisory.json'))['advisory']
    ck('salience advisory covers all issues', len(_adv)>=18, str(len(_adv)))
    _pr=[a['params'] for a in _adv]
    ck('salience regime probs valid (p+q+b~1, all in [0,1])',
       all(0<=x['p']<=1 and 0<=x['q']<=1 and 0<=x['b']<=1 and 0.94<=x['p']+x['q']+x['b']<=1.01 for x in _pr))
else: ck('salience advisory present (run salience_shock.py --advise)', False)
if os.path.exists(B+r'\dashboard_data.json'):
    _dd=_j.load(open(B+r'\dashboard_data.json',encoding='utf-8'))
    _sal=_dd.get('engine',{}).get('sal')
    ck('bundle exports Noble sal params (one per issue layer)', bool(_sal) and len(_sal)==len(_dd['issues']),
       '%s vs %s'%(len(_sal or []),len(_dd['issues'])))
if os.path.exists(B+r'\press_hcoi.json'):
    _hc=_j.load(open(B+r'\press_hcoi.json',encoding='utf-8'))
    _allm=[m2 for c in _hc['communities'] for m2 in c['members']]
    ck('HCoI communities non-trivial and disjoint (Louvain partition)',
       len(_hc['communities'])>=3 and len(_allm)==len(set(_allm)), '%d comms'%len(_hc['communities']))
if os.path.exists(B+r'\press_engagement.json'):
    _pe=_j.load(open(B+r'\press_engagement.json',encoding='utf-8'))
    ck('engagement layer: valid probabilities, no self-edges',
       all(0<e['p']<=1 and e['src']!=e['dst'] for e in _pe['directed']), '%d edges'%len(_pe['directed']))
if os.path.exists(B+r'\stance_dynamics.json'):
    _sd=_j.load(open(B+r'\stance_dynamics.json',encoding='utf-8'))
    ck('stance-dynamics imputation beats coin flip on masked cells (sign agreement > 0.55)',
       _sd['masking_validation']['sign_agreement']>0.55, str(_sd['masking_validation']))
    ck('imputation proposals carry provenance (user adjudicates; matrix untouched)',
       os.path.exists(B+r'\stance_imputation_proposals.csv') and
       'provenance' in open(B+r'\stance_imputation_proposals.csv',encoding='utf-8-sig').readline())
if os.path.exists(B+r'\frame_alignment.json'):
    _fa=_j.load(open(B+r'\frame_alignment.json',encoding='utf-8'))
    ck('frame-alignment rates in [0,1] across issues', all(0<=x['alignment']<=1 for x in _fa['alignment']))


print('\n[GOLD] Golden labels + domain classifier + collector DB')
if os.path.exists(B+'\\cu_golden_labels.csv'):
    import csv as _csv
    _gl=list(_csv.DictReader(open(B+'\\cu_golden_labels.csv',encoding='utf-8-sig')))
    _cls={'support','oppose','neutral','unrelated'}
    ck('golden label set parses, all labels in vocabulary', len(_gl)>=700 and all(r['label'] in _cls for r in _gl), '%d rows'%len(_gl))
    _sp=collections.Counter(r['split'] for r in _gl)
    ck("golden set preserves the user's train/test split and firewalls goldeval",
       _sp.get('train')==456 and _sp.get('test')==87 and _sp.get('goldeval',0)>0, str(dict(_sp)))
    _du=collections.Counter(r['url'] for r in _gl if r['url'])
    ck('golden set has no duplicate URLs', max(_du.values())==1 if _du else False)
    _cf=[r for r in _gl if r['conflict']]
    # two provenance formats are compliant: the arrow ('aug=X->queue=Y', unresolved flag) and the
    # user-adjudicated form 'user-resolved: <label> (was aug=X queue=Y)' (2026-07-03 overnight pass, 36 rows)
    ck('label conflicts are flagged, not silently resolved',
       all(('->' in r['conflict']) or r['conflict'].startswith('user-resolved:') for r in _cf), '%d flagged (user-resolved carries its trail inline)'%len(_cf))
else: ck('golden label set present (run build_golden_labels.py)', False)
if os.path.exists(B+'\\stance_clf_composed.pkl'):
    try:
        import pickle as _pk
        _clf=_pk.load(open(B+'\\stance_clf_composed.pkl','rb'))
        _pr=_clf.predict(['target: Credit Card Mandates. speaker: member. Senator introduces the Credit Card Competition Act to cap interchange.',
                          'target: Credit Unions (direct). speaker: member. Senator celebrates credit union community investment.'])
        ck('domain classifier loads and predicts sane labels', _pr[0]=='oppose' and _pr[1]=='support', str(list(_pr)))
    except Exception as e: ck('domain classifier loads', False, str(e)[:80])
if os.path.exists(B+'\\social_media.db'):
    import sqlite3 as _sq
    _c=_sq.connect(B+'\\social_media.db')
    _ok=_c.execute('PRAGMA integrity_check').fetchone()[0]
    _orph=_c.execute("SELECT COUNT(*) FROM posts WHERE platform IS NULL OR post_id IS NULL").fetchone()[0]
    ck('social_media.db integrity + no orphan posts', _ok=='ok' and _orph==0, '%s / %d orphans'%(_ok,_orph))
_tif=B+r'\tweet_stance_final.csv'
if os.path.exists(_tif):
    with open(_tif,encoding='utf-8-sig',newline='') as _handle:
        _tir=csv.DictReader(_handle); _tifields=set(_tir.fieldnames or []); _tiout=list(_tir)
    ck('tweet stance output preserves image-reader provenance as a distinct signal',
       {'image_read','image_confidence','image_disposition'}.issubset(_tifields),
       'fields=%s'%sorted(_tifields))
    _barr_ids={'511658226','511657485','498264314','458365083'}
    _barr_out={row.get('quorum_id'):row.get('stack_stance') for row in _tiout
               if row.get('quorum_id') in _barr_ids}
    ck('Barr credit-card late-fee tweet gold is applied by exact Quorum ID',
       _barr_out=={qid:'support' for qid in _barr_ids},str(_barr_out))
    _followup_ids={'559107834','559097784','552715581','552700010','545211528','433093684',
                   '591392565','591305285','539036369','536549427','536545303','445704177'}
    _followup_out={row.get('quorum_id'):row.get('stack_stance') for row in _tiout
                   if row.get('quorum_id') in _followup_ids}
    ck('Britt/Lummis user-ratified social evidence is support by exact Quorum ID',
       _followup_out=={qid:'support' for qid in _followup_ids},str(_followup_out))
    _cbdc_out={row.get('quorum_id'):row.get('stack_stance') for row in _tiout
               if row.get('quorum_id') in {'620647802','615403327'}}
    _registry=CO.merged_rows('cu_valence_registry.json','entries',('key',),data_root=B)
    _cbdc=next((row for row in _registry if row.get('key')=='cbdc'),{})
    ck('CBDC polarity is anti-CU and anti-CBDC Davidson posts compose to support',
       _cbdc.get('valence')==-1 and _cbdc_out=={'620647802':'support','615403327':'support'},
       'registry=%s tweets=%s'%(_cbdc,_cbdc_out))
_tirows=CO.merged_rows('tweet_image_annotations.json','annotations',('quorum_id',),data_root=B)
_tibad=sum(1 for row in _tirows if not row.get('quorum_id') or
           row.get('review_status')!='approved' or
           row.get('cu_alignment') not in ('support','oppose','neutral','unrelated') or
           not 0<=float(row.get('confidence',-1))<=1)
ck('approved tweet-image annotations carry exact ID, stance, confidence, and review status',
   _tibad==0,'%d invalid annotations'%_tibad)
_sgr=B+r'\speech_gold_retired.json'
if os.path.exists(_sgr):
    _sgdoc=json.load(open(_sgr)); _sgrows=_sgdoc.get('rows') or []
    ck('retired speech gold is limited to explicit neutral precision negatives',
       _sgdoc.get('schema')=='speech-gold-retired-v1' and
       all(row.get('gold_stance')=='neutral' and row.get('reason') for row in _sgrows),
       '%d retired rows'%len(_sgrows))

print('\n========================================')

# RECEPTIVITY LAYER (2026-07-02): coverage + documented caps.
if os.path.exists(B+r'\member_receptivity.json'):
    _rj=json.load(open(B+r'\member_receptivity.json'))['members']
    ck('receptivity layer covers all members', len(_rj)==533, str(len(_rj)))
    ck('receptivity multipliers within documented caps',
       all(0.30<=v['seed_prob']<=0.90 and 0.50<=v['pressure_mult']<=1.60 and 0.80<=v['edge_mult']<=1.20 for v in _rj.values()))

# W1 guard 2026-07-02: the press-pillar SUMPRODUCT weight row once held stray LET-formula remnants that
# evaluated to 0 on the three TOP-tier categories — every weight must be a positive literal number.
import openpyxl as _px
_wb=_px.load_workbook(B+r'\Moneyball Main Version - Blended.xlsx',data_only=True)
_w535=[_wb['Press Release - Normalized'].cell(535,c).value for c in range(5,24)]
ck('press-pillar weight row 535: 19 numeric positive weights (no formula remnants)',
   all(isinstance(v,(int,float)) and v>0 for v in _w535),str(_w535))
# 2026-07-04 integrity fix guard: the press component must be DYNAMIC (built from the current tone
# csv), never the frozen 07-02 archive (which silently excluded every post-archive article: 197
# members frozen forever). Checks the blend's meta artifact is fresh vs the csv and coverage grew.
_pm_p=B+r'\press_component_meta.json'
if os.path.exists(_pm_p):
    _pm=json.load(open(_pm_p))
    _tone_p=os.path.join(os.path.dirname(B),'Draft Systems','cu_press_stance_tone.csv')
    _fresh=(os.path.getmtime(_tone_p)-_pm.get('built_from_csv_mtime',0))<8*86400 if os.path.exists(_tone_p) else True   # csv may be up to one weekly cycle newer; beyond that the blend stopped ingesting
    ck('media-blend press component is DYNAMIC and fresh (not the frozen archive)',
       _fresh and _pm.get('members_nonzero',0)>=197,
       'members=%s fresh=%s'%(_pm.get('members_nonzero'),_fresh))
    # 2026-07-04: intensity must be NON-DEGENERATE. The first dynamic build clipped |press_value|
    # to a 0.25 floor that 100% of articles saturated -> the component silently became a pure
    # article-count measure (flipped Warren +310 ranks). Guard: recorded pv stats must show real
    # spread (many distinct values, median strictly inside (0, max)).
    _pvs=_pm.get('pv_stats') or {}
    ck('press component intensity non-degenerate (no clip saturation / constant weights)',
       _pvs.get('n_distinct',0)>=20 and 0<_pvs.get('median',0)<_pvs.get('max',0),
       'pv_stats=%s'%_pvs)
    # 2026-07-04 stack wiring guard: when the press_stance_final.csv sidecar exists, the stack must
    # actually drive the signs (>=90% stack vs DEBATE-fallback) and the unrelated->0 mention-noise
    # kill must be active. A silent fall-back to DEBATE (0.571 adversarial-frame) is a regression.
    _sk=_pm.get('stack') or {}
    if _sk.get('sidecar'):
        _tot=_sk.get('stack_signed',0)+_sk.get('debate_fallback',0)
        ck('press stack drives the dynamic component (not silently on DEBATE fallback)',
           _tot>0 and _sk.get('stack_signed',0)/_tot>=0.9 and _sk.get('unrelated_zeroed',0)>0,
           'stack=%s'%_sk)
        ck('user-ratified press gold overrides reach the scoring component',
           _sk.get('gold_overrides',0)>0,'stack=%s'%_sk)
    # 2026-07-16 own-source channel gate: speech, hearing, and interview sidecars were
    # integrity-checked before publication.  Keep those gates observable and fail loud if a
    # later build silently disables a channel, admits learned hearing rows, or double-counts
    # Bluesky (which is already represented in the social component).
    _osc=_pm.get('own_source_channels') or {}
    _validated=('speech','hearing','interview')
    ck('validated own-source press channels are enabled and loaded',
       all((_osc.get(ch) or {}).get('enabled') is True and
           (_osc.get(ch) or {}).get('loaded') is True and
           (_osc.get(ch) or {}).get('signed_gold_rows',
                                    (_osc.get(ch) or {}).get('signed_rows',0))>0
           for ch in _validated),
       'own_source_channels=%s'%_osc)
    _hm=_osc.get('hearing') or {}
    ck('hearing channel remains gold-only',
       _hm.get('tier_policy')=='gold only' and _hm.get('signed_gold_rows',0)>0 and
       _hm.get('non_gold_rows_excluded',0)>0,
       'hearing=%s'%_hm)
    _bm=_osc.get('bluesky') or {}
    ck('Bluesky remains single-counted',
       _bm.get('enabled') is False and _bm.get('loaded') is False and
       'no press-side duplicate' in _bm.get('tier_policy',''),
       'bluesky=%s'%_bm)

    _sidecar_specs=(
        ('speech','speeches_cu.csv','speech_stance_final.csv','speech_gold_queue.csv'),
        ('hearing','hearings_cu.csv','hearing_stance_final.csv','hearing_gold_queue.csv'),
        ('interview','interviews_cu.csv','interview_stance_final.csv','interview_gold_queue.csv'),
    )
    _valid_stances={'support','oppose','neutral'}
    _valid_tiers={'gold','unanimous','majority','review'}
    _sidecars_ok=True; _sidecar_detail=[]
    for _ch,_src_name,_final_name,_gold_name in _sidecar_specs:
        _src_path=os.path.join(B,_src_name); _final_path=os.path.join(B,_final_name)
        _gold_path=os.path.join(B,_gold_name)
        _files_present=all(os.path.exists(p) for p in (_src_path,_final_path,_gold_path))
        _src_rows=[]; _final_rows=[]
        if _files_present:
            with open(_src_path,encoding='utf-8-sig',newline='') as _fh:
                _src_rows=list(csv.DictReader(_fh))
            with open(_final_path,encoding='utf-8-sig',newline='') as _fh:
                _final_rows=list(csv.DictReader(_fh))
        _row_parity=_files_present and len(_src_rows)>0 and len(_src_rows)==len(_final_rows)
        _identity_ok=_row_parity and all(
            all((src.get(field) or '')==(final.get(field) or '')
                for field in ('bioguide','member','date','categories'))
            for src,final in zip(_src_rows,_final_rows)
        )
        _values_ok=_row_parity and all(
            row.get('stance') in _valid_stances and row.get('tier') in _valid_tiers
            for row in _final_rows
        )
        _fresh=_files_present and os.path.getmtime(_final_path)>=max(
            os.path.getmtime(_src_path),os.path.getmtime(_gold_path))
        _channel_ok=_files_present and _row_parity and _identity_ok and _values_ok and _fresh
        _sidecars_ok=_sidecars_ok and _channel_ok
        _sidecar_detail.append('%s:%d/%d identity=%s values=%s fresh=%s'%(
            _ch,len(_src_rows),len(_final_rows),_identity_ok,_values_ok,_fresh))
    ck('own-source channel sidecars preserve source row identity',
       _sidecars_ok,'; '.join(_sidecar_detail))
_pra_p=B+r'\press_domain_routing_audit.json'
if os.path.exists(_pra_p):
    _pra=json.load(open(_pra_p)); _prrows=_pra.get('rows') or []
    _prbad=sum(1 for row in _prrows if not row.get('routing_basis') or
               not isinstance(row.get('effective_categories'),list) or
               ('General Credit Union' in row.get('effective_categories',[]) and
                len(row.get('effective_categories',[]))>1))
    ck('press domain ledger covers every input row with one auditable disposition',
       _pra.get('schema')=='press-domain-routing-audit-v1' and len(_prrows)==2388 and _prbad==0,
       'rows=%d bad=%d'%(len(_prrows),_prbad))
    ck('press routing rejects unsupported declared tags instead of trusting them blindly',
       sum(bool(row.get('discarded_declared')) for row in _prrows)>0,
       'no discarded declared tags recorded')
    _barr_press=[row for row in _prrows if row.get('member_source')=='Andy Barr' and
                 row.get('title')=='House, Senate resolutions would void CFPB credit card late fees rule']
    ck('Barr late-fee CRA press evidence is scored as user-ratified support',
       bool(_barr_press) and all(row.get('gold_override')=='support' for row in _barr_press),
       str(_barr_press[:2]))
else:
    ck('press domain routing audit present',False,'press_domain_routing_audit.json missing')
_lra_p=B+r'\legislation_routing_audit.json'
if os.path.exists(_lra_p):
    _lra=json.load(open(_lra_p)); _lrrows=_lra.get('rows') or []
    _target={('H.R.7733','117'):'CDFI Fund',('S.869','118'):'CDFI Fund',
             ('H.R.7036','118'):'Data Privacy & Cyber',('H.R.7022','117'):'Data Privacy & Cyber',
             ('S.4698','117'):'Data Privacy & Cyber'}
    _lrbad=0
    for (_bill,_congress),_cat in _target.items():
        _seen={row.get('effective_category') for row in _lrrows
               if row.get('bill')==_bill and str(row.get('congress')) in (_congress,_congress+'.0')
               and row.get('effective_category')}
        _lrbad += int(bool(_seen) and _seen!={_cat})
    _ldup=sum((_lra.get('summary') or {}).get('duplicates',{}).values())
    ck('legislation ledger covers every source row and scores one strongest-role event',
       _lra.get('schema')=='legislation-routing-audit-v1' and
       len(_lrrows)==_lra.get('source_rows') and _ldup>0,
       'rows=%d source=%s duplicates=%d'%(len(_lrrows),_lra.get('source_rows'),_ldup))
    ck('pinned CDFI/privacy bills have exactly one effective policy column',
       _lrbad==0 and (_lra.get('summary') or {}).get('effective_categories',{}).get('CDFI Fund',0)>0,
       '%d targeted bill routing errors'%_lrbad)
else:
    ck('legislation routing audit present',False,'legislation_routing_audit.json missing')
# 2026-07-04 nomological fix guard: receptivity must run on the BEHAVIORAL basis (the lexical Big5
# was INDICTED by the pre-registered label-free validation — big5_nomological_prereg.md). A silent
# fall-back to lexical (missing dispositions artifact) is a regression to an indicted instrument.
_rc_p=B+r'\member_receptivity.json'
if os.path.exists(_rc_p):
    _rc=json.load(open(_rc_p))
    _bd_p=B+r'\behavioral_dispositions.json'
    _bd_ok=os.path.exists(_bd_p) and len((json.load(open(_bd_p)) or {}).get('members',{}))>=450
    ck('receptivity runs on the behavioral basis (lexical Big5 indicted 2026-07-04)',
       _rc.get('basis')=='behavioral' and _bd_ok,
       'basis=%s dispositions_ok=%s'%(_rc.get('basis'),_bd_ok))
    # ROUND-7/7b N-slot guard (2026-07-11 promotion, user-ratified, pre-registered gate 5/5): the
    # susceptibility N-slot must remain affective-content N. The 2026-07-12 criterion audit clarifies
    # that E/party residualization is a confound control rather than independent validity evidence; it
    # does not silently overturn the locked promotion. Coverage floor 150.
    _nsl=_rc.get('n_slot'); _nsrc=sum(1 for v in _rc.get('members',{}).values() if v.get('n_source')=='affective')
    ck('receptivity N-slot = affective-content N (round-7/7b; no silent behavioral/text revert)',
       _nsl=='affective' and _nsrc>=150,'n_slot=%s affective_members=%d'%(_nsl,_nsrc))
    # COOP-SLOT guard (audit 2026-07-10, mirrors the N-slot guard): the cooperation-role slot must be the promoted
    # share-based cooperation_v2 (@94f09cf); a silent RECEPT_COOP=v1 revert or a coverage collapse
    # (coop_v2_candidate staleness / c2-floor starvation) would sail through the basis gate above.
    _csl=_rc.get('coop_slot'); _csrc=sum(1 for v in _rc.get('members',{}).values() if v.get('coop_source')=='coop_v2')
    _cv=[v.get('cooperation_v2_z') for v in _rc.get('members',{}).values() if v.get('cooperation_v2_z') is not None]
    ck('receptivity cooperation-role slot = cooperation_v2 (promotion @94f09cf; coverage>=250, |z|<=5, no silent v1 revert)',
       _csl=='v2' and _csrc>=250 and len(_cv)>=250 and all(abs(x)<=5.0 for x in _cv),
       'coop_slot=%s coop_v2_members=%d'%(_csl,_csrc))
    # 2026-07-13 construct-contamination audit (user observed cooperation/A friction): cooperation_v2
    # is a cross-party cosponsorship SHARE; qualified Big-Five A is a package-disjoint low-antagonism
    # interview facet. They must remain separately sourced and must not collapse into one score. The
    # |r|<.30 diagnostic is a contamination alarm (<9% shared variance), not a validity claim.
    _cvp=B+r'\coop_v2_candidate.json'; _avp=B+r'\interview_antagonism.json'
    if os.path.exists(_cvp) and os.path.exists(_avp):
        import numpy as _np, math as _math, csv as _csv
        _cvm=json.load(open(_cvp,encoding='utf-8')).get('members',{}); _avm=json.load(open(_avp,encoding='utf-8')).get('members',{})
        _co_bad=sum(1 for m,d in _cvm.items() if d.get('cooperation_v2_z') is not None and
                    ((_rc.get('members',{}).get(m) or {}).get('coop_source')!='coop_v2' or
                     abs(((_rc.get('members',{}).get(m) or {}).get('cooperation_v2_z') or 0)-d['cooperation_v2_z'])>.002))
        ck('cooperation and low-antagonism A remain separately sourced (zero substitutions)',
           _co_bad==0, '%d cooperation_v2 values substituted/mismatched'%_co_bad)
        _con={d['member']:d for d in json.load(open(B+r'\congestion.json',encoding='utf-8'))}
        _g={d.get('full_name'):d.get('gender') for d in _csv.DictReader(open(DATA_FE+r'\legislators.csv',encoding='utf-8-sig'))}
        _rows=[]
        for _m in set(_cvm)&set(_avm)&set(_con):
            _x=(_cvm.get(_m) or {}).get('cooperation_v2_z')
            if _x is None: continue
            _c=_con[_m]; _p=str(_c.get('party') or '').lower()
            _rows.append([_x,_avm[_m]['agreeableness_z'],1.0 if _p.startswith('r') else 0.0,
                          1.0 if _g.get(_m)=='F' else 0.0,1.0 if _c.get('chamber')=='Senate' else 0.0,
                          _math.log1p(_avm[_m].get('words') or 0)])
        _z=_np.asarray(_rows,float); _raw=float(_np.corrcoef(_z[:,0],_z[:,1])[0,1]) if len(_z)>5 else 9
        _X=_np.c_[_np.ones(len(_z)),_z[:,2:]]
        _rx=_z[:,0]-_X@_np.linalg.lstsq(_X,_z[:,0],rcond=None)[0]; _ry=_z[:,1]-_X@_np.linalg.lstsq(_X,_z[:,1],rcond=None)[0]
        _part=float(_np.corrcoef(_rx,_ry)[0,1]) if len(_z)>5 else 9
        ck('cooperation vs low-antagonism A cross-contamination below alarm threshold',
           len(_z)>=180 and abs(_raw)<.30 and abs(_part)<.30,
           'n=%d raw_r=%+.3f partial_r(party,gender,chamber,interview-volume)=%+.3f'%(len(_z),_raw,_part))
    if _bd_ok:
        _rv=[d['reactivity_z'] for d in json.load(open(_bd_p))['members'].values() if d.get('reactivity_z') is not None]
        ck('reactivity_z disposition sane (coverage >=150, |z|<=5)',
           len(_rv)>=150 and all(abs(v)<=5.0 for v in _rv),'%d covered'%len(_rv))
else:
    # (2026-07-09 fix: this branch previously fired a MISLABELED press-component check name)
    ck('receptivity artifact present (member_receptivity.json)',False,'missing — chain step 1 (receptivity_layer.py) did not run')
# S2-04 guard (audit 2026-07-02): every workbook reader hardcodes range(3,536) for the 533-member roster.
# If a member row is ever added/removed, those readers silently truncate or NUL the tail. Assert on
# R&A col A — TYPED names, immune to the phase-2 openpyxl-save cache death (Blended col A is
# formula-backed and legitimately reads empty between Excel opens).
_roster=len(json.load(open(B+r'\congestion.json')))
_ra_ws=_wb['Rank & Allocation']; _wbrows=sum(1 for r in range(3,536) if _ra_ws.cell(r,1).value)
ck('workbook member rows == roster size (no silent row drift)', _wbrows==_roster, 'workbook %d vs roster %d'%(_wbrows,_roster))
# Roadmap-#9 phase 1 (2026-07-02): the whole raw->normalized->pillar->score->rank chain is reproduced
# in Python (compute_pillars.py) and must match the sheet EXACTLY — catches any silent formula edit.
os.environ['PILLAR_SNAPSHOT_FREEZE']='1'   # the audit must NEVER advance the pillar snapshot vintage (2026-07-17 guard)
try:
    _pp=run('compute_pillars.py')
finally:
    os.environ.pop('PILLAR_SNAPSHOT_FREEZE',None)
ck('python pillar shadow-compute parity (rank math reproducible, 533/533)','PARITY: CLEAN' in _pp,
   _pp.strip().splitlines()[-1] if _pp else 'no output')
# rank-integrity guards (2026-07-17): (a) the pillar snapshot must not silently serve a stale
# leg/inf/demo basis (the 2026-07-16 522-rank incident class); (b) the own-source channel inputs
# must not silently freeze while their channels are ON (extractors are not chain-wired yet —
# collection runs on Paul's 3-week cadence, so the ceiling is cadence+grace = 28 days).
_snp=B+r'\pillar_inputs_snapshot.json'
_sn_stamp=(json.load(open(_snp,encoding='utf-8')).get('stamped') if os.path.exists(_snp) else None)
if _sn_stamp:
    import datetime as _dt2
    _sn_age=(_dt2.datetime.now().astimezone()-_dt2.datetime.fromisoformat(_sn_stamp)).days
    ck('pillar snapshot vintage <= 14 days (stale-basis tripwire)', _sn_age<=14, 'stamped %s (%dd)'%(_sn_stamp,_sn_age))
else:
    ck('pillar snapshot vintage <= 14 days (stale-basis tripwire)', True, 'unstamped pre-guard snapshot (stamps on next compute run)')
_dbm=os.path.getmtime(B+r'\social_media.db') if os.path.exists(B+r'\social_media.db') else None
_chstale=[]
for _cf in ('speeches_cu.csv','hearings_cu.csv','interviews_cu.csv','bluesky_cu.csv'):
    _p2=B+'\\'+_cf
    if _dbm and os.path.exists(_p2) and (_dbm-os.path.getmtime(_p2))>28*86400:
        _chstale.append('%s (%.0fd behind corpus)'%(_cf,(_dbm-os.path.getmtime(_p2))/86400))
ck('own-source channel inputs within 28 days of the corpus (channels are default-ON)', not _chstale, str(_chstale))

if os.path.exists(B+r'\stance_training.jsonl'):
    import hashlib as _hl, json as _js
    _tr=set(); _te=set()
    for _l in open(B+r'\stance_training.jsonl',encoding='utf-8'):
        _r=_js.loads(_l)
        _vs=[_r.get('text',''),_r.get('title',''),_r.get('stance_span','')]
        _hs={_hl.md5(v.strip().lower()[:400].encode()).hexdigest() for v in _vs if v and v.strip()}
        if _r['split']=='train': _tr|=_hs
        elif _r['split']=='test': _te|=_hs
    _ov=len(_tr&_te)
    ck('train/test split-leak (text-hash, incl. augment views) within known bound', _ov<=4,
       '%d collisions (4 known: same-ACU-article rows 75/214; fix = same-text same-split on next label refresh)'%_ov)

# run-5b guards (audit 2026-07-06): the psychology/media-layer artifacts had ZERO audit coverage —
# a truncated big5_llm_opus.json or malformed reference/co-appearance edges would sail through.
# Schema + sanity floors only (freshness is meaningless mid-collector-crawl); each skips if absent.
if os.path.exists(B+r'\big5_llm_opus.json'):
    _b5=json.load(open(B+r'\big5_llm_opus.json',encoding='utf-8')).get('traits',{})
    _bad=sum(1 for d in _b5.values() if not all(k in d for k in ('O','C','E','A','ES')) or not d.get('n_chunks'))
    ck('big5_llm artifact schema (5 traits + n_chunks per member, coverage >=400)',
       len(_b5)>=400 and _bad==0, '%d members, %d malformed'%(len(_b5),_bad))
if os.path.exists(B+r'\member_references.json'):
    _re_=json.load(open(B+r'\member_references.json',encoding='utf-8')).get('edges',[])
    _bad=sum(1 for e in _re_ if not all(k in e for k in ('a','b','n_raw','n_scored','pos','neg','net')) or e['n_scored']<(e['pos']+e['neg']))
    ck('reference-layer edge schema (keys present, scored>=pos+neg, >=1000 edges)',
       len(_re_)>=1000 and _bad==0, '%d edges, %d malformed'%(len(_re_),_bad))
# run-5h guard (user 2026-07-06 'make sure the displays are entirely accurate'): the member-desk
# pillar bars + weights MUST match the rank authority (ranks_python.json) — they silently drifted
# (Crapo influence 1.000 in the authority displayed as 0.222 from a stale workbook column; weights
# 0.40/0.20/0.35/0.05 vs the live 0.25/0.30/0.40/0.05), making the #1 rank look unjustified on the
# desk. This asserts the DISPLAYED decomposition reconstructs the DISPLAYED score for every member.
if os.path.exists(B+r'\member_desk_bundle.json') and os.path.exists(B+r'\ranks_python.json'):
    _bun=json.load(open(B+r'\member_desk_bundle.json',encoding='utf-8')); _ra=json.load(open(B+r'\ranks_python.json',encoding='utf-8'))
    _W=_bun.get('weights',{}); _pcts=_ra.get('pcts',{})
    _bad=[]; _nomatch=0
    for _m in _bun.get('members',[]):
        _nm=_m.get('name'); _p=_pcts.get(_nm)
        if not _p: continue
        # displayed pillars == authority pcts
        if any(abs((_m.get(_dk) or 0)-_p[_ak])>1e-4 for _dk,_ak in (('press_n','press'),('leg_n','leg'),('comm_n','inf'),('demo_n','dem'))):
            _nomatch+=1
        # displayed weights*displayed pillars reconstruct displayed score
        _rc=_W.get('leg',0)*(_m.get('leg_n') or 0)+_W.get('press',0)*(_m.get('press_n') or 0)+_W.get('comm',0)*(_m.get('comm_n') or 0)+_W.get('demo',0)*(_m.get('demo_n') or 0)
        if abs(_rc-(_m.get('score') or 0))>0.003: _bad.append(_nm)
    ck('desk pillar bars match rank authority (no stale-workbook drift)', _nomatch==0, '%d members mismatch authority pcts'%_nomatch)
    ck('desk displayed pillars*weights reconstruct displayed score', len(_bad)==0, '%d members fail reconstruction: %s'%(len(_bad),_bad[:3]))
    # Paul-directed fill contract (2026-07-17): every member renders a Big-Five panel — gold-anchored
    # values primary, labeled estimate layer where gold is absent; estimate slots must carry a
    # provenance label. Guard-only display returns via DESK_B5_FALLBACK=0 (then this check adapts).
    _fill_on=os.environ.get('DESK_B5_FALLBACK','1')!='0'
    if _fill_on:
        _b5empty=[]; _fb_unlabeled=0
        for _m in _bun.get('members',[]):
            _p=_m.get('psy') or {}
            _g=(_p.get('big5_pct') or [None]*5); _f=(_p.get('big5_fb_pct') or [None]*5); _s=(_p.get('big5_fb_src') or [None]*5)
            _mrg=[(_g[_i] if _i<len(_g) and _g[_i] is not None else (_f[_i] if _i<len(_f) else None)) for _i in range(5)]
            if not any(_v is not None for _v in _mrg): _b5empty.append(_m.get('name'))
            for _i in range(5):
                if (_i>=len(_g) or _g[_i] is None) and _i<len(_f) and _f[_i] is not None and not (_i<len(_s) and _s[_i]):
                    _fb_unlabeled+=1
        ck('every member renders a Big-Five panel (gold or labeled estimate; Paul fill 2026-07-17)',
           not _b5empty, str(_b5empty[:6]))
        ck('every estimate slot carries a provenance label', _fb_unlabeled==0, '%d unlabeled'%_fb_unlabeled)
    _fr=json.load(open(B+r'\formal_visibility_roles.json',encoding='utf-8'))
    import influence_roles as _ir
    _frm=_fr.get('members') or {}; _rb=_ra.get('formal_role_influence_bonus') or {}
    _rd=_ra.get('formal_role_display_weight') or {}
    _rn=_ra.get('formal_role_influence_nominal') or {}; _rr=_ra.get('formal_role_influence_eligibility') or {}
    _rbase=_ra.get('influence_raw_base') or {}; _radj=_ra.get('influence_raw_adjusted') or {}
    _cseats={r[0]:str(r[4] or '') for r in _wb['Committee Data'].iter_rows(min_row=2,values_only=True) if r[0]}
    _desk_by={m.get('name'):m for m in _bun.get('members',[])}
    _rolebad=0
    for _nm in _ra.get('rank',{}):
        _entry=_frm.get(_nm) or {}; _nom=float(_entry.get('influence_bonus') or 0.0)
        _want,_reason,_overlap=_ir.eligible_bonus(_entry,_cseats.get(_nm,''))
        _rolebad += int(abs(float(_rn.get(_nm,999))-_nom)>1e-9 or _rr.get(_nm)!=_reason)
        _rolebad += int(abs(float(_rb.get(_nm,999)))>1e-9)
        _rolebad += int(abs(float(_rd.get(_nm,999))-_want)>1e-9)
        _rolebad += int(abs(float(_radj.get(_nm,999))-float(_rbase.get(_nm,0)))>1e-9)
        _dm=_desk_by.get(_nm) or {}
        _rolebad += int(abs(float(_dm.get('formal_role_influence_bonus',999)))>1e-9)
        _rolebad += int(abs(float(_dm.get('formal_role_display_weight',999))-_want)>1e-9)
        _rolebad += int(abs(float(_dm.get('comm_raw',999))-float(_radj.get(_nm,0)))>1e-9)
    _rmeta=_ra.get('formal_role_influence') or {}
    ck('formal office remains titled display context and never changes Influence or rank',
       _fr.get('schema')=='party-elite-influence-roles-v3' and
       _fr.get('valid_until','')>__import__('datetime').date.today().isoformat() and
       len(_frm)>=20 and 'Chuck Grassley' not in _frm and _rolebad==0 and
        _rb.get('Mike Johnson')==0.0 and _rb.get('Tim Scott')==0.0 and _rb.get('John Barrasso')==0.0 and
        _rd.get('Mike Johnson')==3.0 and _rmeta.get('applied_party_elites')==0 and
        _rmeta.get('listed_party_elites')==len(_frm) and _rmeta.get('display_context_party_elites',0)>0,
        '%d listed party elites, %d rank-applied, %d display contexts, %d decomposition errors'%(
        len(_frm),_rmeta.get('applied_party_elites',0),_rmeta.get('display_context_party_elites',0),_rolebad))
    # Blue/Pink remain independently constructed. The equal-weight E remains a frozen comparator;
    # the active OOF-gated E is audited below with dominance and role-adjusted visibility included.
    _psy=[m.get('psy') or {} for m in _bun.get('members',[])]
    _bp=[p for p in _psy if len(p.get('big5_pct') or [])==5 and len(p.get('llm_pct') or [])==5]
    _same_n=sum(1 for p in _bp if p['big5_pct'][4]==p['llm_pct'][4])
    ck('desk broad-gold and raw Pink Neuroticism reads do not collapse to one value',
       len(_bp)>=400 and _same_n<len(_bp), '%d paired members, %d identical N percentiles'%(len(_bp),_same_n))
    _blue=next((_p for _p in (B+r'\big5_blue_validated.json',B+r'\big5_blue_emb.json',B+r'\big5_v2.json') if os.path.exists(_p)),B+r'\big5_v2.json')
    _pink=B+r'\big5_llm_opus.json'
    if os.path.exists(_blue) and os.path.exists(_pink) and os.path.getmtime(_pink)<os.path.getmtime(_blue):
        _unflag=sum(1 for p in _bp if not ((p.get('big5_vintage') or {}).get('pink_older_days')>0))
        ck('desk preserves mixed Blue/Pink corpus-vintage metadata for archived diagnostics',
           _unflag==0, '%d/%d paired members lack the stale-vintage disclosure metadata'%(_unflag,len(_bp)))
    _egp=B+r'\big5_e_gold_anchored.json'; _evp=B+r'\big5_external_gold_validation.json'
    if os.path.exists(_egp) and os.path.exists(_evp):
        _eg=json.load(open(_egp,encoding='utf-8')); _ev=json.load(open(_evp,encoding='utf-8'))
        _egm=_eg.get('members') or {}; _epr=_ev.get('panel_results') or {}
        _ebad=sum(1 for d in _egm.values() if d.get('E') is None or d.get('percentile') is None or
                  d.get('channel_weights')!={'blue_visibility':.5,'pink_observer':.5})
        ck('external Senate gold preserves the discriminant equal-weight Extraversion comparator',
           _eg.get('schema')=='big5-e-gold-anchored-v1' and _eg.get('status')=='ACTIVE_EXTERNAL_GOLD_ANCHORED' and
           _ev.get('verdict')=='PASS' and len(_egm)>=450 and _ebad==0 and
           all((d.get('e_r') or 0)>=.30 and d.get('diagonal_dominant') and (d.get('e_permutation_p_one_sided') or 1)<.05
               for d in _epr.values()),
           '%d E members, %d malformed; panel r=%s'%(
           len(_egm),_ebad,{k:round(v.get('e_r',0),3) for k,v in _epr.items()}))
    _mgp=B+r'\big5_gold_anchored.json'
    if os.path.exists(_mgp):
        _mg=json.load(open(_mgp,encoding='utf-8')); _mgt=_mg.get('traits') or {}; _mge=_mgt.get('E') or {}
        _mgem=_mge.get('members') or {}; _mgew=_mge.get('selected_weights') or {}
        _mgebad=sum(1 for d in _mgem.values() if d.get('E') is None or d.get('percentile') is None)
        _mgpan=(_mge.get('validation') or {}).get('panels') or {}
        ck('multitrait OOF Senate gold activates E with dominance and role-adjusted visibility',
           _mg.get('schema')=='big5-gold-anchored-v2' and set(_mgt)=={'O','C','E','A','N'} and
           _mge.get('status')=='ACTIVE_EXTERNAL_GOLD_ANCHORED' and len(_mgem)>=450 and _mgebad==0 and
           _mgew.get('interaction_dominance',0)>=.1 and _mgew.get('role_adjusted_visibility',0)>=.1 and
           all(d.get('diagonal_dominant') and (d.get('r') or 0)>=.2 and
               (d.get('permutation_p_one_sided') or 1)<.1 for d in _mgpan.values()) and
           all((d.get('validation') or {}).get('predictions_are_out_of_fold') for d in _mgt.values()),
           '%d E members, %d malformed; weights=%s; panel r=%s'%(
           len(_mgem),_mgebad,_mgew,{k:round(v.get('r',0),3) for k,v in _mgpan.items()}))
        _expected_status={'O':'ACTIVE_ONE_PANEL_ASSOCIATION_SUPPORTED',
                          'C':'ACTIVE_ONE_PANEL_ASSOCIATION_SUPPORTED',
                          'E':'ACTIVE_EXTERNAL_GOLD_ANCHORED',
                          'A':'ACTIVE_EXTERNAL_GOLD_ANCHORED',
                          'N':'ACTIVE_ONE_PANEL_ASSOCIATION_SUPPORTED'}
        ck('all broad Big Five traits activate with explicit panel-strength labels',
           {k:v.get('status') for k,v in _mgt.items()}==_expected_status and
           all(any((v.get('association_gates') or {}).values()) for v in _mgt.values()) and
           all(len(v.get('members') or {})>=300 for v in _mgt.values()),
           'statuses=%s; association gates=%s'%(
           {k:v.get('status') for k,v in _mgt.items()},
           {k:v.get('association_gates') for k,v in _mgt.items()}))
    # Source-primary Blue is the completed Round-4 promotion, not the pooled lexical v2 draft.
    _bep=B+r'\big5_blue_emb.json'
    if os.path.exists(_bep):
        _be=json.load(open(_bep,encoding='utf-8')); _bem=_be.get('traits',{}); _bev=_be.get('validation') or {}
        _bars=_bev.get('bars') or {}; _bad=sum(1 for d in _bem.values() if
             not all(k in d for k in ('O','C','E','A','ES','N','tokens','n_source')) or
             abs(float(d['ES'])+float(d['N']))>.002 or d.get('n_source')!='affective_n')
        ck('source-primary Blue promotion preserves frozen MTMM/criterion gate and affective-N slot',
           _bev.get('verdict')=='PASS' and all(_bars.values()) and
           _bev.get('mtmm_same_trait_cross_genre',0)>1.15*_bev.get('mtmm_cross_trait_within_genre',1) and
           set(_be.get('source_genres') or [])=={'CREC','INTERVIEW','social'} and
           len(_bem)>=450 and _bad==0,
           '%d members, %d malformed, verdict=%s, MTMM=%s/%s'%
           (len(_bem),_bad,_bev.get('verdict'),_bev.get('mtmm_same_trait_cross_genre'),
            _bev.get('mtmm_cross_trait_within_genre')))
    _bvp=B+r'\big5_blue_validated.json'
    if os.path.exists(_bvp):
        _bv=json.load(open(_bvp,encoding='utf-8')); _bvm=_bv.get('traits',{}); _vv=_bv.get('validation') or {}; _rr=_bv.get('reliability') or {}
        _diag=_vv.get('diagonal') or {}; _known=_vv.get('known_groups') or {}; _rt=_rr.get('traits') or {}
        _bad=sum(1 for d in _bvm.values() if not all(k in d for k in ('O','C','E','A','N','ES','tokens')) or
                 (d.get('N') is not None and abs(float(d['ES'])+float(d['N']))>.002))
        ck('validated multimethod Big Five preserves all 5 reliability/known-group/discriminant gates',
           len(_bvm)>=450 and _bad==0 and _vv.get('diagonal_pass_n')==5 and _vv.get('known_groups_pass_n')==5 and
           len(_diag)==5 and all(d.get('pass') for d in _diag.values()) and
           len(_known)==5 and all(d.get('pass') for d in _known.values()) and
           _rr.get('pass_n')==5 and _rr.get('package_overlap')==0 and len(_rt)==5 and all(d.get('pass') for d in _rt.values()),
           '%d members, %d malformed, reliability=%s, known=%s, diagonal=%s, package overlap=%s'%
           (len(_bvm),_bad,_rr.get('pass_n'),_vv.get('known_groups_pass_n'),_vv.get('diagonal_pass_n'),_rr.get('package_overlap')))
        _ocebad=sum(1 for d in _bvm.values() if any(d.get(k) is None for k in ('O','C','E')))
        _ncovered=sum(1 for d in _bvm.values() if d.get('N') is not None)
        _coverage=_vv.get('coverage') or {}
        ck('independent O/C/E coverage is not truncated by the narrower affective-N stratum',
           _ocebad==0 and len(_bvm)>_ncovered and
           all(_coverage.get(k)==len(_bvm) for k in ('O','C','E')) and _coverage.get('N')==_ncovered,
           '%d text-basis members, %d O/C/E gaps, %d N-covered; recorded coverage=%s'%
           (len(_bvm),_ocebad,_ncovered,_coverage))
    # Round-9 qualified A facet remains separate from the cooperation engine role.
    _avp=B+r'\interview_antagonism.json'
    if os.path.exists(_avp):
        _av=json.load(open(_avp,encoding='utf-8')); _avm=_av.get('members',{})
        _mal=sum(1 for d in _avm.values() if not all(k in d for k in
                 ('antagonism_z','agreeableness_z','raw_antagonism','half_a','half_b','words'))
                 or abs(d['agreeableness_z']+d['antagonism_z'])>.002)
        ck('validated A-antagonism facet preserves qualified schema and inverse pole',
           _av.get('n')==len(_avm) and len(_avm)>=300 and _av.get('reliability_sb',0)>=.65 and _mal==0
           and (_av.get('source') or {}).get('sha256')=='15FFB00BA199708C17733260205E0580CE766F3B37E0C69A74965B35B4AB0575'
           and 'residualization-only' in ((_av.get('controls') or {}).get('role') or ''),
           '%d members, SB=%s, %d malformed/non-inverse'%
           (len(_avm),_av.get('reliability_sb'),_mal))
        _desk_av={m.get('name'):(m.get('psy') or {}).get('a_validated_pct')
                  for m in _bun.get('members',[]) if (m.get('psy') or {}).get('a_validated_pct') is not None}
        ck('validated A facet propagates to desk without replacing engine/text A',
           len(_desk_av)==len(_avm) and all(0<=v<=1 for v in _desk_av.values())
           and all((m.get('psy') or {}).get('a_validated_scope')=='interview/media-covered stratum only'
                   for m in _bun.get('members',[]) if (m.get('psy') or {}).get('a_validated_pct') is not None),
           '%d artifact members, %d desk values'%(len(_avm),len(_desk_av)))
        _scope=_av.get('scope') or {}; _sel=_av.get('selection_audit') or {}
        ck('validated A claim is sampling-frame scoped (no full-roster percentile/imputation)',
           _scope.get('covered_n')==322 and _scope.get('roster_n')==533
           and _scope.get('imputation')=='forbidden'
           and _sel.get('verdict')=='DEMOGRAPHICALLY SELECTED MEDIA-COVERED STRATUM',
           'scope=%s selection=%s'%(_scope,_sel.get('verdict')))
        _nar=[(m.get('psy') or {}).get('narrative') or {} for m in _bun.get('members',[])]
        _nar_present=[n for n in _nar if n]
        _withheld=sum(1 for m in _bun.get('members',[]) if
                      (m.get('psy') or {}).get('narrative_status')=='withheld_pending_guarded_local_rebuild_and_user_approval')
        _stale=sum(1 for n in _nar if n and 'double-reviewed' not in (n.get('measurement_note') or '').lower())
        _safe=(not _nar_present and _withheld==533) or (len(_nar_present)==533 and _stale==0)
        ck('current narratives cannot silently fall back to legacy E/audio prose', _safe,
           '%d guarded desk narratives, %d explicitly withheld, %d stale provenance notes'%
           (len(_nar_present),_withheld,_stale))
# 2026-07-13 personal-texture catch-net: this display-only artifact must cover every member for whom
# explicit autobiographical evidence was found, retain exact provenance, and propagate byte-for-field
# into the member desk. Coverage must retain explicit evidence for >=90% of candidate-bearing members;
# this ratio gate rewards precision without forcing false positives merely to hit a roster count.
# No low-confidence or source-less inference may enter a profile.
_ptp=B+r'\member_personal_texture.json'
if os.path.exists(_ptp):
    _pt=json.load(open(_ptp,encoding='utf-8')); _ptm=_pt.get('members',{}); _pp=_pt.get('provenance') or {}
    _bad=0; _ni=0; _local_narrative_leak=0
    for _m,_d in _ptm.items():
        _it=_d.get('items') or []; _ni+=len(_it)
        if _d.get('n_items')!=len(_it): _bad+=1
        _bad+=sum(1 for x in _it if not all(x.get(k) not in (None,'') for k in
                 ('category','claim','confidence','evidence','source','classifier')) or x.get('confidence')=='low'
                 or not isinstance(x.get('narrative_eligible'),bool))
        _local_narrative_leak+=sum(1 for x in _it if str(x.get('classifier','')).startswith('local_explicit_')
                                   and x.get('narrative_eligible') is not False)
    ck('personal-texture artifact precision/schema/provenance (no volume reward)',
       len(_ptm)>=.90*_pp.get('candidate_member_count',9999) and _ni>0 and _bad==0 and _local_narrative_leak==0 and
       _pp.get('candidate_count',0)>=_pp.get('semantic_gate_count',1) and
       _pp.get('classified_count',0)>=_pp.get('semantic_gate_count',1),
       '%d members, %d details, %d malformed, %d local-to-narrative leaks, candidates=%s classified=%s gate=%s'%
       (len(_ptm),_ni,_bad,_local_narrative_leak,_pp.get('candidate_count'),_pp.get('classified_count'),_pp.get('semantic_gate_count')))
    if os.path.exists(B+r'\member_desk_bundle.json'):
        _dm={m.get('name'):m.get('personal') for m in json.load(open(B+r'\member_desk_bundle.json',encoding='utf-8')).get('members',[])}
        _mis=sum(1 for m,d in _ptm.items() if _dm.get(m)!=d)
        ck('personal-texture layer propagates exactly to member profiles (display-only)',
           _mis==0 and sum(1 for d in _dm.values() if d)==len(_ptm),
           '%d payload mismatches; artifact=%d desk=%d'%(_mis,len(_ptm),sum(1 for d in _dm.values() if d)))
else:
    ck('personal-texture artifact present (member_personal_texture.json)',False,
       'missing — run personal_texture_layer.py --score --go or --local, then rebuild the desk')
# run-5j guard (2026-07-06): exact-duplicate speech rows (same member+package+text) inflate word
# counts + big5 chunk sampling. The CREC backfill introduced 922 (no pre-insert dedup); collector
# legs now carry stored-package guards. Assert the corpus stays exact-dup-clean (threshold covers
# the small user-gated legacy ABC/CNN set, currently 0 post-dedup).
if os.path.exists(B+r'\social_media.db'):
    import sqlite3 as _sq
    _c=_sq.connect(B+r'\social_media.db',timeout=60); _c.execute('PRAGMA query_only=1')
    _ex=_c.execute("SELECT COALESCE(SUM(c-1),0) FROM (SELECT COUNT(*) c FROM speeches GROUP BY member,package,text HAVING c>1)").fetchone()[0]
    _c.close()
    ck('speech corpus exact-duplicate free (member+package+text)', _ex<=50, '%d exact-dup surplus rows'%_ex)
if os.path.exists(B+r'\co_appearance.json'):
    _co=json.load(open(B+r'\co_appearance.json',encoding='utf-8')).get('edges',[])
    _bad=sum(1 for e in _co if not all(k in e for k in ('a','b','n','last','bipartisan')) or e['n']<1)
    ck('co-appearance edge schema (keys present, n>=1, non-empty)', len(_co)>=100 and _bad==0,
       '%d pairs, %d malformed'%(len(_co),_bad))
# 2026-07-13 production audio guard: zero verified members is a valid fail-closed state. Every
# accepted clip must have two distinct, dated, method-identified target-only reviews, and no display
# producer may fall back to the retired model-only verifier or legacy 88-feature cache.
from audio_rules import MIN_MEMBER_TARGET_SEC as _MIN_AUDIO_MEMBER_SEC
_ar=json.load(open(B+r'\b5_audio_reviews_v2.json',encoding='utf-8'))
_au=json.load(open(B+r'\b5_audio_descriptors_v2.json',encoding='utf-8'))
_abad=0
for _rel,_clip in (_au.get('members') or {}).items():
    if (_clip.get('target_elapsed_sec',0)<_MIN_AUDIO_MEMBER_SEC or _clip.get('source_n',0)<2 or
        _clip.get('program_n',0)<2 or not _clip.get('cross_program')): _abad+=1
    for _row in _clip.get('clips') or []:
        if (_row.get('venue')!='on_air_interview' or not _row.get('outlet') or
            not _row.get('program') or not _row.get('aired_at') or not _row.get('recording_context') or
            _row.get('target_words',0)<50 or _row.get('pitch_deviation_from_speaker_baseline_z') is None or
            not (_row.get('topic') or {}).get('id')): _abad+=1
for _rel,_clip in (_au.get('rejected') or {}).items():
    if not isinstance(_clip,list) or not _clip: _abad+=1
for _rel,_clip in ((_ar.get('clips') or {}).items()):
    _revs=_clip.get('reviews') or []; _ids=[r.get('reviewer') for r in _revs]
    if ((_clip.get('venue')!='on_air_interview' or not _clip.get('outlet') or not _clip.get('program') or
         not _clip.get('aired_at') or not _clip.get('recording_context') or not _clip.get('speech_id') or
         not (_clip.get('topic') or {}).get('id')) or
        len(_revs)!=2 or len(set(_ids))!=2 or any(r.get('identity')!='TARGET' or
        r.get('speaker')!='TARGET_ONLY' or not r.get('reviewed_at') or not r.get('review_method') or
        r.get('transcript_alignment_confirmed') is not True or r.get('interviewer_excluded') is not True
        or r.get('speech_unit_confirmed') is not True or r.get('topic_from_aligned_transcript_confirmed') is not True
        for r in _revs)): _abad+=1
_acon=_ar.get('contract') or {}; _mcon=_au.get('measurement_contract') or {}; _dcon=_au.get('deployment_contract') or {}
ck('audio v2 fail-closed on-air, aligned, program-conditioned contract active',
   _ar.get('schema')=='b5-audio-review-v2' and _au.get('schema')=='b5-audio-descriptors-v2' and
   _au.get('status')=='ACTIVE_FAIL_CLOSED' and _dcon.get('unverified_audio_enters_profiles') is False and
   _dcon.get('on_air_interviews_only') is True and _dcon.get('program_conditioned_percentiles') is True and
   _dcon.get('within_speaker_pitch_deviation') is True and _dcon.get('transcript_identified_topics') is True and
    _dcon.get('audio_may_modify_big_five') is False and _acon.get('transcript_alignment_and_interviewer_exclusion_double_confirmed') is True and
    _acon.get('speech_unit_and_transcript_topic_double_confirmed') is True and
    _acon.get('minimum_confirmed_target_seconds_per_member')==_MIN_AUDIO_MEMBER_SEC and
    _acon.get('minimum_independent_programs_per_member')==2 and
    _acon.get('minimum_speech_units_per_member')==5 and _acon.get('minimum_distinct_topics_per_member')==2 and
    _mcon.get('minimum_target_seconds_per_member')==_MIN_AUDIO_MEMBER_SEC and
    _mcon.get('minimum_speech_units_per_member')==5 and _mcon.get('minimum_distinct_topics_per_member')==2 and
    _mcon.get('minimum_aligned_target_words_per_clip')==50 and
   _mcon.get('percentile_normalization')=='within outlet+program; never raw cross-speaker pitch/energy' and _abad==0,
   '%d accepted clips, %d verified members, %d contract errors'%(_au.get('accepted_clip_n',0),len(_au.get('members') or {}),_abad))
_ncode=open(FE+r'\b5_narratives.py',encoding='utf-8').read()
_dcode=open(FE+r'\make_member_desk.py',encoding='utf-8').read()
_wcode=open(FE+r'\weekly_refresh.py',encoding='utf-8').read()
ck('legacy audio and paid narrative paths retired from unattended/display consumption',
   'b5_audio_cache.json' not in _ncode and 'b5_audio_verify.py' not in _wcode and
   'b5_audio_build.py' not in _wcode and 'b5_narratives.py' not in _wcode and
   '--allow-paid-api' in _ncode and 'member_narratives.json' not in _dcode,
   'weekly/display code must consume guarded v2 artifacts only')
_ci=json.load(open(B+r'\corpus_integrity_report.json',encoding='utf-8'))
_cidb=_ci.get('database') or {}; _cic=_ci.get('accepted_model') or {}; _ciq=_ci.get('quarantined_chrg') or {}
ck('corpus integrity gate protects accepted scores and verifies repaired CHRG attribution',
   _ci.get('schema')=='corpus-integrity-v1' and _ci.get('status')=='PASS' and
   not _ci.get('hard_errors') and _cic.get('cache_exact_match') is True and _cic.get('cache_content_match') is True and
   _cidb.get('exact_duplicate_surplus')==0 and _ciq.get('enters_default_big5') is False and
   (_ciq.get('service_chamber_mismatch_rows',0)+_ciq.get('no_member_service_rows',0))==0 and
   _ciq.get('repair_status')=='source-rebuilt and verified',
   'status=%s hard=%s cache=%s invalid_chrg=%d'%
   (_ci.get('status'),_ci.get('hard_errors'),_cic.get('cache_exact_match'),
    _ciq.get('service_chamber_mismatch_rows',0)+_ciq.get('no_member_service_rows',0)))
_bpanel=json.load(open(B+r'\behavior_panel.json',encoding='utf-8'))
_bdisp=json.load(open(B+r'\behavioral_dispositions.json',encoding='utf-8'))
ck('corrupt legacy CHRG hearing counts cannot re-enter behavioral diligence',
   _bdisp.get('schema')=='behavioral-dispositions-v2' and _bdisp.get('status')=='PASS' and
   'chrg_hearings' not in (_bpanel.get('outcomes') or []) and
   all('chrg_hearings' not in (row or {}) for row in (_bpanel.get('members') or {}).values()) and
   'diligence=roll-call participation+LEP follow-through' in (_bdisp.get('method') or '') and
   'diligence=roll-call participation+LEP follow-through+CHRG' not in (_bdisp.get('method') or '') and
   ((_bdisp.get('integrity') or {}).get('diligence') or {}).get('legacy_chrg_enters_diligence') is False and
   (_bdisp.get('coverage') or {}).get('diligence_z',0)>=525 and
   len(((_bdisp.get('inputs_sha256') or {}).get('lep_criterion.json') or ''))==64,
   'behavior panel/disposition method must keep quarantined CHRG out of C')
# voice-ES display artifact RETIRED 2026-07-12 (user): the +0.208 look was gender-confounded (locked retest
# partial r -0.03). The desk shows behavioral reactivity on the role radar and independently ranked
# Blue/Pink N=-ES on the text radar; no voice-ES schema check remains. Guard against a silent
# regression to the retired artifact:
ck('voice-ES retired from display (b5_es_voice not re-wired into the desk)',
   'b5_es_voice' not in open(FE+r'\make_member_desk.py',encoding='utf-8').read().split('RETIRED')[-1],
   'desk must not consume the gender-confounded voice-ES artifact')
_desk_tpl=open(FE+r'\templates\make_member_desk_00.tpl',encoding='utf-8').read()
ck('desk presents the accepted five-axis display under canonical Big-Five names',
   all(s in _desk_tpl for s in ('Extraversion','Openness','Conscientiousness',
                                'Agreeableness','Neuroticism',
                                'Gold-validated: gold-anchored axes use the OOF Senate-gold composites')) and
   all(len((m.get('psy') or {}).get('big5_pct') or [])==5 for m in _bun.get('members',[])
       if (m.get('psy') or {}).get('big5_pct')),
   'accepted display must carry five canonically named axes and disclose the E scope')
_comm=json.load(open(B+r'\communication_style.json',encoding='utf-8'))
_csrc=_comm.get('members') or {}
_gdoc=json.load(open(B+r'\big5_gold_anchored.json',encoding='utf-8')); _gtraits=_gdoc.get('traits') or {}
_egsrc=((_gtraits.get('E') or {}).get('members') or {})
_gorder=('E','O','C','A','N')
_cbad=0; _egbad=0; _gbad=0
for _row in _bun.get('members',[]):
    _name=_row.get('name'); _psy=_row.get('psy') or {}; _src=_csrc.get(_name) or {}
    if _src:
        _got=_psy.get('communication_style') or {}
        _adj=(_src.get('public_exposure') or {}).get('role_adjusted_percentile')
        _act=((_src.get('facets') or {}).get('activity') or {}).get('percentile')
        _cbad += int((_got.get('public_exposure') or {}).get('role_adjusted_percentile') != _adj)
        _cbad += int(_psy.get('e_public_proxy_pct') != _adj)
        # Public exposure propagates independently.  The Big-Five E slot is the OOF-gated
        # four-channel composite where both primary channels exist and remains missing otherwise.
        if _psy.get('big5_pct'):
            _egbad += int(_psy['big5_pct'][0] != (_egsrc.get(_name) or {}).get('percentile'))
            _gbad += sum(int(_psy['big5_pct'][i] != (((_gtraits.get(trait) or {}).get('members') or {}).get(_name) or {}).get('percentile'))
                         for i,trait in enumerate(_gorder))
        _cbad += int(((_psy.get('behav_pct') or [None,None,None,None])[3]) != _act)
_mike=_csrc.get('Mike Johnson') or {}; _mexp=_mike.get('public_exposure') or {}
ck('observable communication facets and role-adjusted exposure propagate as separate context',
   _comm.get('schema')=='communication-style-v1' and len(_csrc)==533 and _cbad==0 and _egbad==0 and _gbad==0 and
   len(_egsrc)>=450 and
   _mexp.get('formal_role')=='Speaker of the House' and
   _mexp.get('role_adjusted_percentile',1)<_mexp.get('observed_visibility_percentile',0) and
    all(s in _desk_tpl for s in ('Observed communication style','Public exposure context','mean words/turn','Warmth cues',
                                 'positive-energy/excitement cues',
                                 )),
   '%d style members, %d style errors, %d E errors, %d broad-trait errors; Mike raw %.3f adjusted %.3f'%
   (len(_csrc),_cbad,_egbad,_gbad,_mexp.get('observed_visibility_percentile',-1),_mexp.get('role_adjusted_percentile',-1)))
_goldn=((_gtraits.get('N') or {}).get('members') or {})
_golddesk={m.get('name'):((m.get('psy') or {}).get('big5_pct') or [None]*5)[4]
          for m in _bun.get('members',[]) if ((m.get('psy') or {}).get('big5_pct') or [None]*5)[4] is not None}
_profm=(json.load(open(B+r'\profile_basis_v2.json',encoding='utf-8')).get('members') or {})
_narrow_n=sum(1 for d in _profm.values() if ((d.get('personality_facets') or {}).get('Neuroticism_affective_content') or {}).get('z') is not None)
ck('broad one-panel Neuroticism occupies the canonical slot while affective N remains separate',
   len(_golddesk)==len(_goldn)>=450 and all(_golddesk.get(m)==d.get('percentile') for m,d in _goldn.items()) and
   _narrow_n>=450 and 'Neuroticism' in _desk_tpl,
   '%d broad-N source members, %d desk N values, %d retained affective-N facets'%(len(_goldn),len(_golddesk),_narrow_n))
_narr2=os.path.exists(B+r'\member_narratives_v2.json')
_nbad=sum(1 for _row in _bun.get('members',[]) if ((_row.get('psy') or {}).get('narrative') is not None or
          (_row.get('psy') or {}).get('narrative_status')!='withheld_pending_guarded_local_rebuild_and_user_approval')) if not _narr2 else 0
ck('legacy E/audio narratives fail closed pending guarded local rebuild',
   (_narr2 or _nbad==0) and 'member_narratives_v2.json' in _dcode and
   'withheld_pending_guarded_local_rebuild_and_user_approval' in _dcode,
   ('guarded v2 narrative artifact present' if _narr2 else '%d desk rows leaked legacy prose or missed withholding status'%_nbad))
ck('desk suppresses synthetic/default-only dominance values and keeps behavior display N-free',
    all(s in _desk_tpl for s in ('Conversational dominance',
                                 'insufficient observed interview coverage')) and
   'Affective N' not in _desk_tpl and
   "const BLAB=['Cross-party cooperation','Legislative follow-through','Interaction dominance','Interview activity','Policy breadth'];" in _desk_tpl and
   "if _r2.get('dominance_v2') is not None:" in
       open(FE+r'\make_member_desk.py',encoding='utf-8').read(),
   'behavior display must expose coverage, plot five behavioral axes, and omit Affective N')
if os.path.exists(B+r'\behavioral_roles_v2_candidate.json') and os.path.exists(B+r'\member_desk_bundle.json'):
    _r2=json.load(open(B+r'\behavioral_roles_v2_candidate.json',encoding='utf-8')).get('members',{})
    _dpsy={m.get('name'):(m.get('psy') or {}) for m in _bun.get('members',[])}
    _r2bad=0
    for _name,_d in _dpsy.items():
        _src=_r2.get(_name)
        if _src:
            _r2bad += int(abs(float(_d.get('dominance_v2_z',999))-float(_src['dominance_v2']))>.000001)
            _r2bad += int(abs(float(_d.get('expressiveness_v2_z',999))-float(_src['expressiveness_v2']))>.000001)
            _r2bad += int((_d.get('behavior_scope') or {}).get('interaction_turns')!=int(_src['qualifying_turns']))
            _r2bad += int(_d.get('expressiveness_v2_pct') is None or (_d.get('behav_pct') or [None,None,None])[2] is None)
        else:
            _r2bad += int(any(k in _d for k in ('dominance_v2_z','expressiveness_v2_z','expressiveness_v2_pct')))
    _producer=open(FE+r'\make_member_desk.py',encoding='utf-8').read()
    _engine_refs=[]
    # The locked study producer, audio validator, and OOF-gold display/profile estimator are research
    # provenance, not influence/rank engines. Any additional Python consumer is a wiring regression.
    _r2_allowed={'make_member_desk.py','audit_full.py','dashboard_integrity.py','b5_narratives.py','communication_style.py',
                 'b5_profile_basis.py','b5_gold_multitrait_validate.py',
                 'behavioral_roles_v2_study.py','interview_audio_e_validate.py'}
    for _fn in os.listdir(FE):
        if not _fn.endswith('.py') or _fn in _r2_allowed: continue
        if 'behavioral_roles_v2_candidate.json' in open(FE+'\\'+_fn,encoding='utf-8',errors='ignore').read(): _engine_refs.append(_fn)
    ck('qualified conversation roles v2 propagate exactly to the desk and remain engine-isolated',
       len(_r2)>=250 and _r2bad==0 and not _engine_refs and
       all(s in _desk_tpl for s in ('Conversation dynamics v2','Verbal expressiveness')) and
       '_ROLES2_ART' in _producer,
       '%d source members, %d propagation errors, engine refs=%s'%(len(_r2),_r2bad,_engine_refs))
ck('desk preserves missing low-antagonism behavior and its selected-stratum scope',
   all(s in _desk_tpl for s in ('Low-antagonism A facet',
                                'insufficient interview coverage')),
   'low antagonism must remain missing outside the package-qualified stratum, never imputed as A')
_qproc=open(FE+r'\process_quorum.py',encoding='utf-8').read()
ck('Quorum member resolver refuses surname-only identity guesses',
   'Never infer identity from surname alone' in _qproc and 'if len(cands)==1: return cands[0]' not in _qproc,
   'same-surname outsiders must be segregated, not inserted into a legislator corpus')
# V-1 (audit 2026-07-08): vote_category_lean.json is a manually-maintained FROZEN input -- its producer
# vote_by_category.py is in no orchestrator and its cleaned-roll-call CSV is an external file with no in-repo
# producer. Surface the frozen-ness: warn if the vote leans fall >30d behind the newest bills archive, so a
# silently-stale Path V / per_category_stance signal is visible instead of aging unnoticed.
_vcl=B+r'\vote_category_lean.json'
if os.path.exists(_vcl):
    import glob as _g2
    _arch=sorted(_g2.glob(B+r'\congress_bills_archive_*.xlsx'))
    _stale=bool(_arch) and os.path.getmtime(_vcl) < os.path.getmtime(_arch[-1])-30*86400
    ck('vote_category_lean.json fresh vs bills archive (manual frozen input; see V-1)', not _stale,
       'vote leans %s'%(('STALE >30d behind '+os.path.basename(_arch[-1])) if _stale else 'current'))
# say-do corroboration gate coherence (2026-07-07): every capped cell must DISPLAY Ally with its raw
# tier preserved as Champion/Defender; every proven cell must BE Champion/Defender in the matrix.
_g=(CS.get('gate') or {})
if _g:
    _cap=_g.get('capped') or {}; _prov=_g.get('proven') or {}
    _b1=sum(1 for m,cc in _cap.items() for c,raw in cc.items()
            if MAT.get(m,{}).get(c)!='Ally' or raw not in ('Champion','Defender'))
    _b2=sum(1 for m,cs_ in _prov.items() for c in cs_ if MAT.get(m,{}).get(c) not in ('Champion','Defender'))
    ck('say-do gate coherent (capped->Ally w/ raw C/D preserved; proven are C/D cells)', _b1==0 and _b2==0,
       '%d capped cells / %d proven cells; %d+%d incoherent'%(sum(len(v) for v in _cap.values()),
       sum(len(v) for v in _prov.values()),_b1,_b2))
    # Only authored/original-cosponsored legislation and aligned recorded votes qualify as action.
    _cor=_g.get('corrob') or {}
    _b3=sum(1 for m,cc in _cor.items() for c,pth in cc.items()
            if MAT.get(m,{}).get(c) not in ('Champion','Defender') or c in (_cap.get(m) or {})
            or pth not in ('A','V'))
    _cbp=collections.Counter(pth for cc in _cor.values() for pth in cc.values())
    ck('say-do corroboration paths coherent (only A/V cells survive as C/D)', _b3==0 and _cbp['J']==0,
       'corroborated A%d V%d; %d incoherent'%(_cbp['A'],_cbp['V'],_b3))
    _oc=_g.get('overall_capped') or {}; _thr=float(_g.get('overall_delivery_threshold') or 2/3)
    _b4=sum(1 for m,d in _oc.items() if d.get('raw')!='Champion' or d.get('final')!='Ally'
            or not d.get('evidence_available') or not set(d.get('cap_reasons') or []).issubset(
                {'insufficient_top_tier_breadth','top_tier_action_coverage','high_volume_low_followthrough'})
            or not d.get('cap_reasons'))
    ck('overall say-do caps have an auditable fixed-rule reason',_b4==0,
       '%d overall caps, %d incoherent'%(len(_oc),_b4))
    import moneyball_core as _mc
    from tier_util import legislative_press_tiers as _lpt, load_tier_weights as _ltw
    _tier_weights=_ltw()
    ck('user-ratified display-tier weights are 80%% Legislative / 20%% Press',
       all(abs(a-b)<1e-12 for a,b in zip(_tier_weights,(0.8,0.2))),str(_tier_weights))
    _tier,_tier_rank=_lpt(list(MAT))
    # Paul's placement contract as of 2026-07-17 (letters-integration session, 'adjust to the true
    # tier 1s'): Tier 1 is MATH-DRIVEN (pure 80/20 order) except two anchors — Barr (standing #1)
    # and Young Kim ('stays'). The 2026-07-16 Tier-1 pins (Fitzpatrick/Williams/Hagerty/Scott/
    # Torres/Cortez Masto/Timmons/Nunn) and Hill's Tier-2 pin were REMOVED by that directive;
    # the 'should all be in tier 2' pins stand. Mirrors tier_overrides.json placements.
    # 2026-07-18 (Paul: 'replace Ann Wagner with Brian Fitzpatrick ... in the champion ranks' ->
    # 'No, just move him'): an EXPLICIT placement pair, not a math result. Fitzpatrick trails Wagner
    # on the PRESS pillar only (leg .996 vs .955; press .153 vs .500), and his press is understated
    # because CU-positive trade coverage reads neutral in the 2026-06-30-frozen press basis.
    # Wagner's Tier-2 placement is load-bearing: forcing Fitzpatrick alone ejects Tim Scott instead.
    _required_tiers={'Andy Barr':1,'Young Kim':1,'Brian Fitzpatrick':1,
                     'Maxine Waters':2,'Juan Vargas':2,'Nydia Velázquez':2,
                     'Mike Crapo':2,'Scott Fitzgerald':2,'Tim Moore':2,'Ann Wagner':2}
    ck('user-ratified Legislative+Press tier placements are exact',
       all(_tier.get(m)==t for m,t in _required_tiers.items()),
       str({m:_tier.get(m) for m in _required_tiers}))
    _display={m:_mc.overall_display_tier(MAT.get(m,{}),_tier.get(m,0),m,_oc) for m in MAT}
    # 2026-07-17 update: Hill's Tier-2 pin removed ('true tier 1s') — he re-enters Tier 1 on pure
    # math (#2) and MUST be crowned alongside Barr and Young Kim.
    _final_named={m:_display.get(m) for m in ('Young Kim','Andy Barr','French Hill')}
    ck('Tier-1 status crowns named legislative/press leaders as Champion',
       _final_named.get('Young Kim')=='Champion' and _final_named.get('Andy Barr')=='Champion'
       and _final_named.get('French Hill')=='Champion',
       str(_final_named))
    _champ={m for m,v in _display.items() if v=='Champion'}
    _tier1={m for m,t in _tier.items() if t==1}
    ck('overall Champion crown exactly equals Legislative+Press Tier 1',
       _champ==_tier1 and len(_champ)==12,
       '%d crowns vs %d Tier-1; missing %s; extra %s'%(
           len(_champ),len(_tier1),sorted(_tier1-_champ),sorted(_champ-_tier1)))
    if os.path.exists(B+r'\dashboard_data.json'):
        _dash=json.load(open(B+r'\dashboard_data.json',encoding='utf-8'))
        _tc={'Champion':'C','Defender':'F','Ally':'A','Highly Persuadable':'H','Persuadable':'P',
             'Lean Persuadable':'G','True Neutral':'T','Lean Skeptic':'L','Neutral':'N',
             'Skeptic':'S','Highly Skeptical':'K','Opponent':'O','—':'-'}
        _expected_gen=''.join(_tc.get(_display.get(m,'—'),'-') for m in _dash.get('name',[]))
        ck('dashboard overall labels exactly match shared crown law',_dash.get('gen')==_expected_gen,
           'dashboard and kernel label strings differ')
    _wb_display={_ra_ws.cell(r,1).value:_ra_ws.cell(r,7).value for r in range(3,536) if _ra_ws.cell(r,1).value}
    ck('workbook stance column exactly matches shared crown law',
       all(_wb_display.get(m)==v for m,v in _display.items()) and len(_wb_display)==len(_display),
       'workbook %d vs matrix %d rows'%(len(_wb_display),len(_display)))
else:
    ck('say-do gate block present in category_stance.json', False, 'gate key missing — producer predates 2026-07-07 landing')

print('FULL-BODY AUDIT — %d passed, %d FAILED'%(len(P),len(F)))
if F:
    print('  FAILURES:',F)
    sys.exit(1)
else: print('  ALL MODULES CLEAN.')
