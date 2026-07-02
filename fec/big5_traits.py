# -*- coding: utf-8 -*-
"""BIG FIVE TRAIT ESTIMATION FROM MEMBER TEXT — Ramey, Klingler & Hollibaugh (PSRM 2016)
"Measuring Elite Personality Using Speech", synthesized (user 2026-07-02: "copy what these articles
did bar for bar in a synthesized manner, it does not have to be domain specific").

Their pipeline, kept step for step:
  1. Aggregate each legislator's text over a fixed span (theirs: floor speeches by year; ours: the
     member's full tweet/post corpus — Birkenmaier & Lechner (Pol.Analysis 2025) validate social
     posts as personality-cue carriers on par with speeches).
  2. Score the text on closed-vocabulary LIWC/MRC-style category rates (per-1000-word usage).
  3. Standardize the category rates ACROSS THE CORPUS and combine with standardized weights into
     the Big Five (their step: Mairesse et al. 2007 pre-trained standardized SMOreg over the same
     features; the model file is LIWC-license-bound, so we use the published Pennebaker & King
     (1999) / Yarkoni (2010) / Mairesse Table-3 correlates as the standardized linear weights —
     same feature families, same signs, linear instead of SVR).
  4. Scores are z-normalized RELATIVE TO THIS CORPUS (their explicit framing: estimates are only
     meaningful relative to the corpus used — congressional distribution, not population norms).
  5. Confidence: they show CI width blows up below ~1k words/member (Fig 3, median 11k/yr).  We
     carry tokens and flag lowconf below LOWCONF_TOKENS (default 1500); consumers shrink lowconf
     trait z-scores toward 0 (uncertainty -> neutral, never a fabricated personality).
Emotional Stability (ES) is reported Ramey-style; Neuroticism N = -ES.
Output: big5_traits.json  {member: {O,C,E,A,ES,N (z-scores), tokens, nposts, lowconf}}"""
import sys,os,json,re,math,collections
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import openpyxl,pandas as pd
B=r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
LOWCONF_TOKENS=int(os.environ.get('LOWCONF_TOKENS','1500'))
C={m['member']:m for m in json.load(open(B+r'\congestion.json'))}; members=list(C)
# ---------------- corpus: posts (ALL topics; RTs EXCLUDED — an RT is someone else's words) + floor
# speeches / hearing transcripts (speeches table, fetch_congressional_record.py — the papers' native
# medium; Ramey used the Congressional Record itself) ----------------
texts=collections.defaultdict(list)
ws=openpyxl.load_workbook(B+r'\Quorum Processed.xlsx',read_only=True)['Members']
hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
im,ic,irt=hdr.index('member'),hdr.index('Content'),hdr.index('is_rt')
for row in ws.iter_rows(min_row=2,values_only=True):
    m,t=row[im],row[ic]
    if m in C and t and str(row[irt]).lower() not in ('1','true'): texts[m].append(str(t))
sm=pd.read_excel(B+r'\Social Master.xlsx',sheet_name='Corpus',dtype=str)
for _,r in sm[sm['platform']=='bluesky'].iterrows():
    if r['member'] in C and r['text'] and str(r.get('is_rt','')).lower() not in ('1','true'): texts[r['member']].append(str(r['text']))
import sqlite3
nsp=0
try:
    con=sqlite3.connect(B+r'\social_media.db')
    for m,t in con.execute('SELECT member,text FROM speeches'):
        if m in C and t: texts[m].append(t); nsp+=1
except sqlite3.OperationalError: pass
print('corpus: speeches %d turns folded in'%nsp)
# ---------------- closed-vocabulary categories (LIWC-analog; word lists are the published exemplars) ----------------
CAT={
 'i':"i me my mine myself im ive ill id",
 'we':"we us our ours ourselves were weve well",
 'you':"you your yours youre",
 'article':"a an the",
 'negate':"no not never none cannot cant dont wont isnt wasnt arent didnt doesnt wouldnt couldnt shouldnt nothing neither nor",
 'posemo':"good great love happy hope proud thank thanks celebrate honor honored best wonderful glad support win wins success benefit benefits opportunity grateful appreciate excellent congratulations congrats delighted terrific amazing",
 'negemo':"bad wrong hate hurt nasty terrible awful fail failed failure crisis threat worry worried fear afraid angry anger sad loss lose damage harm danger attack disaster outrage shameful disgrace broken devastating reckless",
 'anx':"worried anxious afraid fear fearful nervous risk risky threat threatened threatens uncertain uncertainty alarm alarming danger dangerous crisis vulnerable insecure unstable",
 'anger':"angry anger outrage outrageous furious fight fighting attack attacks hate disgusting shameful corrupt corruption betray betrayal rigged radical disgrace hostile",
 'social':"talk talked share friend friends family families community communities neighbor neighbors together meeting met join joined joining team people folks everyone partnership colleagues constituents neighbors town",
 'family':"family families mother father mom dad son daughter kids children parents grandmother grandfather wife husband brother sister grandparents",
 'achieve':"work working achieve achieved achievement accomplish accomplished succeed success successful win won effort efforts goal goals deliver delivered results earn earned lead leader leadership fought secured secure passed",
 'work':"job jobs work working business businesses economy economic employment worker workers industry wages career manufacturing farmers",
 'insight':"think thinking believe know understand consider considered because reason reasons means thus therefore however although question questions idea ideas analysis examine evidence solution solutions",
 'certain':"always never must certainly absolutely definitely clear clearly undeniable undoubtedly every all guarantee fact facts truth proven commit committed",
 'tentat':"maybe perhaps might may could seems appears possibly somewhat unclear likely probably potential potentially",
}
CATW={k:set(v.split()) for k,v in CAT.items()}
FEATS=list(CAT)+['sixltr','exclam','number']
# ---------------- standardized trait weights (Pennebaker–King / Yarkoni / Mairesse correlates) ----------------
# Signs are the replicated literature directions; magnitudes are the relative correlate strengths.
W={
 'E':{'social':.30,'posemo':.20,'we':.15,'exclam':.10,'tentat':-.15,'negemo':-.10},          # extraverts: social+positive, low hedging
 'ES':{'anx':-.30,'negemo':-.25,'i':-.20,'negate':-.15,'posemo':.10},                        # stability = inverse neuroticism markers
 'A':{'posemo':.25,'family':.20,'social':.15,'we':.15,'anger':-.25},                         # agreeable: warm/communal, low hostility
 'C':{'achieve':.35,'work':.20,'certain':.10,'negate':-.20,'negemo':-.15},                   # conscientious: achievement talk, low negations
 'O':{'sixltr':.30,'insight':.25,'article':.20,'i':-.15,'certain':-.10},                     # open: lexical complexity + cognitive words
}
WORD=re.compile(r"[a-z']+")
# chamber RITUAL is not lexical behavior — floor-management boilerplate (heavier for the majority that
# runs procedure) contaminates formality features; strip the fixed phrases before category counting.
RITUAL=re.compile(r'\b(i yield (back|to|the floor|myself)[^.]*|mr\.? (president|speaker)|madam (president|speaker)|'
 r'i ask unanimous consent[^.]*|the (gentleman|gentlewoman|gentlelady) from [a-z ]+|'
 r'i (rise|reserve the balance of my time)|pursuant to [^.]*|the clerk will [^.]*|'
 r'i suggest the absence of a quorum|the balance of my time)\b',re.I)
def rates(txts):
    cnt=collections.Counter(); n=0; six=0; excl=0; num=0
    for t in txts:
        t=re.sub(r'https?://\S+|\bt\.co/\S+|@\w+','',t.lower())   # URLs + @handles are not lexical behavior
        t=RITUAL.sub(' ',t)
        excl+=t.count('!'); num+=len(re.findall(r'\d',t))
        for w in WORD.findall(t.replace("'","")):
            n+=1
            if len(w)>=6: six+=1
            for c,s in CATW.items():
                if w in s: cnt[c]+=1
    if n==0: return None,0
    r={c:1000.0*cnt[c]/n for c in CAT}; r['sixltr']=1000.0*six/n; r['exclam']=1000.0*excl/n; r['number']=1000.0*num/n
    return r,n
R={}; TOK={}
for m in members:
    r,n=rates(texts.get(m,[]))
    if r: R[m]=r; TOK[m]=n
# corpus-standardize each feature (Ramey step: standardized features -> standardized model)
mu={f:sum(R[m][f] for m in R)/len(R) for f in FEATS}
sd={f:(sum((R[m][f]-mu[f])**2 for m in R)/max(1,len(R)-1))**.5 or 1.0 for f in FEATS}
Z={m:{f:(R[m][f]-mu[f])/sd[f] for f in FEATS} for m in R}
raw={m:{t:sum(w*Z[m][f] for f,w in W[t].items()) for t in W} for m in R}
tmu={t:sum(raw[m][t] for m in raw)/len(raw) for t in W}
tsd={t:(sum((raw[m][t]-tmu[t])**2 for m in raw)/max(1,len(raw)-1))**.5 or 1.0 for t in W}
out={}
for m in members:
    if m in raw:
        tr={t:round((raw[m][t]-tmu[t])/tsd[t],3) for t in W}
        out[m]={**tr,'N':round(-tr['ES'],3),'tokens':TOK[m],'nposts':len(texts.get(m,[])),'lowconf':TOK[m]<LOWCONF_TOKENS}
    else:
        out[m]={'O':0.0,'C':0.0,'E':0.0,'A':0.0,'ES':0.0,'N':0.0,'tokens':0,'nposts':0,'lowconf':True}
json.dump({'traits':out,'method':'Ramey-synth closed-vocab standardized linear (P&K/Yarkoni/Mairesse correlates); z relative to THIS corpus',
           'lowconf_tokens':LOWCONF_TOKENS},open(B+r'\big5_traits.json','w'))
nm=sum(1 for m in out if out[m]['tokens']>0); lc=sum(1 for m in out if out[m]['lowconf'])
tk=sorted(TOK.values())
print('BIG5: %d/%d members with text | median tokens %d | lowconf(<%d tok): %d'%(nm,len(members),tk[len(tk)//2] if tk else 0,LOWCONF_TOKENS,lc))
for t in ('O','C','E','A','ES'):
    top=sorted((m for m in out if not out[m]['lowconf']),key=lambda m:-out[m][t])[:3]
    bot=sorted((m for m in out if not out[m]['lowconf']),key=lambda m:out[m][t])[:3]
    print('  %-2s high: %-45s low: %s'%(t,', '.join(top),', '.join(bot)))
# ---- Ramey Table-1 style external validation: trait vs party (their DV: Common Space conservatism;
# expected signs there: O −, C +, E −, A −, ES +). Party = roster proxy for ideology. ----
import csv as _csv
party={}
for cn in (119,118,117):
    p=B+r'\votes\HS%d_members.csv'%cn
    if os.path.exists(p):
        for row in _csv.DictReader(open(p,encoding='utf-8')):
            b=row.get('bioguide_id'); pc=str(row.get('party_code'))
            if b and b not in party: party[b]={'100':0,'200':1}.get(pc)
pr={m:party.get(C[m].get('bioguide','')) for m in members}
EXP={'O':-1,'C':1,'E':-1,'A':-1,'ES':1}
print('validation (point-biserial r vs Republican, weighted members only; Ramey-expected sign in []):')
val=[m for m in out if not out[m]['lowconf'] and pr.get(m) is not None]
for t in ('O','C','E','A','ES'):
    xs=[out[m][t] for m in val]; ys=[pr[m] for m in val]
    mx=sum(xs)/len(xs); my=sum(ys)/len(ys)
    sx=(sum((x-mx)**2 for x in xs))**.5; sy=(sum((y-my)**2 for y in ys))**.5
    r=sum((x-mx)*(y-my) for x,y in zip(xs,ys))/max(1e-9,sx*sy)
    ok='OK' if r*EXP[t]>0 else 'REVERSED'
    print('  %-2s r=%+.3f [%+d expected]  %s'%(t,r,EXP[t],ok))
print('wrote big5_traits.json')
