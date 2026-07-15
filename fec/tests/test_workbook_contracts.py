from __future__ import annotations

import json
import os
from pathlib import Path
import unittest

import openpyxl

ROOT=Path(__file__).resolve().parents[2]
WORKBOOK=Path(os.environ.get('MONEYBALL_WORKBOOK',ROOT/'outputs'/'workbook_repair'/'Moneyball Audit Workbook.xlsx'))


@unittest.skipUnless(WORKBOOK.exists(),'generated Moneyball Audit Workbook.xlsx is not present')
class WorkbookContracts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.book=openpyxl.load_workbook(WORKBOOK,read_only=False,data_only=False)
        with open(ROOT/'category_stance.json',encoding='utf-8') as handle:
            cls.category=json.load(handle)
        with open(ROOT/'ranks_python.json',encoding='utf-8') as handle:
            cls.ranks=json.load(handle)

    @classmethod
    def tearDownClass(cls):
        cls.book.close()

    def test_legislative_normalization_uses_the_same_member_row(self):
        raw=self.book['Legislative Raw']; normalized=self.book['Legislative Normalized']
        for col,row in ((3,3),(152,535)):
            values=[float(raw.cell(r,col).value or 0) for r in range(3,536)]
            value=values[row-3]; pos=max([0]+values); neg=abs(min([0]+values))
            expected=0 if value==0 else value/pos if value>0 and pos else value/neg if neg else 0
            self.assertAlmostEqual(expected,float(normalized.cell(row,col).value or 0),places=12)
            self.assertEqual(raw.cell(row,1).value,normalized.cell(row,1).value)

    def test_category_stance_sheet_exactly_matches_live_matrix(self):
        ws=self.book['Category Stance']; cats=self.category['cats']; matrix=self.category['matrix']
        self.assertEqual(cats,[ws.cell(2,c).value for c in range(3,21)])
        rows={ws.cell(r,1).value:[ws.cell(r,c).value for c in range(3,21)] for r in range(3,536)}
        self.assertEqual(set(matrix),set(rows))
        for name in matrix:
            self.assertEqual([matrix[name].get(cat,'—') for cat in cats],rows[name],name)

    def test_workbook_contains_rank_authority(self):
        self.assertIn('Rank Authority',self.book.sheetnames)
        ws=self.book['Rank Authority']
        headers=[ws.cell(4,c).value for c in range(1,16)]
        self.assertEqual('Score',headers[1]); self.assertEqual('Rank',headers[2])
        rows={ws.cell(r,1).value:(ws.cell(r,2).value,ws.cell(r,3).value) for r in range(5,538)}
        self.assertEqual(set(self.ranks['rank']),set(rows))
        for name in self.ranks['rank']:
            self.assertAlmostEqual(self.ranks['score'][name],rows[name][0],places=12)
            self.assertEqual(self.ranks['rank'][name],rows[name][1])

    def test_overview_states_the_authority_boundary_and_displacement_repair(self):
        ws=self.book['Overview']
        fields={ws.cell(r,1).value:ws.cell(r,2).value for r in range(3,11)}
        self.assertIn('not a rank or stance authority',fields['Scope'])
        self.assertIn('formal office is display-only',fields['Rank authority'])
        self.assertGreater(fields['Legacy displaced normalized cells found'],0)


if __name__=='__main__':
    unittest.main(verbosity=2)
