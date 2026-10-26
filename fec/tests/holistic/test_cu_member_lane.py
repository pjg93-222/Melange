# -*- coding: utf-8 -*-
"""CR lane (build_cu_member_lane): offline BILLSTATUS -> CU cosponsor network + bills archive.

Fixtures are inline BILLSTATUS xml zipped into the billstatus<c>/BILLSTATUS-<c>-<t>.zip layout. Golden checks
read the real 0625 network / 0624 archive under MONEYBALL_ROOT (paths.ROOT) and skip when absent.
"""
import io
import math
import os
import subprocess
import sys
import zipfile

import pytest

FEC = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, FEC)

import openpyxl  # noqa: E402
from paths import ROOT  # noqa: E402
from network import build_cu_member_lane as L  # noqa: E402

ROSTER = {'S000001': 'Sam Sponsor', 'A000001': 'Ann Orig', 'B000001': 'Bob Reg', 'W000001': 'Wes Withdrawn'}
SECS = ['Interchange', 'CDFI Fund', 'Payments']


def bs(number, sponsor='S000001', cos=(), title='A Bill', actions=(), btype='HR', summary=None):
    c = ''.join('<item><bioguideId>%s</bioguideId><fullName>Rep. %s</fullName><isOriginalCosponsor>%s</isOriginalCosponsor>%s</item>'
                % (b, b, o, '<sponsorshipWithdrawnDate>2025-01-02</sponsorshipWithdrawnDate>' if w else '') for b, o, w in cos)
    a = ''.join('<item><actionDate>%s</actionDate><actionCode>%s</actionCode><type>%s</type><text>%s</text></item>' % x for x in actions)
    sm = ('<summaries><summary><versionCode>00</versionCode><updateDate>2025-01-01</updateDate><text><![CDATA[<p>%s</p>]]></text></summary></summaries>' % summary) if summary else ''
    return ('<billStatus><bill><number>%d</number><type>%s</type><title>%s</title>'
            '<sponsors><item><bioguideId>%s</bioguideId><fullName>Rep. %s [X-ST-1]</fullName></item></sponsors>'
            '<cosponsors>%s</cosponsors><committees><item><name>Financial Services Committee</name><chamber>House</chamber></item></committees>'
            '<actions>%s</actions>%s</bill></billStatus>' % (number, btype, title, sponsor, sponsor, c, a, sm)).encode()


def make_raw(tmp_path, docs):
    """docs: {(congress, type, number): xml bytes} -> raw root"""
    by = {}
    for (c, t, n), x in docs.items():
        by.setdefault((c, t), {})['BILLSTATUS-%d%s%d.xml' % (c, t, n)] = x
    for (c, t), members in by.items():
        d = tmp_path / ('billstatus%d' % c); d.mkdir(exist_ok=True)
        with zipfile.ZipFile(d / ('BILLSTATUS-%d-%s.zip' % (c, t)), 'w') as z:
            for k, v in members.items():
                z.writestr(k, v)
    return str(tmp_path)


def build_fixture(tmp_path, docs, seed=None, **kw):
    raw = make_raw(tmp_path, docs)
    seed = seed or {k: ['Interchange'] for k in docs}
    bills, refused = L.read_bills(raw, list(seed))
    rows = L.archive_rows(seed, bills, ROSTER, SECS, **kw)
    return bills, refused, rows


COS3 = (('A000001', 'True', False), ('B000001', 'False', False), ('W000001', 'False', True))


def test_edge_weights_and_withdrawn_in_C(tmp_path):
    bills, ref, rows = build_fixture(tmp_path, {(119, 'hr', 1): bs(1, cos=COS3)})
    assert not ref
    edges, refused, labels = L.stage3_edges(rows, bills, ROSTER, edge_withdrawn='include')
    f = 1 / math.sqrt(3)   # C counts the withdrawn cosponsor (R's cosponsors page lists it)
    w = {e['To']: e for e in edges}
    assert math.isclose(w['A000001']['Weight'], 1.0 * f) and math.isclose(w['B000001']['Weight'], 0.5 * f)
    assert w['A000001']['NOriginal'] == 1 and w['B000001']['NRegular'] == 1 and w['W000001']['NBills'] == 1
    e2, _, _ = L.stage3_edges(rows, bills, ROSTER, edge_withdrawn='exclude')
    assert math.isclose({e['To']: e for e in e2}['A000001']['Weight'], 1 / math.sqrt(2))


def test_success_and_decay_and_summing(tmp_path):
    passed = ((('2025-02-01', '8000', 'Floor', 'Passed/agreed to in House'),))
    docs = {(118, 'hr', 1): bs(1, cos=(('A000001', 'True', False),), actions=passed),
            (119, 'hr', 2): bs(2, cos=(('A000001', 'True', False),))}
    bills, _, rows = build_fixture(tmp_path, docs)
    edges, _, _ = L.stage3_edges(rows, bills, ROSTER)
    e = edges[0]
    assert (e['From'], e['To'], e['NBills']) == ('S000001', 'A000001', 2)
    assert math.isclose(e['Weight'], 3.0 * 0.8 + 1.0)   # SUCCESS x3 x decay 0.8, plus the plain 119th bill


def test_self_loop_dropped(tmp_path):
    bills, _, rows = build_fixture(tmp_path, {(119, 'hr', 1): bs(1, cos=(('S000001', 'True', False), ('A000001', 'True', False)))})
    edges, _, _ = L.stage3_edges(rows, bills, ROSTER)
    assert [e['To'] for e in edges] == ['A000001']


def test_offroster_sponsor_no_sponsor_row_no_edges_but_cosponsor_rows(tmp_path):
    bills, _, rows = build_fixture(tmp_path, {(119, 'hr', 1): bs(1, sponsor='Z999999', cos=(('A000001', 'True', False),))})
    assert {r['Role'] for r in rows} == {'original-cosponsor', 'cosponsor'}
    edges, refused, _ = L.stage3_edges(rows, bills, ROSTER)
    assert edges == [] and refused == []


def test_withdrawn_dial_and_roles(tmp_path):
    docs = {(119, 'hr', 1): bs(1, cos=COS3)}
    _, _, ex = build_fixture(tmp_path, docs, withdrawn='exclude')
    assert [(r['Role'], r['Member']) for r in ex] == [('sponsor', 'Sam Sponsor'), ('original-cosponsor', 'Ann Orig'),
                                                      ('cosponsor', 'Ann Orig'), ('cosponsor', 'Bob Reg')]
    assert ex[0]['CosponsorCount'] == 2
    _, _, inc = build_fixture(tmp_path, docs, withdrawn='include')
    assert 'Wes Withdrawn' in {r['Member'] for r in inc}
    _, _, no = build_fixture(tmp_path, docs, originals_also_cosponsor=False)
    assert [r['Role'] for r in no].count('cosponsor') == 1


@pytest.mark.parametrize('acts,btype,want', [
    ([], 'hr', 'Introduced'),
    ([('d', '8000', 'Floor', 'Passed/agreed to in House')], 'hr', 'Passed House'),
    ([('d', '8000', 'Floor', 'Passed/agreed to in House'), ('d', '17000', 'Floor', 'Passed/agreed to in Senate')], 'hr', 'Passed Senate'),
    ([('d', '17000', 'Floor', 'Agreed to in Senate')], 'sres', 'Agreed to in Senate'),
    ([('d', '', 'ResolvingDifferences', 'x')], 'hr', 'Resolving Differences'),
    ([('d', '36000', 'BecameLaw', 'Became Public Law No: 119-1.')], 'hr', 'Became Law'),
    ([('d', '', 'President', 'Vetoed by President.'), ('d', '', 'Floor', 'On passage over veto, failed: not agreed to by two-thirds')], 'hr', 'Failed to pass over veto'),
])
def test_tracker_status(acts, btype, want):
    assert L.tracker_status(acts, btype) == want


def test_link_and_name_roundtrip():
    from network import build_cu_placeholder_backfill as BF
    for t in L.PREFIX:
        link = L.detail_link(119, t, 7)
        assert BF.is_bill_link(link) and BF.bill_key(L.bill_name(t, 7), 119) == (119, t, 7)
    assert L.detail_link(117, 'hres', 3) == '/bill/117th-congress/house-resolution/3'
    assert L.decay(117) == pytest.approx(0.64)


def test_archive_sheets_and_closeness_blank(tmp_path):
    bills, _, rows = build_fixture(tmp_path, {(119, 'hr', 1): bs(1, cos=COS3, summary='Hello')})
    assert rows[0]['bill_summary'] == 'Hello' and rows[0]['Committee'] == 'Committees: House - Financial Services'
    out = str(tmp_path / 'o'); os.makedirs(out)
    long = ['Credit Union Capital Requirements and More', 'CDFI Fund']
    L.write_archive(os.path.join(out, 'a.xlsx'), [dict(r, Section=long[0]) for r in rows], long)
    wb = openpyxl.load_workbook(os.path.join(out, 'a.xlsx'))
    assert wb.sheetnames == ['All Bills', ('01 ' + long[0])[:31], '02 CDFI Fund']
    assert wb['02 CDFI Fund'].max_row == 1   # header-only
    edges, _, labels = L.stage3_edges(rows, bills, ROSTER)
    cent = L.centrality(edges, labels)
    L.write_network(os.path.join(out, 'n.xlsx'), edges, cent)
    ws = openpyxl.load_workbook(os.path.join(out, 'n.xlsx'))['Centrality']
    got = {r[0]: r[7] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert got['S000001'] is not None and got['A000001'] is None   # sinks have no out-reach -> blank


def test_xxe_guard(tmp_path):
    bad = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e "y">]><billStatus><bill/></billStatus>'
    bills, refused = L.read_bills(make_raw(tmp_path, {(119, 'hr', 1): bad}), [(119, 'hr', 1)])
    assert bills == {} and 'DOCTYPE' in refused[0][1]


def test_missing_key_is_refusal_not_silent(tmp_path):
    raw = make_raw(tmp_path, {(119, 'hr', 1): bs(1)})
    bills, refused = L.read_bills(raw, [(119, 'hr', 1), (119, 'hr', 99), (118, 's', 5)])
    assert list(bills) == [(119, 'hr', 1)] and len(refused) == 2


def test_root_refusal():
    with pytest.raises(L.Refusal):
        L.assert_out_ok(L.REPO, ROOT)
    with pytest.raises(L.Refusal):
        L.assert_out_ok(ROOT, ROOT)
    L.assert_out_ok(os.path.join(L.REPO, 'data', 'cr_lane', 'x'), ROOT)


def test_build_refuses_root_out_and_missing_inputs(tmp_path, capsys):
    assert L.main(['build', '--out', L.REPO, '--run-id', 'x']) == 3          # root out refused (exit 3, nothing written)
    assert L.main(['build', '--out', str(tmp_path), '--run-id', 'x', '--raw-root', str(tmp_path / 'nope')]) == 3
    assert 'inputs missing' in capsys.readouterr().err
    assert not (tmp_path / 'x').exists()


REF_NET = os.path.join(ROOT, L.REF_NETWORK)
REF_ARC = os.path.join(ROOT, L.REF_ARCHIVE)


@pytest.mark.skipif(not os.path.exists(REF_NET), reason='0625 network not present')
def test_golden_centrality_matches_0625():
    wb = openpyxl.load_workbook(REF_NET, read_only=True)
    it = wb['Edges'].iter_rows(values_only=True); next(it)
    edges = [dict(zip(L.EDGE_HEADERS, r)) for r in it]
    labels = {}
    for e in edges:
        labels[e['From']] = e['From_label']; labels[e['To']] = e['To_label']
    got = L.centrality(edges, labels)
    it = wb['Centrality'].iter_rows(values_only=True); next(it)
    ref = [dict(zip(L.CENT_HEADERS, r)) for r in it]
    assert [g['node'] for g in got] == [r['node'] for r in ref]
    for g, r in zip(got, ref):
        for h in L.CENT_HEADERS[2:]:
            if r[h] is None:
                assert math.isnan(g[h])
            else:
                assert abs(g[h] - r[h]) < 1e-9, (g['node'], h)


@pytest.mark.skipif(not (os.path.exists(REF_NET) and os.path.exists(REF_ARC)), reason='reference files not present')
def test_golden_schema_diff_of_fixture_candidate(tmp_path):
    seed = {(119, 'hr', 1): ['Interchange']}
    bills, _, rows = build_fixture(tmp_path, {(119, 'hr', 1): bs(1, cos=COS3)}, seed=seed)
    order = ['Examination & Supervision', 'Interchange']
    edges, _, labels = L.stage3_edges(rows, bills, ROSTER)
    n, a = str(tmp_path / 'n.xlsx'), str(tmp_path / 'a.xlsx')
    L.write_network(n, edges, L.centrality(edges, labels))
    L.write_archive(a, rows, order)
    res = L.schema_diff(n, a, REF_NET, REF_ARC)
    # sheet-name set differs (fixture has 2 sections, not 17) and nothing else may break the contract
    assert [f for f in res['fail'] if 'sheet names' not in f] == []


def test_variant_report_flags_raw_lastaction_and_counts():
    def row(la, c=119, m='Ann Orig', title='T', summ=None):
        return {'Section': 'Interchange', 'Role': 'cosponsor', 'Member': m, 'Name': 'H.R.1', 'Congress': c,
                'LastAction': la, 'Title': title, 'bill_summary': summ}
    a = [row('Introduced')]; a0 = [row('Referred to the House Committee on Financial Services.', title='T2')]
    rep = L.variant_report({'A0': a0, 'A': a}, a)
    assert rep['A0']['lastaction_invalid'] == 1 and rep['A']['lastaction_invalid'] == 0
    assert rep['A0']['vs_A']['bills_title_differs'] == 1 and rep['A0']['vs_A']['bills_lastaction_differs'] == 1
    assert rep['A']['rows_by_congress'] == {119: 1}


def test_r_faithful_network_counts_per_section_and_opt_in_dedupe(tmp_path):
    docs = {(119, 'hr', 1): bs(1, cos=COS3), (119, 'hr', 2): bs(2, sponsor='Z999999', cos=COS3)}
    bills, _, rows = build_fixture(tmp_path, docs)
    billset = {(119, 'hr', 1): 2, (119, 'hr', 2): 1}
    e2, skipped, _ = L.r_network_edges(billset, bills, ROSTER)
    e1, _, _ = L.r_network_edges(billset, bills, ROSTER, count='dedupe')
    a2 = {e['To']: e for e in e2}['A000001']; a1 = {e['To']: e for e in e1}['A000001']
    assert math.isclose(a2['Weight'], 2 * a1['Weight']) and a2['NBills'] == 2 and a1['NBills'] == 1
    assert skipped == [((119, 'hr', 2), 'sponsor off roster')]   # reported, no edges


def test_r_faithful_uses_current_status_and_bioguide_sponsor(tmp_path):
    passed = (('2025-02-01', '8000', 'Floor', 'Passed/agreed to in House'),)
    bills, _, _ = build_fixture(tmp_path, {(119, 'hr', 1): bs(1, cos=(('A000001', 'True', False),), actions=passed)})
    e, _, _ = L.r_network_edges({(119, 'hr', 1): 1}, bills, ROSTER)
    assert (e[0]['From'], math.isclose(e[0]['Weight'], 3.0)) == ('S000001', True)


def test_archive_set_per_section_multiplier(tmp_path):
    docs = {(119, 'hr', 1): bs(1, cos=(('A000001', 'True', False),))}
    seed = {(119, 'hr', 1): ['Interchange', 'Payments']}
    bills, _, rows = build_fixture(tmp_path, docs, seed=seed)
    per, _, _ = L.network_edges('archive', 'per-section', rows, bills, ROSTER, seed)
    ded, _, _ = L.network_edges('archive', 'dedupe', rows, bills, ROSTER, seed)
    assert math.isclose(per[0]['Weight'], 2 * ded[0]['Weight'])


def test_billset_file_is_tracked_not_ignored():
    p = os.path.join(L.REPO, 'fec', 'network', 'cr_r_network_billset.json')
    assert os.path.isfile(p)
    r = subprocess.run(['git', 'check-ignore', '-q', p], cwd=L.REPO)
    assert r.returncode == 1, '.gitignore (*.json) must not ignore the R bill set'


def test_billset_file_is_r_network_set():
    bs_ = L.load_billset()
    assert len(bs_) == 229 and sum(bs_.values()) == 293          # 311 checkpoint bills - 18 stale (+3 Tim Scott supplied)
    assert all(t in L.PREFIX and c in L.CONGRESSES for c, t, n in bs_)
    assert (119, 's', 2486) in bs_ and (118, 's', 3992) in bs_ and (117, 's', 3382) in bs_


# ---- CR-BF: glossary-bug backfill baked into the R-faithful base
def test_glossary_backfill_adds_lost_bills_once_and_never_double_counts(tmp_path):
    docs = {(119, 'hr', 1): bs(1, cos=COS3), (119, 'hr', 2): bs(2, cos=COS3)}
    bills, _, _ = build_fixture(tmp_path, docs)
    r_set = {(119, 'hr', 1): 2}
    lost = {(119, 'hr', 1): 1, (119, 'hr', 2): 1}               # hr1 is in BOTH sets; hr2 only lost
    merged = L.with_glossary_backfill(r_set, lost)
    assert merged == {(119, 'hr', 1): 2, (119, 'hr', 2): 1}     # R's per-section count wins, no double count
    base, _, _ = L.r_network_edges(r_set, bills, ROSTER)
    full, _, _ = L.r_network_edges(merged, bills, ROSTER)
    w = lambda es: {(e['From'], e['To']): e['Weight'] for e in es}
    one = {k: v / 2 for k, v in w(base).items()}                # weight of one count of hr1
    assert all(math.isclose(w(full)[k], w(base)[k] + one[k]) for k in one)   # hr2 adds exactly one hr1-sized count


def _glossary_golden():
    ref = os.path.join(ROOT, L.REF_ARCHIVE)
    sidecar = os.path.join(ROOT, 'data', 'network', 'cu_placeholder_backfill_20260625.json')
    if not (os.path.exists(ref) and os.path.exists(sidecar)):
        pytest.skip('live 0624 archive / backfill sidecar absent')
    return ref, sidecar


def test_golden_lost_set_is_the_111_bills_of_the_live_sidecar():
    import json
    ref, sidecar = _glossary_golden()
    lost = L.glossary_lost_set(ref)
    doc = json.load(open(sidecar, encoding='utf-8'))
    assert len(lost) == 111 and lost == {L.BF.bill_key(n, c): 1 for n, c, _ in doc['lost_bills']}
    assert len(set(lost) - set(L.load_billset())) == 109        # 2 are already in R's set: kept once


def test_baked_base_not_keyed_in_core_so_sidecar_never_stacks():
    import moneyball_core as MC
    cr = 'congress_cosponsor_network_20261026_cr3.xlsx'
    assert MC.cu_backfill_ties(cr) == {} and MC.cu_dropped_sponsors(cr) == frozenset()
    assert list(MC.CU_PLACEHOLDER_BACKFILL_FILES) == ['congress_cosponsor_network_20260625_205554.xlsx']   # 0625 only, untouched


def test_golden_baked_ties_match_live_sidecar_within_tolerance():
    """The 111 bills rebuilt from CURRENT BILLSTATUS vs the frozen live sidecar (Oct-17 cache + 0624 LastAction)."""
    import json
    ref, sidecar = _glossary_golden()
    raw = os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Melange', 'lda_rebuild', 'raw')
    if not os.path.isdir(raw):
        pytest.skip('BILLSTATUS raw root absent')
    lost = L.glossary_lost_set(ref)
    roster, _ = L.load_roster(ROOT)
    bills, refused = L.read_bills(raw, list(lost))
    assert not refused
    edges, skipped, _ = L.r_network_edges(lost, bills, roster)
    got = {(e['From'], e['To']): e['Weight'] for e in edges}
    side = {(e['from'], e['to']): e['weight'] for e in json.load(open(sidecar, encoding='utf-8'))['edges']}
    assert len(set(got) & set(side)) >= 0.99 * len(side)                      # measured 1,267 of 1,272 pairs
    assert abs(sum(got.values()) - sum(side.values())) <= 0.06 * sum(side.values())   # measured within 5%: BILLSTATUS drift since the frozen sidecar


# ----------------------------------------------------------------------------- FIX83 (CodeRabbit, PR #83)
def test_out_allowed_only_under_live_data_cr_lane_or_outside_root(tmp_path):
    ok = os.path.join(ROOT, 'data', 'cr_lane')
    L.assert_out_ok(ok, ROOT); L.assert_out_ok(os.path.join(ok, 'run1'), ROOT); L.assert_out_ok(str(tmp_path / 'sandbox'), ROOT)
    for bad in (os.path.join(ROOT, 'data'), os.path.join(ROOT, 'sub'), os.path.join(ROOT, 'data', 'cr_lane_x'),
                os.path.join(ROOT, 'data', 'cr_lane', '..'), os.path.join(ROOT, 'data', 'cr_lane', '..', 'engine'),
                os.path.join(ROOT, 'data', 'cr_lane', '..', '..', 'fec')):
        with pytest.raises(L.Refusal):
            L.assert_out_ok(bad, ROOT)


def test_out_refuses_junction_into_live_root(tmp_path):
    link = tmp_path / 'jx'
    try:
        if os.name == 'nt':
            subprocess.run(['cmd', '/c', 'mklink', '/J', str(link), os.path.join(ROOT, 'data')], check=True, capture_output=True)
        else:
            os.symlink(os.path.join(ROOT, 'data'), link)
    except Exception:
        pytest.skip('cannot create link')
    with pytest.raises(L.Refusal):
        L.assert_out_ok(str(link / 'x'), ROOT)


@pytest.mark.parametrize('rid', ['..', '.', 'a/b', r'a\b', r'..\..', 'C:x'])
def test_bad_run_id_refused_nothing_written(tmp_path, rid):
    out = tmp_path / 'o'
    assert L.main(['build', '--out', str(out), '--run-id', rid]) == 3
    assert not out.exists()


def test_build_prechecks_reference_workbooks_even_with_seed_override(tmp_path, monkeypatch, capsys):
    raw = make_raw(tmp_path / 'raw', {})
    seed = tmp_path / 'seed.xlsx'; seed.write_bytes(b'x')
    monkeypatch.setattr(L, 'REF_NETWORK', 'no_such_ref_network.xlsx')
    rc = L.main(['build', '--out', str(tmp_path / 'o'), '--run-id', 'r1', '--raw-root', raw, '--seed-archive', str(seed),
                 '--network-set', 'archive'])
    assert rc == 3 and 'no_such_ref_network.xlsx' in capsys.readouterr().err and not (tmp_path / 'o' / 'r1').exists()


def _pg_world(monkeypatch, ties_ok):
    from network import cr_parity_gate as PG
    E = {('a', 'b'): (1.0, 1, 1, 0)}
    g = {'per': {}, 'stale': [], 'E': E, 'nometa': [], 'ties_by': {}, 'ties': {('a', 'b'): [1.0, 1, 1, 0]} if ties_ok else {}}
    monkeypatch.setattr(PG, 'reconstruct', lambda p: g)
    return PG


def test_write_billset_refused_without_parity(tmp_path, monkeypatch):
    PG = _pg_world(monkeypatch, ties_ok=False)
    out = tmp_path / 'bs.json'
    assert PG.write_billset('x.csv', str(out)) != 0 and not out.exists()


def test_write_billset_written_when_parity_passes(tmp_path, monkeypatch):
    PG = _pg_world(monkeypatch, ties_ok=True)
    out = tmp_path / 'bs.json'
    assert PG.write_billset('x.csv', str(out)) == 0 and out.exists()


def test_parity_refuses_a_non_finite_weight_on_either_side(tmp_path, monkeypatch):
    # CodeRabbit PR #83: a blank reference Weight cell reads as NaN, and abs(NaN - x) > tol is False.
    from network import cr_parity_gate as PG
    for ref_w, tie_w in ((float('nan'), 1.0), (1.0, float('nan')), (float('inf'), 1.0)):
        g = {'per': {}, 'stale': [], 'E': {('a', 'b'): (ref_w, 1, 1, 0)}, 'nometa': [], 'ties_by': {},
             'ties': {('a', 'b'): [tie_w, 1, 1, 0]}}
        monkeypatch.setattr(PG, 'reconstruct', lambda p, g=g: g)
        bad, extra = PG.parity(g)
        assert bad == [('a', 'b')] and extra == [], (ref_w, tie_w)
        out = tmp_path / 'bs.json'
        assert PG.write_billset('x.csv', str(out)) != 0 and not out.exists()


def test_edge_growth_last_action_is_recomputed_for_bills_outside_the_seed():
    from network import cr_edge_growth as EG
    passed = [('2025-02-01', '8000', 'Floor', 'Passed/agreed to in House')]
    bills = {(119, 'hr', 1): {'actions': passed}, (119, 'hr', 2): {'actions': []}}
    la = EG.last_actions(bills)
    assert la[(119, 'hr', 1)] == L.tracker_status(passed, 'hr') and la[(119, 'hr', 2)] == 'Introduced'
