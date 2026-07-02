import openpyxl,sys,json,collections,glob,pandas as pd
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
M=r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx'
wb=openpyxl.load_workbook(M,data_only=True)
nrm=wb['Legislative Score - Normalized']; pr=wb['Press Release - Normalized']
nrow={nrm.cell(r,1).value:r for r in range(3,536) if nrm.cell(r,1).value}
prow={pr.cell(r,1).value:r for r in range(2,535) if pr.cell(r,1).value}
CSJ=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\category_stance.json'))
MAT=CSJ['matrix']; SCO=CSJ['scores']
VC=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\vote_category_lean.json'))
n2b={v:k for k,v in VC['name_of'].items()}; votes=VC['votes']
cats=[('Examination & Supervision',[(3,1)],5),('CU Capital Requirements',[(11,1)],6),('NCUA Budget & Insurance',[(19,1)],7),
('Interchange',[(27,-1)],8),('Payments',[(35,1)],9),('Technology & Emerging',[(43,1)],10),('Fraud Prevention',[(51,1)],11),
('Data Privacy & Cyber',[(59,1)],12),('Charter Modernization',[(67,1)],13),('Defense CU Issues',[(75,1)],14),
('Field of Membership',[(83,1)],15),('CDFI Fund',[(91,1)],17),('CFPB',[(99,1),(145,-1)],19),
('Credit Card Mandates',[(107,-1)],20),('TCPA',[(115,1)],21),('Member Business Lending',[(123,1)],22),('Mortgage & Housing',[(131,1)],23)]
def eng(m,block):
    r=nrow.get(m); return sum((nrm.cell(2,block+o).value or 0)*(nrm.cell(r,block+o).value or 0) for o in (0,2,4)) if r else 0.0
WHO='Richard Durbin'
print('='*70,'\nDURBIN per-category breakdown (leg_signed | press | vote n | combined | TIER)')
for lab,blocks,pc in cats:
    ls=sum(d*eng(WHO,b) for b,d in blocks)
    pv=(pr.cell(prow[WHO],pc).value or 0) if WHO in prow else 0
    vi=votes.get(n2b.get(WHO,''),{}).get(lab); vn=vi['n'] if vi else 0; vl=vi['lean'] if vi else None
    print('  %-26s leg %+6.3f  press %+5.2f  vote %s  -> %+6.3f  %s'%(
        lab,ls,pv,('%.2f(n%d)'%(vl,vn)) if vi else '  -  ',SCO[WHO][lab],MAT[WHO][lab]))
# what BILLS drive Durbin -- from the archive
A=pd.read_excel(sorted(glob.glob(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\congress_bills_archive_*.xlsx'))[-1],sheet_name='All Bills')
du=A[A['Member'].astype(str).str.contains('Durbin',case=False,na=False)]
print('\nDURBIN bills in archive (%d): Section | Role | Bill | Title'%len(du))
for _,r in du.iterrows():
    print('  %-26s %-16s %-9s %s'%(str(r['Section'])[:26],str(r['Role'])[:16],str(r['Name'])[:9],str(r['Title'])[:48]))
# Interchange & Credit Card section bills -- categorization check
for sec in ['Interchange','Credit Card']:
    s=A[A['Section'].astype(str).str.contains(sec,case=False,na=False)]
    bills=s.drop_duplicates('Name')[['Name','Title']]
    print('\n%s section: %d bill-rows, %d distinct bills:'%(sec,len(s),len(bills)))
    for _,r in bills.head(12).iterrows(): print('   %-10s %s'%(str(r['Name'])[:10],str(r['Title'])[:60]))
# locate interchange-routing / credit card competition act
hit=A[A['Title'].astype(str).str.contains('interchange|routing|competition',case=False,na=False)].drop_duplicates('Name')
print('\nBills with interchange/routing/competition in title (-> which Section):')
for _,r in hit.iterrows(): print('   %-10s [%-22s] %s'%(str(r['Name'])[:10],str(r['Section'])[:22],str(r['Title'])[:46]))
# opponent scarcity: tier counts per category
print('\nTIER COUNTS per category (Champion/Ally/Neutral/Skeptic/Opponent/—):')
members=list(MAT.keys())
for lab,_,_ in cats:
    c=collections.Counter(MAT[m][lab] for m in members)
    print('  %-26s C%-4d A%-4d N%-4d S%-3d O%-3d  -%d'%(lab,c['Champion'],c['Ally'],c['Neutral'],c['Skeptic'],c['Opponent'],c['—']))
# known opponents: are they mislabeled champions?
print('\nKNOWN CU-OPPONENTS — champion/ally counts (should be LOW):')
for who in ['Richard Durbin','Maxine Waters','Sherrod Brown','Elizabeth Warren','Jack Reed','Roger Marshall']:
    if who in MAT:
        ch=[c for c,_,_ in cats if MAT[who][c]=='Champion']; al=[c for c,_,_ in cats if MAT[who][c]=='Ally']
        op=[c for c,_,_ in cats if MAT[who][c]=='Opponent']
        print('  %-18s champ=%d %s | ally=%d | opp=%d %s'%(who,len(ch),[x.split(" &")[0][:10] for x in ch],len(al),len(op),op))
