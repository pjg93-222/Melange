import openpyxl, glob, os, sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
WB = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx'
bak = max(glob.glob(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.PRE-STANCESYNC-*.xlsx'), key=os.path.getmtime)
def load(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True); ws = wb['Blended Model']
    h1 = [ws.cell(1, c).value for c in range(1, 29)]; h2 = [ws.cell(2, c).value for c in range(1, 29)]
    data = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(1, 29)] for r in range(3, 536) if ws.cell(r, 1).value}
    wb.close(); return h1, h2, data
h1, h2, A = load(bak); _, _, B = load(WB)
print('headers row1:', [x for x in h1 if x])
print('headers row2:', [x for x in h2 if x])
for c in range(28):
    chg = 0; mx = 0.0; ex = None
    for n in A:
        a = A[n][c]; b = B.get(n, [None]*28)[c]
        if a != b:
            chg += 1
            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                mx = max(mx, abs(a-b))
            elif ex is None:
                ex = (n, a, b)
    if chg:
        print('col%d  %r / %r  changed=%d  max_numeric_diff=%g  ex=%s' % (c+1, h1[c], h2[c], chg, mx, ex))
