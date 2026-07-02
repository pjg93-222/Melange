import openpyxl,sys,json,collections,statistics
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
wb=openpyxl.load_workbook(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx',data_only=True)
pr=wb['Press Release - Normalized']; ra=wb['Rank & Allocation']
members=[ra.cell(r,1).value for r in range(3,536) if ra.cell(r,1).value]
prow={pr.cell(r,1).value:r for r in range(2,535) if pr.cell(r,1).value}
LEG=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\archive_leg.json'))
VC=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\vote_category_lean.json'))
n2b={v:k for k,v in VC['name_of'].items()}; votes=VC['votes']
CS=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\category_stance.json'))
MAT=CS['matrix']; SCO=CS['scores']
cats=[('Examination & Supervision',5),('CU Capital Requirements',6),('NCUA Budget & Insurance',7),('Interchange',8),
('Payments',9),('Technology & Emerging',10),('Fraud Prevention',11),('Data Privacy & Cyber',12),('Charter Modernization',13),
('Defense CU Issues',14),('Field of Membership',15),('CDFI Fund',17),('CFPB',19),('Credit Card Mandates',20),
('TCPA',21),('Member Business Lending',22),('Mortgage & Housing',23)]
print('PER-CATEGORY CONNECTEDNESS + DISTRIBUTION  (leg=archive, press=sheet, vote n>=5)')
print('%-26s | leg+  leg-  | prs+  prs- | vote | CHAMP%%  pos%%  | leg-orientation'%'category')
defensive=[]
for lab,pc in cats:
    leg={m:LEG.get(m,{}).get(lab,0.0) for m in members}
    prs={m:((pr.cell(prow[m],pc).value or 0) if m in prow else 0) for m in members}
    lp=sum(1 for m in members if leg[m]>1e-9); ln=sum(1 for m in members if leg[m]<-1e-9)
    pp=sum(1 for m in members if prs[m]>0.02); pn=sum(1 for m in members if prs[m]<-0.02)
    vv=sum(1 for m in members if (votes.get(n2b.get(m,''),{}).get(lab,{}) or {}).get('n',0)>=5)
    champ=sum(1 for m in members if MAT[m][lab]=='Champion')
    pos=sum(1 for m in members if SCO[m][lab]>0.001); neg=sum(1 for m in members if SCO[m][lab]<-0.001)
    orient='DEFENSIVE (leg mostly anti)' if ln>lp else 'offensive'
    if ln>lp: defensive.append(lab)
    print('%-26s | %4d  %4d | %4d  %4d | %4d | %4d %3d%% %3d%% | %s'%(lab,lp,ln,pp,pn,vv,champ,100*champ//533,100*pos//533,orient))
print('\nDEFENSIVE categories (legislation presence = anti-CU):',defensive)
# For defensive cats: who are the PRESS-based defenders vs the legislative opponents?
for lab,pc in [('Interchange',8),('Credit Card Mandates',20)]:
    prs={m:((pr.cell(prow[m],pc).value or 0) if m in prow else 0) for m in members}
    leg={m:LEG.get(m,{}).get(lab,0.0) for m in members}
    defs=sorted([m for m in members if prs[m]>0.02],key=lambda m:-prs[m])
    opps=sorted([m for m in members if leg[m]<-1e-9],key=lambda m:leg[m])
    print('\n%s  -- PRESS DEFENDERS (press>0): %d ->'%(lab,len(defs)),[ '%s %.2f'%(m.split()[-1],prs[m]) for m in defs[:8]])
    print('%s  -- LEGISLATIVE pushers (leg<0): %d ->'%(lab,len(opps)),[ '%s %.1f'%(m.split()[-1],leg[m]) for m in opps[:8]])
# distribution of combined positive scores per a few cats (to set per-category thresholds)
print('\nPositive-score percentiles per category (p50/p75/p90 among score>0):')
for lab,pc in [('Mortgage & Housing',23),('Member Business Lending',22),('Interchange',8),('Examination & Supervision',5)]:
    ps=sorted(SCO[m][lab] for m in members if SCO[m][lab]>0.001)
    if ps:
        q=lambda f:ps[min(len(ps)-1,int(f*len(ps)))]
        print('  %-26s n+=%-3d  p50 %.2f  p75 %.2f  p90 %.2f  max %.2f'%(lab,len(ps),q(.5),q(.75),q(.9),ps[-1]))
