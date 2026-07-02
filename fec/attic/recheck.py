import csv,sys,re,collections,json
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
BANK={'C00336768','C00104299','C00128512','C00379693','C00034595','C00300178','C00429852','C00004275',
'C00043489','C00364778','C00009639','C00386524','C00008474','C00039305','C00032698','C00035519','C00186064',
'C00018036','C00488882','C00432252','C00648345','C00501429','C00890707','C00438051','C00040535','C00365122',
'C00326595','C00038604','C00410274','C00879890','C00579540','C00589119','C00370015'}  # +Nelnet
fec2mem={}
with open('legislators.csv',encoding='utf-8') as f:
    r=csv.reader(f); h=next(r); fi=h.index('fec_ids');ni=h.index('full_name');bi=h.index('bioguide_id')
    legbio=set()
    for row in r:
        legbio.add(row[bi])
        for cid in re.findall(r'[HSP]\d[A-Z0-9]{2}\d{5}',row[fi]): fec2mem[cid]=(row[ni],row[bi])
candname={}
for line in open('cn.txt',encoding='latin-1'):
    p=line.split('|')
    if len(p)>1: candname[p[0]]=p[1]
bankc=collections.defaultdict(float); unmapped=collections.defaultdict(float)
for line in open('itpas2.txt',encoding='latin-1'):
    p=line.split('|')
    if len(p)<17 or p[0] not in BANK or p[5] not in ('24K','24Z'): continue
    if len(p)>19 and p[19]=='X': continue
    cand=p[16]
    try: amt=float(p[14])
    except: continue
    if cand in fec2mem: bankc[fec2mem[cand][1]]+=amt   # by bioguide
    elif cand: unmapped[cand]+=amt
tot=sum(bankc.values())
print(f'bank total to CURRENT members (w/ Nelnet): ${tot:,.0f} | members receiving: {sum(1 for v in bankc.values() if v>0)}')
print(f'unmapped (to non-current candidates): ${sum(unmapped.values()):,.0f} across {len(unmapped)} candidates')
print('\ntop UNMAPPED recipients (should be challengers/losers, NOT sitting members):')
for cand,amt in sorted(unmapped.items(),key=lambda x:-x[1])[:8]:
    print(f'  {candname.get(cand,cand):34s} ${amt:>8,.0f}  [{cand}]')
# model member coverage: compare model bioguides (Centrality node) to legislators
import openpyxl
cen=openpyxl.load_workbook(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx', data_only=True)['Centrality']
modelbio={cen.cell(r,1).value for r in range(2,cen.max_row+1) if cen.cell(r,1).value}
modelbio={b for b in modelbio if isinstance(b,str) and b and b[0].isalpha()}
missing=[b for b in modelbio if b in legbio is False]
notinleg=[b for b in modelbio if b not in legbio]
print(f'\nmodel network members (bioguide): {len(modelbio)} | not found in legislators-current: {len(notinleg)}')
if notinleg[:10]: print('  e.g.:', notinleg[:10])
