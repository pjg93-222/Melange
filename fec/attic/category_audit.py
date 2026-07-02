import openpyxl,sys,json,collections,statistics
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
P=[];F=[]
def ok(c,m):(P if c else F).append(m)
M=r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx'
wb=openpyxl.load_workbook(M,data_only=True)
nrm=wb['Legislative Score - Normalized']; pr=wb['Press Release - Normalized']; ra=wb['Rank & Allocation']
gen={ra.cell(r,1).value:ra.cell(r,7).value for r in range(3,536) if ra.cell(r,1).value}
members=[nrm.cell(r,1).value for r in range(3,536) if nrm.cell(r,1).value]
nrow={nrm.cell(r,1).value:r for r in range(3,536) if nrm.cell(r,1).value}
prow={pr.cell(r,1).value:r for r in range(2,535) if pr.cell(r,1).value}
VC=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\vote_category_lean.json'))
b2n=VC['name_of']; n2b={v:k for k,v in b2n.items()}; votes=VC['votes']
D=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\category_stance.json'))
J_m=D['matrix']; J_s=D['scores']
cats=[('Examination & Supervision',[(3,1)],5),('CU Capital Requirements',[(11,1)],6),('NCUA Budget & Insurance',[(19,1)],7),
('Interchange',[(27,-1)],8),('Payments',[(35,1)],9),('Technology & Emerging',[(43,1)],10),('Fraud Prevention',[(51,1)],11),
('Data Privacy & Cyber',[(59,1)],12),('Charter Modernization',[(67,1)],13),('Defense CU Issues',[(75,1)],14),
('Field of Membership',[(83,1)],15),('CDFI Fund',[(91,1)],17),('CFPB',[(99,1),(145,-1)],19),
('Credit Card Mandates',[(107,-1)],20),('TCPA',[(115,1)],21),('Member Business Lending',[(123,1)],22),('Mortgage & Housing',[(131,1)],23)]
NEG={'Interchange','Credit Card Mandates'}
def eng(m,block):
    r=nrow.get(m); return sum((nrm.cell(2,block+o).value or 0)*(nrm.cell(r,block+o).value or 0) for o in (0,2,4)) if r else 0.0
def p80(vals):
    a=sorted(abs(x) for x in vals if abs(x)>1e-9); return a[int(0.8*len(a))] if a else 1.0
def clip(x): return max(-1.0,min(1.0,x))
def tier(c):
    return 'Champion' if c>=.40 else 'Ally' if c>=.12 else 'Opponent' if c<=-.40 else 'Skeptic' if c<=-.12 else 'Neutral'
# recompute everything independently
rec={}; recs={}; legP={}; prP={}; voP={}; overlay_changed=0; overlay_cells=0
for lab,blocks,pc in cats:
    legraw={m:sum(d*eng(m,b) for b,d in blocks) for m in members}
    prraw={m:((pr.cell(prow[m],pc).value or 0) if m in prow else 0) for m in members}
    sl=p80(legraw.values()); sp=p80(prraw.values())
    for m in members:
        ln=clip(legraw[m]/sl); pn=clip(prraw[m]/sp)
        vinfo=votes.get(n2b.get(m,''),{}).get(lab); vn=vinfo['n'] if vinfo else 0; hv=bool(vinfo) and vn>=5
        comb=(.55*ln+.35*pn+(.10*clip(vinfo['lean']) if hv else 0))/(1.0 if hv else .90)
        comb_nov=(.55*ln+.35*pn)/.90
        sig=abs(legraw[m])>1e-9 or abs(prraw[m])>1e-9 or hv
        rec[(m,lab)]='—' if not sig else tier(comb); recs[(m,lab)]=round(comb,3)
        legP[(m,lab)]=ln; prP[(m,lab)]=pn; voP[(m,lab)]=clip(vinfo['lean']) if hv else None
        if hv:
            overlay_cells+=1
            if (tier(comb)!=tier(comb_nov)): overlay_changed+=1
# ===== A completeness/bounds =====
ok(len(members)==533,'A1 members = %d'%len(members))
ncell=sum(1 for m in members for c,_,_ in cats)
ok(all(abs(recs[(m,c)])<=1.0001 for m in members for c,_,_ in cats),'A2 all combined scores within [-1,1]')
ok(all((m in J_m and c in J_m[m]) for m in members for c,_,_ in cats),'A3 every member×category present in output (%d cells)'%ncell)
# ===== B reproduce vs stored output =====
mism=[(m,c) for m in members for c,_,_ in cats if rec[(m,c)]!=J_m[m][c]]
ok(len(mism)==0,'B1 independent recompute matches stored matrix exactly (mismatches=%d %s)'%(len(mism),mism[:3]))
sdif=[(m,c) for m in members for c,_,_ in cats if abs(recs[(m,c)]-J_s[m][c])>0.002]
ok(len(sdif)==0,'B2 recomputed scores match stored scores (diffs>0.002 = %d)'%len(sdif))
# ===== C direction correctness =====
# no member is "Champion" on a NEGATIVE category (engaged there = anti-CU)
badchamp=[(m,c) for m in members for c,_,_ in cats if c in NEG and J_m[m][c]=='Champion']
ok(len(badchamp)==0,'C1 no Champion on negative categories Interchange/CreditCard (would be backwards): %d'%len(badchamp))
# any Opponent on a purely-positive category must be press-driven (press<0). CFPB is mixed-direction (Restrain+/Reach-) so legislation can be anti there too.
LEGNEG={'Interchange','Credit Card Mandates','CFPB'}
posopp=[(m,c) for m in members for c,_,_ in cats if c not in LEGNEG and J_m[m][c]=='Opponent']
posopp_press=[t for t in posopp if prP[t]<-0.05]
ok(len(posopp)==0 or len(posopp_press)==len(posopp),'C2 every Opponent on a positive category is press-driven (%d/%d)'%(len(posopp_press),len(posopp)))
# ===== D vote-overlay influence =====
ok(overlay_cells>0,'D1 vote overlay active on %d cells (n>=5); changed tier vs leg+press on %d (%.0f%%)'%(overlay_cells,overlay_changed,100*overlay_changed/max(overlay_cells,1)))
# ===== E consistency with general stance =====
ordv={'Champion':5,'Ally':4,'Neutral':3,'Unengaged':3,'Skeptic':2,'Opponent':1}
net={m:sum(recs[(m,c)] for c,_,_ in cats) for m in members}
xs=[ordv.get(gen.get(m),3) for m in members]; ys=[net[m] for m in members]
def spear(a,b):
    ra_=sorted(range(len(a)),key=lambda i:a[i]); rb=sorted(range(len(b)),key=lambda i:b[i])
    rk=[0]*len(a); rl=[0]*len(a)
    for i,j in enumerate(ra_): rk[j]=i
    for i,j in enumerate(rb): rl[j]=i
    n=len(a); return 1-6*sum((rk[i]-rl[i])**2 for i in range(n))/(n*(n*n-1))
sp_=spear(xs,ys)
ok(sp_>=0.45,'E1 per-category roll-up correlates with general stance (Spearman %.2f, expect >=0.45)'%sp_)
# ===== F divergence diagnosis: general Opponent/Skeptic but category-Champion on >=3 =====
def champs(m): return [c for c,_,_ in cats if J_m[m][c]=='Champion']
div=[m for m in members if gen.get(m) in ('Opponent','Skeptic') and len(champs(m))>=3]
eng_only=0; tot=0
for m in div:
    for c in champs(m):
        tot+=1
        if prP[(m,c)]<=0.05: eng_only+=1   # champion w/ no positive press = pure cosponsorship engagement
fpct=100*eng_only/max(tot,1)
ok(True,'F1 %d general-opponents are category-champions on>=3; of their champ cells %.0f%% are cosponsorship-only (no positive press) -> the known artifact'%(len(div),fpct))
print('AUDIT (category-stance logic) — %d passed, %d FAILED\n'%(len(P),len(F)))
for x in P: print('  [PASS] '+x)
for x in F: print('  [FAIL] '+x)
# illustrative hand-check
print('\nHAND-CHECK cells (member | category | leg | press | vote | combined | tier):')
for m,c in [('Tim Scott','Examination & Supervision'),('Richard Hudson','Interchange'),('Richard Durbin','Credit Card Mandates')]:
    print('   %-15s %-26s leg %+.2f press %+.2f vote %s -> %+.3f %s'%(m,c.split(' &')[0],legP[(m,c)],prP[(m,c)],('%.2f'%voP[(m,c)]) if voP[(m,c)] is not None else ' n/a',recs[(m,c)],J_m[m][c]))
