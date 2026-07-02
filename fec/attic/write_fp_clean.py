import shutil,os,json,sys
import win32com.client as w
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
src=r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx'
bak=r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.PRE-FPCLEAN-20260626.xlsx'
tmp=r'C:\Temp\mb_fpclean.xlsx'
shutil.copy(src,bak); os.makedirs(r'C:\Temp',exist_ok=True); shutil.copy(src,tmp)
cleaned={int(k):v for k,v in json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\fp_clean_cols.json')).items()}
import openpyxl
nrm0=openpyxl.load_workbook(src,data_only=True)['Legislative Score - Normalized']
names=[nrm0.cell(r,1).value for r in range(3,536)]
xl=w.DispatchEx('Excel.Application'); xl.Visible=False; xl.DisplayAlerts=False
wb=xl.Workbooks.Open(tmp); ws=wb.Worksheets('Legislative Score - Normalized')
for col,cmap in cleaned.items():
    vals=[[cmap.get(names[i],None)] for i in range(len(names))]
    ws.Range(ws.Cells(3,col),ws.Cells(2+len(vals),col)).Value=vals
xl.CalculateFull(); wb.Save(); wb.Close(False); xl.Quit(); shutil.copy(tmp,src)
# verify + before/after
wb2=openpyxl.load_workbook(src,data_only=True); bm=wb2['Blended Model']; ra=wb2['Rank & Allocation']
errs=sum(1 for s in wb2.worksheets for row in s.iter_rows() for c in row if isinstance(c.value,str) and c.value in {'#REF!','#DIV/0!','#VALUE!','#N/A','#NAME?'})
alloc=sum((bm.cell(r,21).value or 0) for r in range(3,536)); funded=sum(1 for r in range(3,536) if (bm.cell(r,21).value or 0)>0)
before=json.load(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\fp_before.json'))
after={ra.cell(r,1).value:(ra.cell(r,4).value,bm.cell(r,20).value) for r in range(3,536) if ra.cell(r,1).value}
moved=[(n,before[n][1],after[n][1]) for n in after if n in before and isinstance(before[n][1],(int,float)) and isinstance(after[n][1],(int,float)) and abs(after[n][1]-before[n][1])>=15]
allocch=sum(1 for n in after if n in before)  # placeholder
print('integrity: errors=%d alloc=$%s funded=%d'%(errs,f'{alloc:,.0f}',funded))
print('members moving >=15 rank spots:',len(moved))
for n,b,a in sorted(moved,key=lambda x:abs(x[2]-x[1]),reverse=True)[:8]: print('   %-22s rank %s -> %s'%(n,int(b),int(a)))
print('\nsanity (should be ~stable): Tim Scott %s->%s, Andy Barr %s->%s, Durbin %s->%s'%(
 int(before['Tim Scott'][1]),int(after['Tim Scott'][1]),int(before['Andy Barr'][1]),int(after['Andy Barr'][1]),int(before['Richard Durbin'][1]),int(after['Richard Durbin'][1])))
print('backup',os.path.basename(bak))
