import openpyxl,sys,re,csv,collections,json
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
M=r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx'
wb=openpyxl.load_workbook(M,data_only=True)
nrm=wb['Legislative Score - Normalized']; ra=wb['Rank & Allocation']; bm=wb['Blended Model']
# 15 score-bearing blocks mapped to interaction Sections (CFPB/CDFI excluded - no FP rows), with topic weight
blk={3:('Examination & Supervision',0.5),11:('Credit Union Capital Requirements',0.25),19:('NCUA Budget & Share Insurance Fund',0.5),
27:('Interchange',-2),35:('Payments',0.5),43:('Technology & Emerging Issues',0.5),51:('Fraud Prevention',0.4),
59:('Data Privacy & Cybersecurity',0.5),67:('Charter Modernization & Governance',0.75),75:('Defense Credit Union Issues',0.2),
83:('Field of Membership',0.75),115:('Telephone Consumer Protection Act Issues',0.2),107:('Credit Card & Consumer-Fee Mandates',-3),
123:('Member Business Lending',0.3),131:('Mortgage Lending & Housing',0.4)}
roleoff={'sponsor':0,'original-cosponsor':2,'cosponsor':4}
def billkey(name,cong):
    s=str(name).upper(); ch='H' if s.startswith('H') else ('S' if s.startswith('S') else (re.sub(r'[^HS]','',s)[:1] or '?'))
    return f"{int(float(cong))}_{ch}_{re.sub(chr(92)+'D','',s)}"
fp=set()
for row in csv.DictReader(open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\bill_drop_final.csv')): fp.add(billkey(row['bill'],row['congress']))
# aggregate decay totals & FP-decay per (member, section, role)
tot=collections.defaultdict(float); fpd=collections.defaultdict(float)
for sh,role in [('Sponsor Interactions','sponsor'),('Original Co-Sponsor Interaction','original-cosponsor'),('Co-Sponsor Interactions','cosponsor')]:
    ws=wb[sh]
    for r in range(2,ws.max_row+1):
        sec=ws.cell(r,1).value; mem=ws.cell(r,3).value; nm=ws.cell(r,4).value; cong=ws.cell(r,10).value; dec=ws.cell(r,11).value
        if not sec or not mem: continue
        if not isinstance(dec,(int,float)): dec=1.0
        tot[(mem,sec,role)]+=dec
        if cong is not None and billkey(nm,cong) in fp: fpd[(mem,sec,role)]+=dec
names=[nrm.cell(r,1).value for r in range(3,536)]; nrow={nrm.cell(r,1).value:r for r in range(3,536) if nrm.cell(r,1).value}
cleaned={}   # col -> {member: new_value}
legdelta=collections.defaultdict(float)
sub={}
for sc,(sec,tw) in blk.items():
    for role,off in roleoff.items():
        col=sc+off; sw=nrm.cell(2,col).value or 0
        if sw==0: continue
        cmap={}
        for m in names:
            r=nrow.get(m)
            if not r: continue
            v=nrm.cell(r,col).value or 0; t=tot.get((m,sec,role),0); f=fpd.get((m,sec,role),0)
            share=(f/t) if t>0 else 0.0
            nv=v*(1-share); cmap[m]=round(nv,12)
            legdelta[m]+=tw*sw*(nv-v)
        cleaned[col]=cmap
# before ranks/scores
before={ra.cell(r,1).value:(ra.cell(r,4).value,bm.cell(r,20).value) for r in range(3,536) if ra.cell(r,1).value}
json.dump({str(k):v for k,v in cleaned.items()},open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\data\legislation\fp_clean_cols.json','w'))
json.dump(before,open(r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\fp_before.json','w'))
naff=sum(1 for m in legdelta if abs(legdelta[m])>1e-9)
print('FP bills:',len(fp),'| members with ANY legislative change:',naff,'/533 | cols touched:',len(cleaned))
print('\nbiggest legislative-score DROPS (FP inflation removed):')
for m,d in sorted(legdelta.items(),key=lambda x:x[1])[:12]:
    if d<0: print('   %-24s %+.4f'%(m,d))
print('\nany INCREASES (FP in a negative-weight/anti topic):')
for m,d in sorted(legdelta.items(),key=lambda x:-x[1])[:5]:
    if d>1e-9: print('   %-24s %+.4f'%(m,d))
