# -*- coding: utf-8 -*-
"""Repair the 'Sub Layer - Raw' per-category block C:U (cols 3-21) — the display-plane sibling of the
AG:AW efficiency repair (repair_sublayer_references.py, 2026-07-17). Paul-directed 2026-07-17 ("Fix all
that"). Three defect classes, all from the 07-15 LegNorm row-shift plus original drift:

  (1) DATA ROW off by one: member at SubRaw row r lives at LegNorm row r-1; formulas reference row r
      (each member scored on the alphabetical successor; Zoe Lofgren row 536 reads EMPTY row 536).
  (2) WEIGHT VECTOR wrong row: SUMPRODUCTs weight against LegNorm $3 (Aaron Bean's DATA row post-shift);
      the true weight row is $2.
  (3) CFPB Restrain window truncated (CU:CX, 4 cols) -> canonical 5-col window CU:CY.
  PLUS: col 14 (Credit Union Tax Status) was press-only because no legislative block existed; the block
      now exists (cols 153-160, added 2026-07-17) -> add the standard legislative SUMPRODUCT term
      (EW:FA x EW$2:FA$2), keeping the existing press term and the user-owned row-3 weight untouched.
  Col 16 (CU-Bank Competitive Dynamics) stays press-only: no legislative block exists for it.

RANK-SAFE: SubRaw C:U feeds only 'Sub Layer - Normalized' C:U -> 'Primary Layer' cols 7-13 (display);
the score path consumes Primary col E (efficiency, AG:AW) + Blended col 4 only. Zero rank delta is the
acceptance gate. Blast radius independently confirmed by scout wf_243aaf89 before --apply.

--plan (default): read-only audit of every target cell; prints deviation counts; writes plan JSON.
--apply: backup FIRST (backups/PRE-SUBLAYER-CU-<ts>) -> rewrite -> fullCalcOnLoad -> atomic save.
"""
import io,sys,os,re,json,shutil,datetime,collections
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding='utf-8')
sys.path.insert(0,r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\fec')
import openpyxl
from openpyxl.utils import get_column_letter
from paths import ROOT as B
from moneyball_core import save_workbook_atomic

WB=B+r'\Moneyball Main Version - Blended.xlsx'
PLAN=os.path.join(os.path.dirname(os.path.abspath(__file__)),'repair_sublayer_cu_plan.json')
APPLY='--apply' in sys.argv

# slot -> (category header prefix on LegNorm row 1, first LegNorm col of the 5-col window)
SLOTS={3:('Examination & Supervision',3),4:('Credit Union Capital Requirements',11),
 5:('NCUA Budget',19),6:('Interchange',27),7:('Payments',35),8:('Technology & Emerging',43),
 9:('Fraud Prevention',51),10:('Data Privacy',59),11:('Charter',67),12:('Defense',75),
 13:('Field of Membership',83),15:('CDFI Fund',91),17:('CFPB Restrain',99),19:('Credit Card',107),
 18:('Telephone|TCPA',115),20:('Member Business',123),21:('Mortgage',131),14:('Credit Union Tax Status',153)}
PRESS=re.compile(r"SUMPRODUCT\('Press Release - Normalized'![A-Z]{1,2}\d+,'Press Release - Normalized'![A-Z]{1,2}\$535\)")

lock=os.path.join(os.path.dirname(WB),'~$'+os.path.basename(WB))
if os.path.exists(lock): sys.exit('REFUSING: Excel lock present.')
wb=openpyxl.load_workbook(WB)
sr=wb['Sub Layer - Raw']; ln=wb['Legislative Score - Normalized']

# assert the LegNorm windows are what we think they are (fail loud on drift)
for c,(pre,w0) in SLOTS.items():
    hdr=str(ln.cell(1,w0).value or '')
    ok=any(hdr.lower().startswith(p.lower()) for p in pre.split('|'))
    if not ok: sys.exit('REFUSING: LegNorm col %d header %r does not start with %r (slot %d)'%(w0,hdr,pre,c))
    w2=ln.cell(2,w0).value
    if not (isinstance(w2,(int,float)) and abs(w2-2/7.0)<1e-6):
        sys.exit('REFUSING: LegNorm col %d row-2 weight %r != 2/7 (slot %d)'%(w0,w2,c))

def canonical(r,c):
    pre,w0=SLOTS[c]
    a=get_column_letter(w0); b2=get_column_letter(w0+4)
    leg=("SUMPRODUCT('Legislative Score - Normalized'!%s%d:%s%d,"
         "'Legislative Score - Normalized'!%s$2:%s$2)")%(a,r-1,b2,r-1,a,b2)
    cur=str(sr.cell(r,c).value or '')
    m=PRESS.search(cur)
    press=m.group(0) if m else None
    return leg,press,cur

devs=[]; tax_added=0
for r in range(4,537):
    nm=sr.cell(r,2).value
    if not nm: continue
    own=str(ln.cell(r-1,1).value or '').strip()
    if own!=str(nm).strip():
        sys.exit('REFUSING: row map broken at SubRaw %d (%r) vs LegNorm %d (%r)'%(r,nm,r-1,own))
    for c in SLOTS:
        leg,press,cur=canonical(r,c)
        if c==14:
            if press is None: sys.exit('REFUSING: col14 press term missing at row %d: %r'%(r,cur[:80]))
            want='='+leg+' + '+press
            if cur.strip()!=want: devs.append((r,c,want))
            continue
        if press is None:
            devs.append((r,c,'='+leg)); continue
        want='='+leg+' + '+press
        if cur.strip()!=want: devs.append((r,c,want))
kinds=collections.Counter(get_column_letter(c) for _,c,_ in devs)
print('deviating cells: %d of %d | by slot: %s'%(len(devs),533*len(SLOTS),dict(sorted(kinds.items()))))
if not APPLY:
    json.dump({'generated':datetime.datetime.now().isoformat(timespec='seconds'),
               'count':len(devs),'sample':[{'cell':'%s%d'%(get_column_letter(c),r),'want':w[:200]} for r,c,w in devs[:12]]},
              open(PLAN,'w',encoding='utf-8'),indent=1)
    print('plan -> %s  (read-only; use --apply)'%PLAN)
    wb.close(); sys.exit(0)

ts=datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
bk=os.path.join(B,'backups','Moneyball Main Version - Blended.PRE-SUBLAYER-CU-%s.xlsx'%ts)
shutil.copy2(WB,bk); print('backup ->',os.path.basename(bk))
for r,c,want in devs:
    sr.cell(r,c).value=want
wb.calculation.fullCalcOnLoad=True
save_workbook_atomic(wb,WB)
print('APPLIED: %d cells rewritten (incl. col-14 Tax legislative terms). fullCalcOnLoad set; saved atomically.'%len(devs))
