# -*- coding: utf-8 -*-
"""Full-scale stance audit: for every member compare the assigned overall Stance Tier to a robust
ground-truth press signal (recency-weighted mean of the CORRECTED per-article sentiment — un-amplified,
so a single weight-1 category mislabel can't dominate) plus the legislative pillar. Flags mismatches.
Writes stance_audit.csv. No changes to any model file."""
import openpyxl, sys, io, csv, collections, statistics
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
WB = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\Moneyball Main Version - Blended.xlsx'
def fs(x):
    try: x = float(x)
    except (TypeError, ValueError): return 0
    return 1 if x > 0 else (-1 if x < 0 else 0)

wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
ra = wb['Rank & Allocation']
L = {r: ra.cell(r, 12).value for r in range(3, 10)}
mem = {}
for r in range(3, 536):
    n = ra.cell(r, 1).value
    if n: mem[n] = {'leg': ra.cell(r, 4).value or 0, 'press_wb': ra.cell(r, 3).value or 0,
                    'assigned': ra.cell(r, 7).value}
# corrected per-article sentiment -> robust per-member press lean (recency-wt mean) + counts
ws = wb['Press Release Articles']; H = {h: i for i, h in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))}
art = collections.defaultdict(lambda: [0.0, 0.0, 0, 0, 0])  # name -> [sum grd*rw, sum rw, npos, nneu, nneg]
for r in ws.iter_rows(min_row=2, values_only=True):
    nm = r[H['member']]
    if nm not in mem: continue
    src = r[H['stance_src']]; fsign = fs(r[H['score']]) if src == 'bill' else fs(r[H['stance_sign']])
    mag = abs(r[H['score_graded']] or 0); grd = 0.0 if fsign == 0 else fsign * (mag if mag > 0 else 0.2)
    rw = r[H['recency_weight']] or 1
    a = art[nm]; a[0] += grd * rw; a[1] += rw
    a[2] += grd > 0; a[3] += grd == 0; a[4] += grd < 0
# bill counts
bills = collections.defaultdict(int)
for sh in ['Sponsor Interactions', 'Original Co-Sponsor Interaction', 'Co-Sponsor Interactions']:
    for r in wb[sh].iter_rows(min_row=2, values_only=True):
        if r[2] in mem: bills[r[2]] += 1
wb.close()

L3, L4, L5, L6, L7, L8, L9 = (L[i] for i in range(3, 10))
def rec_stance(leg, gtp):     # workbook formula, but press = robust gt_press (mean scale)
    if leg <= L3 or (gtp <= L4 and leg < L8): return 'Opponent'
    if leg <= L5 or (gtp <= L6 and leg < L8): return 'Skeptic'
    if leg >= L7: return 'Champion'
    if leg >= L8: return 'Ally'
    if leg > L9: return 'Neutral'
    return 'Unengaged'

rows = []
for n, d in mem.items():
    a = art[n]; gtp = (a[0] / a[1]) if a[1] > 0 else 0.0
    d['gtp'] = gtp; d['pos'] = a[2]; d['neu'] = a[3]; d['neg'] = a[4]; d['nbills'] = bills[n]
    d['recommended'] = rec_stance(d['leg'], gtp)
    rows.append((n, d))

mism = [(n, d) for n, d in rows if d['assigned'] != d['recommended']]
print('AUDIT: %d members | assigned vs recommended mismatches: %d' % (len(rows), len(mism)))
# transitions
tr = collections.Counter((d['assigned'], d['recommended']) for n, d in mism)
print('\ntop transitions (assigned -> recommended):')
for (a, b), c in tr.most_common(12): print('   %-10s -> %-10s : %d' % (a, b, c))
# opponents that look wrong: assigned Opponent but gt_press not negative
opp_wrong = [(n, d) for n, d in rows if d['assigned'] == 'Opponent' and d['gtp'] > -0.05]
print('\nassigned OPPONENT but corrected press NOT negative (gtp>-0.05): %d' % len(opp_wrong))
for n, d in sorted(opp_wrong, key=lambda x: -x[1]['gtp'])[:15]:
    print('   %-22s leg=%+.2f press_wb=%+.2f gtp=%+.2f arts(+/0/-)=%d/%d/%d -> rec %s' % (
        n[:22], d['leg'], d['press_wb'], d['gtp'], d['pos'], d['neu'], d['neg'], d['recommended']))
# write full csv
out = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\data\stance\stance_audit.csv'
with io.open(out, 'w', encoding='utf-8', newline='') as f:
    w = csv.writer(f); w.writerow(['member', 'assigned', 'recommended', 'mismatch', 'leg', 'press_workbook',
                                   'gt_press_corrected', 'art_pos', 'art_neu', 'art_neg', 'n_bills'])
    for n, d in sorted(rows, key=lambda x: x[0]):
        w.writerow([n, d['assigned'], d['recommended'], 'YES' if d['assigned'] != d['recommended'] else '',
                    round(d['leg'], 3), round(d['press_wb'], 3), round(d['gtp'], 3),
                    d['pos'], d['neu'], d['neg'], d['nbills']])
print('\nwrote', out)
