import openpyxl, glob, os, sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
B = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WB = B + r'\Moneyball Main Version - Blended.xlsx'
PRIS = max(glob.glob(B + r'\Moneyball Main Version - Blended.PRE-STANCEGUARD-*.xlsx'), key=os.path.getmtime)
def load(p, sheet):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True); ws = wb[sheet]
    h = [ws.cell(2, c).value or ws.cell(1, c).value for c in range(1, ws.max_column+1)]
    d = {ws.cell(r,1).value: [ws.cell(r,c).value for c in range(1, ws.max_column+1)] for r in range(3, 536) if ws.cell(r,1).value}
    wb.close(); return h, d
for sheet in ['Blended Model', 'Rank & Allocation']:
    h, A = load(PRIS, sheet); _, Bd = load(WB, sheet)
    ncols = len(h)
    print('=== %s (%d cols) ===' % (sheet, ncols))
    any_chg = False
    for c in range(ncols):
        chg = 0; mx = 0.0
        for n in A:
            a = A[n][c]; b = Bd.get(n, [None]*ncols)[c]
            if a != b:
                chg += 1
                if isinstance(a,(int,float)) and isinstance(b,(int,float)): mx = max(mx, abs(a-b))
        if chg:
            any_chg = True
            print('  col%d %r  changed=%d  max_numeric_diff=%g' % (c+1, h[c], chg, mx))
    if not any_chg: print('  NO CHANGES (identical to pristine)')
