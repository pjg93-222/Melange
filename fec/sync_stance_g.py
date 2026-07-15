# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.abspath(__file__)));from paths import ROOT as MB_ROOT  # centralized root
"""Col-G-ONLY stance sync: recompute the overall stance ladder from the category_stance matrix with
the ABB26 IMP weights and write ONLY Rank & Allocation col G.

PHASE-2 (system-of-record inversion, user sign-off 2026-07-02): rewritten from Excel COM to openpyxl —
no COM, no Excel process, no kill-orphan, no calc=MANUAL dance. Python owns all derived math
(compute_pillars --write runs right after this in the chain); the workbook's own formulas refresh
whenever the user opens it (fullCalcOnLoad set on save). Values-only write; every formula preserved.
Run after any IMP change or Category Stance edit (run_pipeline does this automatically)."""
import json,sys,collections
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
import openpyxl
from moneyball_core import overall_tier,save_workbook_atomic   # shared stance law + crash-safe workbook replace
B=MB_ROOT
WB=B+r'\Moneyball Main Version - Blended.xlsx'
CS=json.load(open(B+r'\category_stance.json',encoding='utf-8'))
M=CS['matrix']; OVERALL_CAPS=(CS.get('gate') or {}).get('overall_capped') or {}
def agg(n): return overall_tier(M.get(n,{}),n,OVERALL_CAPS)
wb=openpyxl.load_workbook(WB)                       # formulas preserved
ws=wb['Rank & Allocation']; cnt=collections.Counter(); n=0
for r in range(3,536):
    nm=ws.cell(r,1).value
    if not nm: continue
    s=agg(nm); ws.cell(r,7).value=s; cnt[s]+=1; n+=1
wb.calculation.fullCalcOnLoad=True                  # the VIEW self-heals on next Excel open
save_workbook_atomic(wb,WB)
print('wrote %d stance cells to col G (openpyxl, no COM; view recalcs on open)'%n)
print('distribution:',dict(cnt))
print('SAVED.')
