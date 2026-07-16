# -*- coding: utf-8 -*-
"""Build the compact **Moneyball Audit Workbook** (openpyxl port of ``build_audit_workbook.mjs``).

Why this exists
---------------
The historical ``build_audit_workbook.mjs`` / ``verify_audit_workbook.mjs`` builders
``require('@oai/artifact-tool')`` — a package that only exists inside the OpenAI
code-interpreter sandbox.  On the analyst workstation there is no ``package.json``, no
``node_modules``, and ``node`` is not on ``PATH``, so the ``.mjs`` path is dead here.
This module reproduces the same *value contract* with openpyxl and the shipped Python
toolchain, no network access required.

Data flow (mirrors the ``.mjs``):
  * ``extract_workbook_audit_inputs.py`` reads the legacy source workbook (read-only) and
    emits the JSON handoff (``moneyball-workbook-audit-inputs-v1``): the raw & self-normalized
    ``Legislative Score`` grids, per-member agenda rows, source SHA-256, and the count of
    legacy displaced normalized cells.
  * ``ranks_python.json`` is the rank/score authority; ``category_stance.json`` is the CU
    advocacy-frame stance matrix.

Six sheets are written (``Overview``, ``Rank Authority``, ``Category Stance``,
``Legislative Agenda``, ``Legislative Raw``, ``Legislative Normalized``) whose cell VALUES
satisfy the four contracts in ``tests/test_workbook_contracts.py``.

Not reproduced (no contract depends on them): the ``.mjs`` cosmetic styling (fills, borders,
column autofit) and the PNG sheet thumbnails.  Light styling only (bold headers, a title cell,
freeze panes) is applied for readability.
"""
from __future__ import annotations

import argparse
import json
import os
import sys

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from paths import ROOT, WORKBOOK as SOURCE_WORKBOOK  # noqa: E402  (SOURCE, not the audit output)

HANDOFF_SCHEMA = "moneyball-workbook-audit-inputs-v1"
EMDASH = "—"  # U+2014, matches the stance-matrix no-position glyph exactly
OUT_DIR = os.path.join(ROOT, "outputs", "workbook_repair")
OUT_NAME = "Moneyball Audit Workbook.xlsx"
HANDOFF_NAME = "workbook_audit_inputs.json"

RANK_HEADERS = [
    "Member", "Score", "Rank", "Parity Score", "Parity Rank", "Legislative %", "Press %",
    "Influence %", "Demographic %", "Influence Raw", "Office Rank Bonus", "Influence Raw Parity",
    "Display Role Weight", "Role Office", "Display Eligibility",
]


def _read_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def ensure_inputs(handoff_path, source_workbook_path):
    """Produce the JSON handoff by running the read-only extractor when it is absent."""
    if os.path.exists(handoff_path):
        return handoff_path
    from extract_workbook_audit_inputs import extract
    from moneyball_core import dump_atomic
    doc = extract(os.path.abspath(source_workbook_path))
    os.makedirs(os.path.dirname(handoff_path) or ".", exist_ok=True)
    dump_atomic(doc, handoff_path, indent=1)
    return handoff_path


def build(handoff, category, ranks, out_path):
    """Write the audit workbook from the loaded JSON authorities to ``out_path``."""
    if handoff.get("schema") != HANDOFF_SCHEMA:
        raise ValueError("workbook handoff schema mismatch: %r" % (handoff.get("schema"),))
    bold = Font(bold=True)
    title_font = Font(bold=True, size=16)
    italic = Font(italic=True)

    wb = openpyxl.Workbook()

    # -- Overview -----------------------------------------------------------------
    ov = wb.active
    ov.title = "Overview"
    ov["A1"] = "Moneyball Audit Workbook"
    ov["A1"].font = title_font
    ov.merge_cells("A1:H1")
    overview_rows = [
        ["Scope", "Generated review view; legacy source workbook is not a rank or stance authority"],
        ["Members", len(ranks["rank"])],
        ["Rank authority", "ranks_python.json " + EMDASH + " workbook-coherent score and rank; formal office is display-only"],
        ["Category authority", "category_stance.json " + EMDASH + " CU advocacy frame"],
        ["Source workbook", handoff["source_workbook"]],
        ["Source SHA-256", handoff["source_workbook_sha256"]],
        ["Legacy displaced normalized cells found", handoff["legacy_normalized_cell_mismatches"]],
        ["Repair", "Normalization rebuilt from each member's same physical source row"],
    ]
    for i, (key, value) in enumerate(overview_rows):
        r = 3 + i
        ov.cell(r, 1, key).font = bold
        ov.cell(r, 2, value)
    diag = ranks.get("effective_weight_diagnostics", {}) or {}
    nominal = diag.get("nominal_weights", {}) or {}
    for j, head in enumerate(["Pillar", "Nominal weight", "Pillar std. dev.", "Weighted spread", "Realized spread share"]):
        ov.cell(12, 1 + j, head).font = bold
    for i, key in enumerate(nominal):
        r = 13 + i
        ov.cell(r, 1, key)
        ov.cell(r, 2, nominal.get(key))
        ov.cell(r, 3, (diag.get("pillar_stddev") or {}).get(key))
        ov.cell(r, 4, (diag.get("weighted_spread") or {}).get(key))
        ov.cell(r, 5, (diag.get("realized_spread_share") or {}).get(key))

    # -- Rank Authority -----------------------------------------------------------
    ra = wb.create_sheet("Rank Authority")
    ra["A1"] = "Rank Authority"
    ra["A1"].font = title_font
    ra.merge_cells("A1:O1")
    ra["A2"] = ("Shipped Python authority. Formal office is titled display context only and never "
                "modifies Influence or rank.")
    ra["A2"].font = italic
    ra.merge_cells("A2:O2")
    for j, head in enumerate(RANK_HEADERS):
        ra.cell(4, 1 + j, head).font = bold
    ordered = sorted(ranks["rank"], key=lambda n: (ranks["rank"][n], n))
    pcts = ranks.get("pcts", {}) or {}
    for i, name in enumerate(ordered):
        r = 5 + i
        p = pcts.get(name, {}) or {}
        ra.cell(r, 1, name)
        ra.cell(r, 2, ranks["score"].get(name))
        ra.cell(r, 3, ranks["rank"].get(name))
        ra.cell(r, 4, ranks["base_score"].get(name))
        ra.cell(r, 5, ranks["base_rank"].get(name))
        ra.cell(r, 6, p.get("leg"))
        ra.cell(r, 7, p.get("press"))
        ra.cell(r, 8, p.get("inf"))
        ra.cell(r, 9, p.get("dem"))
        ra.cell(r, 10, ranks["influence_raw_base"].get(name))
        ra.cell(r, 11, ranks["formal_role_influence_bonus"].get(name))
        ra.cell(r, 12, ranks["influence_raw_adjusted"].get(name))
        ra.cell(r, 13, ranks["formal_role_display_weight"].get(name))
        ra.cell(r, 14, ranks["formal_role_influence_office"].get(name))
        ra.cell(r, 15, ranks["formal_role_influence_eligibility"].get(name))
    ra.freeze_panes = "A5"

    # -- Category Stance ----------------------------------------------------------
    cats = list(category["cats"])
    matrix = category["matrix"]
    cs = wb.create_sheet("Category Stance")
    cs["A1"] = "Current Category Stance " + EMDASH + " CU Advocacy Frame"
    cs["A1"].font = title_font
    cs.merge_cells("A1:T1")
    for j, head in enumerate(["Member", "General Stance"] + cats):
        cs.cell(2, 1 + j, head).font = bold
    general = {row[0]: row[2] for row in handoff["rank_rows"]}
    for i, name in enumerate(sorted(matrix)):
        r = 3 + i
        cs.cell(r, 1, name)
        cs.cell(r, 2, general.get(name) or EMDASH)
        member = matrix.get(name) or {}
        for j, cat in enumerate(cats):
            cs.cell(r, 3 + j, member.get(cat, EMDASH))
    cs.cell(537, 1, "Positive means aligned with the credit-union advocacy position. Defensive issue "
                    "signs are normalized upstream; this sheet is generated, never hand-pasted.")
    cs["A537"].font = italic
    cs.merge_cells("A537:T537")
    cs.freeze_panes = "C3"

    # -- Legislative Agenda -------------------------------------------------------
    la = wb.create_sheet("Legislative Agenda")
    la["A1"] = "Legislative Agenda " + EMDASH + " Fresh Source View"
    la["A1"].font = title_font
    la.merge_cells("A1:C1")
    for j, head in enumerate(["Member", "Legislative agenda raw", "General Stance"]):
        la.cell(2, 1 + j, head).font = bold
    for i, row in enumerate(handoff["rank_rows"]):
        r = 3 + i
        la.cell(r, 1, row[0])
        la.cell(r, 2, row[1])
        la.cell(r, 3, row[2])
    la.cell(537, 1, "Source workbook SHA-256: " + str(handoff["source_workbook_sha256"]))
    la["A537"].font = italic
    la.merge_cells("A537:C537")
    la.freeze_panes = "A3"

    # -- Legislative Raw / Normalized (bulk grids, start at A1) --------------------
    lr = wb.create_sheet("Legislative Raw")
    for row in handoff["legislative_raw"]:
        lr.append(list(row))
    lr.freeze_panes = "C3"
    ln = wb.create_sheet("Legislative Normalized")
    for row in handoff["legislative_normalized"]:
        ln.append(list(row))
    ln.freeze_panes = "C3"

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path


def build_default(out_path=None, handoff_path=None, source_workbook_path=None):
    """Convenience entry point: ensure the JSON handoff exists, then build the workbook.

    Loads ``ranks_python.json`` and ``category_stance.json`` from ROOT and always (re)writes
    the workbook so callers/tests exercise the current builder rather than a stale artifact.
    """
    out_path = out_path or os.path.join(OUT_DIR, OUT_NAME)
    handoff_path = handoff_path or os.path.join(OUT_DIR, HANDOFF_NAME)
    source_workbook_path = source_workbook_path or SOURCE_WORKBOOK
    ensure_inputs(handoff_path, source_workbook_path)
    handoff = _read_json(handoff_path)
    category = _read_json(os.path.join(ROOT, "category_stance.json"))
    ranks = _read_json(os.path.join(ROOT, "ranks_python.json"))
    return build(handoff, category, ranks, out_path)


def main():
    parser = argparse.ArgumentParser(description="Build the Moneyball Audit Workbook (openpyxl).")
    parser.add_argument("--handoff", default=None,
                        help="workbook_audit_inputs.json path (extracted read-only if absent)")
    parser.add_argument("--source-workbook", default=None,
                        help="legacy source workbook to extract from when --handoff is absent")
    parser.add_argument("--out", default=None, help="output .xlsx path")
    args = parser.parse_args()
    out = build_default(out_path=args.out, handoff_path=args.handoff,
                        source_workbook_path=args.source_workbook)
    print("built audit workbook: %s" % out)


if __name__ == "__main__":
    main()
