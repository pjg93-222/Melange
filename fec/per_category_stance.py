import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.abspath(__file__)));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
import openpyxl, sys, json, collections, math
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
M=MB_ROOT+r'\Moneyball Main Version - Blended.xlsx'
wb=openpyxl.load_workbook(M,data_only=True)
ra=wb['Rank & Allocation']
import moneyball_core as MC
_normM,_w535,_=MC.press_matrix()   # PHASE-2: Python-owned press normalization (raw values + row-535 dial literals; no recalc dependency)
gen={ra.cell(r,1).value:ra.cell(r,7).value for r in range(3,536) if ra.cell(r,1).value}
members=[ra.cell(r,1).value for r in range(3,536) if ra.cell(r,1).value]
LEG=json.load(open(MB_ROOT+r'\archive_leg.json'))
VC=json.load(open(MB_ROOT+r'\vote_category_lean.json'))
n2b={v:k for k,v in VC['name_of'].items()}; votes=VC['votes']
cats=[('Examination & Supervision',5),('CU Capital Requirements',6),('NCUA Budget & Insurance',7),('Interchange',8),
('Payments',9),('Technology & Emerging',10),('Fraud Prevention',11),('Data Privacy & Cyber',12),('Charter Modernization',13),
('Defense CU Issues',14),('Field of Membership',15),('CU Tax Status',16),('CDFI Fund',17),('CFPB',19),('Credit Card Mandates',20),
('TCPA',21),('Member Business Lending',22),('Mortgage & Housing',23)]
# deep-audit fix 2026-07-02: 'CU Tax Status' (Normalized col 16) was added to category_stance.json by hand
# on 2026-07-01 but never to THIS producer list — the next regeneration silently dropped the 18th category
# while the ABB26 IMP dicts (weight 3.0!) kept expecting it. Data now exists (190 press rows + social).
# NOT added to DEFENSIVE (weighting change = user's call; it is a defend-type fight — flag standing).
DEFENSIVE={'Interchange','Credit Card Mandates'}
def weights(cat): return (0.55,0.70,0.20) if cat in DEFENSIVE else (0.55,0.35,0.10)
pw=[_w535[pc] for _,pc in cats]
def p80(vals):
    a=sorted(abs(x) for x in vals if abs(x)>1e-9); return a[int(0.8*len(a))] if a else 1.0
def clip(x): return max(-1.0,min(1.0,x))
press={(m,lab):_normM[pc].get(m,0.0) for m in members for lab,pc in cats}
opl={m:sum(press[(m,lab)]*pw[i] for i,(lab,pc) in enumerate(cats)) for m in members}
neg=sorted(v for v in opl.values() if v<0); ascale=abs(neg[int(0.08*len(neg))]) if neg else 1.0
ANTI_STANCE={'Opponent':1.0,'Skeptic':0.5}
anti={m:max(clip(-opl[m]/(ascale or 1)),ANTI_STANCE.get(gen.get(m),0.0)) for m in members}
PRIOR=0.70
# ---- self-calibrating threshold engine: median + k*IQR, rate clamped to a sane band ----
K=1.25; RFLOOR,RCAP=0.10,0.40; MINN=8
def qtile(s,q):
    return s[min(len(s)-1,max(0,int(round(q*(len(s)-1)))))] if s else None
def shaped_cut(vals,hi,rfloor=RFLOOR,rcap=RCAP,k=K):
    s=sorted(vals); n=len(s)
    if n<MINN: return None,None
    q1=qtile(s,0.25); med=qtile(s,0.5); q3=qtile(s,0.75); iqr=max(q3-q1,1e-9)
    if hi:
        c=med+k*iqr; rate=min(rcap,max(rfloor,sum(1 for x in s if x>=c)/n)); return qtile(s,1-rate),rate
    c=med-k*iqr; rate=min(rcap,max(rfloor,sum(1 for x in s if x<=c)/n)); return qtile(s,rate),rate
# PASS 1: combined score + raw leg/press
comb={}; hassig={}; legR={}; prR={}
for lab,pc in cats:
    WL,WP,WV=weights(lab)
    legraw={m:LEG.get(m,{}).get(lab,0.0) for m in members}; prraw={m:press[(m,lab)] for m in members}
    legc={m:math.copysign(math.sqrt(abs(legraw[m])),legraw[m]) for m in members}   # sqrt-compress heavy tail -> reduce ceiling saturation
    sl=p80(legc.values()); sp=p80(prraw.values())
    for m in members:
        ln=clip(legc[m]/sl); pn=clip(prraw[m]/sp)
        vinfo=votes.get(n2b.get(m,''),{}).get(lab); vn=vinfo['n'] if vinfo else 0; hv=bool(vinfo) and vn>=5
        vc=WV*clip(vinfo['lean']) if hv else 0
        c=clip(WL*ln+WP*pn+vc) if lab in DEFENSIVE else (WL*ln+WP*pn+vc)/(WL+WP+(WV if hv else 0))
        if c>0 and prraw[m]<=0.02 and anti[m]>0: c*=(1-PRIOR*anti[m])
        comb[(lab,m)]=c; legR[(lab,m)]=legraw[m]; prR[(lab,m)]=prraw[m]
        hassig[(lab,m)]=abs(legraw[m])>1e-9 or abs(prraw[m])>1e-9 or hv
# pooled fallbacks for thin categories (no global constants -> pooled empirical)
OFF=[lab for lab,_ in cats if lab not in DEFENSIVE]
POOLp=sorted(comb[(lab,m)] for lab in OFF for m in members if hassig[(lab,m)] and comb[(lab,m)]>0)
POOLn=sorted(comb[(lab,m)] for lab in OFF for m in members if hassig[(lab,m)] and comb[(lab,m)]<0)
pooled_champ=qtile(POOLp,0.88); pooled_opp=qtile(POOLn,0.12); pooled_ally=qtile(POOLp,0.55); pooled_skep=qtile(POOLn,0.45)
POOLlegneg=sorted(legR[(lab,m)] for lab in DEFENSIVE for m in members if legR[(lab,m)]<0)
POOLprspos=sorted(prR[(lab,m)] for lab in DEFENSIVE for m in members if prR[(lab,m)]>0)
# PASS 2: classify
matrix=collections.defaultdict(dict); scores=collections.defaultdict(dict); cnt=collections.Counter(); rep=[]
for lab,pc in cats:
    nt=0
    if lab in DEFENSIVE:
        oppc,_=shaped_cut([legR[(lab,m)] for m in members if legR[(lab,m)]<0],hi=False,rfloor=0.15,rcap=0.6)
        defl,_=shaped_cut([legR[(lab,m)] for m in members if legR[(lab,m)]>0],hi=True)
        defp,_=shaped_cut([prR[(lab,m)] for m in members if prR[(lab,m)]>0],hi=True,rfloor=0.15,rcap=0.5)
        oppc=oppc if oppc is not None else qtile(POOLlegneg,0.5); defl=defl if defl is not None else qtile(POOLlegneg,0.0) and 2.0
        defp=defp if defp is not None else qtile(POOLprspos,0.6) or 0.03
        defl=defl or 2.0
        for m in members:
            leg=legR[(lab,m)]; prs=prR[(lab,m)]
            if not hassig[(lab,m)]: s='—'
            elif leg<=oppc: s='Opponent'
            elif leg<0: s='Skeptic'
            elif leg>=defl or prs>=defp: s='Defender'; nt+=1
            elif leg>0 or prs>=0.4*defp: s='Ally'
            elif prs<=-0.4*defp: s='Skeptic'
            else: s='Neutral'
            matrix[m][lab]=s; scores[m][lab]=round(comb[(lab,m)],3); cnt[s]+=1
        rep.append((lab,'DEF',oppc,defp,nt))
    else:
        P=sorted(comb[(lab,m)] for m in members if hassig[(lab,m)] and comb[(lab,m)]>0)
        Ng=sorted(comb[(lab,m)] for m in members if hassig[(lab,m)] and comb[(lab,m)]<0)
        cc,cr=shaped_cut(P,hi=True,rfloor=0.03,rcap=0.12)     # Champion: very tight elite band (<=12% of positives)
        ac,_=shaped_cut(P,hi=True,rfloor=0.30,rcap=0.60)      # champ+Ally boundary; below -> persuadable Neutral
        oc,_=shaped_cut(Ng,hi=False,rfloor=0.05,rcap=0.25)    # Opponent
        sc,_=shaped_cut(Ng,hi=False,rfloor=0.30,rcap=0.60)    # opp+Skeptic boundary (persuadable edge)
        cc=cc if cc is not None else pooled_champ; ac=ac if ac is not None else pooled_ally
        oc=oc if oc is not None else pooled_opp; sc=sc if sc is not None else pooled_skep
        for m in members:
            c=comb[(lab,m)]
            if not hassig[(lab,m)]: s='—'
            elif c>0: s='Champion' if c>=cc else ('Ally' if c>=ac else 'Neutral')
            elif c<0: s='Opponent' if c<=oc else ('Skeptic' if c<=sc else 'Neutral')
            else: s='Neutral'
            if s=='Champion': nt+=1
            matrix[m][lab]=s; scores[m][lab]=round(c,3); cnt[s]+=1
        rep.append((lab,'off',cc,cr,nt))
# split each category's Neutral band into Highly Persuadable / True Neutral / Lean Skeptic by its OWN tertiles of lean
for lab,pc in cats:
    nz=sorted(scores[m][lab] for m in members if matrix[m][lab]=='Neutral')
    t33=qtile(nz,0.34) if len(nz)>=6 else -0.0001; t67=qtile(nz,0.67) if len(nz)>=6 else 0.0001
    for m in members:
        if matrix[m][lab]=='Neutral':
            c=scores[m][lab]
            matrix[m][lab]='Highly Persuadable' if c>=t67 else ('Lean Skeptic' if c<t33 else 'True Neutral')
cnt=collections.Counter(matrix[m][lab] for m in members for lab,_ in cats)
MC.dump_atomic({'cats':[c[0] for c in cats],'defensive':list(DEFENSIVE),'matrix':{m:matrix[m] for m in members},
           'scores':{m:scores[m] for m in members}},MB_ROOT+r'\category_stance.json')
print('tier distribution:',dict(cnt))
persu=cnt['Highly Persuadable']+cnt['True Neutral']+cnt['Lean Skeptic']+cnt['Skeptic']
print('persuadable middle: HighlyPersuadable %d | TrueNeutral %d | LeanSkeptic %d | Skeptic %d  (total %d) | Champions+Defenders %d'%(
      cnt['Highly Persuadable'],cnt['True Neutral'],cnt['Lean Skeptic'],cnt['Skeptic'],persu,cnt['Champion']+cnt['Defender']))
print('\nSELF-CALIBRATED thresholds per category (derived from its own median+IQR, rate-bounded):')
for lab,ty,cut,rate,nt in rep:
    print('   %-26s %-3s champ/def cut %+.3f  rate %4s%%  n=%3d'%(lab,ty,cut if cut is not None else 0,('%2d'%round(100*rate)) if rate else ' -',nt))
def prof(m): return ([c for c,_ in cats if matrix[m][c] in ('Champion','Defender')],[c for c,_ in cats if matrix[m][c]=='Opponent'])
print('\nVALIDATION:')
for who in ['Richard Durbin','Tim Scott','Andy Barr','Maxine Waters']:
    if who in matrix:
        ch,op=prof(who); print('   %-15s [%s] champ/def=%s opp=%s'%(who,gen.get(who),[c.split(' &')[0][:8] for c in ch],[c.split(' &')[0] for c in op]))
print('Interchange DEFENDERS:',[m.split()[-1] for m in members if matrix[m]['Interchange']=='Defender'])
