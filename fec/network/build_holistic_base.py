# -*- coding: utf-8 -*-
import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""HOLISTIC BASE NETWORK PRODUCER (Phase 2 of the holistic-network v1 build).

Reads the machine-local govinfo BILLSTATUS bulk archive (all bill types, congresses per
dial) and writes `holistic_network_BASE_<date>.xlsx` — the DIRECTED sponsor->cosponsor
base graph over the ALL-BILLS universe, in the extended schema of the design guide §4.5:
the 8 CU-base columns plus per-role weight subtotals WOriginal/WRegular (role dials
factored OUT), so any regular-weight re-ruling is a recompute, never a re-collection.

Weight formula = the CU scraper's Stage-3 formula verbatim (reference snapshot
audits/legfix5/draft_systems_scraper_snapshot_20260726/congress_member_scraper.R —
roles :1773-1774, PASSED_RE :1786, f_bill = succ*decay*(1/sqrt C) :1923-1927,
group-sum :1949-1951), with the D1-ruled role dials (original 1.0 / regular 0.2).
SUCCESS is judged on the bill's LATEST action text (D6-b scraper parity — see the
SUCCESS-SEMANTICS caveat below), C is the full-bill count of DISTINCT cosponsors
(withdrawn included — page-scrape parity; duplicate BILLSTATUS items collapsed and
counted), departed-sponsor policy = accept-the-drop (D6-d; drop share reported against
the TIE-CAPABLE population per congress = COR-5).

SUCCESS-SEMANTICS CAVEAT (PRO-1 adversarial round, 2026-08-17, measured and open):
`latestAction` is the LAST action, not the high-water mark. A bill that passed its origin
chamber and then died in the other chamber shows e.g. 'Received in the Senate and referred
to...' and scores SUCCESS=1.0. Measured over 117-119: PASSED_RE fires on latestAction for
1,083/55,195 bills (1.96%) but on some <actions> item for 5,019 (9.09%); 2,652 of the
37,076 tie-contributing bills are affected, carrying 39,390 of 410,411 ties, and scoring
them would raise unrolled tie-weight ~14.9%. This builder is FAITHFUL to the D6-b ruling
(scraper parity) and therefore inherits the scraper's limitation. Scanning the full
<actions> list is a SPEC CHANGE that would break CU parity -> Paul's ruling, not the
builder's. Recorded in the eval RECORD; do not "fix" it here without re-opening D6-b.

Freeze: audits/holistic_v1_eval/FREEZE_ADDENDUM_2026-08-17.md. Execution home = the
holistic workbranch (RECORD amendment 2026-08-17); paths.py roots this script to its
own tree, so the output is invisible to every production glob until merge — and the
name assert below refuses the production family outright (COR-9).

Dials: HOLISTIC_CONGRESSES=117,118,119  HOLISTIC_DECAY=0.8  HOLISTIC_NEWEST=119
       HOLISTIC_ORIGINAL_W=1.0  HOLISTIC_REGULAR_W=0.2
       HOLISTIC_TYPES=hconres,hjres,hr,hres,s,sconres,sjres,sres
       HOLISTIC_BASE_OUT (path override)  HOLISTIC_CACHE_DIR (bill-record cache)
       HOLISTIC_RAW (archive root; default %LOCALAPPDATA%\\Melange\\lda_rebuild\\raw)

Import-safe: pure functions at module level, all I/O under main() — the Gate-2 fixture
suite imports this module (deliberate deviation from module-level-script house style)."""
import os, re, json, math, zipfile, hashlib, collections, datetime
import xml.etree.ElementTree as ET

# ---- pinned reference constants (scraper parity) ----
SUCCESS_MULT = 3.0
PASSED_RE = re.compile(r'passed|agreed to in|became.*law|public law|to president|resolving differences|veto', re.I)
PRODUCTION_NAME_RE = re.compile(r'^congress_cosponsor_network(?:_AUG)?_?.*\.xlsx$', re.I)  # refuse (COR-9)

def dials():
    return {
        'congresses': tuple(int(c) for c in os.environ.get('HOLISTIC_CONGRESSES', '117,118,119').split(',')),
        'decay': float(os.environ.get('HOLISTIC_DECAY', '0.8')),
        'newest': int(os.environ.get('HOLISTIC_NEWEST', '119')),
        'ow': float(os.environ.get('HOLISTIC_ORIGINAL_W', '1.0')),
        'rw': float(os.environ.get('HOLISTIC_REGULAR_W', '0.2')),
        'types': tuple(t.strip().lower() for t in os.environ.get(
            'HOLISTIC_TYPES', 'hconres,hjres,hr,hres,s,sconres,sjres,sres').split(',')),
    }

def passed_own_chamber(last_action_text):
    """Scraper parity :1786-1787 — regex on the latest action; None/empty -> False."""
    return bool(last_action_text) and bool(PASSED_RE.search(last_action_text))

def bill_factor(passed, age, C, decay=0.8, success_mult=SUCCESS_MULT):
    """f_bill = f_succ * f_decay * sizenorm(C); sizenorm = MULTIPLIED 1/sqrt(C) (:1927).
    age < 0 or None -> full weight (:112). C <= 0 has no ties; return 0.0 defensively."""
    if C is None or C <= 0: return 0.0
    a = 0 if (age is None or age < 0) else age
    f_succ = success_mult if passed else 1.0
    return f_succ * (decay ** a) * (1.0 / math.sqrt(C))

def assemble(bills, roster, ow, rw, decay=0.8, newest=119):
    """bills: iterable of dicts {congress, passed, sponsor, cos: [(bioguide, is_orig, withdrawn), ...]}.
    roster: set of bioguide ids. Returns (edges, stats). edges[(From,To)] = dict with
    WOriginal/WRegular subtotals (role dials factored OUT), NBills/NOriginal/NRegular,
    Weight = ow*WOriginal + rw*WRegular. Rows exist whenever NBills > 0 — a zero role
    dial must NOT drop rows (COR-2 re-weighting basis). Sponsor off roster -> bill
    contributes ZERO ties (D6-d accept-the-drop). C = FULL-BILL count (D6-c), taken
    before any roster filtering. Self-loops dropped (:1931)."""
    edges = {}
    st = collections.Counter()
    drop_by_cg = collections.Counter(); bills_by_cg = collections.Counter()
    tiecap_by_cg = collections.Counter()
    for b in bills:
        cg = b['congress']; bills_by_cg[cg] += 1
        # PRO-1 fix (adversarial round 2026-08-17, verified on BILLSTATUS-118hr6116.xml where
        # G000578/L000596/S001214 each appear twice as (original, withdrawn) then (regular, live)):
        # BILLSTATUS can list one member twice on a bill. Un-deduped, that member's single tie was
        # added to BOTH role subtotals, counted NBills=2 for one bill, and inflated C -- depressing
        # every other tie on the bill through 1/sqrt(C). Keep the LAST item per bioguide (the
        # current sponsorship state, which is the order BILLSTATUS emits) so one member on one bill
        # is exactly one tie of exactly one role. C is then the DISTINCT cosponsor count.
        _seen = {}
        for _c in b['cos']:
            if _c[0] in _seen: st['cosponsor_duplicate_items_collapsed'] += 1
            _seen[_c[0]] = _c
        cos = list(_seen.values())
        C = len(cos)
        if C == 0:
            st['bills_no_cosponsor'] += 1; continue
        tiecap_by_cg[cg] += 1          # COR-5 denominator: bills that COULD carry a tie
        if not b['sponsor'] or b['sponsor'] not in roster:
            st['bills_sponsor_off_roster'] += 1; drop_by_cg[cg] += 1; continue
        f = bill_factor(b['passed'], newest - cg, C, decay)
        st['bills_contributing'] += 1
        for (to, is_orig, withdrawn) in cos:
            if withdrawn: st['ties_withdrawn_included'] += 1
            if not to or to not in roster:
                st['ties_cosponsor_off_roster'] += 1; continue
            if to == b['sponsor']:
                st['ties_self_loop_dropped'] += 1; continue
            e = edges.setdefault((b['sponsor'], to),
                                 {'WOriginal': 0.0, 'WRegular': 0.0, 'NBills': 0,
                                  'NOriginal': 0, 'NRegular': 0})
            e['NBills'] += 1
            if is_orig: e['WOriginal'] += f; e['NOriginal'] += 1
            else:       e['WRegular'] += f; e['NRegular'] += 1
            st['ties_kept'] += 1
    for e in edges.values():
        e['Weight'] = ow * e['WOriginal'] + rw * e['WRegular']
    st['edges'] = len(edges)
    # COR-5 (PRO-1 fix): the drop SHARE must divide by the tie-capable population, not by all
    # bills -- zero-cosponsor bills can never contribute a tie, so including them in the
    # denominator understates the loss (measured: 23.7% -> 29.5% on the 117th).
    share = {cg: (drop_by_cg.get(cg, 0) / tiecap_by_cg[cg] if tiecap_by_cg.get(cg) else 0.0)
             for cg in bills_by_cg}
    return edges, {'counters': dict(st),
                   'bills_by_congress': dict(bills_by_cg),
                   'bills_tie_capable_by_congress': dict(tiecap_by_cg),
                   'sponsor_dropped_by_congress': dict(drop_by_cg),
                   'sponsor_dropped_share_of_tie_capable': share}

def assert_output_name(path):
    base = os.path.basename(path)
    if PRODUCTION_NAME_RE.match(base):
        raise ValueError('COR-9: refusing to write a production-family filename: ' + base)
    if not base.startswith('holistic_network_BASE_'):
        raise ValueError('COR-9: holistic base output must be named holistic_network_BASE_*: ' + base)
    if not base.lower().endswith('.xlsx'):   # PRO-1: prefix alone let a .manifest.json-style name through
        raise ValueError('COR-9: holistic base output must be a .xlsx workbook: ' + base)

# ---- BILLSTATUS extraction (stage 1; cached NDJSON per congress) ----
def _findtext(el, *paths):
    for p in paths:
        v = el.findtext(p)
        if v is not None: return v
    return None

def parse_billstatus_xml(data):
    """One BILLSTATUS xml -> bill record dict (no roster logic here)."""
    root = ET.fromstring(data)
    bill = root.find('bill') if root.find('bill') is not None else root
    cg = _findtext(bill, 'congress')
    la = _findtext(bill, 'latestAction/text')
    sponsor = _findtext(bill, 'sponsors/item/bioguideId')
    cos = []
    for it in bill.findall('cosponsors/item'):
        bg = it.findtext('bioguideId')
        orig = (it.findtext('isOriginalCosponsor') or '').strip().lower() == 'true'
        wd = bool((it.findtext('sponsorshipWithdrawnDate') or '').strip())
        cos.append((bg, orig, wd))
    return {'congress': int(cg) if cg else None, 'passed': passed_own_chamber(la),
            'sponsor': sponsor, 'cos': cos}

def extract_congress(raw_root, cg, types, cache_dir):
    """Parse every BILLSTATUS zip for one congress into an NDJSON cache; resume-safe."""
    os.makedirs(cache_dir, exist_ok=True)
    cache = os.path.join(cache_dir, 'billcache_%d.ndjson' % cg)
    zdir = os.path.join(raw_root, 'billstatus%d' % cg)
    zips = sorted(z for z in os.listdir(zdir) if z.lower().endswith('.zip')
                  and z.replace('BILLSTATUS-%d-' % cg, '').replace('.zip', '').lower() in types)
    expect = 0
    for z in zips:
        with zipfile.ZipFile(os.path.join(zdir, z)) as zf:
            expect += sum(1 for n in zf.namelist() if n.lower().endswith('.xml'))
    if os.path.exists(cache):
        have = sum(1 for _ in open(cache, encoding='utf-8'))
        if have == expect:
            print('  cache hit billcache_%d.ndjson (%d bills)' % (cg, have)); return cache, expect
    n = bad = 0
    with open(cache + '.tmp', 'w', encoding='utf-8') as out:
        for z in zips:
            with zipfile.ZipFile(os.path.join(zdir, z)) as zf:
                for name in zf.namelist():
                    if not name.lower().endswith('.xml'): continue
                    try:
                        rec = parse_billstatus_xml(zf.read(name))
                        rec['src'] = name
                        out.write(json.dumps(rec) + '\n'); n += 1
                    except Exception as ex:
                        bad += 1
                        print('  PARSE FAIL %s: %s' % (name, ex))
            print('  %s done (cum %d bills, %d failures)' % (z, n, bad))
    os.replace(cache + '.tmp', cache)
    if bad: print('  WARNING congress %d: %d parse failures of %d files' % (cg, bad, expect))
    return cache, expect

def main():
    d = dials()
    raw = os.environ.get('HOLISTIC_RAW', os.path.join(os.environ.get('LOCALAPPDATA', ''), r'Melange\lda_rebuild\raw'))
    cache_dir = os.environ.get('HOLISTIC_CACHE_DIR', os.path.join(MB_ROOT, r'data\holistic'))
    out = os.environ.get('HOLISTIC_BASE_OUT',
                         os.path.join(MB_ROOT, 'holistic_network_BASE_%s.xlsx' %
                                      datetime.date.today().strftime('%Y%m%d')))
    assert_output_name(out)
    print('HOLISTIC BASE BUILD | congresses=%s types=%s ow=%.2f rw=%.2f decay=%.2f newest=%d' %
          (d['congresses'], ','.join(d['types']), d['ow'], d['rw'], d['decay'], d['newest']))

    roster_path = os.path.join(MB_ROOT, r'data\engine\congestion.json')
    C = json.load(open(roster_path, encoding='utf-8'))
    bio2name = {m['bioguide']: m['member'] for m in C if m.get('bioguide')}
    roster = set(bio2name)
    print('roster: %d members, %d with bioguide' % (len(C), len(roster)))

    caches = []
    for cg in d['congresses']:
        print('congress %d:' % cg)
        cache, expect = extract_congress(raw, cg, d['types'], cache_dir)
        caches.append((cg, cache, expect))

    def bills():
        for cg, cache, _ in caches:
            for line in open(cache, encoding='utf-8'):
                rec = json.loads(line)
                if rec['congress'] is None: rec['congress'] = cg
                yield rec
    edges, stats = assemble(bills(), roster, d['ow'], d['rw'], d['decay'], d['newest'])

    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Edges'
    ws.append(['From', 'To', 'Weight', 'NBills', 'NOriginal', 'NRegular',
               'From_label', 'To_label', 'WOriginal', 'WRegular'])
    for (a, b), e in sorted(edges.items()):
        ws.append([a, b, e['Weight'], e['NBills'], e['NOriginal'], e['NRegular'],
                   bio2name.get(a, a), bio2name.get(b, b), e['WOriginal'], e['WRegular']])
    wb.save(out)

    members_in = {a for a, _ in edges} | {b for _, b in edges}
    isolates = sorted(m for m in roster if m not in members_in)
    # PRO-1 fix: pin ONLY the zips this build actually read (the dial's type set), not every
    # zip in the directory -- a manifest that pins unread inputs misrepresents provenance.
    zip_pins = {}
    for cg, _, _ in caches:
        zdir = os.path.join(raw, 'billstatus%d' % cg)
        for z in sorted(os.listdir(zdir)):
            if not z.lower().endswith('.zip'): continue
            bt = z.replace('BILLSTATUS-%d-' % cg, '').replace('.zip', '').lower()
            if bt not in d['types']: continue          # not read -> not pinned
            zp = os.path.join(zdir, z)
            zip_pins[z] = {'sha256': hashlib.sha256(open(zp, 'rb').read()).hexdigest(),
                           'mtime_utc': datetime.datetime.fromtimestamp(
                               os.path.getmtime(zp), datetime.timezone.utc).strftime('%Y-%m-%d')}
    missing_types = [(cg, t) for cg, _, _ in caches for t in d['types']
                     if not os.path.exists(os.path.join(raw, 'billstatus%d' % cg,
                                                        'BILLSTATUS-%d-%s.zip' % (cg, t)))]
    if missing_types:
        raise SystemExit('PRO-1 fail-closed: dialled bill types with no zip present: %s '
                         '(a silently dropped type would produce a plausible-but-wrong graph)' % missing_types)
    # PRO-1 fix: the BASE manifest records a canonical content hash OF ITS OWN OUTPUT, so the
    # AUG's base pin can be checked against something reproducible (raw xlsx bytes are not).
    _ws = openpyxl.load_workbook(out, data_only=True)['Edges']
    _rows = sorted(tuple(c.value for c in r) for r in _ws.iter_rows(min_row=2))
    _canon = hashlib.sha256(repr(_rows).encode()).hexdigest()
    manifest = {
        'built_utc': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
        'config': 'holistic', 'vintage': 'BASE', 'dials': {k: (list(v) if isinstance(v, tuple) else v) for k, v in d.items()},
        'formula': 'w = role * SUCCESS(3.0 if passed-latest-action) * decay^age * 1/sqrt(C_fullbill); scraper parity :1773-1788,:1923-1951',
        'scope': {'cosponsorship': 'all-bills'}, 'bpf': 'n/a (base vintage; kernel-side concern)',
        'roster': {'members': len(C), 'with_bioguide': len(roster), 'source': 'data/engine/congestion.json'},
        'stats': stats, 'isolates': {'count': len(isolates), 'bioguides': isolates},
        'inputs': {'billstatus_zips': zip_pins},
        'output': os.path.basename(out), 'output_canonical_rows_sha256': _canon,
    }
    mpath = out.replace('.xlsx', '.manifest.json')
    json.dump(manifest, open(mpath, 'w', encoding='utf-8'), indent=1)
    c = stats['counters']
    print('SUMMARY: %d edges | bills contributing %d / no-cosponsor %d / sponsor-off-roster %d | '
          'ties kept %d / cosponsor-off-roster %d / self-loops %d / withdrawn-included %d | '
          'members present %d / isolates %d' %
          (c.get('edges', 0), c.get('bills_contributing', 0), c.get('bills_no_cosponsor', 0),
           c.get('bills_sponsor_off_roster', 0), c.get('ties_kept', 0),
           c.get('ties_cosponsor_off_roster', 0), c.get('ties_self_loop_dropped', 0),
           c.get('ties_withdrawn_included', 0), len(members_in), len(isolates)))
    print('COR-5 sponsor-dropped share of TIE-CAPABLE bills by congress: %s' %
          {k: '%d/%d = %.1f%%' % (stats['sponsor_dropped_by_congress'].get(k, 0),
                                  stats['bills_tie_capable_by_congress'][k],
                                  100 * stats['sponsor_dropped_share_of_tie_capable'][k])
           for k in sorted(stats['bills_tie_capable_by_congress'])})
    print('  (all-bills denominators, for reference: %s | duplicate cosponsor items collapsed: %d)' %
          (stats['bills_by_congress'], c.get('cosponsor_duplicate_items_collapsed', 0)))
    print('wrote %s + manifest' % out)

if __name__ == '__main__':
    main()
