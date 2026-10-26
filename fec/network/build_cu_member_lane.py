# -*- coding: utf-8 -*-
"""CR LANE - member-level regeneration of the CU cosponsor network and bills archive, OFFLINE.

Rebuilds congress_cosponsor_network_*.xlsx (Edges + Centrality) and congress_bills_archive_*.xlsx
(All Bills + 17 per-section sheets) in the exact schemas of the Draft Systems R scraper, from local
GovInfo BILLSTATUS bulk zips, keyed by bioguide (no name->bioguide step, so no Menendez-class error).
Section membership is a property of the bill and is seeded from the existing R archives.

NETWORK/ZERO-WRITE: reads the live root (MONEYBALL_ROOT) and the BILLSTATUS zips; writes ONLY to
<repo>/data/cr_lane/<run_id>/ and refuses any --out that is a repo/live root. Never promotes: copying
into the root is Paul's call (see CR-design.md section 5). Reuses build_cu_placeholder_backfill's
Stage 3 weighting (backfill_ties). Import has no side effects.

NETWORK DEFAULT = R-FAITHFUL (Paul 2026-10-25): the bill set R's Stage 3 used (fec/network/cr_r_network_billset.json: the
checkpoint bills minus the 18 stale entries, plus the 3 Tim Scott sponsor rows -- the set cr_parity_gate.py reproduces), each
bill counted once PER R DetailLink (per section), re-weighted from BILLSTATUS (real drift only: new cosponsorships and
passed-chamber status; sponsor identity by bioguide). Opt-in: --network-set archive (every archive bill, "method A") and
--count dedupe (one count per bill, "method B"). New bills join R's set only after discovery (a later step, Paul's GO):
append [congress, type, number, n_sections] to the billset file (n_sections = sections the bill is in), or use
--network-set archive. Archive default = variant B (0624 rows for the 117th/118th + BILLSTATUS 119th rows).

  python fec/network/build_cu_member_lane.py build [--network-set r|archive] [--count per-section|dedupe]
  python fec/network/build_cu_member_lane.py diff --candidate-dir DIR
"""
import os as _o, sys as _s
_s.path.insert(0, _o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))))
_s.path.insert(0, _o.path.dirname(_o.path.abspath(__file__)))
import argparse, collections, datetime, hashlib, html, io, json, math, os, re, sys, zipfile
import xml.etree.ElementTree as ET

import build_cu_placeholder_backfill as BF  # stale-ok: script-mode import after explicit own-dir sys.path insert

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
NEWEST_CONGRESS = 119
CONGRESSES = (117, 118, 119)
BILLSET = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cr_r_network_billset.json')
REF_NETWORK = 'congress_cosponsor_network_20260625_205554.xlsx'
REF_ARCHIVE = 'congress_bills_archive_20260624_115822.xlsx'
SEED_ARCHIVES = (REF_ARCHIVE, 'congress_bills_archive_20260825_115305.xlsx')
HEADERS = ['Section', 'Role', 'Member', 'Name', 'Title', 'Sponsor', 'CosponsorCount', 'Committee',
           'LastAction', 'Congress', 'DecayWeight', 'DetailLink', 'bill_summary']
EDGE_HEADERS = ['From', 'To', 'Weight', 'NBills', 'NOriginal', 'NRegular', 'From_label', 'To_label']
CENT_HEADERS = ['node', 'label', 'InfluenceScore', 'WeightedOutDegree', 'WeightedInDegree', 'OutDegree',
                'InDegree', 'Closeness', 'Betweenness', 'Eigenvector', 'PageRank']
ROLES = ('sponsor', 'original-cosponsor', 'cosponsor')
# display prefix <-> BILLSTATUS type <-> congress.gov path pieces
PREFIX = {'hr': 'H.R.', 's': 'S.', 'hres': 'H.Res.', 'sres': 'S.Res.', 'hjres': 'H.J.Res.',
          'sjres': 'S.J.Res.', 'hconres': 'H.Con.Res.', 'sconres': 'S.Con.Res.'}
PATH_KIND = {'hr': ('house', 'bill'), 's': ('senate', 'bill'), 'hres': ('house', 'resolution'),
             'sres': ('senate', 'resolution'), 'hjres': ('house', 'joint-resolution'),
             'sjres': ('senate', 'joint-resolution'), 'hconres': ('house', 'concurrent-resolution'),
             'sconres': ('senate', 'concurrent-resolution')}
BIOGUIDE_RE = re.compile(r'^[A-Z]\d{6}$')
NAME_RE = re.compile(r'^(H\.R\.|S\.|H\.Res\.|S\.Res\.|H\.J\.Res\.|S\.J\.Res\.|H\.Con\.Res\.|S\.Con\.Res\.)\d+$')
TRACKER_VOCAB = {'Introduced', 'Passed House', 'Passed Senate', 'Agreed to in Senate', 'Agreed to in House',
                 'Resolving Differences', 'Became Law', 'Failed to pass over veto', 'To President'}


class Refusal(Exception):
    pass


# ----------------------------------------------------------------------------- pure helpers
def bill_name(btype, number):
    return PREFIX[btype] + str(number)


def detail_link(congress, btype, number):
    ch, kind = PATH_KIND[btype]
    sfx = 'th' if 10 <= congress % 100 <= 20 else {1: 'st', 2: 'nd', 3: 'rd'}.get(congress % 10, 'th')
    return '/bill/%d%s-congress/%s-%s/%d' % (congress, sfx, ch, kind, number)


def decay(congress):
    return 0.8 ** (NEWEST_CONGRESS - int(congress))


def strip_html(s):
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', s or ''))).strip()


def tracker_status(actions, bill_type):
    """High-water tracker label from BILLSTATUS actions [(date, code, type, text)] (R tracker vocabulary).
    Ladder, first hit wins: Became Law > Failed to pass over veto > Resolving Differences > Passed Senate/
    Agreed to in Senate > Passed House > Introduced.  ('To President' is folded into Passed*: PASSED_RE
    treats all of them alike.)"""
    simple = bill_type in ('sres', 'sconres', 'hres', 'hconres')
    txt = ' '.join(a[3] or '' for a in actions)
    types = {a[2] for a in actions}; codes = {a[1] for a in actions}
    if 'BecameLaw' in types or codes & {'36000', '41000'} or re.search(r'\bBecame (Public|Private) Law', txt):
        return 'Became Law'
    if re.search(r'two-thirds|Failed of passage|not sustained|Passed over veto|veto', txt, re.I) \
            and re.search(r'veto', txt, re.I) and re.search(r'failed|not agreed|not sustained', txt, re.I):
        return 'Failed to pass over veto'
    if 'ResolvingDifferences' in types:
        return 'Resolving Differences'
    senate = codes & {'17000', '19000'} or re.search(r'Passed Senate|Agreed to in Senate', txt) or \
        (bill_type in ('s', 'sjres') and 'Passed/agreed to in Senate' in txt)
    house = codes & {'8000', '8500'} or re.search(r'Passed House|Agreed to in House|Passed/agreed to in House', txt)
    if senate:
        return 'Agreed to in Senate' if simple else 'Passed Senate'
    if house:
        return 'Passed House'
    return 'Introduced'


def bill_key_from_name(name, congress):
    return BF.bill_key(name, congress)


# ----------------------------------------------------------------------------- inputs
def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def section_seed(archive_paths):
    """-> (seed {(congress, type, number): [Section,...] template order}, section_order [Section,...]).
    Reads 'All Bills' of each archive (first archive fixes the template order)."""
    import openpyxl
    seed = collections.OrderedDict(); order = []
    for p in archive_paths:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        rows = wb['All Bills'].iter_rows(values_only=True)
        hdr = list(next(rows)); ix = {h: i for i, h in enumerate(hdr)}
        for r in rows:
            if r[ix['Name']] is None:
                continue
            k = bill_key_from_name(r[ix['Name']], r[ix['Congress']]); sec = r[ix['Section']]
            if k is None or not sec:
                continue
            if sec not in order:
                order.append(sec)
            lst = seed.setdefault((k[0], k[1], k[2]), [])
            if sec not in lst:
                lst.append(sec)
        names = wb.sheetnames
        wb.close()
        if p == archive_paths[0] and len(names) > 1:   # template order from the per-section sheet names
            tmpl = []
            for n in names[1:]:
                rest = re.sub(r'^\d+ ', '', n)
                tmpl.append(next((s for s in order if s.startswith(rest)), rest))
            order = tmpl + [s for s in order if s not in tmpl]
    return seed, order


def check_xml(raw):
    up = raw.upper()
    if b'<!DOCTYPE' in up or b'<!ENTITY' in up:
        raise Refusal('DOCTYPE/ENTITY forbidden in BILLSTATUS xml')


def _t(el, path):
    x = el.find(path)
    return (x.text or '').strip() if x is not None and x.text else ''


def parse_billstatus(raw, congress, btype, number):
    """BILLSTATUS xml bytes -> Bill dict (see CR-design section 6)."""
    check_xml(raw)
    b = ET.fromstring(raw).find('bill')
    if b is None:
        raise Refusal('no <bill>')
    sp = b.find('sponsors/item')
    sponsor = {'bioguide': _t(sp, 'bioguideId'), 'fullName': _t(sp, 'fullName')} if sp is not None else None
    cos = []
    for it in b.findall('cosponsors/item'):
        cos.append({'bioguide': _t(it, 'bioguideId'), 'fullName': _t(it, 'fullName'),
                    'original': _t(it, 'isOriginalCosponsor').lower() == 'true',
                    'withdrawn': bool(_t(it, 'sponsorshipWithdrawnDate'))})
    cmts = collections.OrderedDict()
    for it in b.findall('committees/item'):
        n = re.sub(r'\s+Committee$', '', _t(it, 'name')); ch = _t(it, 'chamber')
        if n and n not in cmts.setdefault(ch, []):
            cmts[ch].append(n)
    acts = [(_t(a, 'actionDate'), _t(a, 'actionCode'), _t(a, 'type'), _t(a, 'text')) for a in b.findall('actions/item')]
    summ = ''
    sl = b.findall('summaries/summary')
    if sl:
        best = max(sl, key=lambda s: (_t(s, 'updateDate'), _t(s, 'versionCode')))
        summ = strip_html(_t(best, 'text'))
    return {'congress': congress, 'type': btype, 'number': number, 'name': bill_name(btype, number),
            'title': _t(b, 'title'), 'sponsor': sponsor, 'cos': cos, 'committees': cmts, 'actions': acts,
            'summary': summ or None}


def read_bills(raw_root, keys):
    """-> (bills {key: Bill}, refusals [(key, reason)]). Opens each zip once; missing key = refusal."""
    bills, refused = {}, []
    by_zip = collections.defaultdict(list)
    for k in keys:
        by_zip[(k[0], k[1])].append(k)
    for (c, t), ks in sorted(by_zip.items()):
        zp = os.path.join(raw_root, 'billstatus%d' % c, 'BILLSTATUS-%d-%s.zip' % (c, t))
        if not os.path.exists(zp):
            refused += [(k, 'zip absent: %s' % zp) for k in ks]; continue
        with zipfile.ZipFile(zp) as z:
            names = set(z.namelist())
            for k in ks:
                mem = 'BILLSTATUS-%d%s%d.xml' % k
                if mem not in names:
                    refused.append((k, 'not in BILLSTATUS')); continue
                try:
                    bills[k] = parse_billstatus(z.read(mem), *k)
                except (Refusal, ET.ParseError) as e:
                    refused.append((k, 'parse: %s' % e))
    return bills, refused


# ----------------------------------------------------------------------------- archive rows
def committee_text(cmts):
    if not cmts:
        return ''
    return 'Committees: ' + ' | '.join('%s - %s' % (ch or 'Joint', '; '.join(ns)) for ch, ns in cmts.items())


def archive_rows(seed, bills, roster_by_bio, section_order, withdrawn='exclude', originals_also_cosponsor=True):
    """-> list[dict] in archive order: Section (template), Role, Member (codepoint), Congress desc, type, number."""
    rows = []
    for k, secs in seed.items():
        b = bills.get(k)
        if not b:
            continue
        live = [c for c in b['cos'] if withdrawn == 'include' or not c['withdrawn']]
        base = {'Name': b['name'], 'Title': b['title'], 'Sponsor': (b['sponsor'] or {}).get('fullName') or None,
                'CosponsorCount': sum(1 for c in b['cos'] if not c['withdrawn']),
                'Committee': committee_text(b['committees']) or None,
                'LastAction': tracker_status(b['actions'], b['type']), 'Congress': b['congress'],
                'DecayWeight': decay(b['congress']), 'DetailLink': detail_link(*k), 'bill_summary': b['summary']}
        members = []
        if b['sponsor'] and b['sponsor']['bioguide'] in roster_by_bio:
            members.append(('sponsor', b['sponsor']['bioguide']))
        seen = set()
        for c in live:
            if c['bioguide'] in roster_by_bio and c['bioguide'] not in seen \
                    and c['bioguide'] != (b['sponsor'] or {}).get('bioguide'):
                seen.add(c['bioguide'])
                members.append(('original-cosponsor' if c['original'] else 'cosponsor', c['bioguide']))
                if c['original'] and originals_also_cosponsor:   # R's 'cosponsor' search also returns originals (0624: 2,942 of 2,972)
                    members.append(('cosponsor', c['bioguide']))
        for sec in secs:
            for role, bio in members:
                rows.append(dict(base, Section=sec, Role=role, Member=roster_by_bio[bio]))
    return sort_rows(rows, section_order)


def sort_rows(rows, section_order):
    ro = {r: i for i, r in enumerate(ROLES)}; so = {x: i for i, x in enumerate(section_order)}
    def tn(r):
        k = bill_key_from_name(r['Name'], r['Congress']); return (k[1], k[2])
    return sorted(rows, key=lambda r: (so.get(r['Section'], 99), ro[r['Role']], r['Member'], -r['Congress'], tn(r)))


def _read_rows(path, keep):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    it = wb['All Bills'].iter_rows(values_only=True); hdr = list(next(it))
    out = [d for d in (dict(zip(hdr, r)) for r in it if r[0] is not None) if keep(d)]
    wb.close()
    return out


def variant_b_rows(ref_archive, rows_a, roster_names, section_order):
    """Variant B (Paul 2026-10-25): R's own 0624 rows for the 117th/118th (roster members only) + the
    BILLSTATUS-derived 119th rows of variant A (replacing the degraded 0825 119th rows)."""
    keep = _read_rows(ref_archive, lambda r: int(r['Congress']) < NEWEST_CONGRESS and r['Member'] in roster_names)
    return sort_rows(keep + [r for r in rows_a if r['Congress'] == NEWEST_CONGRESS], section_order)


def variant_a0_rows(ref_archive, aug_archive, roster_names, section_order):
    """Variant A0 (Paul's literal '0624+0825', NO BILLSTATUS rebuild): R's 0624 rows for the 117th/118th plus the
    EXISTING 0825 rows for the 119th (roster members only, raw LastAction strings left as they are)."""
    old = _read_rows(ref_archive, lambda r: int(r['Congress']) < NEWEST_CONGRESS and r['Member'] in roster_names)
    new = _read_rows(aug_archive, lambda r: int(r['Congress']) == NEWEST_CONGRESS and r['Member'] in roster_names)
    return sort_rows(old + new, section_order)


def _flip_fn():
    try:
        import bill_rules as BR
    except Exception:
        return None
    return lambda r: bool(BR.excluded_bill(r['Name'], r['Congress'], r['Title'] or '', r['bill_summary'] or ''))


def variant_report(variants, ref_rows):
    """Report over {'A0':rows,'A':rows,'B':rows}: counts by congress, member coverage, invalid LastAction,
    title/summary diffs and BR.excluded_bill verdict flips (each vs the 0624 reference rows, and vs A)."""
    flip = _flip_fn()
    def per_bill(rows):
        d = {}
        for r in rows:
            d.setdefault((r['Name'], r['Congress']), r)
        return d
    def verdicts(rows):
        return {k: flip(r) for k, r in per_bill(rows).items()} if flip else None
    ref_b = per_bill(ref_rows); ref_v = verdicts(ref_rows)
    def stat(rows):
        la = collections.Counter(r['LastAction'] for r in rows)
        return {'rows': len(rows), 'rows_by_congress': dict(collections.Counter(r['Congress'] for r in rows)),
                'members': len({r['Member'] for r in rows}), 'bills': len(per_bill(rows)),
                'lastaction_invalid': sum(n for k, n in la.items() if k not in TRACKER_VOCAB),
                'lastaction_invalid_examples': [k[:70] for k in la if k not in TRACKER_VOCAB][:3],
                'lastaction_counts': {k: n for k, n in la.items() if k in TRACKER_VOCAB}}
    def vs(rows, other, other_v):
        b, o = per_bill(rows), per_bill(other); both = set(b) & set(o); v = verdicts(rows)
        key = lambda r: (r['Section'], r['Role'], r['Member'], r['Name'], r['Congress'])
        kr, ko = {key(r) for r in rows}, {key(r) for r in other}
        return {'rows_only_here': len(kr - ko), 'rows_only_there': len(ko - kr), 'rows_in_both': len(kr & ko),
                'members_only_here': sorted({r['Member'] for r in rows} - {r['Member'] for r in other}),
                'members_only_there': sorted({r['Member'] for r in other} - {r['Member'] for r in rows}),
                'bills_title_differs': sum(1 for k in both if (b[k]['Title'] or '') != (o[k]['Title'] or '')),
                'bills_summary_differs': sum(1 for k in both if (b[k]['bill_summary'] or '') != (o[k]['bill_summary'] or '')),
                'bills_lastaction_differs': sum(1 for k in both if b[k]['LastAction'] != o[k]['LastAction']),
                'excluded_bill_flips': (sum(1 for k in both if v[k] != other_v[k]) if v and other_v else 'unavailable')}
    out = {name: stat(rows) for name, rows in variants.items()}
    for name, rows in variants.items():
        out[name]['vs_0624'] = vs(rows, ref_rows, ref_v)
    vA = verdicts(variants['A'])
    for name in variants:
        if name != 'A':
            out[name]['vs_A'] = vs(variants[name], variants['A'], vA)
    return out


# ----------------------------------------------------------------------------- network
def load_billset(path=BILLSET):
    """-> {(congress, type, number): n_links} (R's network bill set, see module doc)."""
    d = json.load(open(path, encoding='utf-8'))
    return {(c, t, n): m for c, t, n, m in d['bills']}


def _labels(bills, roster_by_bio):
    labels = dict(roster_by_bio)
    for b in bills.values():
        for c in ([b['sponsor']] if b['sponsor'] else []) + b['cos']:
            if c['bioguide'] and c['bioguide'] not in labels:
                labels[c['bioguide']] = c['fullName']
    return labels


def r_network_edges(billset, bills, roster_by_bio, edge_withdrawn='include', count='per-section'):
    """R-faithful Stage 3: R's bill set x BILLSTATUS cosponsors/status. -> (edges, skipped [(key, reason)], labels)."""
    ties = collections.defaultdict(lambda: [0.0, 0, 0, 0]); skipped = []
    for k, m in sorted(billset.items()):
        b = bills.get(k); sp = (b or {}).get('sponsor')
        if not b or not sp or sp['bioguide'] not in roster_by_bio:
            skipped.append((k, 'no BILLSTATUS' if not b else 'sponsor off roster')); continue
        cos = [(c['bioguide'], c['original']) for c in b['cos'] if edge_withdrawn == 'include' or not c['withdrawn']]
        row = {'Name': b['name'], 'Congress': k[0], 'DecayWeight': decay(k[0]), 'Member': 'x',
               'LastAction': tracker_status(b['actions'], k[1])}
        t, _ = BF.backfill_ties([row] * (m if count == 'per-section' else 1), {k: {'sponsor': sp['bioguide'], 'cos': cos}},
                                {'x': sp['bioguide']})
        for key, v in t.items():
            a = ties[key]
            for i in range(4):
                a[i] += v[i]
    labels = _labels(bills, roster_by_bio)
    edges = [{'From': f, 'To': t, 'Weight': v[0], 'NBills': v[1], 'NOriginal': v[2], 'NRegular': v[3],
              'From_label': labels.get(f, f), 'To_label': labels.get(t, t)} for (f, t), v in sorted(ties.items())]
    return edges, skipped, labels


def stage3_edges(rows, bills, roster_by_bio, edge_withdrawn='include', mult=None):
    """Stage 3 over the archive's sponsor rows (dedup by (Name, Congress)) via backfill_ties.
    -> (edges list[dict] sorted (From,To), refused [(name, congress, reason)], labels {bioguide: label})."""
    roster_bio = {m: b for b, m in roster_by_bio.items()}
    seen, sp = set(), []
    for r in rows:
        if r['Role'] == 'sponsor' and (r['Name'], r['Congress']) not in seen:
            seen.add((r['Name'], r['Congress']))
            k = bill_key_from_name(r['Name'], r['Congress'])
            sp.extend([r] * (mult or {}).get((k[0], k[1], k[2]), 1))   # mult: per-section counting (opt-in archive set)
    cache = {}
    for k, b in bills.items():
        cos = [(c['bioguide'], c['original']) for c in b['cos'] if edge_withdrawn == 'include' or not c['withdrawn']]
        cache[k] = {'sponsor': (b['sponsor'] or {}).get('bioguide'), 'cos': cos}
    ties, refused = BF.backfill_ties(sp, cache, roster_bio)
    labels = _labels(bills, roster_by_bio)
    edges = [{'From': f, 'To': t, 'Weight': v[0], 'NBills': v[1], 'NOriginal': v[2], 'NRegular': v[3],
              'From_label': labels.get(f, f), 'To_label': labels.get(t, t)} for (f, t), v in sorted(ties.items())]
    return edges, refused, labels


def network_edges(netset, count, rows, bills, roster, seed, edge_withdrawn='include', billset=None):
    """Dispatcher for the two bill sets. -> (edges, refused/skipped, labels)."""
    if netset == 'r':
        return r_network_edges(billset if billset is not None else load_billset(), bills, roster, edge_withdrawn, count)
    mult = {k: len(v) for k, v in seed.items()} if count == 'per-section' else None
    return stage3_edges(rows, bills, roster, edge_withdrawn, mult)


def _norm01(v):
    import numpy as np
    v = np.nan_to_num(np.asarray(v, dtype=float), nan=0.0)
    lo, hi = v.min(), v.max()
    return np.zeros_like(v) if hi == lo else (v - lo) / (hi - lo)


def centrality(edges, labels):
    """igraph-parity Centrality rows (design section 2): returns list[dict] in sheet order."""
    import networkx as nx, numpy as np
    G = nx.DiGraph()
    nodes = list(dict.fromkeys([e['From'] for e in edges] + [e['To'] for e in edges]))
    G.add_nodes_from(nodes)
    for e in edges:
        G.add_edge(e['From'], e['To'], weight=e['Weight'], distance=1.0 / e['Weight'])
    n = len(nodes); ix = {v: i for i, v in enumerate(nodes)}
    wout = dict(G.out_degree(weight='weight')); win = dict(G.in_degree(weight='weight'))
    close = {}
    for v in nodes:
        d = nx.single_source_dijkstra_path_length(G, v, weight='distance')
        d.pop(v, None)
        close[v] = 1.0 / sum(d.values()) if d else float('nan')
    btw = nx.betweenness_centrality(G, normalized=False, weight='distance')
    A = np.zeros((n, n))
    for e in edges:
        A[ix[e['From']], ix[e['To']]] = e['Weight']
    # ponytail: dense eig is O(n^3) on ~650 nodes; switch to scipy.sparse.linalg.eigs if n exceeds ~5k
    w, V = np.linalg.eig(A.T)
    vec = np.abs(np.real(V[:, int(np.argmax(np.real(w)))]))
    eig = vec / vec.max() if vec.max() > 0 else vec
    pr = nx.pagerank(G, alpha=0.85, weight='weight', tol=1e-14, max_iter=10000)
    wo = np.array([wout[v] for v in nodes]); bt = np.array([btw[v] for v in nodes])
    infl = np.round(100 * _norm01(0.60 * _norm01(wo) + 0.25 * _norm01(bt) + 0.15 * _norm01(eig)), 1)
    rows = [{'node': v, 'label': labels.get(v, v), 'InfluenceScore': float(infl[i]), 'WeightedOutDegree': wout[v],
             'WeightedInDegree': win[v], 'OutDegree': G.out_degree(v), 'InDegree': G.in_degree(v),
             'Closeness': close[v], 'Betweenness': btw[v], 'Eigenvector': float(eig[i]), 'PageRank': pr[v]}
            for i, v in enumerate(nodes)]
    rows.sort(key=lambda r: (-r['InfluenceScore'], -r['WeightedOutDegree']))
    return rows


# ----------------------------------------------------------------------------- writers
def _cell(v):
    if v is None:
        return None
    if isinstance(v, float):
        if math.isnan(v):
            return None
        if v == int(v) and abs(v) < 1e15:
            return int(v)
    return v


def write_network(path, edges, cent):
    import openpyxl, moneyball_core as MC
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Edges'; ws.append(EDGE_HEADERS)
    for e in edges:
        ws.append([_cell(e[h]) for h in EDGE_HEADERS])
    wc = wb.create_sheet('Centrality'); wc.append(CENT_HEADERS)
    for r in cent:
        wc.append([_cell(r[h]) for h in CENT_HEADERS])
    MC.save_workbook_atomic(wb, path)


def sheet_name(i, section):
    return ('%02d %s' % (i, section))[:31]


def write_archive(path, rows, section_order):
    import openpyxl, moneyball_core as MC
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'All Bills'; ws.append(HEADERS)
    for r in rows:
        ws.append([_cell(r[h]) for h in HEADERS])
    for i, sec in enumerate(section_order, 1):
        s = wb.create_sheet(sheet_name(i, sec)); s.append(HEADERS)
        for r in rows:
            if r['Section'] == sec:
                s.append([_cell(r[h]) for h in HEADERS])
    MC.save_workbook_atomic(wb, path)


# ----------------------------------------------------------------------------- schema diff
def _cls(v):
    if v is None:
        return 'none'
    return 'num' if isinstance(v, (int, float)) and not isinstance(v, bool) else 'str'


def _sheet(wb, name):
    it = wb[name].iter_rows(values_only=True)
    hdr = tuple(next(it)); return hdr, [r for r in it if any(c is not None for c in r)]


def _dtype_map(hdr, rows):
    out = {}
    for i, h in enumerate(hdr):
        out[h] = sorted({_cls(r[i]) for r in rows})
    return out


def schema_diff(candidate_net, candidate_arc, ref_net, ref_arc, roster_members=None):
    """-> {'fail': [...], 'report': {...}}. FAIL on contract breaks, REPORT content deltas."""
    import openpyxl
    fail, rep = [], {}
    cn, rn = openpyxl.load_workbook(candidate_net, read_only=True), openpyxl.load_workbook(ref_net, read_only=True)
    ca, ra = openpyxl.load_workbook(candidate_arc, read_only=True), openpyxl.load_workbook(ref_arc, read_only=True)
    for tag, c, r in (('network', cn, rn), ('archive', ca, ra)):
        if c.sheetnames != r.sheetnames:
            fail.append('%s sheet names/order differ: %s vs %s' % (tag, c.sheetnames, r.sheetnames))
    # network
    ch, crow = _sheet(cn, 'Edges'); rh, rrow = _sheet(rn, 'Edges')
    if ch != rh: fail.append('Edges headers %s != %s' % (ch, rh))
    cd, rd = _dtype_map(ch, crow), _dtype_map(rh, rrow)
    for h in ch:
        if h in rd and not set(cd[h]) <= set(rd[h]) | ({'num'} if set(rd[h]) <= {'num'} else set()):
            fail.append('Edges col %s dtype %s not within reference %s' % (h, cd[h], rd[h]))
    for r in crow:
        if not (BIOGUIDE_RE.match(str(r[0])) and BIOGUIDE_RE.match(str(r[1]))):
            fail.append('Edges non-bioguide key %s->%s' % (r[0], r[1])); break
    if [(r[0], r[1]) for r in crow] != sorted((r[0], r[1]) for r in crow): fail.append('Edges not sorted (From,To)')
    if any(r[3] != r[4] + r[5] for r in crow): fail.append('NBills != NOriginal + NRegular')
    if any(r[0] == r[1] for r in crow): fail.append('self-loop in Edges')
    kh, krow = _sheet(cn, 'Centrality'); rkh, rkrow = _sheet(rn, 'Centrality')
    if kh != rkh: fail.append('Centrality headers %s != %s' % (kh, rkh))
    if [(-r[2], -r[3]) for r in krow] != sorted((-r[2], -r[3]) for r in krow): fail.append('Centrality not sorted')
    if any(not BIOGUIDE_RE.match(str(r[0])) for r in krow): fail.append('Centrality non-bioguide node')
    ce, re_ = {(r[0], r[1]): r for r in crow}, {(r[0], r[1]): r for r in rrow}
    rep['edges'] = {'candidate': len(crow), 'reference': len(re_), 'added': len(set(ce) - set(re_)),
                    'removed': len(set(re_) - set(ce)),
                    'weight_changed_1e-9': sum(1 for k in set(ce) & set(re_) if abs(ce[k][2] - re_[k][2]) > 1e-9)}
    rep['centrality'] = {'candidate_nodes': len(krow), 'reference_nodes': len(rkrow),
                         'top25_candidate': [(r[1], r[2]) for r in krow[:25]],
                         'top25_reference': [(r[1], r[2]) for r in rkrow[:25]]}
    # archive
    ah, arow = _sheet(ca, 'All Bills'); rah, rarow = _sheet(ra, 'All Bills')
    if ah != rah: fail.append('All Bills headers differ')
    ix = {h: i for i, h in enumerate(ah)}
    for r in arow:
        if not NAME_RE.match(str(r[ix['Name']])): fail.append('bad Name %r' % r[ix['Name']]); break
    if any(r[ix['Role']] not in ROLES for r in arow): fail.append('bad Role')
    if any(r[ix['Congress']] not in CONGRESSES for r in arow): fail.append('bad Congress')
    if any(r[ix['DecayWeight']] not in (1, 0.8, 0.8 ** 2) and abs((r[ix['DecayWeight']] or 0) - decay(r[ix['Congress']])) > 1e-12 for r in arow): fail.append('bad DecayWeight')
    if any(r[ix['LastAction']] not in TRACKER_VOCAB for r in arow):
        fail.append('LastAction outside tracker vocabulary: %s' % sorted({r[ix['LastAction']] for r in arow} - TRACKER_VOCAB))
    if roster_members is not None:
        off = sorted({r[ix['Member']] for r in arow} - set(roster_members))
        if off: fail.append('Member not on roster: %s' % off[:5])
    ad, rad = _dtype_map(ah, arow), _dtype_map(rah, rarow)
    for h in ah:
        if h in rad and 'str' in ad[h] and 'str' not in rad[h]: fail.append('All Bills col %s gained str' % h)
        if h in rad and 'num' in ad[h] and 'num' not in rad[h]: fail.append('All Bills col %s gained num' % h)
    ck = lambda rows: {(r[ix['Section']], r[ix['Role']], r[ix['Member']], r[ix['Name']], r[ix['Congress']]) for r in rows}
    cs, rs = ck(arow), ck(rarow)
    if len(cs) != len(arow): fail.append('duplicate (Section,Role,Member,Name,Congress) rows')
    per_sec = collections.Counter(k[0] for k in cs - rs); per_sec_r = collections.Counter(k[0] for k in rs - cs)
    bills = lambda rows: {(r[ix['Name']], r[ix['Congress']]) for r in rows}
    rep['archive'] = {'candidate_rows': len(arow), 'reference_rows': len(rarow), 'rows_added': len(cs - rs),
                      'rows_removed': len(rs - cs), 'bills_added': len(bills(arow) - bills(rarow)),
                      'bills_removed': len(bills(rarow) - bills(arow)),
                      'rows_added_by_section': dict(per_sec), 'rows_removed_by_section': dict(per_sec_r)}
    rep['archive']['excluded_bill_flips'] = _exclusion_flips(arow, rarow, ix)
    return {'fail': fail, 'report': rep}


def _exclusion_flips(arow, rarow, ix):
    try:
        import bill_rules as BR
    except Exception as e:  # report-only helper
        return 'unavailable: %s' % e
    f = lambda r: bool(BR.excluded_bill(r[ix['Name']], r[ix['Congress']], r[ix['Title']] or '', r[ix['bill_summary']] or ''))
    cur = {(r[ix['Name']], r[ix['Congress']]): f(r) for r in arow}
    ref = {(r[ix['Name']], r[ix['Congress']]): f(r) for r in rarow}
    return sum(1 for k in cur if k in ref and cur[k] != ref[k])


# ----------------------------------------------------------------------------- CLI
def assert_out_ok(out, live_root):
    o = os.path.realpath(out)
    for forbidden in {os.path.realpath(live_root), os.path.realpath(REPO)}:
        if o == forbidden:
            raise Refusal('refusing to write to a repo/live root: %s' % o)


def load_roster(root):
    d = json.load(open(os.path.join(root, 'data', 'engine', 'roster.json'), encoding='utf-8'))
    return {m['bioguide']: m['member'] for m in d['members']}, d


def build(a):
    from paths import ROOT
    out_base = a.out or os.path.join(REPO, 'data', 'cr_lane')
    run_id = a.run_id or datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    out = os.path.join(out_base, run_id)
    assert_out_ok(out_base, ROOT); assert_out_ok(out, ROOT)
    raw_root = a.raw_root or os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Melange', 'lda_rebuild', 'raw')
    seeds = a.seed_archive or [os.path.join(ROOT, p) for p in SEED_ARCHIVES]
    need = [raw_root, os.path.join(ROOT, 'data', 'engine', 'roster.json')] + seeds + ([a.billset] if a.network_set == 'r' else [])
    gone = [p for p in need if not os.path.exists(p)]
    if gone:
        raise Refusal('WARNING: inputs missing, nothing built: %s' % gone)
    roster, rdoc = load_roster(ROOT)
    seed, order = section_seed(seeds)
    billset = load_billset(a.billset) if a.network_set == 'r' else {}
    bills, refused = read_bills(raw_root, list(set(seed) | set(billset)))
    if len(refused) > a.allow_refusals:
        raise Refusal('%d missing BILLSTATUS keys (> --allow-refusals %d): %s' % (len(refused), a.allow_refusals, refused[:5]))
    rows = archive_rows(seed, bills, roster, order, a.withdrawn, a.originals_also_cosponsor == 'yes')
    edges, ref3, labels = network_edges(a.network_set, a.count, rows, bills, roster, seed, a.edge_withdrawn, billset or None)
    if a.network_set == 'archive' and len(ref3) > a.allow_refusals:   # R-set skips (off-roster sponsor) are reported, not gating
        raise Refusal('%d sponsor-identity refusals: %s' % (len(ref3), ref3[:5]))
    cent = centrality(edges, labels)
    os.makedirs(out, exist_ok=True)
    npath = os.path.join(out, 'congress_cosponsor_network_%s.xlsx' % run_id)
    apath = os.path.join(out, 'congress_bills_archive_%s_A.xlsx' % run_id)    # variant A (all BILLSTATUS), opt-in
    bpath = os.path.join(out, 'congress_bills_archive_%s.xlsx' % run_id)      # variant B = the default archive
    ref_arc = os.path.join(ROOT, REF_ARCHIVE); names = set(roster.values())
    rows_b = variant_b_rows(ref_arc, rows, names, order)
    rows_a0 = variant_a0_rows(ref_arc, os.path.join(ROOT, SEED_ARCHIVES[1]), names, order)
    a0path = os.path.join(out, 'congress_bills_archive_%s_A0.xlsx' % run_id)
    write_network(npath, edges, cent); write_archive(apath, rows, order); write_archive(bpath, rows_b, order)
    write_archive(a0path, rows_a0, order)
    ref_a = os.path.join(ROOT, REF_ARCHIVE)
    sd = schema_diff(npath, bpath, os.path.join(ROOT, REF_NETWORK), ref_a, roster.values())   # the default archive (B)
    sd_a = schema_diff(npath, apath, os.path.join(ROOT, REF_NETWORK), ref_a, roster.values())
    sd['variant_A_fail'] = sd_a['fail']
    sd_a0 = schema_diff(npath, a0path, os.path.join(ROOT, REF_NETWORK), ref_a, roster.values())   # informational: A0 is expected to FAIL LastAction
    ab = variant_report({'A0': rows_a0, 'A': rows, 'B': rows_b}, _read_rows(ref_arc, lambda r: True))
    ab['A0']['schema_fail_informational'] = sd_a0['fail']
    zips = sorted(os.path.join(dp, f) for dp, _, fs in os.walk(raw_root) for f in fs
                  if f.endswith('.zip') and re.search(r'billstatus11[789]', dp))
    man = {'schema': 'cr-manifest-v1', 'run_id': run_id, 'built_at': datetime.datetime.now().isoformat(),
           'dials': {'network_set': a.network_set, 'count': a.count, 'default_archive': 'B', 'withdrawn': a.withdrawn, 'edge_withdrawn': a.edge_withdrawn, 'originals_also_cosponsor': a.originals_also_cosponsor, 'newest_congress': NEWEST_CONGRESS,
                     'allow_refusals': a.allow_refusals},
           'inputs_sha256': {'roster.json': sha256(os.path.join(ROOT, 'data', 'engine', 'roster.json')),
                             **{os.path.basename(p): sha256(p) for p in seeds},
                             **{os.path.relpath(p, raw_root): sha256(p) for p in zips}},
           'counts': {'seed_bills': len(seed), 'bills_read': len(bills), 'sections': len(order),
                      'archive_rows': len(rows), 'edges': len(edges), 'centrality_nodes': len(cent)},
           'refusals': {'billstatus': [[list(k), r] for k, r in refused], 'sponsor_identity_or_skipped': [list(map(str, x)) for x in ref3]},
           'outputs': {os.path.basename(npath): sha256(npath), os.path.basename(apath): sha256(apath), os.path.basename(bpath): sha256(bpath), os.path.basename(a0path): sha256(a0path)},
           'variants': ab, 'schema_diff': sd, 'acceptance_passed': not sd['fail'] and not sd_a['fail']}
    with open(os.path.join(out, 'cr_manifest.json'), 'w', encoding='utf-8') as fh:   # written LAST
        json.dump(man, fh, indent=1, default=str)
    print('cr lane (archive B): %d archive rows, %d edges, %d nodes, %d refusals, schema fails=%d -> %s'
          % (len(rows_b), len(edges), len(cent), len(refused) + len(ref3), len(sd['fail']), out))
    return 0 if man['acceptance_passed'] else 2


def diff(a):
    from paths import ROOT
    d = a.candidate_dir
    man = json.load(open(os.path.join(d, 'cr_manifest.json'), encoding='utf-8'))
    if not man.get('acceptance_passed'):
        raise Refusal('candidate manifest has acceptance_passed=false')
    npath = next(os.path.join(d, f) for f in os.listdir(d) if f.startswith('congress_cosponsor_network_'))
    apath = next(os.path.join(d, f) for f in sorted(os.listdir(d)) if f.startswith('congress_bills_archive_') and not re.search(r'_(A|A0)\.xlsx$', f))   # default archive (B)
    roster, _ = load_roster(ROOT)
    res = schema_diff(npath, apath, a.ref_network or os.path.join(ROOT, REF_NETWORK),
                      a.ref_archive or os.path.join(ROOT, REF_ARCHIVE), roster.values())
    print(json.dumps(res, indent=1, default=str))
    return 0 if not res['fail'] else 2


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    sp = ap.add_subparsers(dest='cmd', required=True)
    b = sp.add_parser('build')
    b.add_argument('--seed-archive', nargs='*'); b.add_argument('--raw-root'); b.add_argument('--out')
    b.add_argument('--run-id'); b.add_argument('--allow-refusals', type=int, default=0)
    b.add_argument('--withdrawn', choices=('exclude', 'include'), default='exclude')
    b.add_argument('--edge-withdrawn', choices=('exclude', 'include'), default='include')
    b.add_argument('--network-set', choices=('r', 'archive'), default='r')
    b.add_argument('--count', choices=('per-section', 'dedupe'), default='per-section')
    b.add_argument('--billset', default=BILLSET)
    b.add_argument('--originals-also-cosponsor', choices=('yes', 'no'), default='yes')
    d = sp.add_parser('diff')
    d.add_argument('--candidate-dir', required=True); d.add_argument('--ref-network'); d.add_argument('--ref-archive')
    a = ap.parse_args(argv)
    try:
        return build(a) if a.cmd == 'build' else diff(a)
    except Refusal as e:
        print('REFUSED: %s' % e, file=sys.stderr)
        return 3


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.exit(main())
