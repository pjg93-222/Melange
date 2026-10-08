"""Fail-closed coverage gate for a merged Mentat industry database.

Usage:
  python validate_industry_coverage.py DATABASE INDUSTRY CYCLES FANOUT_RECEIPT
      [--output RECEIPT]

The gate is intentionally independent of the collector and workbook generator.
It proves that a successful fanout receipt covers the complete company x cycle
grid, that every FEC API work unit was fetched and stored, that all collection
canaries passed, and that SQLite integrity/foreign-key checks are clean.  It
writes an atomic machine-readable receipt and exits nonzero on any failure.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import sys
from contextlib import closing
from pathlib import Path
from typing import Any, Iterable


CYCLE_MAP = {
    "2021-2022": 2022,
    "2022": 2022,
    "2023-2024": 2024,
    "2024": 2024,
    "2025-2026": 2026,
    "2026": 2026,
}


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def roster_cell(value: object, source_row: int) -> dict[str, Any]:
    original = clean_text(value)
    result: dict[str, Any] = {
        "variants": [],
        "exclusions": [],
        "metadata": {},
    }
    if not original:
        return result
    upper = original.upper()
    if re.match(r"^NOT\s+", upper):
        result["exclusions"] = [
            clean_text(re.sub(r"^NOT\s+", "", original, flags=re.IGNORECASE))
        ]
        return result
    cik = re.fullmatch(r"CIK\s+([0-9]+)", original, flags=re.IGNORECASE)
    if cik:
        result["metadata"] = {"cik": cik.group(1)}
        return result
    instruction = re.fullmatch(
        r"""JUST\s+["']?(.+?)["']?""",
        original,
        flags=re.IGNORECASE,
    )
    if instruction:
        result["metadata"] = {
            "search_instruction": clean_text(instruction.group(1))
        }
        return result
    if re.search(
        r"^(BUT|HOWEVER)\b|\bSO YOU NEED:?$|^THE INSURANCE ACTIVITY\b",
        upper,
    ):
        return result
    confidence = (
        "spot_check"
        if re.search(r"\u26a0|SPOT[- ]?CHECK", upper)
        else (
            "safe"
            if re.search(r"\bSAFE\b|DISTINCTIVE", upper)
            else "normal"
        )
    )
    for piece in re.split(r"\s+/\s+", original):
        variant = clean_text(re.sub(r"\s*\([^)]*\)", "", piece))
        if variant:
            result["variants"].append(
                {
                    "variant": variant,
                    "confidence": confidence,
                    "source_row": source_row,
                }
            )
    return result


def workbook_roster(path: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError(
            "Roster validation requires openpyxl, which is already a Mentat "
            "workbook dependency."
        ) from error

    resolved = path.resolve(strict=True)
    book = openpyxl.load_workbook(
        resolved,
        read_only=True,
        data_only=False,
    )
    try:
        industries: list[dict[str, Any]] = []
        all_checks: list[bool] = []
        for sheet in book.worksheets:
            raw_rows = [
                (row_number, [cell.value for cell in row])
                for row_number, row in enumerate(
                    sheet.iter_rows(),
                    start=1,
                )
            ]
            blocks: list[list[tuple[int, list[object]]]] = []
            if sheet.title == "Crypto":
                blocks = [
                    [(row_number, row)]
                    for row_number, row in raw_rows
                    if any(clean_text(value) for value in row)
                ]
            else:
                current: list[tuple[int, list[object]]] = []
                for row_number, row in raw_rows:
                    if any(clean_text(value) for value in row):
                        current.append((row_number, row))
                    elif current:
                        blocks.append(current)
                        current = []
                if current:
                    blocks.append(current)

            companies: list[dict[str, Any]] = []
            parse_errors: list[str] = []
            for block in blocks:
                parsed = [
                    roster_cell(value, row_number)
                    for row_number, row in block
                    for value in row
                    if clean_text(value)
                ]
                variants: list[dict[str, Any]] = []
                seen_variants: set[str] = set()
                for cell in parsed:
                    for variant in cell["variants"]:
                        key = clean_text(variant["variant"])
                        if key not in seen_variants:
                            seen_variants.add(key)
                            variants.append(variant)
                exclusions = sorted(
                    {
                        exclusion
                        for cell in parsed
                        for exclusion in cell["exclusions"]
                        if exclusion
                    },
                    key=lambda value: (value.upper(), value),
                )
                if not variants:
                    parse_errors.append(
                        f"Rows {block[0][0]}-{block[-1][0]} have no usable variants."
                    )
                    continue
                companies.append(
                    {
                        "company": variants[0]["variant"],
                        "variant_count": len(variants),
                        "variants": variants,
                        "exclusion_count": len(exclusions),
                        "exclusions": exclusions,
                        "source_rows": [item[0] for item in block],
                    }
                )

            company_keys = [company["company"].casefold() for company in companies]
            duplicate_companies = sorted(
                {
                    company["company"]
                    for company in companies
                    if company_keys.count(company["company"].casefold()) > 1
                },
                key=lambda value: (value.upper(), value),
            )
            owners: dict[str, set[str]] = {}
            labels: dict[str, str] = {}
            for company in companies:
                for variant in company["variants"]:
                    key = re.sub(
                        r"[^A-Z0-9]+",
                        "",
                        variant["variant"].upper(),
                    )
                    if not key:
                        continue
                    labels.setdefault(key, variant["variant"])
                    owners.setdefault(key, set()).add(company["company"])
            cross_company_variants = [
                {
                    "variant": labels[key],
                    "companies": sorted(
                        company_owners,
                        key=lambda value: (value.upper(), value),
                    ),
                }
                for key, company_owners in sorted(owners.items())
                if len(company_owners) > 1
            ]
            checks = {
                "has_company_blocks": bool(companies),
                "all_blocks_parse": not parse_errors,
                "company_names_unique": not duplicate_companies,
                "variants_unambiguous": not cross_company_variants,
            }
            all_checks.extend(checks.values())
            industries.append(
                {
                    "industry": sheet.title,
                    "ready": all(checks.values()),
                    "company_count": len(companies),
                    "variant_count": sum(
                        company["variant_count"] for company in companies
                    ),
                    "exclusion_pattern_count": sum(
                        company["exclusion_count"] for company in companies
                    ),
                    "companies_with_exclusions": sum(
                        company["exclusion_count"] > 0
                        for company in companies
                    ),
                    "companies_without_exclusions": sum(
                        company["exclusion_count"] == 0
                        for company in companies
                    ),
                    "checks": checks,
                    "parse_errors": parse_errors,
                    "duplicate_companies": duplicate_companies,
                    "cross_company_variants": cross_company_variants,
                    "companies": companies,
                }
            )
    finally:
        book.close()
    return {
        "status": "succeeded" if industries and all(all_checks) else "failed",
        "mode": "roster",
        "workbook": str(resolved),
        "industry_count": len(industries),
        "industries": [industry["industry"] for industry in industries],
        "readiness": industries,
    }


def parse_cycles(value: str) -> tuple[int, ...]:
    labels = [
        item.strip()
        for item in value.replace(",", "|").split("|")
        if item.strip()
    ]
    unknown = sorted(set(labels).difference(CYCLE_MAP))
    if not labels or unknown:
        raise ValueError(
            "Cycles must use 2021-2022/2022, 2023-2024/2024, or "
            f"2025-2026/2026; invalid={unknown!r}"
        )
    return tuple(sorted({CYCLE_MAP[label] for label in labels}))


def atomic_write_json(path: Path, value: dict[str, Any]) -> None:
    target = path.resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    pending = target.with_name(
        f".{target.name}.{os.getpid()}.pending"
    )
    with pending.open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, indent=2, sort_keys=True)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(pending, target)


def table_exists(connection: sqlite3.Connection, table: str) -> bool:
    row = connection.execute(
        """
        SELECT 1
        FROM sqlite_master
        WHERE type='table' AND name=?
        """,
        (table,),
    ).fetchone()
    return row is not None


def placeholders(values: Iterable[object]) -> str:
    materialized = tuple(values)
    if not materialized:
        raise ValueError("At least one SQL value is required.")
    return ",".join("?" for _ in materialized)


def scalar(
    connection: sqlite3.Connection,
    sql: str,
    parameters: Iterable[object] = (),
) -> int:
    row = connection.execute(sql, tuple(parameters)).fetchone()
    return 0 if row is None else int(row[0])


def as_int(value: Any, field: str) -> int:
    if isinstance(value, bool):
        return int(value)
    try:
        return int(value)
    except (TypeError, ValueError) as error:
        raise ValueError(
            f"Fanout receipt field {field!r} is not an integer: {value!r}"
        ) from error


def truthy(value: Any) -> bool:
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "succeeded"}
    return bool(value)


def load_fanout_receipt(path: Path) -> dict[str, Any]:
    with path.resolve(strict=True).open(encoding="utf-8") as stream:
        value = json.load(stream)
    if not isinstance(value, dict):
        raise ValueError("Fanout receipt must be a JSON object.")
    return value


def quick_check(connection: sqlite3.Connection) -> list[str]:
    return [
        str(row[0])
        for row in connection.execute("PRAGMA quick_check")
    ]


def foreign_key_failures(connection: sqlite3.Connection) -> int:
    return sum(1 for _ in connection.execute("PRAGMA foreign_key_check"))


def expected_companies(
    connection: sqlite3.Connection,
    industry: str,
) -> tuple[str, ...]:
    if not table_exists(connection, "companies"):
        raise RuntimeError("Merged database is missing the companies table.")
    rows = connection.execute(
        """
        SELECT DISTINCT company
        FROM companies
        WHERE industry=?
        ORDER BY company COLLATE NOCASE, company
        """,
        (industry,),
    ).fetchall()
    companies = tuple(str(row[0]).strip() for row in rows if str(row[0]).strip())
    if not companies:
        raise RuntimeError(
            f"Merged database has no company roster for industry {industry!r}."
        )
    return companies


def validate_fanout_receipt(
    receipt: dict[str, Any],
    database: Path,
    expected_task_count: int,
) -> dict[str, Any]:
    status = str(receipt.get("status", "")).strip().lower()
    merged = truthy(receipt.get("merged", False))
    task_count = as_int(receipt.get("task_count", -1), "task_count")
    succeeded = as_int(receipt.get("succeeded", -1), "succeeded")
    failed = as_int(receipt.get("failed", -1), "failed")
    source_value = receipt.get("source_database")
    source_matches = True
    if source_value:
        source_matches = (
            Path(str(source_value)).resolve() == database.resolve()
        )
    checks = {
        "status_succeeded": status == "succeeded",
        "merged": merged,
        "task_count_exact": task_count == expected_task_count,
        "succeeded_exact": succeeded == expected_task_count,
        "failed_zero": failed == 0,
        "source_database_matches": source_matches,
    }
    return {
        "status": status,
        "merged": merged,
        "task_count": task_count,
        "succeeded": succeeded,
        "failed": failed,
        "source_database_matches": source_matches,
        "checks": checks,
    }


def api_progress_grid(
    connection: sqlite3.Connection,
    companies: tuple[str, ...],
    cycles: tuple[int, ...],
) -> tuple[list[dict[str, Any]], int, int]:
    if not table_exists(connection, "api_work_progress"):
        raise RuntimeError(
            "Merged database is missing the api_work_progress table."
        )
    company_sql = placeholders(companies)
    cycle_sql = placeholders(cycles)
    parameters = (*companies, *cycles)
    rows = connection.execute(
        f"""
        SELECT
          company,
          cycle,
          SUM(CASE WHEN kind='employer' THEN 1 ELSE 0 END) AS employer_units,
          SUM(CASE WHEN kind='corporate' THEN 1 ELSE 0 END) AS corporate_units,
          SUM(
            CASE
              WHEN fetched_complete<>1 OR stored_complete<>1 THEN 1
              ELSE 0
            END
          ) AS incomplete_units,
          COUNT(*) AS total_units,
          COALESCE(SUM(rows_seen), 0) AS rows_seen
        FROM api_work_progress
        WHERE company IN ({company_sql})
          AND cycle IN ({cycle_sql})
        GROUP BY company, cycle
        """,
        parameters,
    ).fetchall()
    observed = {
        (str(row[0]), int(row[1])): {
            "company": str(row[0]),
            "cycle": int(row[1]),
            "employer_units": int(row[2]),
            "corporate_units": int(row[3]),
            "incomplete_units": int(row[4]),
            "total_units": int(row[5]),
            "rows_seen": int(row[6]),
        }
        for row in rows
    }
    grid: list[dict[str, Any]] = []
    missing = 0
    incomplete = 0
    for company in companies:
        for cycle in cycles:
            item = observed.get((company, cycle))
            if item is None:
                item = {
                    "company": company,
                    "cycle": cycle,
                    "employer_units": 0,
                    "corporate_units": 0,
                    "incomplete_units": 0,
                    "total_units": 0,
                    "rows_seen": 0,
                }
            complete = (
                item["employer_units"] >= 1
                and item["corporate_units"] >= 1
                and item["incomplete_units"] == 0
            )
            item["complete"] = complete
            if item["total_units"] == 0:
                missing += 1
            elif not complete:
                incomplete += 1
            grid.append(item)
    return grid, missing, incomplete


def scoped_count(
    connection: sqlite3.Connection,
    table: str,
    companies: tuple[str, ...],
    cycles: tuple[int, ...],
) -> tuple[int, int]:
    if not table_exists(connection, table):
        raise RuntimeError(f"Merged database is missing the {table} table.")
    company_sql = placeholders(companies)
    cycle_sql = placeholders(cycles)
    cycle_column = (
        "two_year_transaction_period"
        if table == "contributions"
        else "cycle"
    )
    total = scalar(connection, f"SELECT COUNT(*) FROM {table}")
    in_scope = scalar(
        connection,
        f"""
        SELECT COUNT(*)
        FROM {table}
        WHERE company IN ({company_sql})
          AND {cycle_column} IN ({cycle_sql})
        """,
        (*companies, *cycles),
    )
    return total, in_scope


def validate(
    database: Path,
    industry: str,
    cycles: tuple[int, ...],
    fanout_receipt: Path,
) -> dict[str, Any]:
    resolved = database.resolve(strict=True)
    receipt = load_fanout_receipt(fanout_receipt)
    uri = resolved.as_uri() + "?mode=ro"
    with closing(sqlite3.connect(uri, uri=True, timeout=60)) as connection:
        connection.execute("PRAGMA busy_timeout=60000")
        connection.execute("PRAGMA query_only=ON")
        integrity = quick_check(connection)
        foreign_keys = foreign_key_failures(connection)
        companies = expected_companies(connection, industry)
        expected_tasks = len(companies) * len(cycles)
        fanout = validate_fanout_receipt(
            receipt,
            resolved,
            expected_tasks,
        )
        grid, missing_grid, incomplete_grid = api_progress_grid(
            connection,
            companies,
            cycles,
        )
        canary_failures = scalar(
            connection,
            f"""
            SELECT COUNT(*)
            FROM canary_log
            WHERE company IN ({placeholders(companies)})
              AND cycle IN ({placeholders(cycles)})
              AND passed<>1
            """,
            (*companies, *cycles),
        )
        contribution_total, contribution_scoped = scoped_count(
            connection,
            "contributions",
            companies,
            cycles,
        )
        corporate_total, corporate_scoped = scoped_count(
            connection,
            "corporate_giving",
            companies,
            cycles,
        )

    checks = {
        "sqlite_quick_check": integrity == ["ok"],
        "foreign_key_check": foreign_keys == 0,
        "fanout_receipt": all(fanout["checks"].values()),
        "company_cycle_grid": missing_grid == 0 and incomplete_grid == 0,
        "canaries": canary_failures == 0,
        "contribution_scope": contribution_total == contribution_scoped,
        "corporate_scope": corporate_total == corporate_scoped,
    }
    return {
        "status": "succeeded" if all(checks.values()) else "failed",
        "database": str(resolved),
        "industry": industry,
        "cycles": list(cycles),
        "companies": list(companies),
        "company_count": len(companies),
        "expected_task_count": expected_tasks,
        "checks": checks,
        "sqlite_quick_check": integrity,
        "foreign_key_failures": foreign_keys,
        "fanout": fanout,
        "missing_company_cycles": missing_grid,
        "incomplete_company_cycles": incomplete_grid,
        "canary_failures": canary_failures,
        "contributions": {
            "total": contribution_total,
            "in_scope": contribution_scoped,
        },
        "corporate_giving": {
            "total": corporate_total,
            "in_scope": corporate_scoped,
        },
        "company_cycle_grid": grid,
    }


def parse_arguments(arguments: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("database", type=Path, nargs="?")
    parser.add_argument("industry", nargs="?")
    parser.add_argument("cycles", nargs="?")
    parser.add_argument("fanout_receipt", type=Path, nargs="?")
    parser.add_argument(
        "--roster",
        type=Path,
        help=(
            "Validate and list every industry discovered from Company Name.xlsx."
        ),
    )
    parser.add_argument("--output", type=Path)
    parsed = parser.parse_args(arguments)
    if parsed.roster is None and any(
        value is None
        for value in (
            parsed.database,
            parsed.industry,
            parsed.cycles,
            parsed.fanout_receipt,
        )
    ):
        parser.error(
            "DATABASE INDUSTRY CYCLES FANOUT_RECEIPT are required unless "
            "--roster is used."
        )
    if parsed.roster is not None and any(
        value is not None
        for value in (
            parsed.database,
            parsed.industry,
            parsed.cycles,
            parsed.fanout_receipt,
        )
    ):
        parser.error("--roster cannot be combined with database gate arguments.")
    return parsed


def main(arguments: list[str]) -> int:
    options = parse_arguments(arguments)
    output = options.output
    try:
        if options.roster is not None:
            result = workbook_roster(options.roster)
        else:
            result = validate(
                options.database,
                options.industry.strip(),
                parse_cycles(options.cycles),
                options.fanout_receipt,
            )
    except Exception as error:
        result = {
            "status": "failed",
            "mode": "roster" if options.roster is not None else "coverage",
            "database": (
                str(options.database.resolve())
                if options.database is not None
                else None
            ),
            "industry": (
                options.industry.strip()
                if options.industry is not None
                else None
            ),
            "error": f"{type(error).__name__}: {error}",
        }
    if output is not None:
        atomic_write_json(output, result)
    print(json.dumps(result, sort_keys=True))
    return 0 if result.get("status") == "succeeded" else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
