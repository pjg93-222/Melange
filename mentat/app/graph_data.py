"""Guided report graphs. Read saved report facts without changing classifications."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from functools import lru_cache, wraps
import csv
import json
from pathlib import Path
import posixpath
import re
import threading
import time
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import downstream
import native_scan
from view_cache import ViewCache, file_signature

from report_insights import digest, digest_stamp, key, local_reference, stamp

TIERS = {'Chief': 'Chief', 'Middle Management': 'Middle-Management', 'Employee': 'Employees'}
# The order the employee-level legend reads in, largest level first on every
# saved report so far; levels_by_company keeps each company's levels in it.
LEVEL_ORDER = (TIERS['Employee'], TIERS['Middle Management'], TIERS['Chief'])
CATEGORIES = ['Super PAC', 'Trade Association', 'JFC', 'Company PAC', 'Leadership PAC',
              'Carey', 'Other PAC', 'Candidate', 'Party Committee', 'Unclassified']
AUTHORIZED = {'Principal campaign committee', 'Authorized candidate committee'}
MAX_SOURCE_ROWS = 1000000
# The four recorded activity directions, in the order the console offers
# them. direction_totals is emitted for every one of them, including the
# empty ones: a drawer that showed only the non-empty sides could not say
# that Support is zero, which is the answer the reader came for.
ACTIVITY_DIRECTIONS = ('Contribution / Other', 'Support', 'Oppose', 'Unknown')
SINGLEFLIGHT_WAIT_SECONDS = 120


def singleflight_cache(maxsize):
    """Share an in-progress file read without blocking reads of other files.

    A caller that has waited SINGLEFLIGHT_WAIT_SECONDS for the read in progress
    stops sharing it and reads the file itself, so one stalled read delays the
    requests behind it instead of blocking them forever.
    """
    def decorate(load):
        cached = lru_cache(maxsize=maxsize)(load)
        condition = threading.Condition()
        reading = set()

        @wraps(load)
        def read(file_stamp):
            deadline = time.monotonic() + SINGLEFLIGHT_WAIT_SECONDS
            leader = True
            with condition:
                while file_stamp in reading:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0 or not condition.wait(remaining):
                        # The read in progress is stalled. Read independently, as
                        # plain lru_cache would, and leave its token to its owner.
                        leader = False
                        break
                if leader:
                    reading.add(file_stamp)
            try:
                return cached(file_stamp)
            finally:
                if leader:
                    with condition:
                        reading.remove(file_stamp)
                        condition.notify_all()

        read.cache_clear = cached.cache_clear
        read.cache_info = cached.cache_info
        return read
    return decorate


def cycle(value):
    match = re.fullmatch(r'(?:\d{4}[-–])?(\d{4})', str(value or '').strip())
    return match[1] if match else ''


def amount(value):
    try:
        result = Decimal(str(value))
        if result.is_finite():
            return result
    except (InvalidOperation, ValueError):
        pass
    raise ValueError('A required report amount is missing or nonnumeric; no partial graph was created.')


def money(value):
    return float(round(value, 2))


def table_rows(sheet, header=5, required=()):
    iterator = sheet.iter_rows(min_row=header, values_only=True)
    headers = [key(v) for v in next(iterator, ())]
    if not set(required) <= set(headers):
        raise ValueError(f'{sheet.title}: required report columns are missing.')
    for number, row in enumerate(iterator, 1):
        if number > MAX_SOURCE_ROWS:
            raise ValueError(f'{sheet.title}: report exceeds the local graph reading limit.')
        record = {h: v for h, v in zip(headers, row) if h}
        if any(v is not None and v != '' for v in record.values()):
            yield record


def activity_xml_rows(path, metadata, sheet_name='Downstream Activity', required=None, optional=False,
                      optional_fields=()):
    """Stream selected activity columns; ignore styles, dates and hyperlinks."""
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    relationship = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    required = set(required or {'cycle','source_committee_id','candidate_id','candidate_officeholder',
                'candidate_party','support_oppose','net_included_amount'})
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        sheet = next((s for s in workbook.iter(ns+'sheet') if s.get('name') == sheet_name), None)
        if sheet is None:
            if optional:
                return
            raise ValueError('This report has no saved downstream activity.')
        relations = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        relation = next((r for r in relations if r.get('Id') == sheet.get(relationship)), None)
        if relation is None or relation.get('TargetMode') == 'External':
            raise ValueError('Downstream worksheet is not stored in this report.')
        target = relation.get('Target','')
        target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            with archive.open('xl/sharedStrings.xml') as stream:
                for _,element in ET.iterparse(stream,events=('end',)):
                    if element.tag == ns+'si':
                        strings.append(''.join(t.text or '' for t in element.iter(ns+'t')))
                        element.clear()
        def value(cell):
            if cell.get('t') == 'inlineStr':
                return ''.join(t.text or '' for t in cell.iter(ns+'t'))
            cached = cell.find(ns+'v')
            if cached is None or cached.text is None:
                return None
            return strings[int(cached.text)] if cell.get('t') == 's' else cached.text
        columns = None
        with archive.open(target) as stream:
            for _,element in ET.iterparse(stream,events=('end',)):
                if element.tag != ns+'row':
                    continue
                number = int(element.get('r','0'))
                if number > MAX_SOURCE_ROWS+5:
                    raise ValueError('Downstream report exceeds the local graph reading limit.')
                if number == 3:
                    metadata['note'] = next((value(c) for c in element if c.get('r') == 'A3'), '')
                elif number == 5:
                    headers = {c.get('r','').rstrip('0123456789'):key(value(c)) for c in element if c.tag == ns+'c'}
                    if not required <= set(headers.values()):
                        raise ValueError(f'{sheet_name}: required report columns are missing.')
                    columns = {col:field for col,field in headers.items() if field in required |
                               {'unverified_version_rows','unresolved_version_rows'} | set(optional_fields)}
                elif number > 5:
                    if columns is None:
                        raise ValueError('Downstream Activity: header row is missing.')
                    row = {columns[col]:value(c) for c in element
                           if c.tag == ns+'c' and (col := c.get('r','').rstrip('0123456789')) in columns}
                    if any(v is not None and v != '' for v in row.values()):
                        yield row
                element.clear()


def report_transfers(file_stamp):
    dependencies = file_signature(p for p in transfer_source(file_stamp[0]) if p is not None)
    return _report_transfers((file_stamp, dependencies))


@singleflight_cache(maxsize=4)
def _report_transfers(version):
    file_stamp, dependencies = version
    metadata = {}
    result = []
    required = {'cycle','source_committee','destination_committee','reported_net_transfer','traversal_status'}
    count = 0
    for row in activity_xml_rows(file_stamp[0], metadata, 'Transfer Paths', required, optional=True):
        if not row.get('source_committee'):
            continue
        count += 1
        # Candidate gifts are already in activity. Never expand a terminal gift.
        if row.get('traversal_status') == 'candidate_committee':
            continue
        result.append({'cycle':cycle(row['cycle']), 'source':str(row['source_committee']),
                       'destination':str(row['destination_committee']),
                       'amount':amount(row['reported_net_transfer'])})
    note = str(metadata.get('note') or '')
    expected = re.search(r'([\d,]+)\s+records;', note)
    if 'first ' in note.lower() or (expected and int(expected[1].replace(',','')) != count):
        payload, manifest = transfer_source(file_stamp[0])
        if not payload or not manifest.is_file() or not payload.is_file():
            raise ValueError('Transfer Paths is a partial display; its pinned full report source is unavailable.')
        import contribution_data as source_data
        proof = json.loads(manifest.read_text(encoding='utf-8-sig'))
        if (proof.get('database_sha256') != source_data._workbook_source(file_stamp)
                or not proof.get('data_sha256') or digest(payload) != proof['data_sha256']):
            raise ValueError('Full transfer source does not match the published report fingerprint.')
        payload_stamp = stamp(payload)
        native = native_scan.tables(payload_stamp, [], payload_table='transfer_paths')
        if native is not None:
            table = native['tables']['transfer_paths']
            rows = (dict(zip(table['headers'], row)) for row in table['rows'])
        else:
            rows = source_data._payload_rows(payload, 'transfer_paths')
        unique = {}
        for row in rows:
            if row.get('status') == 'candidate_committee':
                continue
            identity = (cycle(row['cycle']), str(row['source']), str(row['destination']))
            if not identity[0] or not all(re.fullmatch(r'C\d{8}', value) for value in identity[1:]):
                raise ValueError('The full transfer source has invalid committee or cycle identifiers.')
            value = amount(row['amount'])
            if identity in unique and unique[identity] != value:
                raise ValueError('Full report paths disagree on one transfer amount.')
            unique[identity] = value
        if (stamp(payload) != payload_stamp or stamp(file_stamp[0]) != file_stamp
                or file_signature(p for p in transfer_source(file_stamp[0]) if p is not None) != dependencies):
            raise ValueError('Full report source changed while reading transfers.')
        return [dict(cycle=k[0], source=k[1], destination=k[2], amount=v) for k,v in sorted(unique.items())]
    return result


report_transfers.cache_clear = _report_transfers.cache_clear
report_transfers.cache_info = _report_transfers.cache_info


def transfer_source(path):
    """The full export pinned to this workbook, never a newer collection."""
    path = Path(path)
    if not path.name.endswith(' Industry Spending Overview.xlsx'):
        return None, None
    import contribution_data as source_data
    try:
        fingerprint = source_data._workbook_source(stamp(path))
    except (ValueError, OSError, KeyError, zipfile.BadZipFile):
        return None, None
    industry = path.name.removesuffix(' Industry Spending Overview.xlsx')
    folder = source_data.PROJECT / '.cache/industry-report' / source_data._slug(industry)
    return (folder / f'industry_data-{fingerprint[:16]}.json',
            folder / f'manifest-{fingerprint[:16]}.json')


@lru_cache(maxsize=32)
def graph_schema(file_stamp):
    try:
        # Workbook metadata is sufficient. Loading openpyxl here materialized the
        # shared-string table for every report just to list its sheet names.
        with zipfile.ZipFile(file_stamp[0]) as archive:
            tree = ET.fromstring(archive.read('xl/workbook.xml'))
            ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            return [s.get('name') for s in tree.iter(ns+'sheet')
                    if re.fullmatch(r'\d{4}[-–]\d{4}', s.get('name', ''))]
    except ET.ParseError as error:
        raise ValueError('Workbook sheet metadata is malformed.') from error


def committee_id(value):
    match = re.fullmatch(r'(?:FEC:)?(C\d{8})', str(value or ''))
    return match[1] if match else ''


@singleflight_cache(maxsize=8)
def report_facts(file_stamp):
    """Recipient facts are a complete partition; Master/Employee facts repeat them."""
    cycles = graph_schema(file_stamp)
    projection = ('company','employee_tier','amount','include_in_analysis','recipient_category',
                  'fact_type','person_key','contributor','recipient_key','committee_id','recipient')
    scanned = native_scan.tables(file_stamp, [
        {'name':'Employee List','header':5,'columns':['person_key','person','company']},
        {'name':'PAC & Candidate Map','header':5},
        {'name':'Recipient List','header':5,'columns':['cycle','recipient_key','recipient']},
        {'name':'Company List','header':5,'columns':['company']},
        *[{'name':name,'header':4,'columns':projection} for name in cycles]])
    book = None if scanned else openpyxl.load_workbook(file_stamp[0], read_only=True, data_only=True)
    sheetnames = scanned['sheets'] if scanned else book.sheetnames
    def rows(name, header=5, required=()):
        if scanned is None:
            yield from table_rows(book[name], header, required)
            return
        table = scanned['tables'][name]
        if not set(required) <= set(table['headers']):
            raise ValueError(f'{name}: required report columns are missing.')
        for row in table['rows']:
            yield dict(zip(table['headers'], row))
    try:
        maps, names, companies = {}, {}, set()
        people, person_names = [], {}
        if 'Employee List' in sheetnames:
            for row in rows('Employee List', required=('person_key', 'person', 'company')):
                if row.get('person_key') and row.get('person'):
                    person_names[(str(row['company']), str(row['person_key']))] = str(row['person'])
        if 'PAC & Candidate Map' in sheetnames:
            for row in rows('PAC & Candidate Map', required=('cycle', 'committee_id')):
                lookup = (cycle(row.get('cycle')), str(row.get('committee_id') or ''))
                if lookup in maps and maps[lookup] != row:
                    raise ValueError('Conflicting candidate links in this report; direct rankings are unavailable.')
                maps[lookup] = row
        if 'Recipient List' in sheetnames:
            for row in rows('Recipient List', required=('cycle', 'recipient_key', 'recipient')):
                names.setdefault((cycle(row['cycle']), row['recipient_key']), row['recipient'])
        if 'Company List' in sheetnames:
            companies.update(str(r['company']) for r in rows('Company List', required=('company',)) if r.get('company'))
        facts, masters, partitions = {}, defaultdict(Decimal), defaultdict(Decimal)
        sheets = graph_schema(file_stamp)
        if not sheets:
            raise ValueError('This report has no supported cycle receipt tables. Use Advanced for custom worksheet graphs.')
        for sheet_name in sheets:
            year = cycle(sheet_name)
            fact_mode = False
            for row in rows(sheet_name, 4, ('company','employee_tier','amount','include_in_analysis')):
                if str(row.get('include_in_analysis', '')).lower() != 'yes':
                    continue
                company = str(row.get('company') or 'Unknown company')
                companies.add(company)
                tier = str(row.get('employee_tier') or 'Unknown')
                category = str(row.get('recipient_category') or 'Unclassified')
                scope = (year, company, tier, category)
                fact_type = row.get('fact_type')
                if (fact_type == 'Employee' or not fact_type) and row.get('person_key'):
                    person = str(row['person_key'])
                    people.append({'cycle':year, 'company':company, 'tier':tier, 'person_id':person,
                                   'person':person_names.get((company, person), str(row.get('contributor') or person)),
                                   'amount':amount(row['amount'])})
                if fact_type:
                    fact_mode = True
                    if fact_type not in ('Master','Employee','Recipient'):
                        raise ValueError('Unknown report fact type; no partial graph was created.')
                    if fact_type == 'Master':
                        masters[scope] += amount(row['amount'])
                    if fact_type != 'Recipient':
                        continue
                value = amount(row['amount'])
                partitions[scope] += value
                recipient = str(row.get('recipient_key') or row.get('committee_id') or row.get('recipient') or 'Unknown recipient')
                identifier = committee_id(row.get('committee_id')) or committee_id(recipient)
                mapping = maps.get((year, identifier), {})
                candidate = str(mapping.get('candidate_id_s') or '')
                linked = (category == 'Candidate' and mapping.get('recipient_category') == 'Candidate'
                          and mapping.get('relationship') in AUTHORIZED and re.fullmatch(r'[HSP][A-Z0-9]{8}', candidate))
                identity = (*scope, recipient)
                if identity not in facts:
                    facts[identity] = {'cycle':year, 'company':company, 'tier':tier, 'category':category,
                        'recipient_id':identifier or recipient,
                        'recipient':str(mapping.get('recipient_committee') or names.get((year,recipient)) or row.get('recipient') or recipient),
                        'candidate_id':candidate if linked else '',
                        'candidate':str(mapping.get('direct_candidate_sponsor') or candidate) if linked else '',
                        'party':str(mapping.get('direct_party') or '') if linked else '', 'amount':Decimal(0)}
                facts[identity]['amount'] += value
            if fact_mode:
                keys = {s for s in masters.keys() | partitions.keys() if s[0] == year}
                if any(abs(masters[s] - partitions[s]) > Decimal('.01') for s in keys):
                    raise ValueError('Recipient facts do not reconcile to the report totals; no partial graph was created.')
        if stamp(file_stamp[0]) != file_stamp:
            raise ValueError('The workbook changed while it was being read. Refresh Graph.')
        return {'facts':list(facts.values()), 'people':people, 'companies':sorted(companies), 'cycles':sorted(cycle(s) for s in sheets),
                'committees':[{'cycle':year, 'id':identifier, 'label':str(row.get('recipient_committee') or identifier),
                               'category':str(row.get('recipient_category') or 'Unclassified')} for (year,identifier),row in maps.items()],
                'sha256':digest_stamp(file_stamp), 'report':Path(file_stamp[0]).name}
    finally:
        if book is not None:
            book.close()


@singleflight_cache(maxsize=8)
def _report_activity(options):
    file_stamp, include_filings = options
    try:
        metadata = {}
        groups = defaultdict(Decimal)
        unverified, unresolved = defaultdict(int), defaultdict(int)
        filings = defaultdict(list)
        labels = {}
        count = 0
        required = ('cycle','source_committee_id','candidate_id','candidate_officeholder',
                    'candidate_party','support_oppose','net_included_amount')
        scanned = native_scan.tables(file_stamp, [{'name':'Downstream Activity','header':5,
            'columns':[*required,'unverified_version_rows','unresolved_version_rows',
                       *(['filing_records'] if include_filings else [])]}])
        if scanned:
            table = scanned['tables'].get('Downstream Activity')
            if table is None:
                raise ValueError('This report has no saved downstream activity.')
            if not set(required) <= set(table['headers']):
                raise ValueError('Downstream Activity: required report columns are missing.')
            metadata['note'] = table['note']
            source_rows = (dict(zip(table['headers'], row)) for row in table['rows'])
        else:
            source_rows = activity_xml_rows(file_stamp[0], metadata,
                                            optional_fields=('filing_records',) if include_filings else ())
        for row in source_rows:
            count += 1
            identity = (cycle(row['cycle']), str(row['source_committee_id']), str(row['candidate_id'] or ''), str(row['support_oppose'] or 'Unknown'))
            groups[identity] += amount(row['net_included_amount'])
            unverified[identity] += int(row.get('unverified_version_rows') or 0)
            unresolved[identity] += int(row.get('unresolved_version_rows') or 0)
            if include_filings and row.get('filing_records'):
                try:
                    records = json.loads(row['filing_records'])
                except (TypeError, ValueError):
                    records = []
                if isinstance(records, list):
                    filings[identity].extend(record for record in records if isinstance(record, dict))
            labels[identity] = (str(row.get('candidate_officeholder') or identity[2] or 'Unidentified candidate'), str(row.get('candidate_party') or ''))
        note = str(metadata.get('note') or '')
        expected = re.search(r'condensed to ([\d,]+)', note)
        if expected and int(expected[1].replace(',','')) != count:
            raise ValueError('Downstream worksheet is a partial display; full activity rankings are unavailable.')
        if stamp(file_stamp[0]) != file_stamp:
            raise ValueError('The workbook changed while it was being read. Refresh Graph.')
        return [{'cycle':k[0], 'committee':k[1], 'candidate_id':k[2], 'direction':k[3],
                 'candidate':labels[k][0], 'party':labels[k][1], 'amount':v,
                 'filings':filings[k],
                 'unverified_version_rows':unverified[k], 'unresolved_version_rows':unresolved[k]} for k,v in groups.items()]
    except (ET.ParseError, IndexError, KeyError, zipfile.BadZipFile) as error:
        raise ValueError('Saved downstream activity could not be read.') from error


def report_activity(file_stamp, include_filings=False):
    return _report_activity((file_stamp, include_filings))


report_activity.cache_clear = _report_activity.cache_clear
report_activity.cache_info = _report_activity.cache_info


def estimate_receipt(path):
    """Require the full hash-bound CSV, never the 10,000-row worksheet preview."""
    path = Path(path)
    manifest_path = path.with_suffix('.audit') / 'audit-manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
    if manifest.get('workbook_name') != path.name or manifest.get('workbook_sha256') != digest(path):
        raise ValueError('The estimate audit bundle does not match this workbook.')
    receipt = local_reference('allocation_estimates.csv', manifest_path.parent)
    if manifest.get('csv_sha256', {}).get(receipt.name) != digest(receipt):
        raise ValueError('The estimate receipt is missing or changed.')
    return stamp(receipt)


@lru_cache(maxsize=4)
def estimate_rows(file_stamp):
    groups = defaultdict(Decimal)
    labels = {}
    with Path(file_stamp[0]).open(encoding='utf-8-sig',newline='') as stream:
        for i, row in enumerate(csv.DictReader(stream),1):
            if i > MAX_SOURCE_ROWS:
                raise ValueError('Estimate receipt exceeds the local graph reading limit.')
            identity = (cycle(row.get('cycle')), row.get('root',''), row.get('candidate_id',''), row.get('direction','Unknown'))
            if not identity[2]:
                raise ValueError('Estimate receipt lacks candidate identifiers.')
            groups[identity] += amount(row.get('estimated_amount'))
            labels[identity] = row.get('candidate_name') or identity[2]
    return [{'cycle':k[0],'committee':k[1],'candidate_id':k[2],'direction':k[3],
             'candidate':labels[k],'party':'','amount':v} for k,v in groups.items()]


def grouped(rows, id_field, label_field=None):
    totals, labels = defaultdict(Decimal), {}
    for row in rows:
        identifier = row[id_field]
        if not identifier and id_field == 'candidate_id':
            identifier = 'unidentified-candidate'
        if not identifier:
            continue
        totals[identifier] += row['amount']
        labels[identifier] = str(row.get(label_field or id_field) or identifier)
    return [{'id':identifier,'label':labels[identifier],'amount':money(value)}
            for identifier,value in sorted(totals.items(), key=lambda item:(-item[1],str(item[0]))) if value != 0]


def levels_by_company(rows):
    """Each company's own non-Corporate giving split by level, for the "By
    company" bars under Giving by employee level (board M4). The rows are the
    ones `levels` is grouped from, so Corporate never enters; a company's
    amount is its non-Corporate total, unlike companies[].amount, which counts
    Corporate giving when the scope is all giving. Every company lists the
    legend's levels in one order, a missing one at 0, then any other level it
    has (Unknown), so its shares always read in the same place."""
    totals = defaultdict(lambda: defaultdict(Decimal))
    for row in rows:
        totals[row['company']][row['level']] += row['amount']
    result = []
    for company, levels in totals.items():
        order = [*LEVEL_ORDER, *sorted(level for level in levels if level not in LEVEL_ORDER)]
        result.append({'company':company, 'amount':money(sum(levels.values(), Decimal(0))),
                       'levels':[{'id':level,'label':level,'amount':money(levels.get(level, Decimal(0)))}
                                 for level in order]})
    return sorted(result, key=lambda item:(-item['amount'], item['company']))


_views = ViewCache()


def graph_view(paths, filters):
    paths = tuple(map(Path, paths))
    # Saved estimate bundles have their own manifest/hash protocol. Keep that
    # gate on every estimate read rather than treating a view cache as proof.
    if filters.get('mode') == 'estimates':
        return _graph_view(paths, filters)
    def dependencies():
        snapshots = list(downstream.ROOT.glob('*/snapshots.json')) if filters.get('mode') == 'activity' else []
        full_sources = [p for path in paths for p in transfer_source(path) if p is not None] if filters.get('mode') == 'activity' else []
        return file_signature([*paths, *snapshots, *full_sources])
    before = dependencies()
    selection = tuple((k,str(filters.get(k) or '')) for k in
                      ('company','tier','cycle','category','recipient','latest','funding','mode','direction'))
    def build():
        value = _graph_view(paths, filters)
        if dependencies() != before:
            raise ValueError('A source changed while preparing Graph. Refresh.')
        return value
    # Keep order and multiplicity: [A, A] is invalid overlapping evidence,
    # not the same request as [A], even though both have the same file set.
    return _views.get((tuple(str(path.resolve()) for path in paths), before, selection), build)


def _graph_view(paths, filters):
    stamps = [stamp(path) for path in paths]
    reports = [report_facts(version) for version in stamps]
    years = [year for report in reports for year in report['cycles']]
    if len(years) != len(set(years)):
        raise ValueError('These reports overlap election cycles. Choose a run with one report per industry/cycle.')
    facts = [r for report in reports for r in report['facts']]
    companies = sorted({c for report in reports for c in report['companies']})
    selected = {k:str(filters.get(k) or '') for k in ('company','tier','cycle','category','recipient')}
    if filters.get('latest') == '1' and years:
        selected['cycle'] = max(years)
    funding = str(filters.get('funding') or 'all')
    mode = str(filters.get('mode') or 'direct')
    direction = str(filters.get('direction') or 'Support')
    if funding not in ('all','personal','corporate') or mode not in ('direct','activity','estimates'):
        raise ValueError('Choose a valid giving scope and candidate view.')
    if selected['company'] and selected['company'] not in companies:
        raise ValueError('Choose a company in this industry.')
    if selected['cycle'] and selected['cycle'] not in years:
        raise ValueError('Choose an election cycle in this report.')
    if selected['tier'] and selected['tier'] not in TIERS.values() and selected['tier'] != 'Unknown':
        raise ValueError('Choose Chief, Middle-Management or Employees.')
    if selected['tier'] and funding == 'corporate':
        raise ValueError('Employee levels apply to personal giving, not corporate giving.')
    scope = [r for r in facts if (not selected['company'] or r['company'] == selected['company'])
             and (not selected['cycle'] or r['cycle'] == selected['cycle'])
             and (funding == 'all' or (r['tier'] == 'Corporate') == (funding == 'corporate'))
             and (not selected['tier'] or (r['tier'] != 'Corporate' and TIERS.get(r['tier'],'Unknown') == selected['tier']))]
    categories = sorted({r['category'] for r in scope})
    recipients = {r['recipient_id']:r['recipient'] for r in scope if not selected['category'] or r['category'] == selected['category']}
    filtered = [r for r in scope if (not selected['category'] or r['category'] == selected['category'])
                and (not selected['recipient'] or r['recipient_id'] == selected['recipient'])]
    direct = [r for r in filtered if r['candidate_id']]
    organizations = [r for r in filtered if r['category'] != 'Candidate']
    employee_rows = [{**r, 'level':TIERS.get(r['tier'], 'Unknown')}
                     for r in filtered if r['tier'] != 'Corporate']
    unresolved = sum((r['amount'] for r in filtered if r['category'] == 'Candidate' and not r['candidate_id']), Decimal(0))
    candidate_rows = direct
    buckets = None
    available, note = True, 'Direct contributions to uniquely linked principal or authorized candidate committees.'
    ties = []
    if mode != 'direct':
        candidate_rows = []
        try:
            if mode == 'estimates' and (selected['company'] or selected['tier'] or funding != 'all'):
                raise ValueError('Saved estimates cover the whole industry. Select all companies, all levels and all giving to view them.')
            funded = defaultdict(Decimal)
            for r in filtered:
                funded[r['cycle'],r['recipient_id']] += r['amount']
            eligible = {k for k,v in funded.items() if v > 0}
            if direction not in ACTIVITY_DIRECTIONS:
                raise ValueError('Choose a recorded activity direction.')
            buckets = {name: [] for name in ACTIVITY_DIRECTIONS}
            network_keys = eligible
            activity_rows = None
            if mode == 'activity':
                activity_rows, network = downstream.merge_activity(
                    [r for version in stamps for r in report_activity(version)], eligible,
                    report_edges=[r for version in stamps for r in report_transfers(version)])
                network_keys = network['keys']
            rows = (activity_rows if mode == 'activity' else
                    [r for path in paths for r in estimate_rows(estimate_receipt(path))])
            for r in rows:
                if (r['cycle'],r['committee']) in network_keys and r['amount'] != 0 and r['direction'] in buckets:
                    buckets[r['direction']].append(r)
            candidate_rows = buckets[direction]
            note = ('Saved observed activity across reachable committees, counted once per committee/cycle. Routes and activity are limited to saved report coverage. This is not money attributed to the selected company or employee level; company views can share the same PAC activity.'
                    if mode == 'activity' else 'Saved conditional industry estimates, not traced company dollars. Support, opposition and contribution/other activity remain separate.')
            if mode == 'activity' and network['units']:
                note = downstream.NOTE + f" {len(network['units'])} saved committee/cycles; {len(network['missing'])} without a saved snapshot. Snapshots supply whole schedule partitions; workbook activity is used only for committees without snapshots."
                if network.get('measurement_comparisons'):
                    note += ' Snapshot and workbook measurements differ; their amounts are shown separately below and are not reconciled.'
            if mode == 'activity' and not activity_rows and network['missing']:
                raise ValueError('Downstream activity is unavailable: no saved activity for the selected committee/cycles. Missing evidence is not zero spending.')
            if mode == 'activity':
                ties = downstream.relationships(activity_rows, eligible, network)
                uncertain = sum(r.get('unverified_version_rows',0) for r in activity_rows)
                unresolved_versions = sum(r.get('unresolved_version_rows',0) for r in activity_rows)
                if uncertain or unresolved_versions:
                    note += f' Filing relationships: {uncertain} unverified observations retained; {unresolved_versions} unresolved observations in the saved source. Amounts are provisional, not certified unique spending.'
        except (OSError, ValueError, KeyError, AttributeError) as error:
            # The partial buckets go with it: a refused read has no sides to
            # report, and four rows of '0 candidates' would be read as a
            # measurement rather than as the absence of one.
            available, note, buckets = False, str(error), None
    def candidate_groups(rows):
        result = grouped(rows,'candidate_id','candidate')
        parties = defaultdict(set)
        for row in rows:
            if row.get('party'):
                parties[row['candidate_id']].add(row['party'])
        for candidate in result:
            candidate['party'] = ' / '.join(sorted(parties[candidate['id']]))
        return result
    chart_candidates = candidate_groups(candidate_rows) if available else []
    if any(stamp(path) != version for path,version in zip(paths,stamps)):
        raise ValueError('A report changed while building the graph. Refresh Graph.')
    return {'filters':{**selected,'funding':funding,'mode':mode,'direction':direction},
            'options':{'companies':companies,'cycles':sorted(set(years)), 'categories':categories,
                       'recipients':[{'id':i,'label':n} for i,n in sorted(recipients.items(),key=lambda p:p[1])],
                       'unknown_tiers':any(r['tier'] not in TIERS and r['tier'] != 'Corporate' for r in facts)},
            'totals':{'giving':money(sum((r['amount'] for r in filtered),Decimal(0))),
                      'direct':money(sum((r['amount'] for r in direct),Decimal(0))),
                      'organizations':money(sum((r['amount'] for r in organizations),Decimal(0))),
                      'unresolved_candidate_links':money(unresolved),
                      'recipients':len({r['recipient_id'] for r in filtered}), 'receipt_groups':len(filtered)},
            'categories':grouped(filtered,'category'), 'organizations':grouped(organizations,'recipient_id','recipient'),
            'companies':grouped(filtered,'company'),
            'levels':grouped(employee_rows,'level'),
            'levels_by_company':levels_by_company(employee_rows),
            'candidates':chart_candidates, 'candidate_available':available, 'candidate_note':note,
            'downstream_edges':network['edges'] if mode == 'activity' and available else [],
            'pac_candidate_relationships':ties,
            'downstream_sources':downstream.sources(network) if mode == 'activity' and available else [],
            'measurement_comparisons':network.get('measurement_comparisons',[]) if mode == 'activity' and available else [],
            'candidates_by_direction':{name:candidate_groups(rows) for name,rows in buckets.items()}
                if mode == 'activity' and buckets is not None else {},
            **({'direction_totals':[
                {'direction':name,
                 'candidates':len([r for r in grouped(buckets.get(name) or [],'candidate_id','candidate') if r['id'] != 'unidentified-candidate']),
                 'unidentified_amount':money(sum((r['amount'] for r in buckets.get(name) or [] if not r['candidate_id']),Decimal(0))),
                 'amount':money(sum((r['amount'] for r in buckets.get(name) or []),Decimal(0)))}
                for name in ACTIVITY_DIRECTIONS]} if mode == 'activity' and buckets is not None
              else {'direction_totals':[]} if mode == 'activity' else {}),
            'candidate_details':[{'candidate_id':r['candidate_id'],'candidate':r['candidate'],
                                 'company':r.get('company',''), 'tier':TIERS.get(r.get('tier'),r.get('tier','')),
                                 'recipient':r.get('recipient',r.get('committee','')), 'cycle':r['cycle'], 'amount':money(r['amount'])}
                                 for r in candidate_rows] if available else [],
            'reports':[{'name':r['report'],'sha256':r['sha256']} for r in reports],
            'note':'Included net giving from saved report receipts; corporate and personal giving remain distinct. Totals may include unverified attribution or receipt roles and are not certified unique giving. Levels follow existing occupation/title rules, not a verified company hierarchy. Employees includes unrecognized titles. Amounts retain the saved report’s rules; code changes do not rewrite historical workbooks.'}
