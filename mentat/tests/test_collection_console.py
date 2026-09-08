"""Offline upload, launch isolation, telemetry, and HTTP boundary checks."""
import base64
import importlib.util
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch, MagicMock
from urllib.request import Request, urlopen
from urllib.error import HTTPError

APP = Path(__file__).resolve().parents[1] / "app"
spec = importlib.util.spec_from_file_location("collection_console", APP / "server.py")
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


class ConsoleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.key_patch = patch('fec_key_store.key_path', return_value=self.root / 'test-keys.dpapi')
        self.key_patch.start()
        self.addCleanup(self.key_patch.stop)
        self.environment_patch = patch.object(app, 'environment_keys', return_value=[])
        self.environment_patch.start()
        self.addCleanup(self.environment_patch.stop)
        original_read = app.read_json
        source_patch = patch.object(app, 'read_json', side_effect=lambda path:
            {} if Path(path).name == 'active_fec_sources.json' else original_read(path))
        source_patch.start()
        self.addCleanup(source_patch.stop)
        self.console = app.Console(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def upload(self, text="canonical_name,alias,industry\nAcme,Acme Inc,Test\nAcme,ACME CORP,Test\nOther,Other LLC,Test\n"):
        return self.console.upload({"filename":"test.csv", "content":base64.b64encode(text.encode()).decode()})

    def test_repeated_canonical_rows_group_aliases(self):
        value = self.upload()
        self.assertEqual(value["state"], "ready")
        industry = value["readiness"]["readiness"][0]
        self.assertEqual(industry["company_count"], 2)
        self.assertEqual(industry["companies"][0]["variant_count"], 3)

    def test_cross_company_alias_collision_is_not_ready(self):
        value = self.upload("canonical_name,alias\nOne,SHARED\nTwo,SHARED\n")
        self.assertEqual(value["state"], "needs review")
        with self.assertRaises(ValueError):
            self.console.launch({"id":value["id"],"cycles":[2026]})

    def test_crypto_keeps_one_company_per_row(self):
        value = self.upload("canonical_name,alias,industry\nOne,One LLC,Crypto\nOne,One Inc,Crypto\nTwo,Two Inc,Crypto\n")
        self.assertEqual(value["readiness"]["readiness"][0]["company_count"], 2)

    def test_empty_name_rejected(self):
        with self.assertRaisesRegex(ValueError, "canonical_name is required"):
            self.upload("canonical_name,alias\n,hello\n")

    def test_existing_workbook_roundtrip_preserves_roster(self):
        first = self.upload()
        content = (self.root / first["id"] / "roster.xlsx").read_bytes()
        second = self.console.upload({"filename":"roster.xlsx","content":base64.b64encode(content).decode()})
        self.assertEqual(first["readiness"]["readiness"], second["readiness"]["readiness"])

    def test_launch_pins_uploaded_roster_and_keeps_key_out_of_arguments_and_disk(self):
        value = self.upload()
        runner = MagicMock()
        runner.pid = 987654
        runner.poll.return_value = None
        with patch.object(app.subprocess,"Popen", return_value=runner) as launch, patch.object(app.threading,"Thread"):
            self.console.launch({"id":value["id"],"cycles":[2026],"api_key":"private-test-key"})
        command = launch.call_args.args[0]
        self.assertNotIn("private-test-key", repr(command))
        self.assertEqual(launch.call_args.kwargs["env"]["FEC_API_KEYS"], "private-test-key")
        self.assertNotIn(b'private-test-key', (self.root / 'test-keys.dpapi').read_bytes())
        self.assertEqual(app.load_keys(), ['private-test-key'])
        directory = self.root / value["id"]
        self.assertNotIn("private-test-key", (directory / "upload.json").read_text())
        runtime = directory / "runtime/mentat"
        self.assertEqual((runtime / "inputs/Company Name.xlsx").read_bytes(), (directory / "roster.xlsx").read_bytes())
        self.assertEqual(Path(launch.call_args.kwargs["env"]["MENTAT_INPUT_DIR"]), runtime / "inputs")
        with self.assertRaises(ValueError):
            self.console.launch({"id":value["id"],"cycles":[2026],"api_key":"private-test-key"})

    def test_invalid_cycle_and_identifier_cannot_launch(self):
        with self.assertRaises(ValueError):
            self.console.launch({"id":"../outside", "cycles":[2026]})
        value = self.upload()
        with self.assertRaises(ValueError):
            self.console.launch({"id":value["id"], "cycles":[2028]})

    @unittest.skipUnless(Path("C:/Program Files/R/R-4.6.0/bin/Rscript.exe").is_file(), "R integration requires installed R 4.6")
    def test_runner_plan_and_offline_R_ingest_use_uploaded_aliases(self):
        value = self.upload()
        runner = MagicMock()
        runner.pid = 987654
        with patch.object(app.subprocess, "Popen", return_value=runner) as launch, patch.object(app.threading,"Thread"):
            self.console.launch({"id":value["id"],"cycles":[2026],"api_key":"offline-plan-placeholder"})
        kwargs = launch.call_args.kwargs
        plan = subprocess.run(launch.call_args.args[0] + ["-PlanOnly"], cwd=kwargs["cwd"],
                              env=kwargs["env"], capture_output=True, text=True, timeout=60)
        self.assertEqual(plan.returncode, 0, plan.stdout + plan.stderr)
        self.assertEqual(json.loads(plan.stdout)["discovered_industries"], ["Test"])
        runtime = Path(kwargs["cwd"])
        data_root = self.root / "offline-ingest"
        result = subprocess.run([self.console.rscript, str(runtime / "mentat.R"),
                                 "--industry=Test", "--cycles=2026", "--skip-scrape", "--defer-report-files",
                                 "--data-root=" + str(data_root)], cwd=runtime, env=kwargs["env"],
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        connection = sqlite3.connect(data_root / "mentat.db")
        try:
            companies = connection.execute("SELECT DISTINCT company FROM companies ORDER BY company").fetchall()
            aliases = connection.execute("SELECT variant FROM companies WHERE company='Acme'").fetchall()
            self.assertEqual(companies, [("Acme",), ("Other",)])
            self.assertEqual({row[0] for row in aliases}, {"Acme", "Acme Inc", "ACME CORP"})
        finally:
            connection.close()

    def test_live_read_uses_database_without_modifying_it(self):
        database = self.root / "mentat.db"
        with sqlite3.connect(database) as connection:
            connection.execute("CREATE TABLE api_work_progress(company,cycle,kind,rows_seen,records_total,fetched_complete,stored_complete,updated_at)")
            connection.execute("INSERT INTO api_work_progress VALUES ('Acme',2026,'employer',10,20,0,0,'now')")
            connection.execute("CREATE TABLE contributions(sub_id,contributor_name,company,matched_variant,contributor_employer,committee_name,contribution_receipt_date,contribution_receipt_amount,company_match_status)")
            connection.execute("INSERT INTO contributions VALUES ('123','Sample','Acme','ACME INC','Acme Inc','Committee','2026-01-01',100,'accepted')")
        connection.close()
        before = database.read_bytes()
        view = app.lane_preview(database)
        self.assertTrue(view["available"])
        self.assertEqual(view["progress"][0]["rows_seen"], 10)
        self.assertEqual(view["records"][0]["sub_id"], "123")
        self.assertEqual(database.read_bytes(), before)

    def test_missing_database_has_no_fake_zero(self):
        self.assertEqual(app.lane_preview(self.root / "absent.db"), {"available":False})

    def test_report_inspection_detects_missing_directions_and_changed_file(self):
        import openpyxl
        path = self.root / "report.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Beneficiary Summary"
        sheet.cell(5, 1, "Downstream Activity")
        book.save(path)
        first = self.console.inspect_report(path)
        self.assertFalse(first['direction_columns'])
        sheet.cell(5, 2, "Support Activity")
        sheet.cell(5, 3, "Oppose Activity")
        book.save(path)
        book.close()
        second = self.console.inspect_report(path)
        self.assertTrue(second['direction_columns'])
        self.assertNotEqual(first['sha256'], second['sha256'])

    def test_attributed_report_warnings_use_named_columns(self):
        import openpyxl
        path = self.root / 'attributed.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Candidate Funding'
        sheet.cell(5, 1, 'Conditional Non-opposition')
        sheet.cell(5, 2, 'Conditional Opposition')
        quality = book.create_sheet('Data Quality')
        for column, value in enumerate(('Industry', 'Check', 'Severity', 'Passed', 'Detail'), 1):
            quality.cell(5, column, value)
        for column, value in enumerate(('Fixture', 'coverage', 'warn', False, 'Unverified source'), 1):
            quality.cell(6, column, value)
        book.save(path)
        book.close()
        result = self.console.inspect_report(path)
        self.assertTrue(result['direction_columns'])
        self.assertEqual(result['warnings'], [{'check': 'coverage', 'severity': 'warn', 'detail': 'Unverified source'}])

    def test_saved_run_reports_failed_coverage_separately_from_completion(self):
        upload = self.upload()
        directory = self.root / upload['id'] / 'runs' / 'test-run'
        directory.mkdir(parents=True)
        coverage = directory / 'coverage.json'
        app.atomic_json(coverage, {'status':'failed','checks':{'company_cycle_grid':False}})
        receipt = directory / 'fanout.json'
        app.atomic_json(receipt, {'task_count':3,'succeeded':2,'failed':1})
        report = self.root / 'report.xlsx'
        import openpyxl
        book = openpyxl.Workbook()
        book.save(report)
        book.close()
        expected = self.console.inspect_report(report)['sha256'].upper()
        app.atomic_json(directory / 'manifest.json', {'run_id':'test-run'})
        app.atomic_json(directory / 'status.json', {'stage':'completed','industries':[
            {'industry':'Test','status':'succeeded','coverage_receipt':str(coverage),
             'fanout_receipt':str(receipt),'output_workbook':str(report),'workbook_sha256':expected}]})
        result = self.console.snapshot(upload['id'])
        self.assertEqual(result['assurance'][0]['collection'],'failed')
        self.assertTrue(result['reports'][0]['quality']['matches_run'])
        self.assertEqual(result['completed_collection'], {'total':3,'completed':3,'failed':1,'running':0})

    def test_http_rejects_untrusted_posts_and_hosts(self):
        server = app.ThreadingHTTPServer(("127.0.0.1",0),app.Handler)
        server.app = self.console
        worker = threading.Thread(target=server.serve_forever,daemon=True)
        worker.start()
        origin = f"http://127.0.0.1:{server.server_port}"
        try:
            with urlopen(origin + "/api/session") as response:
                self.assertEqual(json.load(response)["token"],self.console.token)
            for headers in ({}, {"X-Mentat-Token":self.console.token,"Origin":"https://untrusted.example"}):
                with self.assertRaises(HTTPError) as error:
                    urlopen(Request(origin + "/api/start", data=b"{}", headers=headers))
                self.assertEqual(error.exception.code,403)
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(origin + "/api/session",headers={"Host":"untrusted.example"}))
            self.assertEqual(error.exception.code,403)
        finally:
            server.shutdown()
            server.server_close()
            worker.join()

    def test_review_is_bound_to_workbook_and_warning_and_preserves_evidence(self):
        directory = self.root / 'fixture'
        directory.mkdir()
        workbook = self.root / 'report.xlsx'
        workbook.write_bytes(b'fixture')
        app.atomic_json(directory / 'status.json', {'industries': [
            {'industry': 'Test', 'status': 'succeeded', 'output_workbook': str(workbook)}]})
        run = {'id': 'test', 'name': 'Fixture', 'state': 'completed', 'directory': directory}
        warning = {'check': 'coverage', 'severity': 'warn', 'detail': 'Old source'}
        quality = {'sha256': 'first', 'warnings': [warning]}
        with patch.object(self.console, 'find', return_value=run), patch.object(self.console, 'inspect_report', return_value=quality):
            first = self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]
            self.console.acknowledge_quality({'id': 'test', 'warning_ids': [first['id']]})
            reviewed = self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]
            self.assertTrue(reviewed['reviewed_at'])
            self.assertEqual(reviewed['detail'], 'Old source')
            self.assertEqual(workbook.read_bytes(), b'fixture')
            # Repeated polling does not mutate cached warnings or their IDs.
            self.assertEqual(first['id'], reviewed['id'])
            self.assertEqual(quality['warnings'], [warning])
            quality['sha256'] = 'new-workbook'
            changed = self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]
            self.assertIsNone(changed['reviewed_at'])
            with self.assertRaisesRegex(ValueError, 'changed'):
                self.console.acknowledge_quality({'id': 'test', 'warning_ids': [first['id']]})
            quality['sha256'] = 'first'
            quality['warnings'] = [{**warning, 'detail': 'New finding'}]
            self.assertIsNone(self.console.snapshot('test')['reports'][0]['quality']['warnings'][0]['reviewed_at'])

    def test_failed_check_cannot_be_acknowledged_away(self):
        with patch.object(self.console, 'snapshot', return_value={'reports': [
            {'quality': {'warnings': [{'id': 'hard-fail', 'severity': 'fail'}]}}]}):
            with self.assertRaises(ValueError):
                self.console.acknowledge_quality({'id': 'test', 'warning_ids': ['hard-fail']})


if __name__ == "__main__":
    unittest.main()
