"""Temporary local-only report server for test_console_browser.cjs."""
import json
from pathlib import Path
import sys
import tempfile
import threading

import openpyxl
from graph_fixture import add_graph_sheets

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
import server as app


class FixtureConsole(app.Console):
    def __init__(self, root):
        self.root = root
        self.token = 'offline-fixture-token'
        self.report_cache = {}
        self.processes = {}
        self.lock = threading.Lock()

    def runs(self):
        return [{'id':'fixture','name':'Local regression fixture','state':'completed',
                 'created_at':1,'directory':self.root}]

    def summary(self):
        return {'runs':[{'id':'fixture','name':'Local regression fixture','state':'completed'}],
                'key_configured':False,'active_key_count':0,'stored_key_count':0}


def main():
    with tempfile.TemporaryDirectory(prefix='mentat-chart-test-') as temporary:
        root = Path(temporary)
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Candidate Funding'
        for row in [['Fixture report'],['Temporary values for UI tests'],[],[],
                    ['Company','Amount','Other Amount','Cycle','Date'],
                    ['A',10,100,2026,'2026-01-01'],['A',30,120,2026,'2026-01-02'],
                    ['B',60,200,2024,'2026-01-03'],['B',None,300,2024,'2026-01-04'],
                    ['<img src=x onerror=alert(1)>',-5,0,2026,'2026-01-05']]:
            sheet.append(row)
        receipts = book.create_sheet('Source Freshness')
        receipts.append(['FEC Dataset','Cycle','Downloaded At','Checked At','Verified Data Through'])
        receipts.append(['cm',2026,'2026-08-01T00:00:00Z','2026-09-08T00:00:00Z','2026-07-31'])
        receipts.append(['oth',2026,None,'2026-09-08T00:00:00Z',None])
        regions = book.create_sheet('Regional sample')
        for row in [['Region','Volume','Cost'], ['North',10,20], ['South',15,30], ['East',25,60]]:
            regions.append(row)
        add_graph_sheets(book)
        report = root / 'Fixture.xlsx'
        book.save(report)
        book.close()
        (root / 'status.json').write_text(json.dumps({'stage':'completed','industries':[
            {'industry':'Fixture','status':'succeeded','output_workbook':str(report)}]}))
        server = app.ThreadingHTTPServer(('127.0.0.1',0), app.Handler)
        server.app = FixtureConsole(root)
        worker = threading.Thread(target=server.serve_forever,daemon=True)
        worker.start()
        print(f'http://127.0.0.1:{server.server_port}',flush=True)
        try:
            sys.stdin.readline()
        finally:
            server.shutdown()
            server.server_close()
            worker.join()


if __name__ == '__main__':
    main()
