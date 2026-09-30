"""Loader validation, then full synthetic end-to-end runs (all workbooks synthetic) with Glass Box emission."""
import json
import math
import os
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

import synth_workbooks
from sgs import glassbox, load, panel, registry, run_all
from sgs.load import LoadError, SheetSpec

N_TESTS = len(registry.REGISTRY)
H5_IDS = [t.id for t in registry.REGISTRY if t.family == "H5"]


def _xlsx(path, header, rows, sheet="data"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


class TestLoader(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.spec = {"data": SheetSpec(("year", "value"), ("study_sector", "year"), study_sector=True)}

    def tearDown(self):
        self.tmp.cleanup()

    def test_rejects_formula_cell(self):
        p = self.dir / "f.xlsx"
        _xlsx(p, ["study_sector", "year", "value"], [["crypto", 2021, 1.0], ["crypto", 2022, "=B2*2"]])
        with self.assertRaises(LoadError) as cm:
            load.read_workbook(p, self.spec)
        self.assertIn("f.xlsx [data] cell C3", str(cm.exception))

    def test_rejects_duplicate_key(self):
        p = self.dir / "d.xlsx"
        _xlsx(p, ["study_sector", "year", "value"], [["crypto", 2021, 1.0], ["crypto", 2021, 2.0]])
        with self.assertRaises(LoadError) as cm:
            load.read_workbook(p, self.spec)
        self.assertIn("d.xlsx [data] key ['study_sector', 'year']", str(cm.exception))

    def test_rejects_missing_column(self):
        p = self.dir / "m.xlsx"
        _xlsx(p, ["study_sector", "year"], [["crypto", 2021]])
        with self.assertRaises(LoadError) as cm:
            load.read_workbook(p, self.spec)
        self.assertIn("m.xlsx [data] column value", str(cm.exception))

    def test_requires_study_sector(self):
        # uniform column name: a sheet that still names the study sector `sector` is rejected, naming the rename
        p = self.dir / "s.xlsx"
        _xlsx(p, ["sector", "year", "value"], [["crypto", 2021, 1.0], ["fintech", 2021, 2.0]])
        with self.assertRaises(LoadError) as cm:
            load.read_workbook(p, self.spec)
        self.assertIn("s.xlsx [data] column study_sector: required column missing", str(cm.exception))
        self.assertIn("`sector` column is present", str(cm.exception))
        q = self.dir / "t.xlsx"  # both present (06/07 keep the source label in `sector`): study_sector is used
        _xlsx(q, ["sector", "study_sector", "year", "value"], [["life_insurance_annuities", "life_insurance", 2021, 1.0]])
        wb = load.read_workbook(q, self.spec)
        self.assertEqual(wb.frames["data"]["study_sector"].tolist(), ["life_insurance"])
        self.assertEqual(len(wb.sha256), 64)


class TestSyntheticEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        base = Path(cls.tmp.name)
        cls.data = synth_workbooks.build(base / "data")  # no 05 in this folder: H5 is NOT BUILT
        cls.out = run_all.run(data_dir=cls.data, results_dir=base / "results", B=999, run_id="analysis-test",
                              log=lambda *a: None)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_panel_shapes_and_statuses(self):
        sy = self.out["tables"]["sector_year"]
        self.assertEqual(len(sy), 9 * 21)
        tel = sy[sy["study_sector"] == "life_insurance"]
        self.assertTrue((tel["contributions_status"] == "not_applicable").all())
        self.assertTrue(sy.loc[sy["year"] == 2026, "contributions_real_m"].isna().all())
        syt = self.out["tables"]["sector_year_tier"]
        self.assertEqual(len(syt), 9 * 21 * 4)
        lv = sy[sy["lobby_amount_undiff_real_m"].notna()]
        self.assertEqual(sorted(lv["study_sector"].unique()), ["commercial_banks", "credit_unions"])

    def test_core_test_estimable_and_recovers_slope(self):
        rows = {r["test_id"]: r for r in self.out["est_rows"]}
        self.assertEqual(rows["C1"]["status"], "ESTIMABLE")
        self.assertEqual(rows["C1"]["window_used"], "2020-2025")
        # 7 year-level parameters need the fallback: complete 2008-2025, most recent 10 years
        self.assertEqual((rows["L1-income"]["window_used"], rows["L1-income"]["window_kind"]), ("2016-2025", "fallback"))
        self.assertIn("complete-case years in 2008-2025: 2008-2025", rows["L1-income"]["fallback_search"])
        c1 = [c for c in self.out["coef_rows"] if c["test_id"] == "C1"]
        self.assertEqual([c["coefficient"] for c in c1], ["growth_previous_year_pct"])
        self.assertLess(abs(c1[0]["estimate"] - 0.3), 0.05)
        self.assertTrue(c1[0]["stratification"].startswith("STRATIFIED") or c1[0]["stratification"].startswith("UNSTRATIFIED"))
        h1 = [c for c in self.out["coef_rows"] if c["test_id"] == "H1-C1"]
        self.assertEqual([c["coefficient"] for c in h1], ["pres_year", "midterm_year"])
        self.assertEqual(rows["DS-House"]["status"], "NOT BUILT")
        self.assertEqual(rows["H4seat-C1"]["status"], "NOT ESTIMABLE")  # election years only: never 10 consecutive

    def test_lobbyview_tests_descriptive_in_fallback(self):
        rows = {r["test_id"]: r for r in self.out["est_rows"]}
        for k in "123":
            r = rows[f"L{k}-LV"]
            # LobbyView observed 2006-2023 (2024 partial): complete 2008-2023 -> most recent 10 = 2014-2023
            self.assertEqual((r["status"], r["label"], r["window_used"], r["window_kind"]),
                             ("ESTIMABLE", "DESCRIPTIVE", "2014-2023", "fallback"))
            self.assertEqual(r["sectors"], "commercial_banks, credit_unions")
        lv = [c for c in self.out["coef_rows"] if c["test_id"] == "L1-LV"]
        self.assertEqual(len(lv), 1)
        self.assertLess(abs(lv[0]["estimate"] - 0.5), 0.05)
        self.assertTrue(math.isnan(lv[0]["boot_p"]))
        self.assertEqual(lv[0]["ci_text"], "n/a (inference not available)")
        self.assertEqual(lv[0]["boot_note"], "2 sectors: wild cluster bootstrap cannot reach p < 0.056; inference not available")

    def test_h5_not_built_without_05(self):
        rows = {r["test_id"]: r for r in self.out["est_rows"]}
        for tid in H5_IDS:
            self.assertEqual(rows[tid]["status"], "NOT BUILT", tid)
            self.assertIn("05_industry_champions.xlsx not found", rows[tid]["reason"])
        self.assertEqual(len(H5_IDS), 13)

    def test_outputs_written_in_study_format(self):
        xl = Path(self.out["xlsx"])
        self.assertTrue(xl.exists() and Path(self.out["report"]).exists() and Path(self.out["estimability"]).exists())
        wb = openpyxl.load_workbook(xl, read_only=True)
        self.assertEqual(wb.sheetnames, ["estimability", "coefficients", "strata", "samples", "sector_year",
                                         "sector_year_tier", "industry_year", "champions", "dictionary", "sources"])
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    self.assertFalse(isinstance(v, str) and v.startswith("="))
        wb.close()
        est = pd.read_excel(xl, sheet_name="estimability")
        self.assertEqual(len(est), N_TESTS)
        self.assertEqual(N_TESTS, 139)  # 129 + the 10 PREVIEW tests (C1-PV..C3-PV x 3 specifications, H1-PV)
        src = pd.read_excel(xl, sheet_name="sources")
        self.assertIn("05_industry_champions.xlsx", set(src["file"]))

    def test_glassbox_event_shape(self):
        with tempfile.TemporaryDirectory() as d:
            eid = glassbox.emit("run.start", "running", "start", "test", "analysis-test", events_dir=Path(d))
            files = os.listdir(d)
            self.assertEqual(files, [f"{eid}.json"])
            ev = json.loads((Path(d) / files[0]).read_text(encoding="utf-8"))
            self.assertEqual(list(ev), ["schema_version", "id", "timestamp_utc", "run_id", "correlation_id", "parent_id",
                                        "lane", "source", "kind", "status", "stage", "summary", "detail", "cwd",
                                        "duration_ms", "bytes", "metadata"])
            self.assertEqual((ev["lane"], ev["source"]), ("claude", "sector-growth-analysis"))


class TestSyntheticEndToEndWithChampions(unittest.TestCase):
    """A synthetic 05 with an old receipt: H5 is read at run time and estimated."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        base = Path(cls.tmp.name)
        data = synth_workbooks.build(base / "data")
        synth_workbooks.write_champions(data, synth_workbooks.champions_frame(slope=5000.0, noise=500.0),
                                        age_minutes=30)
        cls.out = run_all.run(data_dir=data, results_dir=base / "results", B=999, run_id="analysis-test-h5",
                              only=["H5", "C1"], log=lambda *a: None)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_h5_estimated_and_reported(self):
        rows = {r["test_id"]: r for r in self.out["est_rows"]}
        self.assertEqual(rows["H5-pooled"]["status"], "ESTIMABLE")
        self.assertEqual(rows["H5-within"]["status"], "ESTIMABLE")
        self.assertEqual(rows["H5-pooled"]["window_used"], "cycles 2022, 2024, 2026")
        self.assertEqual(rows["H5-pooled"]["dv_column"], "direct_giving_real_2025_usd")
        self.assertIn("employee channel only", rows["H5-pooled"]["flag"])
        self.assertEqual(rows["H5-pooled-no2026"]["status"], "NOT ESTIMABLE")
        self.assertEqual(rows["H5-pooled-trade"]["status"], "NOT BUILT")
        self.assertIn("outcome-contaminated: money affects margins", rows["H5-general_margin"]["flag"])
        coefs = {c["test_id"]: c for c in self.out["coef_rows"] if c["test_id"].startswith("H5")}
        self.assertLess(abs(coefs["H5-within"]["estimate"] - 5000.0), 600.0)
        self.assertEqual(coefs["H5-pooled"]["unit_of_analysis"], "board x champion x cycle")
        self.assertTrue(any(s.get("board") for s in self.out["strata_rows"]))
        md = Path(self.out["report"]).read_text(encoding="utf-8")
        self.assertIn("### H5-pooled (EXTENSION): ESTIMABLE", md)
        self.assertIn("stratum board bank", md)
        src = [s for s in self.out["sources"] if s["file"] == "05_industry_champions.xlsx"][0]
        self.assertEqual(len(src["sha256"]), 64)
        champs = self.out["tables"]["champions"]
        self.assertGreater(len(champs), 0)


class TestDemShareStub(unittest.TestCase):
    def test_not_built(self):
        with self.assertRaises(panel.NotBuilt):
            panel.build_dem_share()


if __name__ == "__main__":
    unittest.main()
