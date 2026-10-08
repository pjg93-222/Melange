# -*- coding: utf-8 -*-
"""Fail-closed identity/category checks for generated dashboard artifacts."""
import hashlib, json, os
import openpyxl
from paths import ROOT, WORKBOOK, p
import moneyball_core as MC
import stance_rules as SR
import demographic_potential as DP


class IntegrityError(RuntimeError):
    pass


OPTIONAL_DISTRICT_SOURCES = (
    'district_industry.json',
    'district_industry_4digit.json',
    'district_community.json',
    'district_employers.json',
    'district_financial.json',
    'district_hmda.json',
)


def _sha(obj):
    raw=json.dumps(obj,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode('utf-8')
    return hashlib.sha256(raw).hexdigest()


def _file_sha(path):
    """Hash file contents, not mtimes, so OneDrive/Excel touches are not data changes."""
    if not os.path.exists(path):
        return None
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):
            h.update(chunk)
    return h.hexdigest()


def _check(condition, message, errors):
    if not condition: errors.append(message)


def preflight(member_names, category_stance, dashboard_data, archive_leg):
    """Validate the identity and category joins before any desk artifact is written."""
    errors=[]; warnings=[]; roster=MC.members(); rset=set(roster); cats=category_stance.get('cats') or []
    _check(len(roster)==len(rset), 'canonical roster contains duplicate names', errors)
    _check(len(member_names)==len(set(member_names)), 'member desk source contains duplicate names', errors)
    _check(set(member_names)==rset, 'member desk source roster differs from canonical roster', errors)
    snames=dashboard_data.get('name') or []
    _check(len(snames)==len(set(snames)), 'dashboard_data contains duplicate names', errors)
    _check(set(snames)==rset, 'dashboard_data roster differs from canonical roster', errors)
    _check(set(category_stance.get('matrix') or {})==rset, 'stance matrix roster differs from canonical roster', errors)
    _check(set(category_stance.get('scores') or {})==rset, 'stance score roster differs from canonical roster', errors)
    _check(set(archive_leg).issubset(rset), 'legislation archive contains unknown/current-roster-mismatched names', errors)
    for name in roster:
        _check(set((category_stance.get('matrix') or {}).get(name,{}) or {})==set(cats),
               'stance categories missing/misaligned for '+name, errors)
        _check(set((category_stance.get('scores') or {}).get(name,{}) or {})==set(cats),
               'stance score categories missing/misaligned for '+name, errors)
    issues=dashboard_data.get('issues') or []
    stance_rows=dashboard_data.get('stance') or []
    _check(len(stance_rows)==len(snames), 'dashboard stance row count does not match dashboard names', errors)
    _check(all(len(s)==len(issues) for s in stance_rows), 'dashboard stance strings do not align to issue count', errors)

    wb=openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=True)
    ra=wb['Rank & Allocation']; wn=[ra.cell(r,1).value for r in range(3,536) if ra.cell(r,1).value]
    # Rows 2:534 are the 533-member roster; row 535 is the literal "Weights" footer.
    rw=wb['Press Release - Raw']; mn=[rw.cell(r,1).value for r in range(2,535) if rw.cell(r,1).value]
    wb.close()
    _check(len(wn)==len(set(wn)) and set(wn)==rset, 'Rank & Allocation roster is duplicated or misaligned', errors)
    _check(len(mn)==len(set(mn)) and set(mn)==rset, 'Press Release - Raw roster is duplicated or misaligned', errors)

    # General CU must be conserved exactly when distributed.
    alloc=((category_stance.get('media') or {}).get('general_credit_union') or {})
    bad_conservation=[]
    for name,a in alloc.items():
        if name not in rset or abs(float(a.get('general',0))-sum(float(x) for x in (a.get('to') or {}).values()))>1e-8:
            bad_conservation.append(name)
    _check(not bad_conservation, 'General CU allocation is not identity-safe/conserved: '+', '.join(bad_conservation[:5]), errors)
    if not alloc: warnings.append('no nonzero General Credit Union values were available to disperse')

    report={'status':'pass' if not errors else 'fail','errors':errors,'warnings':warnings,
            'counts':{'roster':len(roster),'categories':len(cats),'general_cu_members':len(alloc)},
            'roster_order_sha256':_sha(roster),'category_order_sha256':_sha(cats)}
    if errors: raise IntegrityError('; '.join(errors[:8]))
    return report


def postflight(bundle, preflight_report):
    """Validate the completed desk bundle and enforce deterministic regeneration."""
    errors=[]; warnings=[]; roster=MC.members(); rset=set(roster); cats=bundle.get('cats') or []
    members=bundle.get('members') or []; names=[m.get('name') for m in members]
    _check(len(names)==len(set(names)) and set(names)==rset, 'desk bundle roster is duplicated or misaligned', errors)
    allowed=set(cats)|set(bundle.get('art_extra_cats') or [])
    n2b=MC.name2bio(); social_members=set(); article_total=bill_total=0; missing_bio=[]
    for m in members:
        name=m.get('name'); article_total+=len(m.get('articles') or []); bill_total+=len(m.get('bills') or [])
        _check(len(m.get('stances') or '')==len(cats), 'desk stance string misaligned for '+str(name), errors)
        expected=n2b.get(name); actual=(m.get('bio') or {}).get('bioguide')
        if expected and actual: _check(actual==expected, 'bioguide mismatch for '+str(name), errors)
        elif expected and not actual: missing_bio.append(name)
        for a in m.get('articles') or []:
            _check(set(a.get('cats') or []).issubset(allowed), 'unknown media category for '+str(name), errors)
            if a.get('source')=='social': social_members.add(name)
        for b in m.get('bills') or []:
            _check(set(b.get('areas') or []).issubset(allowed), 'unknown bill category for '+str(name), errors)
    potential_meta=bundle.get('demographic_potential_meta') or {}
    potential_codes={row.get('sector') for row in potential_meta.get('eligible_sectors') or []
                     if isinstance(row,dict)}
    potential_lq_min=float(potential_meta.get('lq_min',1.25))
    potential_count=0
    potential_markets={}
    score_weights=potential_meta.get('score_weights') or {}
    _check(score_weights==DP.SCORE_WEIGHTS,
           'demographic-potential weights differ from the pinned formula',
           errors)
    _check(abs(sum(float(value) for value in score_weights.values())-1.0)<1e-9,
           'demographic-potential weights do not sum to 100%',errors)
    _check(potential_meta.get('score_version')==DP.SCORE_VERSION,
           'demographic-potential score version is not current',errors)
    _check(potential_meta.get('weight_basis')==DP.WEIGHT_BASIS and
           potential_meta.get('not_empirically_optimized') is True,
           'demographic-potential weight provenance is missing or changed',
           errors)
    component_definitions=potential_meta.get('component_definitions') or []
    _check([row.get('key') for row in component_definitions
            if isinstance(row,dict)]==list(DP.SCORE_WEIGHTS),
           'demographic-potential component formula catalog is incomplete',
           errors)
    source_catalog=potential_meta.get('source_catalog') or {}
    _check(source_catalog==DP.SOURCE_CATALOG and all(
        str((row or {}).get('official_url') or '').startswith('https://')
        for row in source_catalog.values()),
        'demographic-potential official source catalog is incomplete',errors)
    sensitivity_audit=potential_meta.get('sensitivity_audit') or {}
    _check(sensitivity_audit.get('version')==DP.SENSITIVITY_VERSION and
           sensitivity_audit.get('scenario_count')==9,
           'demographic-potential sensitivity audit is missing',errors)
    community_meta=bundle.get('community_meta') or {}
    _check(
        community_meta.get('schema')=='district-community-acs-v1' and
        community_meta.get('vintage')==2024 and
        community_meta.get('survey')=='ACS 1-year',
        'ACS community metadata is missing or not the pinned 2024 vintage',
        errors,
    )
    _check(
        set(community_meta.get('profile_families') or {})==
        {'DP03','DP04','DP05'} and
        all(str(url).startswith('https://')
            for url in (community_meta.get('official_urls') or {}).values()),
        'ACS community table/source catalog is incomplete',
        errors,
    )
    structural_meta=potential_meta.get('structural_adjustment') or {}
    _check(
        structural_meta.get('version')==
        DP.STRUCTURAL_ADJUSTMENT_VERSION and
        structural_meta.get('status')==
        'descriptive_companion_not_causal',
        'demographic-potential structural adjustment is missing',
        errors,
    )
    structurally_adjusted_markets={}
    for m in members:
        potential=m.get('demographic_potential') or {}
        if not potential.get('qualified'):
            _check(potential.get('score') is None,
                   'unqualified member has a demographic-potential score for '+
                   str(m.get('name')),errors)
            continue
        potential_count+=1
        industry=m.get('industry') or {}
        market=(industry.get('kind'),industry.get('geo'))
        prior=potential_markets.get(market)
        if prior is not None:
            _check(
                (prior.get('score'),prior.get('rank'))==
                (potential.get('score'),potential.get('rank')),
                'members in the same demographic market have different '
                'scores/ranks for '+str(market),errors)
            _check(
                (
                    (prior.get('structural_adjustment') or {}).get('score'),
                    (prior.get('structural_adjustment') or {}).get('rank'),
                ) == (
                    (potential.get('structural_adjustment') or {}).get(
                        'score'),
                    (potential.get('structural_adjustment') or {}).get(
                        'rank'),
                ),
                'members in the same demographic market have different '
                'structurally adjusted scores/ranks for '+str(market),
                errors,
            )
        else:
            potential_markets[market]=potential
        signals=potential.get('signals') or []
        _check(bool(signals),'qualified demographic potential has no signals for '+str(m.get('name')),errors)
        for signal in signals:
            code=signal.get('sector')
            lq=signal.get('lq')
            shift=signal.get('competitive_shift')
            _check(code in potential_codes,
                   'unknown demographic-potential sector for '+str(m.get('name')),errors)
            _check(isinstance(lq,(int,float)) and not isinstance(lq,bool) and lq>=potential_lq_min,
                   'demographic-potential LQ below threshold for '+str(m.get('name')),errors)
            _check(isinstance(shift,(int,float)) and not isinstance(shift,bool) and shift>0,
                   'demographic-potential shift is not positive for '+str(m.get('name')),errors)
            _check(signal.get('crosscheck') in ('passed','not_covered'),
                   'invalid demographic-potential cross-check for '+str(m.get('name')),errors)
        score=potential.get('score')
        components=potential.get('components') or {}
        available=potential.get('available') or {}
        coverage=potential.get('evidence_coverage')
        sensitivity=potential.get('sensitivity') or {}
        _check(isinstance(score,(int,float)) and not isinstance(score,bool) and 0<=score<=100,
               'invalid demographic-potential score for '+str(m.get('name')),errors)
        _check(set(components)==set(score_weights),
               'demographic-potential components do not match weights for '+
               str(m.get('name')),errors)
        _check(set(available)==set(score_weights),
               'demographic-potential availability does not match weights for '+
               str(m.get('name')),errors)
        _check(all(isinstance(value,(int,float)) and not isinstance(value,bool)
                   and 0<=value<=1 for value in components.values()),
               'demographic-potential component is outside [0,1] for '+
               str(m.get('name')),errors)
        if set(components)==set(score_weights) and isinstance(score,(int,float)):
            reconstructed=round(100*sum(
                float(score_weights[key])*float(components[key])
                for key in score_weights),1)
            _check(abs(reconstructed-float(score))<=0.01,
                   'demographic-potential score does not reconstruct for '+
                   str(m.get('name')),errors)
        _check(isinstance(coverage,(int,float)) and not isinstance(coverage,bool)
               and 0<=coverage<=1,
               'invalid demographic-potential evidence coverage for '+
               str(m.get('name')),errors)
        _check(isinstance(potential.get('rank'),int) and
               potential.get('rank')>=1,
               'invalid demographic-potential rank for '+str(m.get('name')),
               errors)
        _check(
            isinstance(sensitivity.get('rank_min'),int) and
            isinstance(sensitivity.get('rank_max'),int) and
            sensitivity['rank_min']<=potential.get('rank')<=
            sensitivity['rank_max'],
            'invalid demographic-potential sensitivity range for '+
            str(m.get('name')),errors)
        large=((m.get('employers') or {}).get('large') or {})
        if large.get('complete') is False:
            _check(
                available.get('large_employer_reach') is False,
                'incomplete CBP large-employer evidence was scored as '
                'available for '+str(m.get('name')),errors)
        _check(potential.get('tier') in {
            label for _,label in DP.TIER_THRESHOLDS},
            'invalid demographic-potential tier for '+str(m.get('name')),
            errors)
        structural=potential.get('structural_adjustment') or {}
        if structural.get('available'):
            structurally_adjusted_markets.setdefault(market,structural)
            community=m.get('community') or {}
            _check(
                community.get('complete') is True and
                community.get('vintage')==2024,
                'structurally adjusted member lacks complete 2024 ACS '
                'context for '+str(m.get('name')),errors)
            _check(
                structural.get('version')==
                DP.STRUCTURAL_ADJUSTMENT_VERSION and
                isinstance(structural.get('score'),(int,float)) and
                0<=structural.get('score')<=100 and
                isinstance(structural.get('rank'),int) and
                structural.get('rank')>=1,
                'invalid structurally adjusted score for '+
                str(m.get('name')),errors)
            controls=structural.get('controls') or {}
            _check(
                set(controls)==set(DP.STRUCTURAL_NUMERIC_CONTROLS)|{
                    'region'},
                'structural control vector is incomplete for '+
                str(m.get('name')),errors)
            adjusted_components=dict(components)
            reach_details=structural.get('components') or {}
            _check(
                set(reach_details)==set(DP.STRUCTURAL_TARGETS),
                'structural reach-component audit is incomplete for '+
                str(m.get('name')),errors)
            for target in DP.STRUCTURAL_TARGETS:
                adjusted=(reach_details.get(target) or {}).get(
                    'adjusted_percentile')
                if isinstance(adjusted,(int,float)):
                    adjusted_components[target]=adjusted
            if (
                set(adjusted_components)==set(score_weights) and
                isinstance(structural.get('score'),(int,float))
            ):
                reconstructed=round(100*sum(
                    float(score_weights[key])*
                    float(adjusted_components[key])
                    for key in score_weights
                ),1)
                _check(
                    abs(reconstructed-float(structural['score']))<=0.01,
                    'structurally adjusted score does not reconstruct for '+
                    str(m.get('name')),errors)
    _check(potential_count==int(potential_meta.get('qualified_members',-1)),
           'demographic-potential metadata count does not match member stickers',errors)
    _check(len(potential_markets)==int(
        potential_meta.get('qualified_markets',-1)),
        'demographic-potential metadata count does not match qualifying markets',
        errors)
    _check(all(potential.get('market_count')==len(potential_markets)
               for potential in potential_markets.values()),
           'demographic-potential member market counts are inconsistent',errors)
    _check(
        len(structurally_adjusted_markets)==int(
            structural_meta.get('adjusted_markets',-1)
        ),
        'structural-adjustment metadata count does not match markets',
        errors,
    )
    _check(all(
        structural.get('market_count')==len(
            structurally_adjusted_markets)
        for structural in structurally_adjusted_markets.values()
    ),'structurally adjusted market counts are inconsistent',errors)
    source_display=((bundle.get('integrity') or {}).get('source_display') or {})
    # DESK_OUT preview runs (make_member_desk no-clobber redirect) must not overwrite the LIVE
    # integrity state/report/reconciliation artifacts — route this module's writes (and the
    # regeneration-state read) to the same directory.
    _outdir=os.environ.get('DESK_OUT')
    _op=(lambda rel: os.path.join(_outdir, rel)) if _outdir else p
    _check(source_display.get('social_displayed',0)>0, 'social source loaded zero visible items', errors)
    def _recon(source,parts):
        accounted=sum(int(source_display.get(k,0) or 0) for k in parts)
        return {'source_rows':int(source),'accounted_rows':accounted,
                'difference':int(source)-accounted,'components':{k:int(source_display.get(k,0) or 0) for k in parts},
                'status':'pass' if int(source)==accounted else 'fail'}
    reconciliation={
        'press':_recon(source_display.get('press_source_rows',0),(
            'press_displayed','press_unknown_member_excluded','press_billstatus_rows_routed_to_legislation')),
        'social':_recon(source_display.get('social_source_rows',0),(
            'social_displayed','social_unverified_handle_excluded','social_unknown_member_excluded',
            'social_unrelated_excluded','social_offtopic_excluded')),
        'legislation':_recon((source_display.get('bill_source_rows',0) or 0)+
                              (source_display.get('bill_manual_rows',0) or 0),(
            'bills_displayed_unique','bill_duplicate_rows_collapsed','bill_unknown_member_excluded',
            'bill_rules_excluded','bill_override_excluded','bill_manual_excluded')),
        # media footprint (2026-07-16): every interview episode / hearing appearance / bluesky post is
        # displayed, summarized beyond the display cap, or carries one explicit exclusion reason.
        'interview_footprint':_recon(source_display.get('interview_episode_rows',0),(
            'interview_appearances_displayed','interview_beyond_display_cap',
            'interview_unknown_member_excluded')),
        # hearing footprint draws on TWO declared sources (2026-07-19): the hearing_antag appearance
        # rows, plus graded CU passages whose hearing has no antag row at all (hearing_antag records
        # the antagonism pass, not every speaker, so a committee chair can be graded while owning no
        # antag row). Both are displayed or accounted; DESK_MEDIA_STANCE=0 makes the second term 0.
        'hearing_footprint':_recon((source_display.get('hearing_appearance_rows',0) or 0)+
                                   (source_display.get('hearing_graded_rows_added',0) or 0),(
            'hearing_appearances_displayed','hearing_beyond_display_cap',
            'hearing_unknown_member_excluded')),
        'bluesky_footprint':_recon(source_display.get('bluesky_post_rows',0),(
            'bluesky_footprint_displayed','bluesky_beyond_display_cap',
            'bluesky_unverified_handle_posts_excluded','bluesky_unknown_member_excluded',
            'bluesky_empty_text_excluded','bluesky_scored_rows_routed_to_social')),
    }
    for source,row in reconciliation.items():
        _check(row['status']=='pass','%s source rows are not fully reconciled (%+d)'%(source,row['difference']),errors)
    _check(article_total==(source_display.get('press_displayed',0) or 0)+(source_display.get('social_displayed',0) or 0)
           +(source_display.get('interview_appearances_displayed',0) or 0)
           +(source_display.get('hearing_appearances_displayed',0) or 0)
           +(source_display.get('bluesky_footprint_displayed',0) or 0),
           'desk article total does not match reconciled press + social + footprint display rows',errors)
    _check(bill_total==(source_display.get('bills_displayed_unique',0) or 0),
           'desk bill total does not match reconciled unique legislation rows',errors)
    MC.dump_atomic({'status':'pass' if all(v['status']=='pass' for v in reconciliation.values()) else 'fail',
                    'rule':'Every source row must be displayed or assigned one explicit exclusion/routing reason.',
                    'sources':reconciliation,'exclusion_samples':source_display.get('exclusion_samples') or {}},
                   _op('dashboard_exclusion_reconciliation.json'),indent=2,ensure_ascii=False)
    if missing_bio: warnings.append('%d members lack desk bioguide enrichment (identity not guessed)'%len(missing_bio))

    # Strip self-referential validation and volatile freshness metadata before hashing the semantic
    # product.  ``generated_at`` and source mtimes are provenance/display metadata, not member data;
    # including them makes every same-input rebuild fail merely because it ran in a later second.
    projection={k:v for k,v in bundle.items() if k not in ('integrity','freshness')}
    output_hash=_sha(projection)
    # Fingerprint every source consumed by make_member_desk plus its routing code
    # and templates.  Content hashes prevent timestamp-only Excel/OneDrive touches
    # from defeating the same-input/same-output corruption check.
    source_files=[
        'Moneyball Main Version - Blended.xlsx','Social Master.xlsx','tweet_stance_final.csv',
        'manual_bill_evidence.csv','bill_overrides.csv','archive_leg.json',
        'behavioral_dispositions.json','behavioral_roles_v2_candidate.json',
        'big5_blue_emb.json','big5_blue_validated.json','big5_llm_opus.json',
        'big5_senate_expert_gold.json','big5_external_gold_validation.json','big5_e_gold_anchored.json',
        'big5_gold_anchored.json',
        'big5_traits.json','big5_v2.json','category_stance.json','communication_style.json',
        'coop_v2_candidate.json','dashboard_data.json','electoral_history.json',
        'electoral_plays.json','interview_antagonism.json','kb_benchmark.json',
        'lep_criterion.json','member_desk_manual.json','member_desk_src.json',
        'member_narratives_v2.json','member_personal_texture.json','member_receptivity.json',
        'member_references.json','model_bio.json','optimize_result.json',
        'pillar_inputs_snapshot.json','press_date_overrides.json','press_stance_overrides.json',
        'profile_basis_v2.json','ranks_python.json','signed_coalitions.json',
        'staff_directory.json','staff_interactions.json','staff_state.json','formal_visibility_roles.json',
        os.path.join('fec','legislators.csv'),os.path.join('fec','make_member_desk.py'),
        os.path.join('fec','media_rules.py'),os.path.join('fec','bill_rules.py'),
        os.path.join('fec','member_identity.py'),os.path.join('fec','dashboard_integrity.py'),
        os.path.join('fec','demographic_potential.py'),
        'DEMOGRAPHIC_POTENTIAL_METHOD.md',
        os.path.join('fec','moneyball_core.py'),os.path.join('fec','stance_rules.py'),
        os.path.join('fec','per_category_stance.py'),os.path.join('fec','compute_pillars.py'),
        os.path.join('fec','influence_roles.py'),
        os.path.join('fec','templates','make_member_desk_00.tpl'),
        # media-footprint source (2026-07-16): the DB plus its WAL sidecar — recent commits live in
        # the -wal until checkpoint, and hashing the main file alone would call changed data "same".
        'social_media.db','social_media.db-wal',
    ]
    source_state={rel:_file_sha(p(rel)) for rel in source_files}
    # Match make_member_desk.L(): root copy wins, then fec/ fallback; if
    # neither exists, _file_sha records None.  These optional district layers
    # all change the semantic desk bundle and therefore must participate in
    # the same-input/same-output fingerprint.
    for rel in OPTIONAL_DISTRICT_SOURCES:
        root_path=os.path.join(ROOT,rel)
        fallback_path=os.path.join(ROOT,'fec',rel)
        selected_path=root_path if os.path.exists(root_path) else fallback_path
        source_state[rel]=_file_sha(selected_path)
    # the footprint display dials shape the bundle, so they are inputs to the same-input check
    source_state['env:DESK_FOOTPRINT']=os.environ.get('DESK_FOOTPRINT','1')
    source_state['env:DESK_MEDIA_CAP']=os.environ.get('DESK_MEDIA_CAP','50')
    source_state['env:DESK_MEDIA_STANCE']=os.environ.get('DESK_MEDIA_STANCE','1')
    input_hash=_sha(source_state); state_path=_op('dashboard_integrity_state.json')
    prior=json.load(open(state_path,encoding='utf-8')) if os.path.exists(state_path) else None
    if prior and prior.get('input_fingerprint')==input_hash:
        _check(prior.get('output_semantic_sha256')==output_hash,
               'same dashboard inputs produced a different semantic bundle', errors)
    report={'status':'pass' if not errors else 'fail','errors':errors,'warnings':warnings,
            'preflight':preflight_report,'counts':{'members':len(members),'articles':article_total,
                'bills':bill_total,'members_with_social':len(social_members),
                'demographic_potential':potential_count,
                'demographic_potential_markets':len(potential_markets)},
            'input_fingerprint':input_hash,'output_semantic_sha256':output_hash,
            'regeneration_match': None if not prior or prior.get('input_fingerprint')!=input_hash else not errors,
            'missing_bioguide_enrichment':missing_bio,'exclusion_reconciliation':reconciliation}
    if errors: raise IntegrityError('; '.join(errors[:8]))
    MC.dump_atomic(report,_op('dashboard_integrity_report.json'),indent=2,ensure_ascii=False)
    MC.dump_atomic({'input_fingerprint':input_hash,'output_semantic_sha256':output_hash},state_path,indent=2)
    return report
