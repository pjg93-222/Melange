# -*- coding: utf-8 -*-
import openpyxl, collections, json, sys, pandas as pd
B = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

C = json.load(open(B + r'\congestion.json', encoding='utf-8'))
name2bio = {m['member']: m.get('bioguide') for m in C}
bio2name = {v: k for k, v in name2bio.items() if v}

wb = openpyxl.load_workbook(B + r'\congress_cosponsor_network_20260625_205554.xlsx', data_only=True)
ws = wb['Edges']
E = {}
for r in range(2, ws.max_row + 1):
    a, b, w, nb = [ws.cell(r, c).value for c in range(1, 5)]
    if a and b: E[(a, b)] = nb
print('network directed edges (bioguide space):', len(E))

A = pd.read_excel(B + r'\congress_bills_archive_20260624_115822.xlsx', sheet_name='All Bills')
A['key'] = A['Name'].astype(str).str.strip() + '|' + A['Congress'].astype(str)
A['rl'] = A['Role'].astype(str).str.lower()
spon = A[A['rl'] == 'sponsor']
b2s = {k: set(g['Member'].astype(str)) for k, g in spon.groupby('key')}
cos = A[A['rl'].isin(['cosponsor', 'original-cosponsor'])]

# DIRECTION TEST on bills where sponsor is a roster member
fwd = rev = 0
for _, r in cos.iterrows():
    k = r['key']; m = str(r['Member']); mb = name2bio.get(m)
    if not mb: continue
    for s in b2s.get(k, ()):
        sb = name2bio.get(s)
        if not sb or sb == mb: continue
        if (sb, mb) in E: fwd += 1
        if (mb, sb) in E: rev += 1
print('DIRECTION: sponsor->cosponsor hits=%d | cosponsor->sponsor hits=%d' % (fwd, rev))

# Restrict to bills whose sponsor row exists AND sponsor is on roster
onroster_bills = {k for k, ss in b2s.items() if any(name2bio.get(s) for s in ss)}
print('\nbills w/ sponsor row:', len(b2s), '| sponsor on roster:', len(onroster_bills))

hit = collections.Counter(); miss = collections.Counter(); meta = {}
for _, r in cos.iterrows():
    k = r['key']
    if k not in onroster_bills: continue
    m = str(r['Member']); mb = name2bio.get(m)
    if not mb: continue            # cosponsor off-roster -> archive wouldn't have the row anyway
    meta.setdefault(k, (r['Congress'], r['Section'], r['LastAction'], r['Title']))
    for s in b2s.get(k, ()):
        sb = name2bio.get(s)
        if not sb or sb == mb: continue
        if (sb, mb) in E: hit[k] += 1
        else: miss[k] += 1

allb = set(hit) | set(miss)
zero = sorted(k for k in allb if hit[k] == 0)
full = [k for k in allb if miss[k] == 0]
print('\nBILLS (sponsor on roster, cosponsor on roster) :', len(allb))
print('  fully in network:', len(full), '| ZERO in network:', len(zero), '| partial:', len(allb)-len(full)-len(zero))
print('  archive pairs missing from network:', sum(miss.values()), 'of', sum(hit.values())+sum(miss.values()))
print('\n  ZERO by Congress:', dict(sorted(collections.Counter(meta[k][0] for k in zero).items())))
print('  ZERO by LastAction:', dict(collections.Counter(meta[k][2] for k in zero).most_common()))
print('\n  TOP ZERO bills by missing pair count:')
for k in sorted(zero, key=lambda x: -miss[x])[:15]:
    print('    %-16s C%s  miss=%-4d %-42s | %s' % (k.split('|')[0], meta[k][0], miss[k], str(meta[k][3])[:42], meta[k][1][:28]))
