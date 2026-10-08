# -*- coding: utf-8 -*-
"""N5 step 1 (READ-ONLY): schema probe of every artifact the classification needs."""
import os, sys, json, csv, zipfile, collections
import openpyxl
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
ROOT = r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems"
P = lambda *a: os.path.join(ROOT, *a)

print('=' * 100)
print('1. WORKBOOK role sheets')
print('=' * 100)
wb = openpyxl.load_workbook(P('Moneyball Main Version - Blended.xlsx'), read_only=True, data_only=True)
for sn in ('Sponsor Interactions', 'Original Co-Sponsor Interaction', 'Co-Sponsor Interactions'):
    ws = wb[sn]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else '' for h in next(it)]
    n = 0
    last = None
    for r in it:
        n += 1
        last = r
    print('  %-34s max_row=%s max_col=%s data_rows=%d' % (sn, ws.max_row, ws.max_column, n))
    print('     header (%d cols): %s' % (len(hdr), hdr))
    print('     LAST data row  : %s' % (list(last)[:14],))
wb.close()

print()
print('=' * 100)
print('2. SELECTED + OLD archive "All Bills"')
print('=' * 100)
for f in ('congress_bills_archive_20260825_115305.xlsx', 'congress_bills_archive_20260624_115822.xlsx'):
    w = openpyxl.load_workbook(P(f), read_only=True, data_only=True)
    print('  %s sheets=%s' % (f, w.sheetnames))
    ws = w['All Bills']
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else '' for h in next(it)]
    n = sum(1 for _ in it)
    print('     data_rows=%d header=%s' % (n, hdr))
    w.close()

print()
print('=' * 100)
print('3. Sentinel Bill Ledger')
print('=' * 100)
w = openpyxl.load_workbook(P('Sentinel Bill Ledger.xlsx'), read_only=True, data_only=True)
print('  sheets=%s' % w.sheetnames)
ws = w['Sentinel Bill Ledger']
it = ws.iter_rows(values_only=True)
hdr = [str(h) if h is not None else '' for h in next(it)]
rows = [dict(zip(hdr, r)) for r in it]
w.close()
print('  data_rows=%d' % len(rows))
print('  header=%s' % hdr)
for col in ('Record Type', 'Review Status', 'Scoring Eligible', 'Role', 'Source', 'Evidence Type'):
    if col in hdr:
        print('  %-18s -> %s' % (col, dict(collections.Counter(
            str(r.get(col)) for r in rows).most_common(12))))
print('  sample row:', json.dumps({k: (str(v)[:60] if v is not None else None)
                                   for k, v in rows[0].items()}, ensure_ascii=False)[:900])

print()
print('=' * 100)
print('4. legislation_routing_audit.json')
print('=' * 100)
RA = json.load(open(P('legislation_routing_audit.json'), encoding='utf-8'))
print('  top keys: %s' % list(RA.keys()))
for k in ('schema', 'source', 'source_rows', 'selected_events',
          'sentinel_ledger_rows', 'sentinel_verified_rows'):
    print('  %-22s %s' % (k, RA.get(k)))
print('  summary.dispositions = %s' % json.dumps(RA['summary']['dispositions'], indent=2))
print('  rows = %d' % len(RA['rows']))
keyfreq = collections.Counter()
for r in RA['rows']:
    keyfreq.update(r.keys())
print('  row field frequency: %s' % dict(keyfreq.most_common()))
byd = collections.defaultdict(list)
for r in RA['rows']:
    byd[r['disposition']].append(r)
for d in sorted(byd):
    ex = byd[d][0]
    print('  --- %-46s n=%-6d sample=%s' % (d, len(byd[d]),
          json.dumps({k: (str(v)[:40] if v is not None else None) for k, v in ex.items()},
                     ensure_ascii=False)[:420]))

print()
print('=' * 100)
print('5. bill_repository_live.json  (per-bill BILLSTATUS cache: cosponsorship DATES)')
print('=' * 100)
BL = json.load(open(P('bill_repository_live.json'), encoding='utf-8'))
print('  top keys: %s' % list(BL.keys())[:20])
bills = BL.get('bills') or {}
print('  bills = %d ; updated_at=%s' % (len(bills), BL.get('updated_at')))
k0 = next(iter(bills))
b0 = bills[k0]
print('  sample key=%r  field keys=%s' % (k0, list(b0.keys())))
for f in ('sponsor', 'cosponsors', 'original_cosponsors'):
    v = b0.get(f)
    if isinstance(v, list) and v:
        print('     %s[0] = %s' % (f, json.dumps(v[0], ensure_ascii=False)[:300]))
    elif isinstance(v, dict):
        print('     %s = %s' % (f, json.dumps(v, ensure_ascii=False)[:300]))
print('  introduced/date-ish fields on sample:', {f: b0.get(f) for f in
      ('introduced_date', 'introducedDate', 'latest_action', 'checked_at', 'updated_at')
      if f in b0})

print()
print('=' * 100)
print('6. xlsx zip entry timestamps (workbook + selected archive)')
print('=' * 100)
for f in ('Moneyball Main Version - Blended.xlsx',
          'congress_bills_archive_20260825_115305.xlsx'):
    z = zipfile.ZipFile(P(f))
    print('  %s' % f)
    for i in z.infolist():
        if 'sheet' in i.filename or 'core.xml' in i.filename or 'workbook.xml' == os.path.basename(i.filename):
            print('     %-40s %s  %d bytes' % (i.filename, i.date_time, i.file_size))
    try:
        core = z.read('docProps/core.xml').decode('utf-8', 'replace')
        print('     docProps/core.xml = %s' % core[:600])
    except KeyError:
        print('     docProps/core.xml absent')
    z.close()

print()
print('=' * 100)
print('7. manual_bill_evidence.csv + congestion roster + movements/admissions/collector sizes')
print('=' * 100)
with open(P('manual_bill_evidence.csv'), encoding='utf-8-sig', newline='') as fh:
    man = list(csv.DictReader(fh))
print('  manual rows=%d fields=%s' % (len(man), list(man[0].keys())))
cong = json.load(open(P('data', 'engine', 'congestion.json'), encoding='utf-8'))
print('  congestion roster=%d' % len(cong))
for rel in ('data/legislation/official_membership_movements.json',
            'data/legislation/sentinel_admissions.json',
            'data/members/collector_roster_updates.ndjson',
            'data/members/sentinel_verified_membership_history.ndjson',
            'data/legislation/membership_withdrawals.csv',
            'data/members/primacy_act_memberships.ndjson',
            'data/members/registry_anchor_memberships.ndjson',
            'data/members/bank_anchor_memberships.ndjson'):
    p = P(*rel.split('/'))
    print('  %-62s %s' % (rel, ('%d bytes' % os.path.getsize(p)) if os.path.exists(p) else 'MISSING'))
