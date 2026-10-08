# Lane C / Q1 : topic/category key sets across every surface. READ-ONLY. stdlib + openpyxl.
# NOTHING under fec/ is imported. The blocks/BLOCKCAT literals below are TRANSCRIBED from
# fec/moneyball_core.py:480-485 (quoted in REPORT.md) so the workbook read matches the producer.
import json, os
import openpyxl

ROOT = r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems"
P = lambda *a: os.path.join(ROOT, *a)

# --- transcribed from fec/moneyball_core.py:480-485 ---
BLOCKS = [3,11,19,27,35,43,51,59,67,75,83,91,99,145,115,107,123,131,153]
BLOCKCAT = {3:'Examination & Supervision',11:'CU Capital Requirements',19:'NCUA Budget & Insurance',27:'Interchange',
 35:'Payments',43:'Technology & Emerging',51:'Fraud Prevention',59:'Data Privacy & Cyber',67:'Charter Modernization',
 75:'Defense CU Issues',83:'Field of Membership',91:'CDFI Fund',99:'CFPB',145:'CFPB',115:'TCPA',
 107:'Credit Card Mandates',123:'Member Business Lending',131:'Mortgage & Housing',153:'CU Tax Status'}
# --- transcribed from fec/bill_rules.py:16-23 ---
BR_CATEGORIES = ('Examination & Supervision','CU Capital Requirements','NCUA Budget & Insurance','Interchange','Payments',
 'Technology & Emerging','Fraud Prevention','Data Privacy & Cyber','Charter Modernization','Defense CU Issues',
 'Field of Membership','CU Tax Status','CDFI Fund','CFPB','Credit Card Mandates','TCPA','Member Business Lending',
 'Mortgage & Housing')

def load(rel):
    with open(P(rel), encoding='utf-8') as fh:
        return json.load(fh)

TARGET = 'CU Tax Status'
TARGET_LABEL = 'Credit Union Tax Status'

def has_target(keys):
    hits = [k for k in keys if 'tax status' in str(k).lower()]
    return hits

rows = []   # (surface, n_keys, target_present, target_key_literal)

print("=" * 100)
print("Q1  CATEGORY / TOPIC KEY SETS -- one surface per section")
print("=" * 100)

# 1. fec/bill_rules.py CATEGORIES (the canonical category tuple)
print("\n-- [code] fec/bill_rules.py:16-23  CATEGORIES --")
print("   n = %d" % len(BR_CATEGORIES))
for c in BR_CATEGORIES: print("   ", c)
rows.append(('code: bill_rules.CATEGORIES', len(BR_CATEGORIES), TARGET in BR_CATEGORIES, TARGET))

# 2. category_stance.json
cs = load('category_stance.json')
print("\n-- category_stance.json  ['cats'] --")
print("   n = %d" % len(cs['cats']))
for c in cs['cats']: print("   ", c)
print("   ['defensive'] =", cs['defensive'])
print("   matrix members =", len(cs['matrix']))
_mk = sorted({k for m in cs['matrix'].values() for k in m})
print("   distinct keys inside matrix rows: n=%d  identical to ['cats']? %s" % (_mk and len(_mk), sorted(cs['cats']) == _mk))
rows.append(('category_stance.json cats', len(cs['cats']), TARGET in cs['cats'], TARGET))

# 3. per_topic_alloc.json  (== per_topic_flip.py output; fec/engine/per_topic_flip.py:94-95)
fl = load('per_topic_alloc.json')
print("\n-- per_topic_alloc.json  (per_topic_flip.py output) --")
print("   n = %d" % len(fl))
for k in fl: print("   %-45s funded_members=%d  total=$%s" % (k, len(fl[k]), format(sum(fl[k].values()), ',')))
rows.append(('per_topic_alloc.json (flip out)', len(fl), bool(has_target(fl)), (has_target(fl) or [None])[0]))

# 4. per_topic_coalition.json
co = load(r'data\lobbying\per_topic_coalition.json')
print("\n-- data\\lobbying\\per_topic_coalition.json --")
print("   n = %d" % len(co))
for k in co: print("   %-45s seeds=%d  total=$%s" % (k, len(co[k]), format(sum(co[k].values()), ',')))
rows.append(('per_topic_coalition.json', len(co), bool(has_target(co)), (has_target(co) or [None])[0]))

# 5. budget_frontier.json + electoral
for rel in (r'data\money\budget_frontier.json', r'data\money\budget_frontier_electoral.json'):
    bf = load(rel)
    topics = bf.get('topics')
    keys = list(topics.keys()) if isinstance(topics, dict) else [t.get('topic') or t.get('label') for t in topics]
    print("\n-- %s  ['topics'] --" % rel)
    print("   container type = %s   n = %d" % (type(topics).__name__, len(keys)))
    for k in keys: print("   ", k)
    rows.append((rel, len(keys), bool(has_target(keys)), (has_target(keys) or [None])[0]))

# 6. dashboard_data.json
dd = load('dashboard_data.json')
print("\n-- dashboard_data.json --")
print("   ['issues']  (rendered dashboard columns)  n = %d" % len(dd['issues']))
for i, k in enumerate(dd['issues']): print("   [%2d] %s" % (i, k))
print("   ['engine']['categories'] (canonical category per column)  n = %d" % len(dd['engine']['categories']))
for i, k in enumerate(dd['engine']['categories']): print("   [%2d] %r" % (i, k))
print("   len(flip)=%d len(coal)=%d len(frontier)=%d len(stats)=%d" % (
    len(dd['flip']), len(dd['coal']), len(dd['frontier']), len(dd['stats'])))
print("   honors columns: 'honors' rows=%d cols=%d ; honors_meta cats n/a here" % (
    len(dd.get('honors') or []), len((dd.get('honors') or [[]])[0])))
rows.append(('dashboard_data.json issues', len(dd['issues']), bool(has_target(dd['issues'])), (has_target(dd['issues']) or [None])[0]))
ec = [c for c in dd['engine']['categories'] if c]
rows.append(('dashboard_data.json engine.categories', len(ec), TARGET in ec, TARGET))

# 7. workbook Legislative Score - Normalized topic blocks
wb = openpyxl.load_workbook(P('Moneyball Main Version - Blended.xlsx'), read_only=True, data_only=True)
print("\n-- workbook sheet names --")
for s in wb.sheetnames: print("   ", s)
nrm = wb['Legislative Score - Normalized']
# read_only sheets need a materialized grid
grid = []
for r_i, row in enumerate(nrm.iter_rows(min_row=1, max_row=2, values_only=True), 1):
    grid.append(list(row))
hdr, w2 = grid[0], grid[1]
def cell(rowvals, col):
    return rowvals[col-1] if col-1 < len(rowvals) else None
print("\n-- workbook 'Legislative Score - Normalized' topic blocks (row 1 header at each block col) --")
print("   sheet max_column = %s" % nrm.max_column)
print("   %-5s %-45s %-28s %s" % ("col", "row-1 header (raw)", "BLOCKCAT[col]", "row-2 sub-weights (o=0,2,4)"))
labels = {}
for sc in BLOCKS:
    raw = cell(hdr, sc)
    lab = str(raw).replace(' - Sponsor', '').replace(' Sponsor', '').strip()
    labels[sc] = lab
    sw = [cell(w2, sc + o) for o in (0, 2, 4)]
    print("   %-5d %-45s %-28s %s" % (sc, repr(raw)[:45], BLOCKCAT[sc], sw))
print("\n   producer label expression (per_topic_flip.py:56, per_topic_coalition.py:42, budget_frontier.py:61):")
print("     labels={sc:str(nrm.cell(1,sc).value).replace(' - Sponsor','').replace(' Sponsor','').strip() for sc in blocks}")
print("   -> derived label set, n=%d:" % len(set(labels.values())))
for sc in BLOCKS: print("     col %-4d -> %r" % (sc, labels[sc]))
rows.append(('workbook Legislative-Normalized blocks', len(BLOCKS), any('tax status' in v.lower() for v in labels.values()),
             next((v for v in labels.values() if 'tax status' in v.lower()), None)))
print("\n   label-set == per_topic_alloc.json key set? %s" % (sorted(set(labels.values())) == sorted(fl.keys())))
print("   labels only in workbook:", sorted(set(labels.values()) - set(fl.keys())))
print("   keys only in per_topic_alloc:", sorted(set(fl.keys()) - set(labels.values())))

# 8. archive_leg.json (which categories actually carry legislative evidence)
al = load('archive_leg.json')
cats_in_leg = {}
for mem, d in al.items():
    for c, v in d.items():
        cats_in_leg.setdefault(c, [0, 0.0])
        cats_in_leg[c][0] += 1
        cats_in_leg[c][1] += v
print("\n-- archive_leg.json : categories carrying legislative evidence --")
print("   members = %d ; distinct categories = %d" % (len(al), len(cats_in_leg)))
for c in sorted(cats_in_leg, key=lambda k: -cats_in_leg[k][0]):
    print("   %-30s members=%3d  signed_sum=%+.2f" % (c, cats_in_leg[c][0], cats_in_leg[c][1]))
print("   CATEGORIES with NO legislative evidence:", sorted(set(BR_CATEGORIES) - set(cats_in_leg)))
rows.append(('archive_leg.json categories', len(cats_in_leg), TARGET in cats_in_leg, TARGET))

print("\n" + "=" * 100)
print("Q1 SUMMARY -- does the surface carry 'CU Tax Status' today?")
print("=" * 100)
print("%-46s %6s  %-8s %s" % ("surface", "n_keys", "carries?", "key literal"))
for s, n, present, lit in rows:
    print("%-46s %6s  %-8s %s" % (s, n, "YES" if present else "NO", lit))
