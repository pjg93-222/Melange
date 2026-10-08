# Lane C / Q2b : size of the workbook-role-sheet vs selected-archive fork. READ-ONLY.
# The desk's displayed bill ledger is built from the 3 workbook role sheets + manual csv
# (fec/surfaces/make_member_desk.py:2092-2139).  The stance matrix is built from the selected
# archive + sentinel-verified + manual (fec/legislation/build_archive_legislative.py:50,90,100).
# This measures the set difference on (member, bill, congress, role).
import os, csv, collections, re, sys, json
import openpyxl
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
ROOT = r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems"
P = lambda *a: os.path.join(ROOT, *a)
norm = lambda b: re.sub(r'\s+', '', str(b or '')).upper()
def _congress(v):
    try: return str(int(float(v)))
    except (TypeError, ValueError): return str(v or '').strip()
def canonical_role(role):
    v = str(role or '').strip().lower()
    if 'original' in v: return 'original-cosponsor'
    if 'cosponsor' in v or 'co-spons' in v: return 'cosponsor'
    if 'sponsor' in v: return 'sponsor'
    return v or 'other'

def xl(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]; it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else '' for h in next(it)]
    out = [dict(zip(hdr, r)) for r in it]; wb.close(); return out

ARCH = xl(P('congress_bills_archive_20260825_115305.xlsx'), 'All Bills')
SENT = [r for r in xl(P('Sentinel Bill Ledger.xlsx'), 'Sentinel Bill Ledger')
        if str(r.get('Scoring Eligible') or '').strip().lower() in {'true','1','yes'}
        and str(r.get('Review Status') or '').strip().lower() == 'verified'
        and str(r.get('Record Type') or '').strip().lower() == 'membership']
with open(P('manual_bill_evidence.csv'), encoding='utf-8-sig', newline='') as fh:
    MANUAL = list(csv.DictReader(fh))
ROLE = {}
wbm = openpyxl.load_workbook(P('Moneyball Main Version - Blended.xlsx'), read_only=True, data_only=True)
for sn in ('Sponsor Interactions', 'Original Co-Sponsor Interaction', 'Co-Sponsor Interactions'):
    ws = wbm[sn]; it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else '' for h in next(it)]
    ROLE[sn] = [dict(zip(hdr, r)) for r in it]
wbm.close()

def key(member, bill, cong, role):
    return (str(member or '').strip(), norm(bill), _congress(cong), canonical_role(role))

desk = collections.Counter()
for sn, rows in ROLE.items():
    for r in rows:
        desk[key(r.get('Member'), r.get('Name'), r.get('Congress'), r.get('Role'))] += 1
matrix = collections.Counter()
for r in ARCH:
    matrix[key(r.get('Member'), r.get('Name'), r.get('Congress'), r.get('Role'))] += 1
for r in SENT:
    matrix[key(r.get('Member'), r.get('Bill'), r.get('Congress'), r.get('Role'))] += 1
manual = {key(r.get('Member'), r.get('Bill'), r.get('Congress'), r.get('Role')) for r in MANUAL}

print("=" * 100)
print("Q2b  FORK SIZE : workbook role sheets (desk ledger) vs selected archive+sentinel (stance matrix)")
print("=" * 100)
print("workbook role-sheet rows          = %d  (distinct member/bill/congress/role keys = %d)"
      % (sum(desk.values()), len(desk)))
print("archive+sentinel rows             = %d  (distinct keys = %d)" % (sum(matrix.values()), len(matrix)))
print("manual_bill_evidence.csv keys     = %d  (read by BOTH sides)" % len(manual))
only_desk = set(desk) - set(matrix) - manual
only_matrix = set(matrix) - set(desk) - manual
print()
print("keys ONLY in the workbook role sheets (desk shows, stance matrix never sees) = %d" % len(only_desk))
print("keys ONLY in archive+sentinel        (stance matrix scores, desk never shows) = %d" % len(only_matrix))
print()
print("  only-in-workbook by role:", dict(collections.Counter(k[3] for k in only_desk)))
print("  only-in-workbook by Congress:", dict(sorted(collections.Counter(k[2] for k in only_desk).items())))
print("  only-in-workbook distinct members: %d ; distinct bills: %d"
      % (len({k[0] for k in only_desk}), len({k[1] for k in only_desk})))
print("  top 25 only-in-workbook bills by key count:")
for b, n in collections.Counter(k[1] for k in only_desk).most_common(25):
    print("     %-16s keys=%d" % (b, n))
print()
print("  only-in-archive by role:", dict(collections.Counter(k[3] for k in only_matrix)))
print("  only-in-archive by Congress:", dict(sorted(collections.Counter(k[2] for k in only_matrix).items())))
print("  top 25 only-in-archive bills by key count:")
for b, n in collections.Counter(k[1] for k in only_matrix).most_common(25):
    print("     %-16s keys=%d" % (b, n))
print()
print("  sample 30 only-in-workbook keys:")
for k in sorted(only_desk)[:30]:
    print("     ", k)
print()
print("-- Tim Scott specifically --")
ts_desk = {k for k in desk if k[0] == 'Tim Scott'}
ts_mat = {k for k in matrix if k[0] == 'Tim Scott'}
print("   workbook keys=%d  archive+sentinel keys=%d  only-workbook=%d  only-archive=%d"
      % (len(ts_desk), len(ts_mat), len(ts_desk - ts_mat - manual), len(ts_mat - ts_desk - manual)))
for k in sorted(ts_desk - ts_mat - manual): print("      ONLY-WORKBOOK", k)
for k in sorted(ts_mat - ts_desk - manual): print("      ONLY-ARCHIVE ", k)
print()
print("-- S.J.Res.18 / H.J.Res.59 key parity --")
for b in ('S.J.RES.18', 'H.J.RES.59'):
    dk = {k for k in desk if k[1] == b}; mk = {k for k in matrix if k[1] == b}
    print("   %-12s workbook keys=%-4d archive keys=%-4d only-workbook=%-4d only-archive=%d"
          % (b, len(dk), len(mk), len(dk - mk), len(mk - dk)))
    for k in sorted(dk - mk)[:12]: print("        ONLY-WORKBOOK", k)
