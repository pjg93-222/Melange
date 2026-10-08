# Lane C / schema probe: sheet names + headers + row counts for every bill source. READ-ONLY.
import os, csv, json
import openpyxl
ROOT = r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems"
P = lambda *a: os.path.join(ROOT, *a)

for rel in ("congress_bills_archive_20260825_115305.xlsx",
            "congress_bills_archive_20260624_115822.xlsx",
            "Sentinel Bill Ledger.xlsx"):
    print("=" * 90)
    print(rel)
    wb = openpyxl.load_workbook(P(rel), read_only=True, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        n = 0
        for _ in ws.iter_rows(min_row=2, values_only=True):
            n += 1
        print("  sheet %-28s data_rows=%-7d headers=%s" % (sn, n, list(hdr)))
    wb.close()

print("=" * 90)
print("Moneyball Main Version - Blended.xlsx : role sheets")
wb = openpyxl.load_workbook(P("Moneyball Main Version - Blended.xlsx"), read_only=True, data_only=True)
for sn in ("Sponsor Interactions", "Original Co-Sponsor Interaction", "Co-Sponsor Interactions",
           "Success Sponsor", "Success Original Co-Sponsor", "CFPB Bills"):
    if sn not in wb.sheetnames:
        print("  sheet %-34s ABSENT" % sn); continue
    ws = wb[sn]
    rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
    n = 0
    for _ in ws.iter_rows(min_row=2, values_only=True):
        n += 1
    print("  sheet %-34s data_rows=%-7d" % (sn, n))
    for i, r in enumerate(rows, 1):
        print("     row%-2d %s" % (i, [None if v is None else (str(v)[:36]) for v in r[:14]]))
wb.close()

print("=" * 90)
print("data\\legislation\\bill_directions.csv")
with open(P("data", "legislation", "bill_directions.csv"), encoding="utf-8-sig", newline="") as fh:
    rd = csv.DictReader(fh)
    print("  fieldnames =", rd.fieldnames)
    rows = list(rd)
print("  rows =", len(rows))
for r in rows[:5]:
    print("   ", r)

print("=" * 90)
print("manual_bill_evidence.csv")
with open(P("manual_bill_evidence.csv"), encoding="utf-8-sig", newline="") as fh:
    rd = csv.DictReader(fh)
    print("  fieldnames =", rd.fieldnames)
    rows = list(rd)
print("  rows =", len(rows))
for r in rows[:3]:
    print("   ", r)

print("=" * 90)
print("data\\legislation\\bill_overrides.csv")
with open(P("data", "legislation", "bill_overrides.csv"), encoding="utf-8-sig", newline="") as fh:
    print(fh.read())

print("=" * 90)
print("legislation_routing_audit.json keys / summary")
with open(P("legislation_routing_audit.json"), encoding="utf-8") as fh:
    ra = json.load(fh)
print("  top keys:", [k for k in ra])
print("  schema =", ra.get("schema"), "| source =", ra.get("source"))
print("  source_rows =", ra.get("source_rows"), "selected_events =", ra.get("selected_events"))
print("  sentinel_ledger_rows =", ra.get("sentinel_ledger_rows"), "sentinel_verified_rows =", ra.get("sentinel_verified_rows"))
print("  summary.dispositions =", json.dumps(ra["summary"]["dispositions"], indent=2))
print("  summary.duplicates =", ra["summary"]["duplicates"])
print("  rows =", len(ra["rows"]))
print("  example row =", json.dumps(ra["rows"][0], indent=2)[:900])
