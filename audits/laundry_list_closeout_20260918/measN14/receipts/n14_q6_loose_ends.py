# -*- coding: utf-8 -*-
"""N14 loose ends.  READ-ONLY.  Four facts the first five receipts leave open:

  1. WHICH shared keys the desk drops for an unresolved member -- the 57 -> 56 and 493 -> 491 gap
     between the raw disagreement sets (Q1) and what the desk would actually DISPLAY (Q5).
  2. The label-set VINTAGE question: the xlsx mirrors of cu_stance_label_set are older than the
     csv, and on the 57 the older xlsx sometimes agrees with the SHEET.  Broken out per xlsx
     sheet ('To Rework' vs 'Done (locked)') against both sides.
  3. cu_golden_labels.csv (2026-07-04, 760 owner golds -- NOT in the brief's named list, and the
     NEWEST owner press-label record on the box) against the 57.
  4. The LIVE-root CU_Press_Stance_Tone.xlsx anomaly: 3,619 release urls where the Draft copy has
     2,375, and on 7 of the 57 it carries a value matching NEITHER source.
  5. The 37 title disagreements: which side is mojibake.
"""
import collections
import io
import json
import os

import openpyxl

import n14_common as C
from members.member_identity import Resolver

C.banner('N14 LOOSE ENDS')

q1 = json.load(io.open(os.path.join(C.HERE, 'n14_q1_disagreement_sets.json'), encoding='utf-8'))
the57 = q1['stance_disagreement_rows']


def url_key(value):
    return C.PS.release_url_key(value)


def norm_title(value):
    return ' '.join(str(value or '').split()).casefold()


# ---- 1. the dropped keys ----------------------------------------------------------------
bundle = json.load(io.open(os.path.join(C.LIVE, 'member_desk_bundle.json'), encoding='utf-8'))
NAMES = [m.get('name') for m in bundle.get('members') or []]
BIO = {m.get('name'): m.get('bioguide') for m in bundle.get('members') or []}
resolver = Resolver(NAMES, BIO)
rows, ph = C.sheet_rows()
sheet_first, _ = C.sheet_index_by_key(rows, ph)
tone_rows = C.PS.load_tone_rows()
shared = {}
for row in tone_rows:
    key = C.DPS.row_key(row.get('url'), row.get('bioguide'), row.get('member'))
    if key[0] and key in sheet_first:
        shared.setdefault(key, row)
print('-- 1. SHARED KEYS THE DESK DROPS (unresolved member) --')
dropped = []
for key in shared:
    srow = sheet_first[key]
    if not resolver.resolve(srow[ph['member']], srow[ph['bioguide']]):
        dropped.append(key)
print('   shared keys %d | dropped for an unresolved member %d' % (len(shared), len(dropped)))
in57 = {(r['release_key'], r['identity']) for r in the57}
src_flat = set()
for key in shared:
    srow = sheet_first[key]
    if (str(srow[ph['stance_src']] or '').strip() == 'debate'
            and (shared[key].get('stance_src') or '').strip() == 'human_label'):
        src_flat.add(key)
for key in dropped:
    srow = sheet_first[key]
    print('   %-28s bioguide=%-10s in the 57: %-3s in the 493 human_label set: %-3s'
          % (str(srow[ph['member']] or '')[:28], str(srow[ph['bioguide']] or ''),
             'YES' if key in in57 else 'no', 'YES' if key in src_flat else 'no'))
    print('      title %s' % str(srow[ph['title']] or '')[:88])
print('   => the 57 disagreements the DESK would actually display: %d'
      % len(in57 - set(dropped)))
print('   => the 493 stance_src keys the DESK would actually display: %d'
      % len(src_flat - set(dropped)))

# ---- 2. the label-set vintage ------------------------------------------------------------
print('\n-- 2. cu_stance_label_set VINTAGE: csv vs the two xlsx mirrors --')
for label, path in (('Draft csv ', os.path.join(C.DRAFT, 'cu_stance_label_set.csv')),
                    ('Draft xlsx', os.path.join(C.DRAFT, 'cu_stance_label_set.xlsx')),
                    ('LIVE  xlsx', os.path.join(C.LIVE, 'cu_stance_label_set.xlsx'))):
    print('   %s  mtime %s  sha256 %s' % (label, C.mtime(path), C.sha256_file(path)[:16]))
ls_rows = C.read_csv_rows(os.path.join(C.SCRATCH, 'cu_stance_label_set.csv'))
ls_by_url = {}
for r in ls_rows:
    uk = url_key(r.get('url'))
    if uk:
        ls_by_url.setdefault(uk, r)
keys57 = {}
for rec in the57:
    keys57[url_key(rec['url'])] = rec
    keys57.setdefault(url_key(rec['csv_url']), rec)
vintage = {}
for label, filename in (('Draft cu_stance_label_set.xlsx', 'labelset_draft.xlsx'),
                        ('LIVE  cu_stance_label_set.xlsx', 'labelset_live.xlsx')):
    wb = openpyxl.load_workbook(os.path.join(C.SCRATCH, filename), read_only=True,
                                data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        hidx = {str(h).strip().lower().replace(' ', '_'): i for i, h in enumerate(header) if h}
        if 'url' not in hidx or 'human_label' not in hidx:
            continue
        tally = collections.Counter()
        drift = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            uk = url_key(str(row[hidx['url']] or ''))
            if not uk:
                continue
            value = str(row[hidx['human_label']] or '').strip().lower()
            csv_row = ls_by_url.get(uk)
            if csv_row and value != (csv_row.get('human_label') or '').strip().lower():
                drift += 1
            rec = keys57.get(uk)
            if rec is None:
                continue
            if value == rec['sheet_stance']:
                tally['xlsx agrees with the SHEET'] += 1
            elif value == rec['csv_stance']:
                tally['xlsx agrees with the CSV'] += 1
            else:
                tally['xlsx agrees with NEITHER'] += 1
        print('   %-32s [%-14s] of the 57: %s   | rows whose human_label DIFFERS from the'
              ' csv: %d' % (label, name[:14], dict(tally.most_common()), drift))
        vintage['%s::%s' % (label.strip(), name)] = dict(tally)
        vintage['%s::%s::drift_vs_csv' % (label.strip(), name)] = drift
    wb.close()
print('   The csv (2026-06-30T18:56) is the LATEST label-set artifact; both xlsx copies predate')
print('   it.  A row where the older xlsx agrees with the SHEET and the newer csv does not is a')
print('   release whose label MOVED between the two owner passes.')

# ---- 3. cu_golden_labels.csv -------------------------------------------------------------
print('\n-- 3. cu_golden_labels.csv (2026-07-04, the NEWEST owner press-label record found) --')
gold_path = os.path.join(C.LIVE, 'data', 'stance', 'cu_golden_labels.csv')
gl = C.read_csv_rows(gold_path)
gl_by_url = {}
for r in gl:
    uk = url_key(r.get('url'))
    if uk:
        gl_by_url.setdefault(uk, r)
print('   rows %d | url-keyed %d | mtime %s' % (len(gl), len(gl_by_url), C.mtime(gold_path)))
tally = collections.Counter()
detail = []
for rec in the57:
    row = gl_by_url.get(url_key(rec['url'])) or gl_by_url.get(url_key(rec['csv_url']))
    if row is None:
        tally['not in cu_golden_labels.csv'] += 1
        continue
    value = (row.get('label') or '').strip().lower()
    if value == rec['csv_stance'] and value == rec['sheet_stance']:
        tally['both'] += 1
    elif value == rec['csv_stance']:
        tally['gold == CSV stance'] += 1
    elif value == rec['sheet_stance']:
        tally['gold == SHEET stance'] += 1
    else:
        tally['gold == neither (%s)' % value] += 1
    detail.append({'member': rec['member'], 'title': rec['title'], 'gold': value,
                   'sheet': rec['sheet_stance'], 'csv': rec['csv_stance'],
                   'split': row.get('split'), 'source': row.get('source')})
print('   %s' % dict(tally.most_common()))
for d in detail:
    if d['gold'] not in (d['csv'],):
        print('      %-22s %-48s gold=%-9s sheet=%-8s csv=%-8s split=%s'
              % (d['member'][:22], d['title'][:48], d['gold'], d['sheet'], d['csv'],
                 d['split']))

# ---- 4. the LIVE tone xlsx anomaly -------------------------------------------------------
print('\n-- 4. CU_Press_Stance_Tone.xlsx: Draft copy vs LIVE-root copy --')
for label, filename in (('Draft', 'tone_draft.xlsx'), ('LIVE ', 'tone_live.xlsx')):
    wb = openpyxl.load_workbook(os.path.join(C.SCRATCH, filename), read_only=True,
                                data_only=True)
    ws = wb['Releases']
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    hidx = {str(h).strip().lower().replace(' ', '_'): i for i, h in enumerate(header) if h}
    n = 0
    urls = set()
    dates = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        n += 1
        uk = url_key(str(row[hidx['url']] or ''))
        if uk:
            urls.add(uk)
        value = row[hidx['date']]
        text = str(value)[:10] if str(value)[:4].isdigit() else ''
        if text:
            dates.append(text)
    dates.sort()
    print('   %s Releases: %d rows, %d distinct release urls, dates %s .. %s'
          % (label, n, len(urls), dates[0] if dates else '-', dates[-1] if dates else '-'))
    wb.close()
print('   The LIVE-root copy holds ~1,240 more releases than the Draft copy that shares the')
print('   base tone CSV\'s mtime.  UNKNOWN: which run wrote the LIVE-root copy -- no repo script')
print('   names it (grep over fec/ for CU_Press_Stance_Tone returns no writer).')

# ---- 5. the 37 title disagreements -------------------------------------------------------
print('\n-- 5. THE 37 TITLE DISAGREEMENTS: WHICH SIDE IS MOJIBAKE? --')
MOJI = ('â€', 'Ã', 'â\x80', 'Â')
sheet_moji = csv_moji = neither_moji = 0
example = []
for key, crow in shared.items():
    srow = sheet_first[key]
    sv, ev = C.norm_field('title', srow[ph['title']], crow.get('title'))
    if sv == ev:
        continue
    s_bad = any(m in sv for m in MOJI)
    c_bad = any(m in ev for m in MOJI)
    if s_bad and not c_bad:
        sheet_moji += 1
    elif c_bad and not s_bad:
        csv_moji += 1
    else:
        neither_moji += 1
    if len(example) < 5:
        example.append((sv[:64], ev[:64]))
print('   title disagreements: %d' % (sheet_moji + csv_moji + neither_moji))
print('   SHEET side carries mojibake, CSV clean : %d' % sheet_moji)
print('   CSV side carries mojibake, SHEET clean : %d' % csv_moji)
print('   neither/both                           : %d' % neither_moji)
for sv, ev in example:
    print('      sheet=%r' % sv)
    print('      csv  =%r' % ev)

C.dump('n14_q6_loose_ends.json', {
    'schema': 'n14-loose-ends-v1',
    'shared_keys': len(shared), 'dropped_unresolved_member': len(dropped),
    'dropped_keys': [list(k) for k in dropped],
    'stance_disagreements_the_desk_displays': len(in57 - set(dropped)),
    'stance_src_keys_the_desk_displays': len(src_flat - set(dropped)),
    'label_set_vintage': vintage,
    'cu_golden_labels': {'rows': len(gl), 'url_keyed': len(gl_by_url),
                         'tally_on_57': dict(tally), 'rows_detail': detail},
    'title_mojibake': {'sheet_bad': sheet_moji, 'csv_bad': csv_moji,
                       'neither': neither_moji},
})
print('DONE N14 LOOSE ENDS')
