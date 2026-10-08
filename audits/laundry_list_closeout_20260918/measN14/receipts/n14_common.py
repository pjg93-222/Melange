# -*- coding: utf-8 -*-
"""N14 measurement lane -- shared read-only loaders.  READ-ONLY: writes nothing outside
audits/laundry_list_closeout_20260918/measN14/receipts and the session scratchpad.

Imports ONLY the two side-effect-free modules (press_stream, surfaces.desk_press_source) with
MONEYBALL_ROOT pointed at LIVE, exactly as fixN1/n1_desk_press_union.py established.  No
production pipeline script is imported or executed: make_member_desk.py and
stance/build_media_blend.py are read as TEXT and the relevant statements are lifted, because
importing either executes it and overwrites live artifacts.
"""
from __future__ import annotations

import ast
import csv
import datetime
import hashlib
import io
import json
import os
import sys

LIVE = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
DRAFT = r'C:\Users\pgoff\OneDrive\Desktop\Draft Systems'
HERE = os.path.dirname(os.path.abspath(__file__))
PACKET = os.path.dirname(HERE)
WT = os.path.abspath(os.path.join(PACKET, os.pardir, os.pardir, os.pardir))
SCRATCH = (r'C:\Users\pgoff\AppData\Local\Temp\claude'
           r'\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-'
           r'unfinished-work-handoff-a3bbad'
           r'\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\measN14')

os.environ['MONEYBALL_ROOT'] = LIVE
sys.path.insert(0, os.path.join(WT, 'fec'))
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

import press_stream as PS                     # side-effect-free reader module
import surfaces.desk_press_source as DPS      # side-effect-free selector module

WB_SCRATCH = os.path.join(SCRATCH, 'wb.xlsx')     # the workbook COPY; LIVE is never opened
DESK_SRC = os.path.join(WT, 'fec', 'surfaces', 'make_member_desk.py')
BLEND_SRC = os.path.join(WT, 'fec', 'stance', 'build_media_blend.py')

# The 13 columns make_member_desk.load_press_and_bills reads with pv(r, k).
DESK_READ = ('member', 'bioguide', 'url', 'stance_src', 'score', 'stance_sign',
             'score_graded', 'title', 'stance', 'stance_conf', 'date', 'categories',
             'relevance')
SIGN = {'support': 1.0, 'oppose': -1.0, 'neutral': 0.0}


def banner(text):
    print('=' * 100)
    print('%s   run %s' % (text, datetime.datetime.now().isoformat(timespec='seconds')))
    print('=' * 100)


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as handle:
        for block in iter(lambda: handle.read(1 << 20), b''):
            digest.update(block)
    return digest.hexdigest()


def mtime(path):
    if not os.path.exists(path):
        return 'MISSING'
    return datetime.datetime.fromtimestamp(os.path.getmtime(path)).isoformat(timespec='seconds')


def read_csv_rows(path, encoding='utf-8-sig'):
    if not path or not os.path.exists(path):
        return []
    with io.open(path, encoding=encoding, newline='') as handle:
        return list(csv.DictReader(handle))


def sheet_rows():
    """(rows, header_index) from the workbook COPY in SCRATCH, read_only."""
    import openpyxl
    if not os.path.exists(WB_SCRATCH):
        raise SystemExit('copy the workbook to %s first' % WB_SCRATCH)
    wb = openpyxl.load_workbook(WB_SCRATCH, read_only=True, data_only=True)
    pw = wb['Press Release Articles']
    header = next(pw.iter_rows(min_row=1, max_row=1, values_only=True))
    ph = {h: i for i, h in enumerate(header)}
    rows = list(pw.iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows, ph


def sheet_index_by_key(rows, ph):
    """release key -> (first sheet row under it, count of sheet rows under it).

    setdefault, exactly as fixN1/n1_desk_press_union.py section 5 does, so the number this
    reproduces is the number that receipt reported.
    """
    first = {}
    count = {}
    for row in rows:
        url = str(row[ph['url']] or '')
        if url.startswith('billstatus:'):
            continue
        key = DPS.row_key(url, row[ph['bioguide']], row[ph['member']])
        if key[0]:
            first.setdefault(key, row)
            count[key] = count.get(key, 0) + 1
    return first, count


def norm_field(field, sheet_value, csv_value):
    """The fixN1 section-5 normalisation, verbatim, so the counts reproduce byte for byte."""
    sv, ev = sheet_value, csv_value
    if field == 'date':
        sv = str(sv)[:10] if str(sv)[:4].isdigit() else ''
        ev = DPS.engine_date(ev)
        ev = str(ev)[:10] if str(ev)[:4].isdigit() else ''
    elif field in ('score', 'stance_sign'):
        sv = str(sv or 0).strip()
        ev = str(int(float(ev))) if str(ev).strip() not in ('', 'None') else '0'
        sv = str(int(float(sv))) if sv not in ('', 'None') else '0'
    elif field in ('score_graded', 'stance_conf'):
        sv = '%.6f' % float(sv or 0)
        try:
            ev = '%.6f' % float(ev or 0)
        except (TypeError, ValueError):
            ev = str(ev)
    else:
        sv, ev = str(sv or '').strip(), str(ev or '').strip()
    return sv, ev


def lift_desk_pieces():
    """The desk's press constants/helpers and its press loop, lifted as TEXT.

    make_member_desk.py executes at import and rewrites the live desk, so it is never
    imported.  This is the fec/tests/stance/test_crosspost_scoring.py pattern that
    fixN1/n1_desk_press_union.py used.
    """
    text = io.open(DESK_SRC, encoding='utf-8').read()
    tree = ast.parse(text)
    want_funcs = {'graded_label', '_fsign', 'norm_cats', '_press_corr'}
    want_names = {'LBL2CAT', '_GRADED'}
    nodes = [n for n in tree.body
             if (isinstance(n, ast.FunctionDef) and n.name in want_funcs)
             or (isinstance(n, ast.Assign)
                 and any(isinstance(t, ast.Name) and t.id in want_names for t in n.targets))]
    got_funcs = {n.name for n in nodes if isinstance(n, ast.FunctionDef)}
    assert got_funcs == want_funcs, sorted(got_funcs)
    lines = text.splitlines()
    start = next(i for i, ln in enumerate(lines) if ln.strip() == 'for r in _press_rows:')
    end = next(i for i, ln in enumerate(lines)
               if i > start and ln.strip().startswith('# --- social'))
    loop_src = '\n'.join(ln[4:] for ln in lines[start:end])
    assert '_press_corr(title' in loop_src and 'PDATE' in loop_src
    return nodes, loop_src, (start + 1, end)


def desk_overrides():
    """PRESSOV as the desk builds it, both overlay positions.

    make_member_desk.py:169 uses config_overlay.merged_rows(..., data_root=B) with
    CODE_ROOT = the tree the code runs from.  In PRODUCTION that tree IS the live root, so the
    two candidate paths collapse to one and only the live file is read.  Run from a worktree
    the worktree's tracked copy is overlaid and WINS by identity.  Both are returned; every
    number is reported under the production position and the worktree position is reported
    beside it.
    """
    import config_overlay as CO
    live_only = CO.merged_rows('data/stance/press_stance_overrides.json', 'corrections',
                               ('url', 'title_contains'), data_root=LIVE, code_root=LIVE)
    with_overlay = CO.merged_rows('data/stance/press_stance_overrides.json', 'corrections',
                                  ('url', 'title_contains'), data_root=LIVE, code_root=WT)
    return live_only, with_overlay


def press_corr_factory(pressov):
    """make_member_desk._press_corr, verbatim (lines 171-176): url EXACT, title substring."""
    def _press_corr(title, url):
        t = (title or '').lower()
        for c in pressov:
            if ((c.get('url') and c['url'] == url)
                    or (c.get('title_contains') and c['title_contains'].lower() in t)):
                return c
        return None
    return _press_corr


def press_gold_factory(pov):
    """build_media_blend._press_gold, verbatim (lines 340-347): url rstrip('/'), substring."""
    def _press_gold(url, title):
        ul = (url or '').strip().rstrip('/')
        tl = (title or '').lower()
        for g in pov:
            st = (g.get('stance') or '').strip().lower()
            if st not in SIGN:
                continue
            if g.get('url') and (g.get('url') or '').strip().rstrip('/') == ul:
                return st
            if g.get('title_contains') and (g.get('title_contains') or '').lower() in tl:
                return st
        return None
    return _press_gold


def stack_map():
    """build_media_blend lines 331-337: press_stance_final.csv -> {turn: (stack_label, tier)}."""
    out = {}
    path = os.path.join(LIVE, 'press_stance_final.csv')
    for row in read_csv_rows(path):
        key = (PS.canonical_url(row.get('url'), keep_fragment=True),
               (row.get('bioguide') or row.get('member') or '').strip().upper())
        out[key] = ((row.get('stack_label') or '').strip().lower(),
                    (row.get('tier') or '').strip().lower())
    return out, path


TIER_W = {'unanimous': 1.0, 'majority': 0.7, 'review': 0.35}


def dump(name, payload):
    path = os.path.join(HERE, name)
    with io.open(path, 'w', encoding='utf-8') as handle:
        json.dump(payload, handle, indent=1, sort_keys=True, default=str)
    print('wrote %s' % name)
    return path
