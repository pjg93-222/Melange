# -*- coding: utf-8 -*-
"""N14 Q3.  For each of the 57, which side equals the OWNER'S OWN recorded label?  READ-ONLY.

Establishes facts only.  No label is proposed and no label is judged: the single question is
which SOURCE carries the owner's own most recent recorded value.

Owner-record files consulted (named by the lane brief; nothing else in Draft Systems is read):
  cu_stance_label_set.csv  (Draft)     human_label column -- the adjudicated label set
  cu_stance_label_set.xlsx (Draft, LIVE root)                the xlsx mirrors
  press_stance_overrides.json          Paul-gold corrections, url-exact or title substring
  press_overrides_from_dashboard.json  the dashboard-exported override queue
  cfpb_stance_review_queue_2026-07-16.csv   PAUL_GOLD column
  CU_Press_Stance_Tone.xlsx / cu_press_releases_scored.xlsx  (the R pipeline's own presentation)

Also: quantifies the stance_src flattening, and says what score_graded IS in each source and
which source the RANKING ENGINE actually reads when it scores.
"""
import collections
import io
import json
import os
import re

import openpyxl

import n14_common as C

C.banner('N14 Q3  THE OWNER RECORD')

q1 = json.load(io.open(os.path.join(C.HERE, 'n14_q1_disagreement_sets.json'), encoding='utf-8'))
the57 = q1['stance_disagreement_rows']
print('the 57 read back from n14_q1_disagreement_sets.json: %d rows' % len(the57))


def norm_title(value):
    return ' '.join(str(value or '').split()).casefold()


# ---- A. cu_stance_label_set.csv: the owner's adjudicated label set -----------------------
LS_PATH = os.path.join(C.SCRATCH, 'cu_stance_label_set.csv')
ls_rows = C.read_csv_rows(LS_PATH)
print('\n-- A. cu_stance_label_set.csv (Draft, mtime %s) --'
      % C.mtime(os.path.join(C.DRAFT, 'cu_stance_label_set.csv')))
print('   rows %d   columns %s' % (len(ls_rows), list(ls_rows[0]) if ls_rows else []))
hl = collections.Counter((r.get('human_label') or '(blank)').strip().lower() for r in ls_rows)
pl = collections.Counter((r.get('pre_label') or '(blank)').strip().lower() for r in ls_rows)
print('   human_label census : %s' % dict(hl.most_common()))
print('   pre_label census   : %s' % dict(pl.most_common()))
print('   split census       : %s'
      % dict(collections.Counter((r.get('split') or '?') for r in ls_rows).most_common()))
print('   rows where human_label != pre_label: %d'
      % sum(1 for r in ls_rows
            if (r.get('human_label') or '').strip().lower()
            != (r.get('pre_label') or '').strip().lower()))
ls_by_key = {}
ls_by_title = collections.defaultdict(list)
for r in ls_rows:
    key = C.DPS.row_key(r.get('url'), None, r.get('member'))
    if key[0]:
        ls_by_key.setdefault(key, []).append(r)
    ls_by_title[norm_title(r.get('title'))].append(r)
print('   distinct release keys in the label set: %d' % len(ls_by_key))

# ---- B. the override ledgers ------------------------------------------------------------
live_ov, wt_ov = C.desk_overrides()
print('\n-- B. press_stance_overrides.json -- %d corrections (production position: LIVE code root)'
      % len(live_ov))
print('   worktree-overlay position: %d corrections (Q2b proved the two are label-identical)'
      % len(wt_ov))
desk_corr = C.press_corr_factory(live_ov)     # make_member_desk._press_corr, url EXACT
engine_gold = C.press_gold_factory(live_ov)   # build_media_blend._press_gold, url rstrip('/')

dash = json.load(io.open(os.path.join(C.LIVE, 'data', 'stance',
                                      'press_overrides_from_dashboard.json'),
                         encoding='utf-8')).get('corrections') or []
dash_by_url = {(c.get('url') or '').strip().rstrip('/'): c for c in dash if c.get('url')}
print('   press_overrides_from_dashboard.json: %d rows' % len(dash))

# ---- C. the CFPB review queue -----------------------------------------------------------
cq = C.read_csv_rows(os.path.join(C.LIVE, 'data', 'stance',
                                  'cfpb_stance_review_queue_2026-07-16.csv'))
gold_filled = [r for r in cq if (r.get('PAUL_GOLD') or '').strip()]
print('\n-- C. cfpb_stance_review_queue_2026-07-16.csv --')
print('   rows %d | PAUL_GOLD filled %d | queue_reason census %s'
      % (len(cq), len(gold_filled),
         dict(collections.Counter((r.get('queue_reason') or '?') for r in cq).most_common())))
print('   id shape sample: %s' % [((r.get('id') or '')[:44]) for r in cq[:4]])
print('   ids that look like a press URL: %d'
      % sum(1 for r in cq if str(r.get('id') or '').startswith('http')))
cq_by_url = {}
for r in cq:
    ident = str(r.get('id') or '')
    if ident.startswith('http'):
        key = C.DPS.row_key(ident, None, r.get('member'))
        if key[0]:
            cq_by_url.setdefault(key, r)
print('   press-URL-keyed queue rows: %d' % len(cq_by_url))

# ---- D. the R pipeline's own presentation workbooks -------------------------------------
print('\n-- D. THE R PIPELINE\'S OWN PRESENTATION WORKBOOKS (case-insensitive headers) --')
keys57 = {(r['release_key'], r['identity']): r for r in the57}
r_workbooks = {}
for label, filename in (('Draft CU_Press_Stance_Tone.xlsx', 'tone_draft.xlsx'),
                        ('LIVE  CU_Press_Stance_Tone.xlsx', 'tone_live.xlsx'),
                        ('Draft cu_press_releases_scored.xlsx', 'scored_draft.xlsx'),
                        ('LIVE  cu_press_releases_scored.xlsx', 'scored_live.xlsx'),
                        ('Draft cu_stance_label_set.xlsx', 'labelset_draft.xlsx'),
                        ('LIVE  cu_stance_label_set.xlsx', 'labelset_live.xlsx')):
    path = os.path.join(C.SCRATCH, filename)
    if not os.path.exists(path):
        print('   %-38s MISSING' % label)
        continue
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        hidx = {str(h).strip().lower().replace(' ', '_'): i
                for i, h in enumerate(header) if h}
        stance_col = next((c for c in ('stance', 'human_label', 'label') if c in hidx), None)
        if 'url' not in hidx or stance_col is None:
            continue
        by_key = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            url = str(row[hidx['url']] or '')
            if not url or url.startswith('billstatus:'):
                continue
            key = C.DPS.row_key(url, row[hidx['bioguide']] if 'bioguide' in hidx else None,
                                row[hidx['member']] if 'member' in hidx else '')
            if key[0]:
                by_key.setdefault(key, row)
        agree_sheet = agree_csv = neither = 0
        found = 0
        for key, rec in keys57.items():
            row = by_key.get(key)
            if row is None:
                continue
            found += 1
            stance = str(row[hidx[stance_col]] or '').strip().lower()
            if stance == rec['sheet_stance']:
                agree_sheet += 1
            elif stance == rec['csv_stance']:
                agree_csv += 1
            else:
                neither += 1
        print('   %-38s [%-22s] keyed %5d | col %-11s | of 57 found %2d ->'
              ' SHEET %2d  CSV %2d  neither %2d'
              % (label, name[:22], len(by_key), stance_col, found,
                 agree_sheet, agree_csv, neither))
        r_workbooks['%s::%s' % (label.strip(), name)] = {
            'rows_keyed': len(by_key), 'stance_column': stance_col, 'found_of_57': found,
            'matches_sheet': agree_sheet, 'matches_csv': agree_csv, 'matches_neither': neither}
    wb.close()

# ---- E. the per-release adjudication ----------------------------------------------------
print('\n' + '-' * 100)
print('-- E. FOR EACH OF THE 57: WHICH SIDE EQUALS THE OWNER\'S OWN RECORD? --')
verdicts = collections.Counter()
by_source = collections.Counter()
detail = []
for rec in the57:
    key = (rec['release_key'], rec['identity'])
    owner = None
    source = None
    # precedence: the most SPECIFIC and most RECENT owner record first
    hit = engine_gold(rec['url'], rec['title']) or engine_gold(rec['csv_url'], rec['csv_title'])
    if hit:
        owner, source = hit, 'press_stance_overrides.json'
    if owner is None:
        dash_hit = dash_by_url.get((rec['url'] or '').strip().rstrip('/'))
        if dash_hit and (dash_hit.get('stance') or '').lower() in C.SIGN:
            owner, source = dash_hit['stance'].lower(), 'press_overrides_from_dashboard.json'
    if owner is None and key in cq_by_url:
        gold = (cq_by_url[key].get('PAUL_GOLD') or '').strip().lower()
        if gold in C.SIGN:
            owner, source = gold, 'cfpb_stance_review_queue_2026-07-16.csv'
    if owner is None:
        cands = ls_by_key.get(key) or []
        if not cands:
            cands = [r for r in ls_by_title.get(norm_title(rec['title']), [])]
        labels = {(r.get('human_label') or '').strip().lower() for r in cands}
        labels = {x for x in labels if x in C.SIGN}
        if len(labels) == 1:
            owner, source = labels.pop(), 'cu_stance_label_set.csv:human_label'
        elif len(labels) > 1:
            owner, source = 'AMBIGUOUS:%s' % '/'.join(sorted(labels)), \
                            'cu_stance_label_set.csv:human_label'
    if owner is None:
        verdict = 'no owner record found'
    elif owner.startswith('AMBIGUOUS'):
        verdict = 'owner record ambiguous'
    else:
        s_hit = owner == rec['sheet_stance']
        c_hit = owner == rec['csv_stance']
        verdict = ('both' if s_hit and c_hit else
                   'CSV matches the owner record' if c_hit else
                   'sheet matches the owner record' if s_hit else 'neither')
    verdicts[verdict] += 1
    by_source[source or '(none)'] += 1
    detail.append(dict(rec, owner_label=owner, owner_source=source, verdict=verdict))

print('\n%-34s %6s' % ('verdict', 'count'))
for verdict, n in verdicts.most_common():
    print('%-34s %6d' % (verdict, n))
print('\n%-44s %6s' % ('owner record came from', 'count'))
for source, n in by_source.most_common():
    print('%-44s %6d' % (source, n))

print('\nthe rows whose owner record does NOT come from the label set:')
for d in detail:
    if d['owner_source'] and d['owner_source'] != 'cu_stance_label_set.csv:human_label':
        print('   %-24s %-58s owner=%-8s src=%s -> %s'
              % (d['member'][:24], d['title'][:58], d['owner_label'],
                 d['owner_source'][:34], d['verdict']))
print('\nthe rows with NO owner record at all:')
none_rows = [d for d in detail if d['owner_source'] is None]
for d in none_rows:
    print('   %-24s %-62s sheet=%-8s csv=%s'
          % (d['member'][:24], d['title'][:62], d['sheet_stance'], d['csv_stance']))
print('   (%d rows)' % len(none_rows))

# ---- F. stance_src: quantify the flattening --------------------------------------------
print('\n' + '-' * 100)
print('-- F. stance_src: THE human_label -> debate FLATTENING --')
rows, ph = C.sheet_rows()
sheet_src = collections.Counter()
for row in rows:
    url = str(row[ph['url']] or '')
    if url.startswith('billstatus:'):
        sheet_src['billstatus(bill)'] += 1
        continue
    sheet_src[str(row[ph['stance_src']] or '(blank)')] += 1
base_src = collections.Counter((r.get('stance_src') or '(blank)')
                               for r in C.read_csv_rows(C.PS.BASE_TONE_PATH))
refresh_src = collections.Counter((r.get('stance_src') or '(blank)')
                                  for r in C.read_csv_rows(C.PS.REFRESH_TONE_PATH))
print('   sheet press rows by stance_src        : %s' % dict(sheet_src.most_common()))
print('   BASE tone CSV rows by stance_src      : %s' % dict(base_src.most_common()))
print('   REFRESH tone CSV rows by stance_src   : %s' % dict(refresh_src.most_common()))
print("   sheet rows carrying 'human_label'     : %d" % sheet_src.get('human_label', 0))
print('   overlap keys where sheet=debate and csv=human_label: %d keys / %d engine rows'
      % (q1['disagreement_by_release_key']['stance_src'],
         q1['disagreement_by_engine_row']['stance_src']))
print("   base CSV human_label rows NOT reached by the overlap: %d"
      % (base_src.get('human_label', 0)
         - q1['disagreement_by_release_key']['stance_src']))
print('   of the %d human_label keys in the overlap, %d also disagree on the STANCE VALUE'
      % (q1['disagreement_by_release_key']['stance_src'], len(the57)))
print('   -> %d carry the owner-adjudicated label AND the same value as the model'
      % (q1['disagreement_by_release_key']['stance_src'] - len(the57)))

# ---- G. score_graded ---------------------------------------------------------------------
print('\n' + '-' * 100)
print('-- G. WHAT score_graded IS IN EACH SOURCE, AND WHAT THE ENGINE READS --')
blend = io.open(C.BLEND_SRC, encoding='utf-8').read()
print("   'score_graded' occurrences in fec/stance/build_media_blend.py: %d"
      % blend.count('score_graded'))
print("   'press_value'  occurrences in fec/stance/build_media_blend.py: %d"
      % blend.count('press_value'))
os.chdir(os.path.join(C.WT, 'fec'))
readers = []
for root, dirs, files in os.walk('.'):
    dirs[:] = [d for d in dirs
               if d not in ('attic', '_quarantine_prereorg_strays_20260810', '__pycache__')]
    for name in files:
        if name.endswith('.py'):
            path = os.path.join(root, name)
            text = io.open(path, encoding='utf-8', errors='replace').read()
            if 'score_graded' in text:
                readers.append((path.replace('\\', '/'), text.count('score_graded')))
print('   every non-attic repo file naming score_graded:')
for path, n in sorted(readers):
    print('      %-56s %d' % (path, n))
print('\n   -> build_media_blend.py (the ranking engine\'s press component) NEVER reads')
print('      score_graded.  It signs a release from press_stance_final.csv stack_label, falls')
print('      back to the CSV stance column, then applies press_stance_overrides; INTENSITY is')
print('      abs(press_value) (line 401), decayed by date (line 406).  So score_graded is a')
print('      DISPLAY column only -- make_member_desk line 1684 (mag) and 1701 (graded_label).')

# which grading run produced each side's score_graded
graded_zero_csv = 0
graded_zero_sheet = 0
by_src = collections.Counter()
base_by_key = {}
for r in C.read_csv_rows(C.PS.BASE_TONE_PATH):
    key = C.DPS.row_key(r.get('url'), r.get('bioguide'), r.get('member'))
    if key[0]:
        base_by_key.setdefault(key, r)
sheet_by_key, _ = C.sheet_index_by_key(rows, ph)
both = set(base_by_key) & set(sheet_by_key)
for key in both:
    br, sr = base_by_key[key], sheet_by_key[key]
    sv, ev = C.norm_field('score_graded', sr[ph['score_graded']], br.get('score_graded'))
    if sv != ev:
        by_src['csv stance_src=%s' % (br.get('stance_src') or '?')] += 1
        if float(ev) == 0.0:
            graded_zero_csv += 1
        if float(sv) == 0.0:
            graded_zero_sheet += 1
print('\n   score_graded disagreement on the BASE-CSV overlap (%d keys), by the CSV\'s'
      ' stance_src: %s' % (len(both), dict(by_src.most_common())))
print('   csv side exactly 0: %d | sheet side exactly 0: %d' % (graded_zero_csv, graded_zero_sheet))
base_graded = collections.Counter()
for r in C.read_csv_rows(C.PS.BASE_TONE_PATH):
    try:
        base_graded['zero' if float(r.get('score_graded') or 0) == 0.0 else 'nonzero'] += 1
    except (TypeError, ValueError):
        base_graded['unparseable:%r' % r.get('score_graded')] += 1
sheet_graded = collections.Counter()
for row in rows:
    if str(row[ph['url']] or '').startswith('billstatus:'):
        continue
    try:
        sheet_graded['zero' if float(row[ph['score_graded']] or 0) == 0.0 else 'nonzero'] += 1
    except (TypeError, ValueError):
        sheet_graded['unparseable'] += 1
print('   BASE CSV score_graded: %s' % dict(base_graded.most_common()))
print('   SHEET    score_graded: %s' % dict(sheet_graded.most_common()))

C.dump('n14_q3_owner_records.json', {
    'schema': 'n14-q3-owner-records-v1',
    'label_set': {'rows': len(ls_rows), 'release_keys': len(ls_by_key),
                  'human_label_census': dict(hl), 'pre_label_census': dict(pl)},
    'override_ledger_rows': len(live_ov), 'dashboard_queue_rows': len(dash),
    'cfpb_queue_rows': len(cq), 'cfpb_queue_paul_gold_filled': len(gold_filled),
    'cfpb_queue_press_url_keyed': len(cq_by_url),
    'r_pipeline_workbooks': r_workbooks,
    'verdicts': dict(verdicts), 'owner_record_sources': dict(by_source),
    'rows': detail,
    'stance_src': {'sheet': dict(sheet_src), 'base_csv': dict(base_src),
                   'refresh_csv': dict(refresh_src)},
    'score_graded': {'engine_reads_it': False,
                     'blend_score_graded_mentions': blend.count('score_graded'),
                     'blend_press_value_mentions': blend.count('press_value'),
                     'repo_readers': dict(readers),
                     'disagreement_by_csv_stance_src': dict(by_src),
                     'base_csv_census': dict(base_graded),
                     'sheet_census': dict(sheet_graded)},
})
print('DONE N14 Q3')
