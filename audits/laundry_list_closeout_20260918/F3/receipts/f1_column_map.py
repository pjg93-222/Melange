# -*- coding: utf-8 -*-
# F3 r1: the COLUMN MAP of the stance path.
#   'Press Release - Raw' headers E..X -> press_matrix() COLS -> per_category_stance cats[]
#   plus row-535 press-pillar weights, the ORIG gate, and the General-CU exclusion ledger.
# READ-ONLY.  Nothing under LIVE\fec is imported.  Workbook read from a SCRATCH copy with
# production semantics (data_only=True, NOT read_only -- mirrors LIVE fec/moneyball_core.py:91).
import os, sys, json, shutil, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import f_lib as F
import openpyxl
from openpyxl.utils import get_column_letter

LIVE    = F.LIVE
SCRATCH = r"C:\Users\pgoff\AppData\Local\Temp\claude\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\F3"
HERE    = os.path.dirname(os.path.abspath(__file__))

WB_LIVE = os.path.join(LIVE, "Moneyball Main Version - Blended.xlsx")
WB_COPY = os.path.join(SCRATCH, "wb_copy.xlsx")
if not os.path.exists(WB_COPY) or os.path.getmtime(WB_COPY) < os.path.getmtime(WB_LIVE):
    shutil.copy2(WB_LIVE, WB_COPY)

print("=" * 108)
print("F3 r1  COLUMN MAP: Raw sheet -> press_matrix -> per_category_stance    run %s"
      % datetime.datetime.now().isoformat(timespec="seconds"))
print("=" * 108)
print("workbook %s" % WB_LIVE)
print("   mtime %s  size %d B"
      % (datetime.datetime.fromtimestamp(os.path.getmtime(WB_LIVE)).isoformat(timespec="seconds"),
         os.path.getsize(WB_LIVE)))
print("   lock at LIVE root present: %s" % os.path.exists(os.path.join(LIVE, ".pipeline_run.lock")))

probs = F.verify_transcriptions()
print("\ntranscription self-check: %s" % ("OK" if not probs else "FAILED -> " + "; ".join(probs)))
if probs:
    sys.exit(2)

MR       = F.media_rules()
RAW2ST   = MR["RAW_TO_STANCE"]
GENERAL  = MR["GENERAL"]
COMPET   = MR["COMPETITIVE"]
PATTERNS = MR["RAW_PATTERNS"]
STANCE_COL = {pc: lab for lab, pc in F.CATS}

wb = openpyxl.load_workbook(WB_COPY, data_only=True)
rw = wb["Press Release - Raw"]
pn = wb["Press Release - Normalized"]
ra = wb["Rank & Allocation"]

# LIVE fec/moneyball_core.py:93  rows 2..534 ; LIVE fec/stance/build_media_blend.py:286-287 rows 2..535
rrow       = {rw.cell(r, 1).value: r for r in range(2, 535) if rw.cell(r, 1).value}
mrow_blend = {rw.cell(r, 1).value: r for r in range(2, 536) if rw.cell(r, 1).value}
print("\nraw-sheet member rows   press_matrix range(2,535) -> %d    blend range(2,536) -> %d"
      % (len(rrow), len(mrow_blend)))
extra = set(mrow_blend) - set(rrow)
print("   written-by-blend but NOT-read-by-press_matrix: %s" % (sorted(extra) if extra else "none"))

print("\n%-4s %-3s %-42s %-26s %-9s %s"
      % ("col", "ltr", "Raw-sheet header", "stance column", "w535", "fate in the stance path"))
print("-" * 132)
rows_out = []
for c in F.ALL_COLS:
    hname = rw.cell(1, c).value
    w535  = pn.cell(535, c).value
    if c in STANCE_COL:
        fate, scol = "SCORED as media evidence", STANCE_COL[c]
    elif c in F.COLS:
        fate, scol = "normalized, then NEVER READ (absent from cats[])", "(none)"
    else:
        fate, scol = "NOT normalized; exclusion ledger only", "(none)"
    print("%-4d %-3s %-42s %-26s %-9s %s"
          % (c, get_column_letter(c), hname, scol,
             ("%.4f" % w535) if isinstance(w535, (int, float)) else str(w535), fate))
    rows_out.append({"col": c, "letter": get_column_letter(c), "header": hname,
                     "stance_column": scol, "w535": w535, "fate": fate,
                     "alias_target": RAW2ST.get(hname)})

print("\n-- alias / list consistency --")
print("   RAW_TO_STANCE entries (media_rules.py:11-30)      : %d" % len(RAW2ST))
print("   per_category_stance cats[] (per_category:28-31)   : %d" % len(F.CATS))
print("   RAW_PATTERNS keys == blend HDRMAP (blend.py:50/53): %d  <- columns the blend WRITES" % len(PATTERNS))
print("   press_matrix COLS (moneyball_core.py:94)          : %d" % len(F.COLS))
hdrs = {rw.cell(1, c).value for c in F.ALL_COLS}
print("   alias targets == cats[] names                    : %s"
      % (set(RAW2ST.values()) == {c for c, _ in F.CATS}))
print("   RAW_PATTERNS keys == sheet headers E..X          : %s" % (set(PATTERNS) == hdrs))
if set(PATTERNS) != hdrs:
    print("      pattern-only %s / sheet-only %s"
          % (sorted(set(PATTERNS) - hdrs), sorted(hdrs - set(PATTERNS))))
print("   headers with NO alias entry                      : %s"
      % sorted(h for h in hdrs if h not in RAW2ST))

# --- ORIG gate: LIVE fec/stance/build_media_blend.py:296, 360, 406 ---------------------
PRO  = os.path.join(LIVE, "data", "stance", "press_raw_original.json")
orig = json.load(open(PRO, encoding="utf-8"))
print("\n-- ORIG gate  (blend.py:360  `_cats=[c for c in _route['categories'] if c in ORIG]`) --")
print("   press_raw_original.json keys: %d   (mtime %s)"
      % (len(orig), datetime.datetime.fromtimestamp(os.path.getmtime(PRO)).isoformat(timespec="seconds")))
for h in sorted(hdrs, key=str):
    print("      %-42s in ORIG: %s" % (h, h in orig))
miss = sorted(h for h in hdrs if h not in orig)
print("   headers that are NOT ORIG keys -> press routed there is dropped BEFORE the sheet: %s" % miss)
print("   ORIG keys that are not sheet headers: %s" % sorted(set(orig) - hdrs))

# --- nonzero census over the roster per_category_stance uses ---------------------------
members = [F.nfc(ra.cell(r, 1).value) for r in range(3, 536) if ra.cell(r, 1).value]
print("\nRank & Allocation roster (per_category_stance.py:17): %d members" % len(members))

raw_all = {c: {m: (rw.cell(rrow[m], c).value or 0.0) for m in rrow} for c in F.ALL_COLS}
rawnfc  = {c: {F.nfc(m): v for m, v in raw_all[c].items()} for c in F.ALL_COLS}
print("\n%-4s %-42s %-8s %-8s %-8s %s" % ("col", "header", "nonzero", "pos", "neg", "nonzero on R&A roster"))
print("-" * 120)
census = {}
for c in F.ALL_COLS:
    v = rawnfc[c]
    nz  = sum(1 for x in v.values() if abs(float(x or 0.0)) > 1e-12)
    pos = sum(1 for x in v.values() if float(x or 0.0) > 1e-12)
    neg = sum(1 for x in v.values() if float(x or 0.0) < -1e-12)
    onr = sum(1 for m in members if abs(float(v.get(m, 0.0) or 0.0)) > 1e-12)
    print("%-4d %-42s %-8d %-8d %-8d %d" % (c, rw.cell(1, c).value, nz, pos, neg, onr))
    census[c] = {"header": rw.cell(1, c).value, "nonzero": nz, "pos": pos, "neg": neg, "nonzero_on_roster": onr}

print("\n-- columns that carry media but reach NO stance cell --")
for c in [x for x in F.COLS if x not in STANCE_COL]:
    print("   col %-2d %-40s %d members nonzero | row535 press-pillar weight %s | NO stance column"
          % (c, rw.cell(1, c).value, census[c]["nonzero_on_roster"], pn.cell(535, c).value))
print("   col 24 %-40s %d members nonzero | never normalized -> exclusion ledger"
      % (rw.cell(1, 24).value, census[24]["nonzero_on_roster"]))

led = F.general_exclusion_ledger(raw_all[24])
print("\n-- general_exclusion_ledger(raw_all[24]) reproduced (media_rules.py:192-201) --")
print("   ledger members %d ; equals col-24 nonzero census: %s" % (len(led), len(led) == census[24]["nonzero"]))
print("   all applied_to_policy_domains == False: %s"
      % all(v["applied_to_policy_domains"] is False for v in led.values()))

# --- what category_stance.json itself records for the media block ----------------------
CS = os.path.join(LIVE, "category_stance.json")
cs = json.load(open(CS, encoding="utf-8"))
gce = cs.get("media", {}).get("general_credit_union_exclusions", {})
print("\n-- category_stance.json media block --")
print("   cats recorded            : %d" % len(cs["cats"]))
print("   general_credit_union_exclusions members: %d" % len(gce))
print("   matches ledger rebuilt from TODAY's workbook: %s" % (len(gce) == len(led)))
print("   NOTE category_stance.json mtime %s vs workbook mtime %s"
      % (datetime.datetime.fromtimestamp(os.path.getmtime(CS)).isoformat(timespec="seconds"),
         datetime.datetime.fromtimestamp(os.path.getmtime(WB_LIVE)).isoformat(timespec="seconds")))

json.dump({"columns": rows_out, "census": {str(k): v for k, v in census.items()},
           "roster_n": len(members), "headers_missing_from_orig": miss,
           "general_ledger_n": len(led), "artifact_gce_n": len(gce),
           "raw_nonzero_members_by_col": {str(c): sorted(m for m, v in rawnfc[c].items()
                                                         if abs(float(v or 0.0)) > 1e-12)
                                          for c in F.ALL_COLS}},
          open(os.path.join(HERE, "f1_column_map.json"), "w", encoding="utf-8"), indent=1, sort_keys=True)
wb.close()
print("\nDONE F3 r1")
