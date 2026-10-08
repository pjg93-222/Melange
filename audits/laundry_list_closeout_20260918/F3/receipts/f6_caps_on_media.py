# -*- coding: utf-8 -*-
# F3 r6: the CAPS and WEIGHTS that act on media evidence inside the stance producer, and how many
# media-bearing cells they actually move.  Closes the last part of M1/M2.
# READ-ONLY.  Nothing under LIVE\fec is imported.
import os, sys, json, shutil, collections, datetime
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import f_lib as F
import openpyxl

LIVE = F.LIVE
SCRATCH = r"C:\Users\pgoff\AppData\Local\Temp\claude\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\F3"
CATS = F.CATS

print("=" * 108)
print("F3 r6  CAPS and WEIGHTS acting on media inside the stance producer   run %s"
      % datetime.datetime.now().isoformat(timespec="seconds"))
print("=" * 108)

wb = openpyxl.load_workbook(os.path.join(SCRATCH, "wb_copy.xlsx"), data_only=True)
rw, ra = wb["Press Release - Raw"], wb["Rank & Allocation"]
rrow = {rw.cell(r, 1).value: r for r in range(2, 535) if rw.cell(r, 1).value}
members = [F.nfc(ra.cell(r, 1).value) for r in range(3, 536) if ra.cell(r, 1).value]
rawnfc = {c: {F.nfc(m): (rw.cell(rrow[m], c).value or 0.0) for m in rrow} for c in F.COLS}
norm = {c: F.normalize_signed_extrema(rawnfc[c]) for c in F.COLS}
press = {(m, lab): norm[pc].get(m, 0.0) for m in members for lab, pc in CATS}
has_press = {k: abs(v) > 1e-9 for k, v in press.items()}
wb.close()

CS = json.load(open(os.path.join(LIVE, "category_stance.json"), encoding="utf-8"))

print("\n(1) WEIGHT the media term carries inside a cell")
print("    LIVE fec/engine/pillar_policy.py:22-25  LABEL_EVIDENCE_WEIGHTS")
print("       standard : legislation 0.55 | media 0.10 | vote 0.10")
print("       defensive: legislation 0.55 | media 0.10 | vote 0.20")
print("    LIVE fec/stance/per_category_stance.py:36  DEFENSIVE = %s" % sorted(F.DEFENSIVE))
print("    LIVE fec/stance/stance_rules.py:26  MEDIA_SCALE = 1.0")
print("    LIVE fec/stance/per_category_stance.py:61  pn=clip(prraw[m]/SR.MEDIA_SCALE)")
print("       -> the normalised media value enters DIRECTLY on [-1,1], no further scaling")
print("    LIVE fec/stance/stance_rules.py:59-62  combine_evidence = weighted mean over PRESENT")
print("       sources only; an absent source is NOT in the denominator.  So in a cell where media")
print("       is the ONLY present source, its effective weight is 1.0, not 0.10.")
media_only = [(m, lab) for m in members for lab, _ in CATS
              if has_press[(m, lab)] and CS["signal"].get(m, {}).get(lab)]
# decompose: media-only == has_press and the cell's score equals the media term alone.  We cannot
# recompute has_leg/has_vote without the leg/vote inputs, so read them.
LEGR = json.load(open(os.path.join(LIVE, "legislation_routing_audit.json"), encoding="utf-8"))
LEG_SIGNAL = {(F.nfc(r["member"]), r["effective_category"]) for r in LEGR["rows"]
              if r.get("disposition") == "selected"}
VC = json.load(open(os.path.join(LIVE, "vote_category_lean.json"), encoding="utf-8"))
n2b = {F.nfc(v): k for k, v in VC["name_of"].items()}
votes = VC["votes"]
def has_vote(m, lab):
    vi = votes.get(n2b.get(m, ""), {}).get(lab)
    return bool(vi) and (vi.get("n", 0) or 0) >= 5          # per_category_stance.py:62
rows = []
for m in members:
    for lab, _ in CATS:
        hp, hl, hv = has_press[(m, lab)], (m, lab) in LEG_SIGNAL, has_vote(m, lab)
        rows.append((m, lab, hp, hl, hv))
comb = collections.Counter((r[2], r[3], r[4]) for r in rows)
print("\n    evidence-presence census over %d cells  (press, legislation, vote):" % len(rows))
for k in sorted(comb, key=lambda x: -comb[x]):
    tag = "MEDIA IS THE ONLY SOURCE -> effective media weight 1.0" if k == (True, False, False) else ""
    print("       press=%-5s leg=%-5s vote=%-5s  %6d  %s" % (k[0], k[1], k[2], comb[k], tag))
mo = comb[(True, False, False)]
print("    media-only cells: %d of %d cells that carry media (%.1f%%)"
      % (mo, sum(1 for r in rows if r[2]), 100.0 * mo / max(1, sum(1 for r in rows if r[2]))))
sig_true = sum(1 for m in members for lab, _ in CATS if CS["signal"].get(m, {}).get(lab))
rebuilt = sum(1 for r in rows if r[2] or r[3] or r[4])
print("    rebuilt signal (press OR leg OR vote): %d   artifact 'signal' True: %d   MATCH: %s"
      % (rebuilt, sig_true, rebuilt == sig_true))
disagree = [(r[0], r[1], r[2], r[3], r[4], CS["signal"].get(r[0], {}).get(r[1]))
            for r in rows if (r[2] or r[3] or r[4]) != bool(CS["signal"].get(r[0], {}).get(r[1]))]
print("    cells where the rebuild disagrees with the artifact: %d" % len(disagree))
for d in disagree[:8]:
    print("       %-24s %-26s press=%s leg=%s vote=%s artifact=%s" % d)

print("\n(2) THE ANTI-STANCE PRIOR -- a media-driven DOWN-weight")
print("    LIVE fec/stance/per_category_stance.py:47-51, 67")
print("       opl[m]=sum(press[(m,lab)]*w535[cat])            <- built ENTIRELY from media")
print("       anti0[m]=clip(-opl[m]/ascale)                   <- ascale = |8th-percentile negative|")
print("       PRIOR=0.70")
print("       :67  if c>0 and prraw[m]<=0.02 and anti[m]>0: c*=(1-PRIOR*anti[m])")
print("    -> a member whose OVERALL media tone is anti-CU has every positive cell with weak media")
print("       (<=0.02) shrunk by up to 70%.  This is the one place media in category X changes the")
print("       score in category Y.  It uses the row-535 press-pillar weights as the cross-category")
print("       mixer, INCLUDING none for col 18 (col 18 is absent from cats[], so its 0.03 weight")
print("       never enters opl either).")
pw = {}
wbp = openpyxl.load_workbook(os.path.join(SCRATCH, "wb_copy.xlsx"), data_only=True)
pn = wbp["Press Release - Normalized"]
for lab, pc in CATS:
    pw[lab] = pn.cell(535, pc).value or 0.0
wbp.close()
opl = {m: sum(press[(m, lab)] * pw[lab] for lab, _ in CATS) for m in members}
neg = sorted(v for v in opl.values() if v < 0)
ascale = abs(neg[int(0.08 * len(neg))]) if neg else 1.0
def clip(x): return max(-1.0, min(1.0, x))
anti0 = {m: clip(-opl[m] / (ascale or 1)) for m in members}
n_anti = sum(1 for m in members if anti0[m] > 0)
print("    members with anti0 > 0 (media-tone-negative)            : %d of %d" % (n_anti, len(members)))
print("    ascale (|8th pct of negative opl|)                      : %.9g" % ascale)
shrinkable = sum(1 for m in members for lab, _ in CATS
                 if anti0[m] > 0 and press[(m, lab)] <= 0.02)
print("    (member,cat) cells eligible for the <=0.02 shrink        : %d" % shrinkable)
print("    max shrink factor applied                              : %.2f (1 - 0.70*1.0)" % (1 - 0.70))

print("\n(3) THE SAY-DO CORROBORATION GATE -- a cap, and it consumes NO media")
print("    LIVE fec/moneyball_core.py:723-739  cap_uncorroborated(member, cm)")
print("       Path A = bill_corroboration()  (moneyball_core.py:625-697, workbook bill role rows)")
print("       Path V = vote_corroboration()  (moneyball_core.py:699-722, vote_category_lean.json)")
print("       :733-738  a Champion/Defender cell with neither A nor V is LOWERED to 'Ally'")
print("    NEITHER path reads media.  The gate therefore never ADDS media evidence; it only removes")
print("    top-tier labels that media evidence helped produce.")
gate = CS["gate"]
capped = gate["capped"]
n_capped = sum(len(v) for v in capped.values())
print("    cells capped Champion/Defender -> Ally : %d across %d members" % (n_capped, len(capped)))
mo_capped = 0; any_media_capped = 0
for m, cc in capped.items():
    for lab in cc:
        if has_press.get((F.nfc(m), lab)):
            any_media_capped += 1
            if not ((F.nfc(m), lab) in LEG_SIGNAL) and not has_vote(F.nfc(m), lab):
                mo_capped += 1
print("    ... of those, cells that CARRY media evidence            : %d" % any_media_capped)
print("    ... of those, cells where media is the ONLY evidence     : %d" % mo_capped)
print("    corroborated cells by path: %s"
      % dict(collections.Counter(p for cc in gate["corrob"].values() for p in cc.values())))
print("    overall crowns capped by the aggregate rule (gate.overall_capped): %d" % len(gate["overall_capped"]))
print("    proven badges (authored bill advanced past introduction) : %d members" % len(gate["proven"]))

print("\n(4) THE LABEL CURVE -- applies to media-bearing cells in declared domains only")
print("    LIVE fec/stance/per_category_stance.py:78-79  SR.curve_label_scores(scores,signal,CURVED_CATEGORIES)")
print("    curved categories (stance_rules.py:28-):", sorted(CS["curve"]["categories"]))
print("    method:", CS["curve"]["method"][:96])
print("    empty_domains_inherit: %s   raw_and_adjusted_scores_unchanged: %s"
      % (CS["curve"]["empty_domains_inherit"], CS["curve"]["raw_and_adjusted_scores_unchanged"]))

print("\n(5) DECAY applied to media BEFORE it reaches the sheet (all in build_media_blend.py)")
print("    press      0.5**(age_years/2.0)   blend.py:394   (2.0 yr half-life)")
print("    press w/o parseable date -> recency_weight column, default 0.5   blend.py:396-397")
print("    social     0.5**(age_years/0.5)   blend.py:120   (0.5 yr half-life, SOCIAL_HL_YEARS)")
print("    speech     0.5**(age_years/2.0)   blend.py:168")
print("    hearing    0.5**(age_years/2.0)   blend.py:211")
print("    interview  0.5**(age_years/2.0)   blend.py:253")
print("    intensity caps: press min(3.0,|press_value|) blend.py:389 (no floor);"
      " social max(0.25,min(1.5,margin)) blend.py:117")
print("    retweets:  inten *= RT_W=0.5      blend.py:36/119")
print("    channel weights onto the press side: speech 1.0 / hearing 0.85 / interview 1.0"
      " (blend.py:461-465)")
print("    press attribution weight: own .gov domain 1.0, third-party PRESS_THIRD_W=1.0"
      " (blend.py:347/401-402)")

json.dump({"evidence_combinations": {str(k): v for k, v in comb.items()},
           "media_only_cells": mo, "rebuilt_signal": rebuilt, "artifact_signal": sig_true,
           "signal_disagreements": len(disagree),
           "anti_members": n_anti, "ascale": ascale, "shrinkable_cells": shrinkable,
           "gate_capped_cells": n_capped, "gate_capped_with_media": any_media_capped,
           "gate_capped_media_only": mo_capped,
           "overall_capped": len(gate["overall_capped"])},
          open(os.path.join(HERE, "f6_caps_on_media.json"), "w", encoding="utf-8"),
          indent=1, sort_keys=True, default=str)
print("\nDONE F3 r6")
