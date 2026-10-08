# -*- coding: utf-8 -*-
# F3 r4: (A) prove the Raw sheet is unchanged since the 2026-09-17 stance run
#        (B) CELL-level closure: which (member, stance category) cells carry media evidence,
#            and the buckets where item contributions are annihilated before has_press
#        (C) M5 publication parity: category_stance.json matrix vs dashboard_data.json STANCE
# READ-ONLY.  Nothing under LIVE\fec is imported.
import os, sys, json, shutil, collections, datetime
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import f_lib as F
import openpyxl

LIVE = F.LIVE
SCRATCH = r"C:\Users\pgoff\AppData\Local\Temp\claude\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\F3"
MR = F.media_rules()
RAW2ST = MR["RAW_TO_STANCE"]
ST2RAW = {v: k for k, v in RAW2ST.items()}
CATS = F.CATS

print("=" * 108)
print("F3 r4  CELL-level closure + publication parity   run %s"
      % datetime.datetime.now().isoformat(timespec="seconds"))
print("=" * 108)

WB_LIVE = os.path.join(LIVE, "Moneyball Main Version - Blended.xlsx")
WB_COPY = os.path.join(SCRATCH, "wb_copy.xlsx")
if not os.path.exists(WB_COPY) or os.path.getmtime(WB_COPY) < os.path.getmtime(WB_LIVE):
    shutil.copy2(WB_LIVE, WB_COPY)
wb = openpyxl.load_workbook(WB_COPY, data_only=True)
rw, ra = wb["Press Release - Raw"], wb["Rank & Allocation"]
rrow = {rw.cell(r, 1).value: r for r in range(2, 535) if rw.cell(r, 1).value}
members = [F.nfc(ra.cell(r, 1).value) for r in range(3, 536) if ra.cell(r, 1).value]
raw_all = {c: {m: (rw.cell(rrow[m], c).value or 0.0) for m in rrow} for c in F.ALL_COLS}
rawnfc = {c: {F.nfc(m): v for m, v in raw_all[c].items()} for c in F.ALL_COLS}

CS = json.load(open(os.path.join(LIVE, "category_stance.json"), encoding="utf-8"))

# ================================================================= (A) provenance proof
print("\n(A) IS THE RAW SHEET STILL THE ONE THE STANCE RUN READ?")
print("    category_stance.json  mtime %s" % datetime.datetime.fromtimestamp(
      os.path.getmtime(os.path.join(LIVE, "category_stance.json"))).isoformat(timespec="seconds"))
print("    workbook              mtime %s   (later: a downstream step re-saved other sheets)"
      % datetime.datetime.fromtimestamp(os.path.getmtime(WB_LIVE)).isoformat(timespec="seconds"))
print("    TEST: per_category_stance.py:156 stores MC.press_general_exclusions() = the VERBATIM col-24")
print("          values.  Byte-compare those stored floats against col 24 of the workbook today.")
gce = CS["media"]["general_credit_union_exclusions"]
col24 = {m: float(v or 0.0) for m, v in rawnfc[24].items() if abs(float(v or 0.0)) > 1e-12}
same = sum(1 for m in gce if m in col24 and repr(float(gce[m]["general"])) == repr(col24[m]))
print("    stored entries %d | today's col-24 nonzero %d | exact float match %d"
      % (len(gce), len(col24), same))
print("    only-in-artifact %s | only-in-workbook %s"
      % (sorted(set(gce) - set(col24))[:5] or "none", sorted(set(col24) - set(gce))[:5] or "none"))
RAW_UNCHANGED = (len(gce) == len(col24) == same)
print("    => Raw-sheet col 24 is byte-identical to what the stance run read: %s" % RAW_UNCHANGED)
print("       (col 24 is written by the same blend pass, same run, as cols 5..23 -- blend.py:443-470)")

# ================================================================= (B) cell-level closure
print("\n(B) CELL-LEVEL CLOSURE: 533 members x 18 stance categories = %d cells" % (len(members) * len(CATS)))
print("    RULE moneyball_core.py:98  norm={c:MR.normalize_signed_extrema(raw_all[c]) for c in COLS}")
print("    RULE per_category_stance.py:46  press={(m,lab):_normM[pc].get(m,0.0) ...}  <- .get default 0.0")
print("    RULE per_category_stance.py:60  has_press=abs(prraw[m])>1e-9   <- the media-evidence test")
norm = {c: F.normalize_signed_extrema(rawnfc[c]) for c in F.COLS}
press = {(m, lab): norm[pc].get(m, 0.0) for m in members for lab, pc in CATS}
has_press = {k: abs(v) > 1e-9 for k, v in press.items()}
n_cells = len(press)
n_hp = sum(1 for v in has_press.values() if v)
print("    cells with media evidence (has_press True) : %d of %d (%.1f%%)"
      % (n_hp, n_cells, 100.0 * n_hp / n_cells))
print("    cells with NO media evidence              : %d" % (n_cells - n_hp))
missing_row = [m for m in members if m not in rrow and m not in {F.nfc(x) for x in rrow}]
print("    roster members absent from the Raw sheet (would silently read 0.0 via .get): %d %s"
      % (len(missing_row), missing_row[:5]))

print("\n    per-category cell counts:")
print("    %-26s %-9s %-9s %-9s %s" % ("stance category", "has_press", "pos", "neg", "raw-sheet nonzero"))
for lab, pc in CATS:
    hp = sum(1 for m in members if has_press[(m, lab)])
    po = sum(1 for m in members if press[(m, lab)] > 1e-9)
    ne = sum(1 for m in members if press[(m, lab)] < -1e-9)
    rz = sum(1 for m in members if abs(float(rawnfc[pc].get(m, 0.0) or 0.0)) > 1e-12)
    flag = "" if hp == rz else "   <-- %d cell(s) lost between sheet and has_press" % (rz - hp)
    print("    %-26s %-9d %-9d %-9d %d%s" % (lab, hp, po, ne, rz, flag))

print("\n    ANNIHILATION BUCKETS -- where an item's contribution vanishes before has_press:")
print("    RULE blend.py:468-469  nv_rounded=round(nv,6); rw.cell(...).value=nv_rounded")
print("         -> any blended cell with |nv| < 5e-7 is written as 0.0 and is indistinguishable")
print("            from 'no media at all'.  The sheet keeps no record of the discarded magnitude.")
tiny = sum(1 for lab, pc in CATS for m in members
           if 0 < abs(float(rawnfc[pc].get(m, 0.0) or 0.0)) < 1e-6)
print("    cells written at |value| < 1e-6 (survivors just above the rounding floor): %d" % tiny)
lost = sum(1 for lab, pc in CATS for m in members
           if abs(float(rawnfc[pc].get(m, 0.0) or 0.0)) > 1e-12 and not has_press[(m, lab)])
print("    cells nonzero on the sheet but has_press False (killed by the 1e-9 test): %d" % lost)
print("    NOTE  a cell whose positive and negative items cancel to 0.0 is recorded by")
print("          per_category_stance.py:60 as ABSENT media evidence, not as measured-zero.")
print("          The producer keeps no counter for this class, so the count of items whose")
print("          contribution was cancelled away CANNOT be recovered from the artifacts.")

# signal reconciliation: signal = has_leg or has_press or hv (per_category_stance.py:69)
sig = CS["signal"]
n_sig = sum(1 for m in members for lab, _ in CATS if sig.get(m, {}).get(lab))
hp_not_sig = [(m, lab) for m in members for lab, _ in CATS
              if has_press[(m, lab)] and not sig.get(m, {}).get(lab)]
print("\n    RECONCILE vs category_stance.json 'signal' (= has_leg OR has_press OR has_vote):")
print("    cells with signal True: %d   cells with has_press True: %d" % (n_sig, n_hp))
print("    has_press True but signal False: %d  <- MUST be 0 (has_press implies signal)" % len(hp_not_sig))
for t in hp_not_sig[:6]:
    print("       %s / %s   press=%.12g" % (t[0], t[1], press[t]))
print("    => media evidence is never dropped between press_matrix and the signal ledger: %s"
      % (len(hp_not_sig) == 0))

# ================================================================= (C) M5 parity
print("\n(C) M5 PUBLICATION PARITY: category_stance.json -> dashboard_data.json")
DD = json.load(open(os.path.join(LIVE, "dashboard_data.json"), encoding="utf-8"))
TC = {'Champion': 'C', 'Defender': 'F', 'Ally': 'A', 'Highly Persuadable': 'H', 'Persuadable': 'P',
      'Lean Persuadable': 'G', 'True Neutral': 'T', 'Lean Skeptic': 'L', 'Neutral': 'N',
      'Skeptic': 'S', 'Highly Skeptical': 'K', 'Opponent': 'O', '—': '-'}   # bundle.py:119
LBLCAT = {'Examination & Supervision': 'Examination & Supervision',
          'Credit Union Capital Requirements': 'CU Capital Requirements',
          'NCUA Budget & Share Insurance Fund': 'NCUA Budget & Insurance', 'Interchange': 'Interchange',
          'Payments': 'Payments', 'Technology & Emerging Issues': 'Technology & Emerging',
          'Fraud Prevention': 'Fraud Prevention', 'Data Privacy & Cybersecurity': 'Data Privacy & Cyber',
          'Charter Modernization & Governance': 'Charter Modernization',
          'Defense Credit Union Issues': 'Defense CU Issues', 'Field of Membership': 'Field of Membership',
          'CDFI Fund': 'CDFI Fund', 'CFPB Restrain': 'CFPB', 'CFPB Reach (-)': 'CFPB',
          'Telephone Consumer Protection Act Issues': 'TCPA',
          'Credit Card & Consumer-Fee Mandates': 'Credit Card Mandates',
          'Member Business Lending': 'Member Business Lending',
          'Mortgage Lending & Housing': 'Mortgage & Housing',
          'Credit Union Tax Status': 'CU Tax Status'}          # bundle.py:22-30
DISPLAY_LABEL = {'CFPB Restrain': 'CFPB Restrain (pro-CU; limits agency reach)'}   # bundle.py:31
DROPPED = {'CFPB Reach (-)'}                                                       # bundle.py:32

FL = json.load(open(os.path.join(LIVE, "per_topic_alloc.json"), encoding="utf-8"))
source = [l for l in FL.keys() if l not in DROPPED]                # bundle.py:37
display = [DISPLAY_LABEL.get(l, l) for l in source]                # bundle.py:38
issue_cat = [LBLCAT.get(l, l) for l in source]                     # bundle.py:39
print("    per_topic_alloc.json labels %d -> source %d (dropped display clones: %s)"
      % (len(FL), len(source), sorted(DROPPED)))
print("    bundle 'issues' in the artifact: %d   last entry %r" % (len(DD["issues"]), DD["issues"][-1]))
print("    reconstructed display list + 'Overall CU stance' == artifact 'issues': %s"
      % (display + ["Overall CU stance"] == DD["issues"]))
print("    issue_cat one-to-one over %d categories: %s" % (len(set(issue_cat)), len(issue_cat) == len(set(issue_cat))))
print("    issue_cat set == category_stance 'cats' set: %s" % (set(issue_cat) == set(CS["cats"])))

names = DD["name"]; stance_str = DD["stance"]; gen_str = DD["gen"]
print("    dashboard members %d   category_stance matrix members %d   roster %d"
      % (len(names), len(CS["matrix"]), len(members)))
print("    every dashboard member present in the matrix: %s"
      % all(n in CS["matrix"] for n in names))
print("    stance string length: %s  (expect %d issues incl. the appended overall char)"
      % (sorted({len(s) for s in stance_str}), len(issue_cat) + 1))

mism = []; checked = 0
for k, n in enumerate(names):
    row = CS["matrix"].get(n, {})
    s = stance_str[k]
    for i, cat in enumerate(issue_cat):
        exp = TC.get(row.get(cat, "—"), "-")
        got = s[i]
        checked += 1
        if exp != got:
            mism.append((n, cat, row.get(cat), exp, got))
print("\n    PER-CATEGORY CELLS COMPARED: %d  (%d members x %d categories)"
      % (checked, len(names), len(issue_cat)))
print("    MISMATCHES: %d" % len(mism))
for t in mism[:12]:
    print("       %-24s %-26s artifact_label=%-18s expect=%s got=%s" % t)

# the 19th char: 'Overall CU stance' = GEN (overall_display_tier), NOT a per-category label
ov_bad = sum(1 for k in range(len(names)) if stance_str[k][len(issue_cat)] != gen_str[k])
print("\n    19th char ('Overall CU stance', appended at bundle.py:665-667) == gen[k] for all members: %s"
      % (ov_bad == 0))
print("    that char is overall_display_tier (moneyball_core.py:558-572) -- the board-top-13 Champion")
print("    crown -- NOT a per-category label, so it is reported separately and not counted above.")
gen_vs_raw = sum(1 for k, n in enumerate(names)
                 if gen_str[k] == 'C' and 'Champion' not in CS["matrix"].get(n, {}).values())
print("    members displayed overall 'C' with no Champion cell in the matrix (crown substitution): %d" % gen_vs_raw)

# Defender: is any Defender label altered on publication?
d_cells = [(n, c) for n in names for c in issue_cat if CS["matrix"].get(n, {}).get(c) == "Defender"]
d_shown = sum(1 for n, c in d_cells if stance_str[names.index(n)][issue_cat.index(c)] == "F")
print("\n    Defender cells in category_stance.json: %d ; published as 'F': %d ; altered: %d"
      % (len(d_cells), d_shown, len(d_cells) - d_shown))

# label vocabulary coverage
vocab = collections.Counter(v for n in CS["matrix"] for v in CS["matrix"][n].values())
unmapped = [v for v in vocab if v not in TC]
print("    label vocabulary in the matrix: %s" % dict(vocab))
print("    labels with NO entry in the bundle's TC decode map (would publish as '-'): %s"
      % (unmapped or "none"))

json.dump({"raw_unchanged": RAW_UNCHANGED, "cells": n_cells, "has_press": n_hp,
           "signal_true": n_sig, "has_press_not_signal": len(hp_not_sig),
           "sheet_nonzero_but_no_has_press": lost, "cells_below_1e6": tiny,
           "parity_checked": checked, "parity_mismatches": len(mism),
           "parity_mismatch_detail": mism[:200], "defender_cells": len(d_cells),
           "defender_published_F": d_shown, "overall_char_ok": ov_bad == 0,
           "crown_substitutions": gen_vs_raw, "vocab": dict(vocab)},
          open(os.path.join(HERE, "f4_cells_and_parity.json"), "w", encoding="utf-8"),
          indent=1, sort_keys=True, default=str)
wb.close()
print("\nDONE F3 r4")
