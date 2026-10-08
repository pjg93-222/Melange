# -*- coding: utf-8 -*-
"""LANE B step 2 (part 2): RE-DERIVE every mention window by copying the extraction logic of
fec/members/member_reference_layer.py AS TEXT.  Nothing under fec/ is imported or executed.

Sources of the copied logic (LIVE == worktree, sha256 272A987D... verified identical):
  fec/members/member_reference_layer.py : 43-175 (norm, pools, regexes, resolve_ref, extract)
  fec/moneyball_core.py                 : 810-825 (NICK_SETS, name_variants)
  fec/members/reference_cache.py        : 9-44   (key, legacy_key, bind_legacy)   [LIVE-only file]
  fec/press_stream.py                   : 68-118 (canonical_url, tone_turn_key, load_tone_rows)

Writes ONLY to SCRATCH.  Proof of re-derivation = reproducing the artifact's own per-edge
n_raw and per-source counters.
"""
import collections, csv, hashlib, io, json, os, re, sqlite3, sys, unicodedata
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

LIVE = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
DRAFT = os.path.abspath(os.path.join(LIVE, os.pardir, 'Draft Systems'))
SCRATCH = (r'C:\Users\pgoff\AppData\Local\Temp\claude'
           r'\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad'
           r'\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\laneB')
B = LIVE
DYAD_CAP = 4                      # member_reference_layer.py:37, artifact records dyad_cap=4
MODEL = 'claude-fable-5'          # member_reference_layer.py:38
CACHE = B + r'\data\members\member_reference_cache_fable.jsonl'
BASELINE = B + r'\data\members\member_reference_cache_baseline_v2.json'
BASE_TONE_PATH = os.path.join(DRAFT, 'cu_press_stance_tone.csv')
REFRESH_TONE_PATH = os.path.join(B, 'data', 'stance', 'press_refresh_tone.csv')
REFRESH_MENTIONS_PATH = os.path.join(B, 'data', 'press', 'official_mentions.jsonl')
REFRESH_RECEIPT_PATH = os.path.join(B, 'data', 'stance', 'press_refresh_feed_receipt.json')

C = {m['member']: m for m in json.load(io.open(B + r'\data\engine\congestion.json', encoding='utf-8'))}
print('roster members (congestion.json): %d' % len(C))

# ---------------- moneyball_core.py:810-825 (copied) ----------------
NICK_SETS = [{'bernard', 'bernie'}, {'charles', 'chuck', 'charlie'}, {'christopher', 'chris'}, {'richard', 'dick', 'rich', 'rick'},
             {'michael', 'mike'}, {'james', 'jim', 'jimmy'}, {'robert', 'bob', 'rob', 'bobby'}, {'thomas', 'tom', 'tommy'}, {'daniel', 'dan', 'danny'},
             {'joseph', 'joe', 'joey'}, {'timothy', 'tim'}, {'patrick', 'pat'}, {'william', 'bill', 'will', 'billy'}, {'edward', 'ed', 'ted', 'eddie'},
             {'theodore', 'ted'}, {'elizabeth', 'liz', 'beth', 'betty'}, {'katherine', 'kathryn', 'katie', 'kate', 'kathy'}, {'margaret', 'maggie', 'meg'},
             {'benjamin', 'ben'}, {'alexander', 'alex'}, {'anthony', 'tony'}, {'andrew', 'andy', 'drew'}, {'matthew', 'matt'}, {'jonathan', 'jon'},
             {'nicholas', 'nick'}, {'gregory', 'greg'}, {'lawrence', 'larry'}, {'gerald', 'jerry'}, {'ronald', 'ron', 'ronny'}, {'donald', 'don'},
             {'kenneth', 'ken'}, {'stephen', 'steven', 'steve'}, {'david', 'dave'}, {'samuel', 'sam'}, {'frederick', 'fred'}, {'susan', 'sue', 'suzy'},
             {'deborah', 'debbie', 'deb'}, {'cynthia', 'cindy'}, {'jacquelyn', 'jacqueline', 'jackie'}, {'raphael', 'rafael', 'raffi'},
             {'joshua', 'josh'}, {'zachary', 'zach', 'zack'}, {'abigail', 'abby'}, {'rebecca', 'becca', 'becky'}, {'jefferson', 'jeff'},
             {'jeffrey', 'jeff'}, {'mitchell', 'mitch'}, {'randall', 'randy'}, {'terrence', 'terri', 'terry'}, {'vincent', 'vince'}]
_NICK = {}
for _ns in NICK_SETS:
    for _a in _ns:
        _NICK.setdefault(_a, set()).update(_ns)
def name_variants(first):
    return _NICK.get(first, set()) | {first}

# ---------------- member_reference_layer.py:43-73 (copied) ----------------
def norm(s):
    s = unicodedata.normalize('NFD', s); s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z ]", '', s.lower()).strip()

SUR = collections.defaultdict(list)
for m in C: SUR[norm(m).split()[-1]].append(m)
SUR_SEN = collections.defaultdict(list); SUR_HOU = collections.defaultdict(list)
for m in C:
    (SUR_SEN if (C[m].get('chamber') == 'Senate') else SUR_HOU)[norm(m).split()[-1]].append(m)
FULLSET = {}
for m in C:
    p = norm(m).split()
    if len(p) >= 2:
        for v in name_variants(p[0]): FULLSET.setdefault((v, p[-1]), m)
BARE2MEM = {fn + ' ' + ln: mem for (fn, ln), mem in FULLSET.items()}
BARE_RE = re.compile(r'\b(?:' + '|'.join(sorted((re.escape(fn) + r'\s+' + re.escape(ln)
                     for fn, ln in FULLSET), key=len, reverse=True)) + r')\b')
TITLE_RE = r'(?:senator|sen\.|representative|rep\.|congressman|congresswoman|speaker|leader|chairman|chairwoman|ranking member)'
RITUAL = re.compile(r'\b(i yield|yield(s|ing)? (to|back)|i thank (the|my|senator|rep)|thank the (gentleman|gentlewoman|gentlelady|chair)|'
                    r'(gentleman|gentlewoman|gentlelady) from|i recognize|is recognized|i join .{0,40} in (recogniz|celebrat|honor|wish)|'
                    r'colloquy|unanimous consent|i ask .{0,30} to join|pursuant to|the clerk|quorum|i reserve)\b', re.I)
EVALUATIVE = re.compile(r'\b(wrong|right|great|dangerous|disastrous|brilliant|leadership|fail|failed|courage|proud|disgrace|'
                        r'attack|defend|oppose|support|fight|betray|hero|corrupt|honest|lie|lying|truth|gut|destroy|save|'
                        r'protect|reckless|smart|effective|partner|friend|mentor|extremis|radical|principled|bipartisan)\b', re.I)

# ---------------- member_reference_layer.py:75-120 (copied verbatim, + form tagging) ----------------
def resolve_ref(text_l, speaker):
    out = []
    for m in re.finditer(r'\b(' + TITLE_RE + r')\s+([a-zÀ-ɏ\'\-]+(?:\s[a-zÀ-ɏ\'\-]+)?)', text_l):
        ttl = m.group(1)
        raw_name = m.group(2)
        possessive_surname = raw_name.split()[0].endswith("'s")
        cand = norm(re.sub(r"'s\b", '', raw_name))
        parts = cand.split()
        if not parts: continue
        pool = (SUR_SEN if ttl.startswith('sen') else
                SUR_HOU if ttl.startswith(('rep', 'congress')) else SUR)
        hit = None
        if len(parts) == 1:
            cs = pool.get(parts[0], [])
            if len(cs) == 1: hit = cs[0]
        else:
            hit = FULLSET.get((parts[0], parts[-1]))
            if not hit and (possessive_surname or parts[1] in {'on', 'for', 'and', 'has', 'is', 'was', 'will', 'would', 'who', 'said',
                                        'the', 'to', 'of', 'with', 'in', 'as', 'at', 'about', 'from', 'that',
                                        'this', 'because', 'joins', 'joined', 'introduced', 'announced', 'supports'}):
                cs = pool.get(parts[0], [])
                if len(cs) == 1: hit = cs[0]
            if not hit:
                cs = pool.get(parts[-1], [])
                if len(cs) == 1:
                    rf = norm(cs[0]).split()[0]
                    if rf.startswith(parts[0]) or parts[0].startswith(rf) or (name_variants(parts[0]) & name_variants(rf)):
                        hit = cs[0]
        if hit and hit != speaker: out.append((hit, m.start(), 'TITLE', m.group(0)))
    for m in BARE_RE.finditer(text_l):
        mem = BARE2MEM.get(re.sub(r'\s+', ' ', m.group(0)))
        if mem and mem != speaker: out.append((mem, m.start(), 'BARE', m.group(0)))
    seen = set(); res = []
    for mem, pos, form, lit in out:
        k = (mem, pos // 80)
        if k not in seen: seen.add(k); res.append((mem, pos, form, lit))
    return res

# ---------------- press_stream.py:68-118 (copied) ----------------
from urllib.parse import urlsplit, urlunsplit
def canonical_url(value, keep_fragment=True):
    value = (value or '').strip()
    if not value: return ''
    try:
        parts = urlsplit(value)
        host = parts.netloc.lower()
        if host.startswith('www.'): host = host[4:]
        path = parts.path.rstrip('/') or '/'
        fragment = parts.fragment if keep_fragment else ''
        return urlunsplit((parts.scheme.lower() or 'https', host, path, parts.query, fragment))
    except ValueError:
        return value.rstrip('/').lower()

def tone_turn_key(row):
    return canonical_url(row.get('url'), keep_fragment=True), (row.get('bioguide') or row.get('member') or '').strip().upper()

def _read_csv(path):
    if not path or not os.path.exists(path): return []
    with io.open(path, encoding='utf-8-sig', newline='') as h: return list(csv.DictReader(h))

def _sha256_file(path):
    d = hashlib.sha256()
    with open(path, 'rb') as h:
        for blk in iter(lambda: h.read(1 << 20), b''): d.update(blk)
    return d.hexdigest()

# replicate validate_refresh_receipt() but REPORT instead of raising
def check_receipt():
    notes = []
    if not os.path.isfile(REFRESH_RECEIPT_PATH): return ['receipt MISSING -> load_tone_rows() would RAISE']
    rc = json.load(io.open(REFRESH_RECEIPT_PATH, encoding='utf-8'))
    if rc.get('schema') != 'press-refresh-engine-feed-v1' or rc.get('status') != 'ACTIVE_TYPED_SUPPLEMENT':
        notes.append('receipt not active: schema=%r status=%r' % (rc.get('schema'), rc.get('status')))
    if rc.get('promotion_mode') != 'CUMULATIVE_ACCEPTED_RUN_LEDGER' or not rc.get('accepted_runs'):
        notes.append('receipt not a cumulative accepted-run ledger')
    exp = {'data/stance/press_refresh_tone.csv': REFRESH_TONE_PATH,
           'data/stance/press_refresh_fulltext.jsonl': os.path.join(B, 'data', 'stance', 'press_refresh_fulltext.jsonl'),
           'data/press/official_mentions.jsonl': REFRESH_MENTIONS_PATH}
    h = rc.get('outputs_sha256') or {}
    if set(h) != set(exp): notes.append('receipt does not cover the complete output set: %s' % sorted(h))
    for rel, p in exp.items():
        if not os.path.isfile(p): notes.append('MISSING output %s' % rel)
        elif _sha256_file(p) != h.get(rel): notes.append('HASH MISMATCH %s' % rel)
    return notes or ['receipt VALID -> load_tone_rows() proceeds']

def load_tone_rows():
    rows, seen = [], set()
    for sp in (BASE_TONE_PATH, REFRESH_TONE_PATH):
        for row in _read_csv(sp):
            if (row.get('source') or '').strip().upper() == 'OFFICIAL_MENTION': continue
            k = tone_turn_key(row)
            if not k[0] or k in seen: continue
            seen.add(k); row['_tone_src_file'] = os.path.basename(sp); rows.append(row)
    return rows

print()
print('=== press refresh receipt gate (press_stream.validate_refresh_receipt) ===')
for n in check_receipt(): print('   %s' % n)

# ---------------- member_reference_layer.py:122-175 extract() (copied; + diagnostics) ----------------
def extract():
    refs = []
    con = sqlite3.connect('file:%s?mode=ro' % (B + r'\social_media.db').replace('\\', '/'), uri=True, timeout=180)
    con.execute('PRAGMA query_only=ON')
    # NOTE: production SELECTs (member,date,source,text); rowid+package added for diagnostics only.
    for rid, spk, date, src, pkg, text in con.execute(
            'SELECT rowid, member, date, source, package, text FROM speeches'):
        if spk not in C or not text: continue
        tl = ' ' + text.lower() + ' '
        for mem, pos, form, lit in resolve_ref(tl, spk):
            w = text[max(0, pos - 170):pos + 200]
            if RITUAL.search(w) and not EVALUATIVE.search(w): continue
            refs.append({'a': spk, 'b': mem, 'src': src, 'date': date or '',
                         'w': re.sub(r'\s+', ' ', w).strip(),
                         '_pos': pos, '_form': form, '_lit': lit, '_sid': rid,
                         '_pkg': pkg or '', '_tlen': len(text), '_raw_w': w})
    con.close()
    n_sp = len(refs); print('   speeches pass: %d' % n_sp)
    for r in load_tone_rows():
        spk = r.get('member'); t = (r.get('title') or '')
        if spk not in C or not t: continue
        tl = ' ' + t.lower() + ' '
        for mem, pos, form, lit in resolve_ref(tl, spk):
            refs.append({'a': spk, 'b': mem, 'src': 'PRESS', 'date': r.get('date') or '', 'w': t[:300],
                         '_pos': pos, '_form': form, '_lit': lit, '_sid': '', '_pkg': '',
                         '_url': r.get('url') or '', '_title': t, '_tonefile': r.get('_tone_src_file', ''),
                         '_tlen': len(t), '_raw_w': t[:300]})
    n_pr = len(refs) - n_sp; print('   PRESS pass: %d' % n_pr)
    if os.path.exists(REFRESH_MENTIONS_PATH):
        for line in io.open(REFRESH_MENTIONS_PATH, encoding='utf-8'):
            try: r = json.loads(line)
            except ValueError: continue
            spk = r.get('host_member'); mem = r.get('mentioned_member')
            if spk not in C or mem not in C or spk == mem: continue
            window = (r.get('evidence') or r.get('title') or '')[:1200]
            if RITUAL.search(window) and not EVALUATIVE.search(window): continue
            refs.append({'a': spk, 'b': mem, 'src': 'PRESS_OFFICIAL_MENTION', 'date': r.get('date') or '',
                         'w': window, '_pos': -1, '_form': 'TYPED', '_lit': '', '_sid': '', '_pkg': '',
                         '_url': r.get('url') or r.get('host_url') or '', '_title': r.get('title') or '',
                         '_mentionrow': {k: r.get(k) for k in ('url', 'host_url', 'source', 'domain', 'host_domain', 'outlet')},
                         '_tlen': len(window), '_raw_w': window})
    n_om = len(refs) - n_sp - n_pr; print('   PRESS_OFFICIAL_MENTION pass: %d' % n_om)
    try:
        import openpyxl
        con2 = sqlite3.connect('file:%s?mode=ro' % (B + r'\social_media.db').replace('\\', '/'), uri=True, timeout=180)
        h2m = {}
        bio2name = {C[m]['bioguide']: m for m in C if C[m].get('bioguide')}
        for bio, tw in con2.execute("SELECT bioguide,twitter FROM members WHERE twitter IS NOT NULL AND twitter!=''"):
            if bio in bio2name: h2m[tw.lower().lstrip('@')] = bio2name[bio]
        con2.close()
        ws = openpyxl.load_workbook(B + r'\Quorum Processed.xlsx', read_only=True)['Members']
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        im, ic, irt, idt = hdr.index('member'), hdr.index('Content'), hdr.index('is_rt'), hdr.index('dt') if 'dt' in hdr else None
        for row in ws.iter_rows(min_row=2, values_only=True):
            spk = row[im]; t = str(row[ic] or '')
            if spk not in C or not t or str(row[irt]).lower() in ('1', 'true'): continue
            for h in re.findall(r'@(\w+)', t):
                mem = h2m.get(h.lower())
                if mem and mem != spk:
                    refs.append({'a': spk, 'b': mem, 'src': 'X', 'date': str(row[idt] or '') if idt is not None else '',
                                 'w': t[:300], '_pos': -1, '_form': 'HANDLE', '_lit': '@' + h, '_sid': '',
                                 '_pkg': '', '_tlen': len(t), '_raw_w': t[:300]})
    except Exception as e:
        print('  (X mention pass skipped: %s)' % str(e)[:200])
    print('   X pass: %d' % (len(refs) - n_sp - n_pr - n_om))
    return refs

print()
print('=== extract() re-derivation ===')
refs = extract()
print('TOTAL re-derived mentions: %d' % len(refs))

# ---------------- reference_cache.py key/legacy_key/bind_legacy (copied) ----------------
def ckey(row): return 'v2:' + hashlib.sha256(json.dumps(['reference-window-v2', row['a'], row['b'], row['w'][:400]], ensure_ascii=False).encode('utf-8')).hexdigest()
def legacy_key(row): return hashlib.md5((row['a'] + '|' + row['b'] + '|' + row['w'][:200]).encode('utf-8', 'replace')).hexdigest()

cache = {}
for ln in io.open(CACHE, encoding='utf-8'):
    try:
        d = json.loads(ln)
        if d['k'] not in cache: cache[d['k']] = d['out']
    except Exception: pass
print('cache distinct keys loaded: %d' % len(cache))
baseline = json.load(io.open(BASELINE, encoding='utf-8'))
assert baseline.get('schema') == 'reference-cache-baseline-v2'
actual = set(cache)
bound = 0
for new, old in baseline['bindings'].items():
    if new not in cache and old in cache:
        cache[new] = cache[old]; bound += 1
legacy_ambiguous = set(baseline['ambiguous']) - actual
print('baseline bindings applied: %d   legacy_ambiguous (live): %d' % (bound, len(legacy_ambiguous)))
print('cache keys after binding  : %d' % len(cache))

# ---------------- aggregation, mirroring lines 228-233, 252-264, 318-346 ----------------
bykey = collections.defaultdict(list)
for r in refs: bykey[(r['a'], r['b'], r['src'])].append(r)
target_all = []
for k, v in bykey.items():
    v.sort(key=lambda r: r['date'], reverse=True)
    target_all.extend(v[:DYAD_CAP])
for r in refs: r['_k'] = ckey(r)
tgt_ids = set(id(r) for r in target_all)
for r in refs: r['_is_target'] = id(r) in tgt_ids
for r in refs: r['_cached'] = r['_k'] in cache
# todo_all: newest DYAD_CAP *cached* windows per dyad/source (line 318)
todo_all = [r for group in bykey.values() for r in [x for x in group if x['_k'] in cache][:DYAD_CAP]]
todo_ids = set(id(r) for r in todo_all)
for r in refs: r['_counted_scored'] = id(r) in todo_ids

print()
print('=' * 78)
print('STEP 2 : TARGET CONTEXTS')
print('=' * 78)
print('target contexts (newest %d per a->b x source)   : %d' % (DYAD_CAP, len(target_all)))
n_tc = sum(1 for r in target_all if r['_cached'])
print('  scored (cache hit incl. legacy binding)       : %d' % n_tc)
print('  UNSCORED                                      : %d' % (len(target_all) - n_tc))
pend = sum(1 for r in target_all if (r['_k'] not in cache) or (r['_k'] in legacy_ambiguous))
print('  pending (uncached OR legacy-ambiguous)        : %d  [drives receipt deferred_contexts]' % pend)
print('todo_all (newest %d CACHED per dyad/source)      : %d   == artifact sum n_scored?' % (DYAD_CAP, len(todo_all)))
print()
print('BY SOURCE  (target contexts)')
print('  %-24s %8s %8s %8s %7s' % ('source', 'target', 'scored', 'unscored', 'pct_un'))
bs = collections.defaultdict(lambda: [0, 0])
for r in target_all:
    bs[r['src']][0] += 1
    if r['_cached']: bs[r['src']][1] += 1
for s in sorted(bs, key=lambda x: -bs[x][0]):
    t, sc = bs[s]
    print('  %-24s %8d %8d %8d %6.1f%%' % (s, t, sc, t - sc, 100.0 * (t - sc) / max(1, t)))
print()
print('BY SOURCE  (ALL raw mentions, and how many are counted into n_scored)')
print('  %-24s %8s %8s' % ('source', 'raw', 'in_n_scored'))
br = collections.Counter(); bsc = collections.Counter()
for r in refs:
    br[r['src']] += 1
    if r['_counted_scored']: bsc[r['src']] += 1
for s in sorted(br, key=lambda x: -br[x]):
    print('  %-24s %8d %8d' % (s, br[s], bsc[s]))

# ---------------- C-SPAN ----------------
print()
print('=' * 78)
print('C-SPAN WINDOWS (INTERVIEW rows whose package is a c-span.org URL)')
print('=' * 78)
cs = [r for r in refs if 'c-span' in r['_pkg'].lower() or 'cspan' in r['_pkg'].lower()]
cst = [r for r in cs if r['_is_target']]
print('c-span raw mentions          : %d' % len(cs))
print('c-span target contexts       : %d' % len(cst))
print('  of those, SCORED (cached)  : %d' % sum(1 for r in cst if r['_cached']))
print('  of those, UNSCORED         : %d' % sum(1 for r in cst if not r['_cached']))
print('c-span mentions counted into n_scored : %d' % sum(1 for r in cs if r['_counted_scored']))
print('distinct c-span programs contributing : %d' % len({r['_pkg'] for r in cs}))
print('c-span edges (a,b pairs)              : %d' % len({(r['a'], r['b']) for r in cs}))

# ---------------- PROOF: reproduce the artifact ----------------
print()
print('=' * 78)
print('PROOF OF RE-DERIVATION vs shipped member_references.json')
print('=' * 78)
art = json.load(io.open(B + r'\member_references.json', encoding='utf-8'))
A = {(e['a'], e['b']): e for e in art['edges']}
mine_raw = collections.Counter(); mine_src = collections.defaultdict(collections.Counter)
for r in refs:
    mine_raw[(r['a'], r['b'])] += 1
    mine_src[(r['a'], r['b'])][r['src']] += 1
mine_scored = collections.Counter()
for r in todo_all: mine_scored[(r['a'], r['b'])] += 1

print('artifact edges                   : %d' % len(A))
print('my dyads with >=1 mention        : %d' % len(mine_raw))
print('my dyads that are artifact edges : %d' % len(set(mine_raw) & set(A)))
print('artifact edges I did NOT produce : %d' % len(set(A) - set(mine_raw)))
print('my dyads absent from artifact     : %d  (expected: formality-only dyads, line 342)' % len(set(mine_raw) - set(A)))
ok = sum(1 for k in A if mine_raw.get(k) == A[k]['n_raw'])
print()
print('n_raw EXACT MATCH            : %d / %d  = %.3f%%' % (ok, len(A), 100.0 * ok / len(A)))
oks = sum(1 for k in A if dict(mine_src[k]) == (A[k].get('sources') or {}))
print('per-source counter EXACT MATCH: %d / %d  = %.3f%%' % (oks, len(A), 100.0 * oks / len(A)))
okn = sum(1 for k in A if mine_scored.get(k, 0) == A[k]['n_scored'])
print('n_scored EXACT MATCH          : %d / %d  = %.3f%%' % (okn, len(A), 100.0 * okn / len(A)))
print('sum n_raw    mine %d  vs artifact %d   (delta %+d)' % (sum(mine_raw.values()), sum(e['n_raw'] for e in art['edges']), sum(mine_raw.values()) - sum(e['n_raw'] for e in art['edges'])))
print('sum n_scored mine %d  vs artifact %d   (delta %+d)' % (sum(mine_scored.values()), sum(e['n_scored'] for e in art['edges']), sum(mine_scored.values()) - sum(e['n_scored'] for e in art['edges'])))
tot_src = collections.Counter()
for k in mine_src:
    for s, v in mine_src[k].items(): tot_src[s] += v
print('my per-source totals  : %s' % dict(tot_src.most_common()))
art_src = collections.Counter()
for e in art['edges']:
    for s, v in (e.get('sources') or {}).items(): art_src[s] += v
print('artifact per-source   : %s' % dict(art_src.most_common()))
print('per-source totals identical: %s' % (dict(tot_src) == dict(art_src)))
mism = [(k, mine_raw.get(k), A[k]['n_raw'], dict(mine_src[k]), A[k].get('sources')) for k in A if mine_raw.get(k) != A[k]['n_raw']]
print()
print('first 12 n_raw mismatches (mine vs artifact):')
for k, a, b, sa, sb in mism[:12]:
    print('   %-46s mine=%-5s art=%-5s  mine_src=%s art_src=%s' % (str(k)[:46], a, b, sa, sb))

# ---------------- dump for P1..P4 ----------------
outp = os.path.join(SCRATCH, 'refs.jsonl')
with io.open(outp, 'w', encoding='utf-8') as f:
    for r in refs:
        f.write(json.dumps(r, ensure_ascii=False) + '\n')
print()
print('wrote %s (%d rows)' % (outp, len(refs)))
with io.open(os.path.join(SCRATCH, 'artifact_edges.json'), 'w', encoding='utf-8') as f:
    json.dump({'edges': art['edges']}, f)
print('DONE 02c')
