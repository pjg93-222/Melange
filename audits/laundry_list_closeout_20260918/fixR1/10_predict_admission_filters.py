# -*- coding: utf-8 -*-
"""FIX R1: predicted effect of the three member-reference admission filters, measured on the
LIVE inputs. READ-ONLY -- nothing under fec/ is imported or executed except the new PURE
helper fec/members/reference_filters.py, and nothing is written outside this receipts dir.

This is laneB/receipts/02c_rederive.py extended. The extraction is the same copied-as-text
logic (that script reproduced the shipped artifact 8,017 of 8,017 edges on n_raw and per-source
counters); the aggregation is extended to the full class tallies so the edge SET, the signed
tallies and the n_scored>=2 drawing floor can all be measured, and then the three filters are
applied through the shipped helper at every dial position.

Sources of the copied logic (verified identical in LIVE and the worktree):
  fec/members/member_reference_layer.py : norm, pools, regexes, resolve_ref, extract, aggregation
  fec/moneyball_core.py                 : NICK_SETS, name_variants
  fec/members/reference_cache.py        : key, legacy_key, bind_legacy   [LIVE-only file]
  fec/press_stream.py                   : canonical_url, tone_turn_key, load_tone_rows
The ONLY imported module is the new helper, so what is measured here is the shipped rule."""
import collections, csv, hashlib, io, json, os, re, sqlite3, sys, unicodedata
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

LIVE = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WORK = (r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\.claude\worktrees'
        r'\unfinished-work-handoff-a3bbad')
DRAFT = os.path.abspath(os.path.join(LIVE, os.pardir, 'Draft Systems'))
OUT = os.path.join(WORK, 'audits', 'laundry_list_closeout_20260918', 'fixR1')
sys.path.insert(0, os.path.join(WORK, 'fec'))
from members import reference_filters as RF     # the shipped pure helper, nothing else

B = LIVE
DYAD_CAP = 4
CACHE = B + r'\data\members\member_reference_cache_fable.jsonl'
BASELINE = B + r'\data\members\member_reference_cache_baseline_v2.json'
BASE_TONE_PATH = os.path.join(DRAFT, 'cu_press_stance_tone.csv')
REFRESH_TONE_PATH = os.path.join(B, 'data', 'stance', 'press_refresh_tone.csv')
REFRESH_MENTIONS_PATH = os.path.join(B, 'data', 'press', 'official_mentions.jsonl')

C = {m['member']: m for m in json.load(io.open(B + r'\data\engine\congestion.json', encoding='utf-8'))}
print('roster members (congestion.json): %d' % len(C))

# ---------------- moneyball_core.py (copied) ----------------
NICK_SETS = [{'bernard', 'bernie'}, {'charles', 'chuck', 'charlie'}, {'christopher', 'chris'}, {'richard', 'dick', 'rich', 'rick'},
             {'michael', 'mike'}, {'james', 'jim', 'jimmy'}, {'robert', 'bob', 'rob', 'bobby'}, {'thomas', 'tom', 'tommy'}, {'daniel', 'dan', 'danny'},
             {'joseph', 'joe', 'joey'}, {'timothy', 'tim'}, {'patrick', 'pat'}, {'william', 'bill', 'will', 'billy'}, {'edward', 'ed', 'ted', 'eddie'},
             {'theodore', 'ted'}, {'elizabeth', 'liz', 'beth', 'betty'}, {'katherine', 'kathryn', 'katie', 'kate', 'kathy'}, {'margaret', 'maggie', 'meg'},
             {'benjamin', 'ben'}, {'alexander', 'alex'}, {'anthony', 'tony'}, {'andrew', 'andy', 'drew'}, {'matthew', 'matt'}, {'jonathan', 'jon'},
             {'nicholas', 'nick'}, {'gregory', 'greg'}, {'lawrence', 'larry'}, {'gerald', 'jerry'}, {'ronald', 'ron', 'ronny'}, {'donald', 'don'},
             {'kenneth', 'ken'}, {'stephen', 'steven', 'steve'}, {'david', 'dave'}, {'samuel', 'sam'}, {'frederick', 'fred'}, {'susan', 'sue', 'suzy'},
             {'deborah', 'debbie', 'deb'}, {'cynthia', 'cindy'}, {'jacquelyn', 'jacqueline', 'jackie'}, {'raphael', 'rafael', 'raffi'},
             {'joshua', 'josh'}, {'zachary', 'zach', 'zack'}, {'abigail', 'abby'}, {'rebecca', 'becca', 'becky'}, {'jefferson', 'jeff'},
             {'jeffrey', 'jeff'}, {'mitchell', 'mitch'}, {'randall', 'randy'}, {'terrence', 'terri', 'terry'}, {'vincent', 'vince'}]
_NICK = {}
for _ns in NICK_SETS:
    for _a in _ns:
        _NICK.setdefault(_a, set()).update(_ns)
def name_variants(first):
    return _NICK.get(first, set()) | {first}

# ---------------- member_reference_layer.py (copied) ----------------
def norm(s):
    s = unicodedata.normalize('NFD', s); s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z ]", '', s.lower()).strip()

SUR = collections.defaultdict(list)
for m in C: SUR[norm(m).split()[-1]].append(m)
SUR_SEN = collections.defaultdict(list); SUR_HOU = collections.defaultdict(list)
for m in C:
    (SUR_SEN if (C[m].get('chamber') == 'Senate') else SUR_HOU)[norm(m).split()[-1]].append(m)
FULLSET = {}
for m in C:
    p = norm(m).split()
    if len(p) >= 2:
        for v in name_variants(p[0]): FULLSET.setdefault((v, p[-1]), m)
BARE2MEM = {fn + ' ' + ln: mem for (fn, ln), mem in FULLSET.items()}
BARE_RE = re.compile(r'\b(?:' + '|'.join(sorted((re.escape(fn) + r'\s+' + re.escape(ln)
                     for fn, ln in FULLSET), key=len, reverse=True)) + r')\b')
TITLE_RE = r'(?:senator|sen\.|representative|rep\.|congressman|congresswoman|speaker|leader|chairman|chairwoman|ranking member)'
SURNAMES = sorted({norm(m).split()[-1] for m in C}, key=len, reverse=True)
RITUAL = re.compile(r'\b(i yield|yield(s|ing)? (to|back)|i thank (the|my|senator|rep)|thank the (gentleman|gentlewoman|gentlelady|chair)|'
                    r'(gentleman|gentlewoman|gentlelady) from|i recognize|is recognized|i join .{0,40} in (recogniz|celebrat|honor|wish)|'
                    r'colloquy|unanimous consent|i ask .{0,30} to join|pursuant to|the clerk|quorum|i reserve)\b', re.I)
EVALUATIVE = re.compile(r'\b(wrong|right|great|dangerous|disastrous|brilliant|leadership|fail|failed|courage|proud|disgrace|'
                        r'attack|defend|oppose|support|fight|betray|hero|corrupt|honest|lie|lying|truth|gut|destroy|save|'
                        r'protect|reckless|smart|effective|partner|friend|mentor|extremis|radical|principled|bipartisan)\b', re.I)
# The producer's own surname pool, handed to the roster filter exactly as the wiring does.
ROSTER_SUR = frozenset(SURNAMES)
assert ROSTER_SUR == RF.roster_surnames(list(C)), 'helper surname pool != producer pool'

def resolve_ref(text_l, speaker):
    out = []
    for m in re.finditer(r'\b(' + TITLE_RE + r')\s+([a-zÀ-ɏ\'\-]+(?:\s[a-zÀ-ɏ\'\-]+)?)', text_l):
        ttl = m.group(1)
        raw_name = m.group(2)
        possessive_surname = raw_name.split()[0].endswith("'s")
        cand = norm(re.sub(r"'s\b", '', raw_name))
        parts = cand.split()
        if not parts: continue
        pool = (SUR_SEN if ttl.startswith('sen') else
                SUR_HOU if ttl.startswith(('rep', 'congress')) else SUR)
        hit = None
        if len(parts) == 1:
            cs = pool.get(parts[0], [])
            if len(cs) == 1: hit = cs[0]
        else:
            hit = FULLSET.get((parts[0], parts[-1]))
            if not hit and (possessive_surname or parts[1] in {'on', 'for', 'and', 'has', 'is', 'was', 'will', 'would', 'who', 'said',
                                        'the', 'to', 'of', 'with', 'in', 'as', 'at', 'about', 'from', 'that',
                                        'this', 'because', 'joins', 'joined', 'introduced', 'announced', 'supports'}):
                cs = pool.get(parts[0], [])
                if len(cs) == 1: hit = cs[0]
            if not hit:
                cs = pool.get(parts[-1], [])
                if len(cs) == 1:
                    rf = norm(cs[0]).split()[0]
                    if rf.startswith(parts[0]) or parts[0].startswith(rf) or (name_variants(parts[0]) & name_variants(rf)):
                        hit = cs[0]
        if hit and hit != speaker: out.append((hit, m.start()))
    for m in BARE_RE.finditer(text_l):
        mem = BARE2MEM.get(re.sub(r'\s+', ' ', m.group(0)))
        if mem and mem != speaker: out.append((mem, m.start()))
    seen = set(); res = []
    for mem, pos in out:
        k = (mem, pos // 80)
        if k not in seen: seen.add(k); res.append((mem, pos))
    return res

# ---------------- press_stream.py (copied) ----------------
from urllib.parse import urlsplit, urlunsplit
def canonical_url(value, keep_fragment=True):
    value = (value or '').strip()
    if not value: return ''
    try:
        parts = urlsplit(value)
        host = parts.netloc.lower()
        if host.startswith('www.'): host = host[4:]
        path = parts.path.rstrip('/') or '/'
        fragment = parts.fragment if keep_fragment else ''
        return urlunsplit((parts.scheme.lower() or 'https', host, path, parts.query, fragment))
    except ValueError:
        return value.rstrip('/').lower()

def tone_turn_key(row):
    return canonical_url(row.get('url'), keep_fragment=True), (row.get('bioguide') or row.get('member') or '').strip().upper()

def _read_csv(path):
    if not path or not os.path.exists(path): return []
    with io.open(path, encoding='utf-8-sig', newline='') as h: return list(csv.DictReader(h))

def load_tone_rows():
    rows, seen = [], set()
    for sp in (BASE_TONE_PATH, REFRESH_TONE_PATH):
        for row in _read_csv(sp):
            if (row.get('source') or '').strip().upper() == 'OFFICIAL_MENTION': continue
            k = tone_turn_key(row)
            if not k[0] or k in seen: continue
            seen.add(k); row['_tone_src_file'] = os.path.basename(sp); rows.append(row)
    return rows

# ---------------- extract(): the producer's four passes, + the url the gate reads ----------------
def extract():
    refs = []
    con = sqlite3.connect('file:%s?mode=ro' % (B + r'\social_media.db').replace('\\', '/'), uri=True, timeout=180)
    con.execute('PRAGMA query_only=ON')
    for spk, date, src, text in con.execute('SELECT member, date, source, text FROM speeches'):
        if spk not in C or not text: continue
        tl = ' ' + text.lower() + ' '
        for mem, pos in resolve_ref(tl, spk):
            w = text[max(0, pos - 170):pos + 200]
            if RITUAL.search(w) and not EVALUATIVE.search(w): continue
            refs.append({'a': spk, 'b': mem, 'src': src, 'date': date or '',
                         'w': re.sub(r'\s+', ' ', w).strip(), '_raw_w': w})
    con.close()
    n_sp = len(refs); print('   speeches pass: %d' % n_sp)
    for r in load_tone_rows():
        spk = r.get('member'); t = (r.get('title') or '')
        if spk not in C or not t: continue
        tl = ' ' + t.lower() + ' '
        for mem, pos in resolve_ref(tl, spk):
            refs.append({'a': spk, 'b': mem, 'src': 'PRESS', 'date': r.get('date') or '', 'w': t[:300],
                         'url': r.get('url') or '', '_raw_w': t[:300], '_title': t})
    n_pr = len(refs) - n_sp; print('   PRESS pass: %d' % n_pr)
    if os.path.exists(REFRESH_MENTIONS_PATH):
        for line in io.open(REFRESH_MENTIONS_PATH, encoding='utf-8'):
            try: r = json.loads(line)
            except ValueError: continue
            spk = r.get('host_member'); mem = r.get('mentioned_member')
            if spk not in C or mem not in C or spk == mem: continue
            window = (r.get('evidence') or r.get('title') or '')[:1200]
            if RITUAL.search(window) and not EVALUATIVE.search(window): continue
            refs.append({'a': spk, 'b': mem, 'src': 'PRESS_OFFICIAL_MENTION', 'date': r.get('date') or '',
                         'w': window, '_raw_w': window})
    n_om = len(refs) - n_sp - n_pr; print('   PRESS_OFFICIAL_MENTION pass: %d' % n_om)
    try:
        import openpyxl
        con2 = sqlite3.connect('file:%s?mode=ro' % (B + r'\social_media.db').replace('\\', '/'), uri=True, timeout=180)
        h2m = {}
        bio2name = {C[m]['bioguide']: m for m in C if C[m].get('bioguide')}
        for bio, tw in con2.execute("SELECT bioguide,twitter FROM members WHERE twitter IS NOT NULL AND twitter!=''"):
            if bio in bio2name: h2m[tw.lower().lstrip('@')] = bio2name[bio]
        con2.close()
        ws = openpyxl.load_workbook(B + r'\Quorum Processed.xlsx', read_only=True)['Members']
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        im, ic, irt, idt = hdr.index('member'), hdr.index('Content'), hdr.index('is_rt'), hdr.index('dt') if 'dt' in hdr else None
        for row in ws.iter_rows(min_row=2, values_only=True):
            spk = row[im]; t = str(row[ic] or '')
            if spk not in C or not t or str(row[irt]).lower() in ('1', 'true'): continue
            for h in re.findall(r'@(\w+)', t):
                mem = h2m.get(h.lower())
                if mem and mem != spk:
                    refs.append({'a': spk, 'b': mem, 'src': 'X', 'date': str(row[idt] or '') if idt is not None else '',
                                 'w': t[:300], '_raw_w': t[:300]})
    except Exception as e:
        print('  (X mention pass skipped: %s)' % str(e)[:200])
    print('   X pass: %d' % (len(refs) - n_sp - n_pr - n_om))
    return refs

# ---------------- reference_cache.py key/bind_legacy (copied) ----------------
def ckey(row): return 'v2:' + hashlib.sha256(json.dumps(['reference-window-v2', row['a'], row['b'], row['w'][:400]], ensure_ascii=False).encode('utf-8')).hexdigest()

print()
print('=== extract() re-derivation ===')
REFS = extract()
print('TOTAL re-derived mentions: %d' % len(REFS))

CACHE_MAP = {}
for ln in io.open(CACHE, encoding='utf-8'):
    try:
        d = json.loads(ln)
        if d['k'] not in CACHE_MAP: CACHE_MAP[d['k']] = d['out']
    except Exception: pass
baseline = json.load(io.open(BASELINE, encoding='utf-8'))
assert baseline.get('schema') == 'reference-cache-baseline-v2'
bound = 0
for new, old in baseline['bindings'].items():
    if new not in CACHE_MAP and old in CACHE_MAP:
        CACHE_MAP[new] = CACHE_MAP[old]; bound += 1
print('cache keys %d (baseline bindings applied %d)' % (len(CACHE_MAP), bound))
for r in REFS: r['_k'] = ckey(r)

ART = json.load(io.open(B + r'\member_references.json', encoding='utf-8'))
AE = {(e['a'], e['b']): e for e in ART['edges']}
print('shipped artifact edges: %d   (dyad_cap %s)' % (len(AE), ART.get('dyad_cap')))
print('shipped artifact already carries an admission_filters block: %s' % ('admission_filters' in ART))


def run(press_own_domain, roster_filter, identical_once):
    """The producer's pipeline at one dial position, using the SHIPPED helper for the filters."""
    refs, excl = RF.apply_admission_filters(REFS, ROSTER_SUR,
                                            press_own_domain=press_own_domain,
                                            roster_filter=roster_filter)
    bykey = collections.defaultdict(list)
    for r in refs: bykey[(r['a'], r['b'], r['src'])].append(r)
    target_all = []
    for k, v in bykey.items():
        v.sort(key=lambda r: r['date'], reverse=True)
        target_all.extend(v[:DYAD_CAP])
    n_target_pre = len(target_all)
    unc_pre = sum(1 for r in target_all if r['_k'] not in CACHE_MAP)
    target_all, dup_target = RF.collapse_identical(target_all, identical_once)
    unc_post = sum(1 for r in target_all if r['_k'] not in CACHE_MAP)
    todo_all = [r for group in bykey.values() for r in [x for x in group if x['_k'] in CACHE_MAP][:DYAD_CAP]]
    n_todo_pre = len(todo_all)
    todo_all, dup_scored = RF.collapse_identical(todo_all, identical_once)
    edges = collections.defaultdict(lambda: {'n_raw': 0, 'n_scored': 0, 'pos': 0, 'neg': 0, 'neu': 0,
                                             'form': 0, 'sources': collections.Counter()})
    for r in refs:
        e = edges[(r['a'], r['b'])]; e['n_raw'] += 1; e['sources'][r['src']] += 1
    for r in todo_all:
        o = CACHE_MAP.get(r['_k'])
        if not o: continue
        e = edges[(r['a'], r['b'])]; e['n_scored'] += 1
        c = o['cls']
        if c == 'positive': e['pos'] += 1
        elif c == 'negative': e['neg'] += 1
        elif c == 'neutral-substantive': e['neu'] += 1
        else: e['form'] += 1
    out = {}
    for ab, e in edges.items():
        sub = e['pos'] + e['neg'] + e['neu']
        if sub == 0: continue                # formality-only dyads never enter the layer
        out[ab] = {'n_raw': e['n_raw'], 'n_scored': e['n_scored'], 'pos': e['pos'], 'neg': e['neg'],
                   'neu': e['neu'], 'form': e['form'], 'sources': dict(e['sources']),
                   'net': round((e['pos'] - e['neg']) / max(1, sub), 3)}
    return {'edges': out, 'excluded': dict(excl), 'dup_target': dup_target, 'dup_scored': dup_scored,
            'mentions': len(refs), 'target_pre': n_target_pre, 'target_post': len(target_all),
            'todo_pre': n_todo_pre, 'todo_post': len(todo_all),
            'paid_queue_uncached_before': unc_pre, 'paid_queue_uncached_after': unc_post}


def drawn(res): return sum(1 for e in res['edges'].values() if e['n_scored'] >= 2)
def signed(res): return sum(1 for e in res['edges'].values() if e['pos'] + e['neg'] > 0)


print()
print('=' * 92)
print('STEP 0  PARITY: all three dials at 0 must reproduce the shipped artifact exactly')
print('=' * 92)
OFF = run(False, False, False)
print('re-derived mentions %d | edges %d | shipped edges %d' % (OFF['mentions'], len(OFF['edges']), len(AE)))
print('dyads I produce that are not shipped edges : %d  (expected 0)' % len(set(OFF['edges']) - set(AE)))
print('shipped edges I do not produce             : %d  (expected 0)' % len(set(AE) - set(OFF['edges'])))
for field in ('n_raw', 'n_scored', 'pos', 'neg', 'neu', 'net'):
    okc = sum(1 for ab in AE if ab in OFF['edges'] and OFF['edges'][ab][field] == AE[ab][field])
    print('  %-8s EXACT MATCH : %d / %d = %.3f%%' % (field, okc, len(AE), 100.0 * okc / len(AE)))
oks = sum(1 for ab in AE if ab in OFF['edges'] and OFF['edges'][ab]['sources'] == (AE[ab].get('sources') or {}))
print('  %-8s EXACT MATCH : %d / %d = %.3f%%' % ('sources', oks, len(AE), 100.0 * oks / len(AE)))
print('sum n_raw    mine %d vs artifact %d' % (sum(e['n_raw'] for e in OFF['edges'].values()),
                                               sum(e['n_raw'] for e in ART['edges'])))
print('sum n_scored mine %d vs artifact %d' % (sum(e['n_scored'] for e in OFF['edges'].values()),
                                               sum(e['n_scored'] for e in ART['edges'])))
print('all dials 0 excluded nothing: %s   (census %s)' % (not OFF['excluded'], OFF['excluded']))
print('drawn (n_scored>=2) %d | signed %d' % (drawn(OFF), signed(OFF)))

print()
print('=' * 92)
print('STEP 1  ALL THREE FILTERS ON (the shipped default)')
print('=' * 92)
ON = run(True, True, True)
print('per-filter exclusion census (mentions dropped before n_raw / the queue):')
for k in RF.EXCLUSION_KEYS[:2]:
    print('   %-26s %6d' % (k, ON['excluded'].get(k, 0)))
print('   %-26s %6d   (scored-queue windows collapsed; n_raw keeps them)'
      % ('identical_window_repeat', ON['dup_scored']))
print('   %-26s %6d   (paid target contexts collapsed)' % ('  ...of which queued', ON['dup_target']))
print('   paid queue uncached contexts: %d -> %d  (contexts --score would have paid for: %+d)'
      % (ON['paid_queue_uncached_before'], ON['paid_queue_uncached_after'],
         ON['paid_queue_uncached_after'] - ON['paid_queue_uncached_before']))
print()
print('mentions (n_raw population) : %6d -> %6d   (%+d)' % (OFF['mentions'], ON['mentions'], ON['mentions'] - OFF['mentions']))
print('scored windows (sum n_scored): %6d -> %6d   (%+d)'
      % (sum(e['n_scored'] for e in OFF['edges'].values()), sum(e['n_scored'] for e in ON['edges'].values()),
         sum(e['n_scored'] for e in ON['edges'].values()) - sum(e['n_scored'] for e in OFF['edges'].values())))
print('EDGES                       : %6d -> %6d   (%+d)' % (len(OFF['edges']), len(ON['edges']), len(ON['edges']) - len(OFF['edges'])))
print('DRAWN edges (n_scored>=2)   : %6d -> %6d   (%+d)' % (drawn(OFF), drawn(ON), drawn(ON) - drawn(OFF)))
print('SIGNED edges (pos+neg>0)    : %6d -> %6d   (%+d)' % (signed(OFF), signed(ON), signed(ON) - signed(OFF)))

gone = sorted(set(OFF['edges']) - set(ON['edges']))
new = sorted(set(ON['edges']) - set(OFF['edges']))
fell = sorted(ab for ab in set(OFF['edges']) & set(ON['edges'])
              if OFF['edges'][ab]['n_scored'] >= 2 > ON['edges'][ab]['n_scored'])
rose = sorted(ab for ab in set(OFF['edges']) & set(ON['edges'])
              if ON['edges'][ab]['n_scored'] >= 2 > OFF['edges'][ab]['n_scored'])
print()
print('edges REMOVED ENTIRELY : %d      (an edge disappears only when it loses ALL signed/neutral support)' % len(gone))
for ab in gone[:25]:
    e = OFF['edges'][ab]
    print('   %-28s -> %-28s n_raw=%-4d n_scored=%-3d net=%+0.2f sources=%s' %
          (ab[0][:28], ab[1][:28], e['n_raw'], e['n_scored'], e['net'], e['sources']))
if len(gone) > 25: print('   ... and %d more' % (len(gone) - 25))
print('edges that APPEAR (cap refill promoted a previously uncapped window): %d' % len(new))
for ab in new[:25]:
    e = ON['edges'][ab]
    print('   %-28s -> %-28s n_raw=%-4d n_scored=%-3d net=%+0.2f sources=%s' %
          (ab[0][:28], ab[1][:28], e['n_raw'], e['n_scored'], e['net'], e['sources']))
print()
print('edges that DROP BELOW the consumers\' n_scored>=2 floor : %d' % len(fell))
for ab in fell[:25]:
    a, b = OFF['edges'][ab], ON['edges'][ab]
    print('   %-26s -> %-26s n_scored %d -> %d | net %+0.2f -> %+0.2f | sources %s' %
          (ab[0][:26], ab[1][:26], a['n_scored'], b['n_scored'], a['net'], b['net'], b['sources']))
if len(fell) > 25: print('   ... and %d more' % (len(fell) - 25))
print('edges that RISE to the floor : %d' % len(rose))
for ab in rose[:10]:
    a, b = OFF['edges'][ab], ON['edges'][ab]
    print('   %-26s -> %-26s n_scored %d -> %d' % (ab[0][:26], ab[1][:26], a['n_scored'], b['n_scored']))
print()
print('net movement among edges present in both: %d edges change net, %d change n_scored'
      % (sum(1 for ab in set(OFF['edges']) & set(ON['edges']) if OFF['edges'][ab]['net'] != ON['edges'][ab]['net']),
         sum(1 for ab in set(OFF['edges']) & set(ON['edges']) if OFF['edges'][ab]['n_scored'] != ON['edges'][ab]['n_scored'])))
print('audit_full floors after the change: edges>=1000 %s | every edge n_scored>=pos+neg %s'
      % (len(ON['edges']) >= 1000,
         all(e['n_scored'] >= e['pos'] + e['neg'] for e in ON['edges'].values())))

print()
print('=' * 92)
print('STEP 2  EACH DIAL TURNED OFF ON ITS OWN (isolating one filter at a time)')
print('=' * 92)
CONFIGS = [('all three ON (default)', True, True, True),
           ('REF_PRESS_OWN_DOMAIN=0', False, True, True),
           ('REF_ROSTER_FILTER=0', True, False, True),
           ('REF_IDENTICAL_ONCE=0', True, True, False),
           ('all three OFF (= shipped artifact)', False, False, False)]
print('%-36s %8s %8s %8s %8s %10s %10s %10s' % ('config', 'mentions', 'edges', 'drawn', 'signed',
                                                'excl_press', 'excl_rost', 'excl_dupe'))
RES = {}
for label, p, r_, i in CONFIGS:
    res = run(p, r_, i); RES[label] = res
    print('%-36s %8d %8d %8d %8d %10d %10d %10d'
          % (label, res['mentions'], len(res['edges']), drawn(res), signed(res),
             res['excluded'].get('press_third_party_domain', 0),
             res['excluded'].get('roster_list_window', 0), res['dup_scored']))
print()
print('marginal contribution of each filter (default minus the config with that filter off):')
for label in ('REF_PRESS_OWN_DOMAIN=0', 'REF_ROSTER_FILTER=0', 'REF_IDENTICAL_ONCE=0'):
    a, b = RES[label], RES['all three ON (default)']
    print('   %-24s edges %+d | drawn %+d | signed %+d | n_scored %+d'
          % (label.split('=')[0], len(b['edges']) - len(a['edges']), drawn(b) - drawn(a), signed(b) - signed(a),
             sum(e['n_scored'] for e in b['edges'].values()) - sum(e['n_scored'] for e in a['edges'].values())))

print()
print('=' * 92)
print('STEP 3  CROSS-CHECKS ON THE IMPLEMENTATION CHOICES')
print('=' * 92)
bad = sum(1 for r in REFS if RF.is_roster_list(r['_raw_w'], ROSTER_SUR) != RF.is_roster_list(r['w'], ROSTER_SUR))
print('roster detector on the RAW slice vs the producer\'s collapsed window: %d of %d disagree'
      % (bad, len(REFS)))
print('  (lane B measured on _raw_w; the producer holds w. 0 disagreements = the graded precision transfers)')
hits = [r for r in REFS if RF.is_roster_list(r['w'], ROSTER_SUR)]
print('roster-list mentions repo-wide: %d   by source %s'
      % (len(hits), dict(collections.Counter(r['src'] for r in hits))))
press = [r for r in REFS if r['src'] == 'PRESS']
third = [r for r in press if not RF.own_official_domain(r.get('url'))]
print('PRESS mentions %d | third-party %d | hosts %s'
      % (len(press), len(third),
         dict(collections.Counter((r.get('url') or '').split('/')[2] if '//' in (r.get('url') or '') else '?'
                                  for r in third))))
pom = [r for r in REFS if r['src'] == 'PRESS_OFFICIAL_MENTION']
print('PRESS_OFFICIAL_MENTION mentions %d -- out of scope for the gate, unchanged by it: %s'
      % (len(pom), RES['all three ON (default)']['excluded'].get('press_third_party_domain', 0) == len(third)))
# Does collapsing the PAID target set ever change which window ends up scored?  It can only do so
# when two windows are equal after whitespace normalisation but hash to DIFFERENT cache keys.
grp = collections.defaultdict(list)
for r in REFS: grp[(r['a'], r['b'], RF.normalized_window(r['w']))].append(r)
split = [k for k, v in grp.items() if len({x['_k'] for x in v}) > 1]
print('normalised-identical windows on one edge that hash to DIFFERENT cache keys: %d groups' % len(split))
print('  (0 => collapsing the paid target set cannot change which window is classified)')
print('Magaziner edges: shipped %d -> filtered %d'
      % (sum(1 for ab in OFF['edges'] if 'Magaziner' in ab[0] + ab[1]),
         sum(1 for ab in ON['edges'] if 'Magaziner' in ab[0] + ab[1])))
for ab in sorted(ab for ab in OFF['edges'] if 'Magaziner' in ab[0] + ab[1]):
    e = OFF['edges'][ab]
    print('   %-26s -> %-26s n_raw=%-3d n_scored=%-3d net=%+0.2f %s  [kept: %s]'
          % (ab[0][:26], ab[1][:26], e['n_raw'], e['n_scored'], e['net'], e['sources'], ab in ON['edges']))
print()
print('the artifact block the producer will write at the default:')
print(json.dumps(RF.admission_filters_block(
    {'REF_PRESS_OWN_DOMAIN': True, 'REF_ROSTER_FILTER': True, 'REF_IDENTICAL_ONCE': True},
    dict(ON['excluded'], identical_window_repeat=ON['dup_scored'])), indent=1, sort_keys=True))

with io.open(os.path.join(OUT, '10_predict_admission_filters.json'), 'w', encoding='utf-8') as f:
    json.dump({'shipped_edges': len(AE),
               'configs': {label: {'mentions': RES[label]['mentions'], 'edges': len(RES[label]['edges']),
                                   'drawn': drawn(RES[label]), 'signed': signed(RES[label]),
                                   'sum_n_scored': sum(e['n_scored'] for e in RES[label]['edges'].values()),
                                   'excluded': RES[label]['excluded'],
                                   'identical_window_repeat': RES[label]['dup_scored']}
                           for label in RES},
               'edges_removed': [list(ab) for ab in gone],
               'edges_added': [list(ab) for ab in new],
               'edges_below_floor': [list(ab) for ab in fell],
               'edges_above_floor': [list(ab) for ab in rose]}, f, ensure_ascii=False, indent=1)
print()
print('wrote 10_predict_admission_filters.json')
print('DONE 10')
