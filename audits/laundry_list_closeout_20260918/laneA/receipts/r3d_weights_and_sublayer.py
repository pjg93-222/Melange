# -*- coding: utf-8 -*-
"""LANE A / R3d -- the press pillar weight row, and which sub-layer each press column feeds.

R3c: Rank & Allocation!C{row} = SUMPRODUCT(INDEX('Press Release - Normalized'!$E$2:$W$534,...),
'Press Release - Normalized'!$E$535:$W$535).  The range stops at W, so column X
('General Credit Union') is consumed by nothing.  Column R ('CU-Bank Competitive Dynamics')
IS inside the range.  This reads the weight row and maps 'Sub Layer - Raw' headers to the
press column each one names, to establish whether a scoring column has any displayed home.
"""
import os, sys, re, collections, datetime
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WBP = os.path.join(ROOT, 'Moneyball Main Version - Blended.xlsx')
print('=' * 100)
print('R3d  press pillar weight row + sub-layer map   run', datetime.datetime.now().isoformat(timespec='seconds'))
print('=' * 100)

MEDIA = ['Examination & Supervision', 'Credit Union Capital Requirements',
         'NCUA Budget & Share Insurance Fund', 'Interchange', 'Payments',
         'Technology & Emerging Issues', 'Fraud Prevention', 'Data Privacy & Cybersecurity',
         'Charter Modernization & Governance', 'Defense Credit Union Issues',
         'Field of Membership', 'Credit Union Tax Status', 'CDFI Fund',
         'CU-Bank Competitive Dynamics', 'CFPB Issues & Structural Reform',
         'Credit Card & Consumer-Fee Mandates', 'Telephone Consumer Protection Act',
         'Member Business Lending', 'Mortgage Lending & Housing', 'General Credit Union']

# ---- A. weight row 535 (values) ---------------------------------------------------------------
wbv = openpyxl.load_workbook(WBP, read_only=True, data_only=True)
pn = wbv['Press Release - Normalized']
row535 = None
rowlabel = None
for r in pn.iter_rows(min_row=535, max_row=535, values_only=True):
    row535 = r
for r in pn.iter_rows(min_row=534, max_row=534, min_col=1, max_col=1, values_only=True):
    rowlabel = r[0]
print("\n-- A. 'Press Release - Normalized' weight row 535 (the SUMPRODUCT's second vector) --")
print('   row 534 col A (last member?) = %r' % rowlabel)
print('   row 535 col A = %r' % (row535[0] if row535 else None))
print('   %-4s %-38s %s' % ('col', 'header', 'weight in row 535'))
inrange, outrange = [], []
for i, name in enumerate(MEDIA):
    idx = 5 + i
    L = get_column_letter(idx)
    w = row535[idx - 1] if row535 and len(row535) >= idx else None
    tag = 'IN  E:W' if L <= 'W' else 'OUT of E:W'
    print('   %-4s %-38s %-18r %s' % (L, name, w, tag))
    (inrange if L <= 'W' else outrange).append((L, name, w))
print('\n   columns inside the pillar range E:W : %d' % len(inrange))
print('   columns OUTSIDE it (never summed)   : %d  -> %s'
      % (len(outrange), [n for _, n, _ in outrange]))
nonzero = [(L, n, w) for L, n, w in inrange if isinstance(w, (int, float)) and abs(w) > 0]
zero = [(L, n, w) for L, n, w in inrange if not (isinstance(w, (int, float)) and abs(w) > 0)]
print('   of the in-range columns, weight != 0 : %d' % len(nonzero))
print('   of the in-range columns, weight == 0 or blank: %d  -> %s'
      % (len(zero), [n for _, n, _ in zero]))

# ---- B. Sub Layer - Raw header map ------------------------------------------------------------
sl = wbv['Sub Layer - Raw']
slh = next(sl.iter_rows(min_row=1, max_row=1, values_only=True))
print("\n-- B. 'Sub Layer - Raw' headers (row 1) --")
for i, h in enumerate(slh[:40], start=1):
    if h is not None:
        print('   col %-3s %s' % (get_column_letter(i), h))
wbv.close()

# ---- C. which press column each Sub Layer - Raw column names ---------------------------------
wbf = openpyxl.load_workbook(WBP, read_only=False, data_only=False)
slf = wbf['Sub Layer - Raw']
pat = re.compile(r"'?Press Release - Normalized'?!\$?([A-Z]{1,3})")
bycol = collections.defaultdict(collections.Counter)
for row in slf.iter_rows(min_row=2, max_row=8):
    for c in row:
        v = c.value
        if isinstance(v, str) and 'Press Release - Normalized' in v:
            for m in pat.findall(v):
                bycol[c.column][m] += 1
print("\n-- C. Sub Layer - Raw column -> press column it pulls (rows 2-8 sample) --")
NAME = {get_column_letter(5 + i): n for i, n in enumerate(MEDIA)}
for col in sorted(bycol):
    hdr = slh[col - 1] if len(slh) >= col else None
    pulls = sorted(bycol[col])
    print('   %-4s %-34s pulls press col %-6s (%s)'
          % (get_column_letter(col), str(hdr)[:34], ','.join(pulls),
             ', '.join(NAME.get(p, '?') for p in pulls)))
pulled = {p for c in bycol for p in bycol[c]}
print('\n   press columns NOT pulled by any Sub Layer - Raw column: %s'
      % sorted({get_column_letter(5 + i) for i in range(20)} - pulled))
for L in sorted({get_column_letter(5 + i) for i in range(20)} - pulled):
    print('      %s  %s' % (L, NAME[L]))
wbf.close()
print('\nDONE R3d')
