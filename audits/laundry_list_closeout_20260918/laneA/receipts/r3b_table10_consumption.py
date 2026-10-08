# -*- coding: utf-8 -*-
"""LANE A / R3b -- the media blend columns are consumed by TABLE name, not sheet name.

build_media_blend's own docstring: "Insertion point = 'Press Release - Raw' (Table10): its
LET-normalized sheet and the Rank & Allocation SUMPRODUCT consume whatever sits there."
R3 found zero `'Press Release - Raw'!` references, so the consumers must use the structured
table reference.  This resolves Table10, dumps the formulas that read it, and reports which
of the 20 category columns each consumer actually names.  Read-only.
"""
import os, sys, io, re, collections, datetime
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl

ROOT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WBP = os.path.join(ROOT, 'Moneyball Main Version - Blended.xlsx')
print('=' * 100)
print('R3b  Table10 / normalized-sheet consumption   run', datetime.datetime.now().isoformat(timespec='seconds'))
print('=' * 100)
wb = openpyxl.load_workbook(WBP, read_only=False, data_only=False)

# ---- A. tables defined on every sheet ---------------------------------------------------------
print('\n-- A. worksheet tables --')
for ws in wb.worksheets:
    try:
        tabs = dict(ws.tables)
    except Exception as exc:
        print('   %-34s <tables unreadable: %s>' % (ws.title, exc)); continue
    if tabs:
        for nm, t in tabs.items():
            try:
                ref = t.ref
            except Exception:
                ref = '?'
            print('   %-34s table %-14s ref %s' % (ws.title, nm, ref))

print('\n-- B. defined names --')
try:
    for nm, dn in list(wb.defined_names.items()):
        print('   %-28s -> %s' % (nm, getattr(dn, 'value', dn)))
except Exception as exc:
    print('   <defined names unreadable: %s>' % exc)

# ---- C. every formula mentioning Table10 ------------------------------------------------------
print('\n-- C. formulas naming Table10 --')
cnt = collections.Counter()
samples = collections.defaultdict(list)
colrefs = collections.Counter()
colpat = re.compile(r'Table10\[\[?([^\]\[]+)\]?\]')
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or 'Table10' not in v:
                continue
            cnt[ws.title] += 1
            for m in colpat.findall(v):
                colrefs[m.strip()] += 1
            if len(samples[ws.title]) < 4:
                samples[ws.title].append('%s!%s = %s' % (ws.title, c.coordinate, v[:600]))
for sh, n in sorted(cnt.items(), key=lambda kv: -kv[1]):
    print('   %-34s %6d formula cells' % (sh, n))
print('\n   Table10 column names named inside those formulas:')
for k, v in sorted(colrefs.items(), key=lambda kv: -kv[1]):
    print('      %-40s %6d' % (k, v))
print('\n   samples:')
for sh in sorted(samples):
    for s in samples[sh]:
        print('      %s' % s)

# ---- D. 'Press Release - Normalized' header + first data-row formulas -------------------------
print("\n-- D. sheet 'Press Release - Normalized' --")
ws = wb['Press Release - Normalized']
hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
print('   header row: %s' % [h for h in hdr[:30]])
for r in range(2, 4):
    vals = []
    for c in range(1, 26):
        v = ws.cell(r, c).value
        if v is not None:
            vals.append('%s%d=%s' % (openpyxl.utils.get_column_letter(c), r, str(v)[:300]))
    print('   row %d:' % r)
    for v in vals[:26]:
        print('      %s' % v)

# ---- E. 'Rank & Allocation' formulas that reference the press pillar --------------------------
print("\n-- E. 'Rank & Allocation' formulas naming the normalized press sheet --")
ra = wb['Rank & Allocation']
seen = 0
for row in ra.iter_rows(min_row=1, max_row=8):
    for c in row:
        v = c.value
        if isinstance(v, str) and ('Press Release' in v or 'Table10' in v):
            print('   %s = %s' % (c.coordinate, v[:700]))
            seen += 1
            if seen > 12:
                break
    if seen > 12:
        break
if not seen:
    print('   (no Press Release reference in rows 1-8; scanning the whole sheet)')
    hit = collections.Counter()
    for row in ra.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and ('Press Release' in v or 'Table10' in v):
                hit[re.sub(r'\d+', '', c.coordinate)] += 1
                if sum(hit.values()) <= 6:
                    print('   %s = %s' % (c.coordinate, v[:700]))
    print('   per-column tally: %s' % dict(hit))
wb.close()
print('\nDONE R3b')
