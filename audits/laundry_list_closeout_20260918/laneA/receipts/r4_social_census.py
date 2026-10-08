# -*- coding: utf-8 -*-
"""LANE A / R4 -- SOCIAL channel census, both consumers (read-only).

Social Master.xlsx 'Corpus' is read by TWO independent implementations with DIFFERENT filter
chains and different denominators:

  ENGINE  fec/stance/build_media_blend.py:58-120   (pandas; feeds the ranking engine)
  DESK    fec/surfaces/make_member_desk.py:1721-1776 (openpyxl; feeds cu-member-desk.html)

Both chains are replicated here row by row, every drop attributed to its quoted rule, and the
silent (uncounted-by-the-producer) drops are named.  Nothing under fec\ is imported: hash-
verified copies of media_rules.py and member_identity.py are used.
"""
import os, sys, io, csv, json, collections, datetime
LIB = r'C:\Users\pgoff\AppData\Local\Temp\claude\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\laneA\lib'
sys.path.insert(0, LIB)
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl
import media_rules as MR
from member_identity import Resolver

ROOT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
SM = os.path.join(ROOT, 'Social Master.xlsx')
print('=' * 100)
print('R4  SOCIAL CHANNEL CENSUS   run', datetime.datetime.now().isoformat(timespec='seconds'))
print('=' * 100)
print('Social Master.xlsx  %d bytes  mtime %s' % (os.path.getsize(SM),
      datetime.datetime.fromtimestamp(os.path.getmtime(SM)).isoformat(sep=' ')))

roster = json.load(io.open(os.path.join(ROOT, 'data', 'engine', 'congestion.json'), encoding='utf-8'))
resolver = Resolver([x['member'] for x in roster], {x['member']: x.get('bioguide') for x in roster})
print('roster members: %d' % len(roster))

STACK_PATH = os.path.join(ROOT, 'data', 'stance', 'tweet_stance_final.csv')
raw_stack = []
with io.open(STACK_PATH, encoding='utf-8-sig', newline='') as fh:
    raw_stack = list(csv.DictReader(fh))
print('tweet_stance_final.csv rows: %d  mtime %s' % (len(raw_stack),
      datetime.datetime.fromtimestamp(os.path.getmtime(STACK_PATH)).isoformat(sep=' ')))

# ENGINE's deterministic collapse (build_media_blend.py:74-83)
grp = collections.defaultdict(list)
for r in raw_stack:
    grp[(r['member'], MR.normalized_text(r['text']))].append(r)
TR = {'unanimous': 3, 'majority': 2, 'review': 1}
STACK = {}
collapsed = 0
for k, rows in grp.items():
    if len(rows) == 1:
        STACK[k] = rows[0]
        continue
    collapsed += 1
    top = collections.Counter(x['stack_stance'] for x in rows).most_common()
    lab = top[0][0] if (len(top) == 1 or top[0][1] > top[1][1]) else 'neutral'
    cands = [x for x in rows if x['stack_stance'] == lab] or rows
    STACK[k] = max(cands, key=lambda x: TR.get(x.get('tier', ''), 0))
print('engine STACK keys: %d   (keys that needed the majority-vote collapse: %d)'
      % (len(STACK), collapsed))
# DESK's last-write-wins index (make_member_desk.py:1718-1720)
DSTACK = {}
for r in raw_stack:
    DSTACK[(r.get('member') or '', MR.normalized_text(r.get('text')))] = r
print('desk  stack keys: %d   (last-write-wins, no collapse)' % len(DSTACK))

SIGN = {'support': 1.0, 'oppose': -1.0, 'neutral': 0.0}
TIERW = {'unanimous': 1.0, 'majority': 0.7, 'review': 0.35}

wb = openpyxl.load_workbook(SM, read_only=True, data_only=True)
print('\nSocial Master sheets: %s' % wb.sheetnames)
ws = wb['Corpus']
hdr_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
H = {h: i for i, h in enumerate(hdr_row)}
print("'Corpus' columns (%d): %s" % (len(hdr_row), list(hdr_row)))
rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()
print('Corpus data rows: %d' % len(rows))


def g(r, k):
    return r[H[k]] if k in H else None


# ---- cell-type diagnostics (the engine reads this sheet through pandas dtype=str) -------------
print('\n-- cell-type diagnostics (divergence risk for the pandas-vs-openpyxl substitution) --')
for col in ('stance', 'verified_handle', 'date', 'stance_margin', 'is_rt', 'platform', 'source'):
    if col not in H:
        print('   %-16s <column absent>' % col)
        continue
    kinds = collections.Counter(type(g(r, col)).__name__ for r in rows)
    vals = collections.Counter(repr(g(r, col)) for r in rows)
    print('   %-16s types=%s  distinct=%d  top=%s'
          % (col, dict(kinds), len(vals), [v for v, _ in vals.most_common(4)]))

print('\n-- platform / source breakdown of all %d Corpus rows --' % len(rows))
for col in ('platform', 'source'):
    if col in H:
        c = collections.Counter(str(g(r, col) or '') for r in rows)
        print('   %s: %s' % (col, dict(sorted(c.items(), key=lambda kv: -kv[1]))))

# ================= ENGINE CHAIN =================
print('\n' + '-' * 100)
print('ENGINE CHAIN  fec/stance/build_media_blend.py:58-120')
print('-' * 100)
print("RULE :59  corp=corp[(corp['stance'].fillna('')!='')&(corp['verified_handle'].astype(str)=='1')]")
e_blank_stance = sum(1 for r in rows if str(g(r, 'stance') or '').strip() == '')
e_unverified = sum(1 for r in rows
                   if str(g(r, 'stance') or '').strip() != '' and str(g(r, 'verified_handle')) != '1')
corp = [r for r in rows
        if str(g(r, 'stance') or '').strip() != '' and str(g(r, 'verified_handle')) == '1']
print('   Corpus rows in                                      : %6d' % len(rows))
print("   dropped: stance blank/NaN                           : %6d" % e_blank_stance)
print("   dropped: verified_handle != '1' (on stance-bearing)  : %6d" % e_unverified)
print('   ENGINE `len(corp)` -- what the engine calls "source rows": %6d' % len(corp))

nover = ndupe = nunknown = nlinkonly = 0
n_zero_sign = n_baddate = n_nocat = nassigned = 0
zero_sign_by_label = collections.Counter()
seen = set()
soc_members = set()
eng_cat = collections.Counter()
for r in corp:
    st = g(r, 'stance')
    inten_override = None
    s = STACK.get((g(r, 'member'), MR.normalized_text(g(r, 'text'))))
    if s:
        st = s['stack_stance']
        inten_override = TIERW.get(s['tier'], 0.5)
        nover += 1
    sg = SIGN.get(str(st or '').strip().lower() if isinstance(st, str) else st, 0.0)
    if not isinstance(st, str):
        sg = 0.0
    else:
        sg = SIGN.get(st, 0.0)
    if sg == 0.0:
        n_zero_sign += 1
        zero_sign_by_label[str(st)] += 1
        continue
    member = resolver.resolve(g(r, 'member'))
    if not member:
        nunknown += 1
        continue
    event = (member, MR.normalized_text(g(r, 'text')))
    dupe = event in seen
    if dupe:
        ndupe += 1
    seen.add(event)
    if not event[1]:
        nlinkonly += 1
        continue
    try:
        datetime.date.fromisoformat(str(g(r, 'date'))[:10])
    except Exception:
        n_baddate += 1
        continue
    hits = MR.route_raw_categories(g(r, 'text') or '')
    if not hits:
        n_nocat += 1
        continue
    nassigned += 1
    soc_members.add(member)
    for c in hits:
        eng_cat[c] += 1

print('\n   per-row disposition of those %d rows:' % len(corp))
print("   RULE :96-97  sg=SIGN.get(_st,0.0); if sg==0.0: continue          -> %6d  *** NOT COUNTED BY THE ENGINE ***" % n_zero_sign)
print('        breakdown by label: %s' % dict(zero_sign_by_label.most_common()))
print("   RULE :98-99  _member=_resolver.resolve(...); if not _member: nunknown+=1; continue -> %6d  (printed)" % nunknown)
print("   RULE :106    elif not _event[1]: nlinkonly+=1; continue          -> %6d  (printed)" % nlinkonly)
print("   RULE :107-108 try: age=...; except Exception: continue           -> %6d  *** NOT COUNTED BY THE ENGINE ***" % n_baddate)
print("   RULE :116-117 hits=MR.route_raw_categories(...); if not hits: continue -> %6d  *** NOT COUNTED BY THE ENGINE ***" % n_nocat)
print('   SCORED / category-assigned (`nassigned`)                         -> %6d' % nassigned)
print('   members carrying social signal (`len(soc)`)                      -> %6d' % len(soc_members))
print('   stack overrides applied (`nover`)                                -> %6d' % nover)
print('   exact cross-post duplicates seen (`ndupe`, RETAINED at CROSSPOST_DEDUP=0) -> %6d' % ndupe)
tot = n_zero_sign + nunknown + nlinkonly + n_baddate + n_nocat + nassigned
print('\n   CLOSES? %d = %d zero-sign + %d unresolved + %d link-only + %d bad-date + %d no-category + %d scored -> %s'
      % (len(corp), n_zero_sign, nunknown, nlinkonly, n_baddate, n_nocat, nassigned,
         'YES' if tot == len(corp) else 'NO  unaccounted=%d' % (len(corp) - tot)))
print('\n   what the engine PRINTS: "%d source rows -> %d scored/category-assigned | %d exact cross-post'
      ' duplicates retained | %d link-only posts skipped | %d unresolved identities excluded | %d members"'
      % (len(corp), nassigned, ndupe, nlinkonly, nunknown, len(soc_members)))
print('   UNPRINTED by the engine: %d zero-sign + %d bad-date + %d no-category = %d rows'
      % (n_zero_sign, n_baddate, n_nocat, n_zero_sign + n_baddate + n_nocat))
print('   ALSO UNPRINTED: the pre-filter drop at :59 -- %d blank-stance + %d unverified-handle = %d rows'
      % (e_blank_stance, e_unverified, e_blank_stance + e_unverified))
print('\n   engine social category assignments (a row splits across its hits):')
for k, v in sorted(eng_cat.items(), key=lambda kv: -kv[1]):
    print('      %-38s %6d' % (k, v))

# ================= DESK CHAIN =================
print('\n' + '-' * 100)
print('DESK CHAIN  fec/surfaces/make_member_desk.py:1721-1776')
print('-' * 100)
d = collections.Counter()
d_seen = {}
d_disp_scored = 0
d_excl = collections.Counter()
d_also = 0
desk_cat = collections.Counter()
CROSSPOST_DEDUP = False
for r in rows:
    d['social_source_rows'] += 1
    if str(g(r, 'verified_handle') or '').strip() != '1':
        d['social_unverified_handle_excluded'] += 1
        continue
    nm = resolver.resolve(g(r, 'member'), None)
    if not nm:
        d['social_unknown_member_excluded'] += 1
        continue
    text = str(g(r, 'text') or '').strip()
    hit = DSTACK.get((str(g(r, 'member') or ''), MR.normalized_text(text)))
    stance = str((hit or {}).get('stack_stance') or g(r, 'stance') or 'neutral').lower()
    if stance == 'unrelated':
        d['social_unrelated_excluded'] += 1
        continue
    raw_cats = MR.route_raw_categories(text)
    if not raw_cats:
        d['social_offtopic_excluded'] += 1
        continue
    key = (nm, MR.normalized_text(text))
    duplicate = key in d_seen
    plat = str(g(r, 'platform') or 'social')
    pdate = str(g(r, 'date') or '')[:10]
    if not duplicate:
        d_seen[key] = (plat, pdate)
    sign = {'support': 1, 'oppose': -1, 'neutral': 0}.get(stance, 0)
    if duplicate:
        d['social_duplicate_visible'] += 1
    elif sign:
        d['social_scored_unique'] += 1
    # social_scoring(sign, key[1], duplicate, d_seen.get(key)) with dedup=False
    if not sign:
        scored, excl, also = False, 'neutral', ''
    elif not key[1]:
        scored, excl, also = False, 'link-only post: no scorable text', ''
    else:
        scored, excl = True, ''
        also = 'also posted on ...' if (duplicate and d_seen.get(key)) else ''
    d['social_displayed'] += 1
    if scored:
        d_disp_scored += 1
    if excl:
        d_excl[excl] += 1
    if also:
        d_also += 1
    for c in raw_cats:
        desk_cat[c] += 1

print("RULE :1729  if str(sv(row,'verified_handle') or '').strip() != '1': ...excluded  -> %6d"
      % d['social_unverified_handle_excluded'])
print("RULE :1732  nm=canon(...); if not nm: ...social_unknown_member_excluded         -> %6d"
      % d['social_unknown_member_excluded'])
print("RULE :1739  if stance == 'unrelated': ...social_unrelated_excluded              -> %6d"
      % d['social_unrelated_excluded'])
print("RULE :1745  if not raw_cats: ...social_offtopic_excluded                        -> %6d"
      % d['social_offtopic_excluded'])
print('DISPLAYED (social_displayed)                                                   -> %6d'
      % d['social_displayed'])
print('   of which scored=True                                                        -> %6d' % d_disp_scored)
print('   of which scored=False by rule:')
for k, v in d_excl.most_common():
    print('      %-40s %6d' % (k, v))
print('   rows carrying an "also posted on" note (D-034)                              -> %6d' % d_also)
print('   social_duplicate_visible (census counter)                                   -> %6d'
      % d['social_duplicate_visible'])
print('   social_scored_unique (census counter)                                       -> %6d'
      % d['social_scored_unique'])
tot = (d['social_unverified_handle_excluded'] + d['social_unknown_member_excluded']
       + d['social_unrelated_excluded'] + d['social_offtopic_excluded'] + d['social_displayed'])
print('\nCLOSES? %d in = %d unverified + %d unknown-member + %d unrelated + %d off-topic + %d displayed -> %s'
      % (d['social_source_rows'], d['social_unverified_handle_excluded'],
         d['social_unknown_member_excluded'], d['social_unrelated_excluded'],
         d['social_offtopic_excluded'], d['social_displayed'],
         'YES' if tot == d['social_source_rows'] else 'NO  unaccounted=%d' % (d['social_source_rows'] - tot)))
print('\ndesk social category assignments:')
for k, v in sorted(desk_cat.items(), key=lambda kv: -kv[1]):
    print('   %-38s %6d' % (k, v))

print('\n-- ENGINE vs DESK denominator gap --')
print('   desk counts every Corpus row as a source row      : %6d' % d['social_source_rows'])
print("   engine's `len(corp)` is already post-filter        : %6d" % len(corp))
print('   gap (rows the engine never calls "source rows")    : %6d' % (len(rows) - len(corp)))
print('\nDONE R4')
