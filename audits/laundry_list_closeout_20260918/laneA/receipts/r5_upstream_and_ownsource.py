# -*- coding: utf-8 -*-
"""LANE A / R5 -- social upstream provenance + the three own-source channels (read-only).

PART 1  Quorum Master.xlsx -> process_quorum.py -> Quorum Processed.xlsx / quorum_member_tweets.csv
        -> build_social_master.py -> Social Master.xlsx 'Corpus'.  Establishes the TRUE rows-in for
        the social channel and the true-duplicate rule that is applied there.
PART 2  speech / hearing / interview own-source channels as build_media_blend consumes them,
        including the staleness gate and the gold-only hearing rule.
"""
import os, sys, io, csv, json, collections, datetime
LIB = r'C:\Users\pgoff\AppData\Local\Temp\claude\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\laneA\lib'
sys.path.insert(0, LIB)
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl
import media_rules as MR
from member_identity import Resolver

ROOT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
print('=' * 100)
print('R5  SOCIAL UPSTREAM + OWN-SOURCE CHANNELS   run', datetime.datetime.now().isoformat(timespec='seconds'))
print('=' * 100)


def stat(p):
    if not os.path.exists(p):
        return 'ABSENT'
    return '%d bytes  mtime %s' % (os.path.getsize(p),
                                   datetime.datetime.fromtimestamp(os.path.getmtime(p)).isoformat(sep=' '))


def rd(p):
    if not os.path.exists(p):
        return []
    with io.open(p, encoding='utf-8-sig', newline='') as fh:
        return list(csv.DictReader(fh))


# ================= PART 1: SOCIAL UPSTREAM =================
print('\n' + '-' * 100)
print('PART 1  SOCIAL CHANNEL PROVENANCE')
print('-' * 100)
QM = os.path.join(ROOT, 'Quorum Master.xlsx')
QP = os.path.join(ROOT, 'Quorum Processed.xlsx')
QMT = os.path.join(ROOT, 'data', 'stance', 'quorum_member_tweets.csv')
for lbl, p in (('Quorum Master.xlsx (raw licensed export)', QM),
               ('Quorum Processed.xlsx', QP),
               ('data/stance/quorum_member_tweets.csv', QMT),
               ('Social Master.xlsx', os.path.join(ROOT, 'Social Master.xlsx'))):
    print('   %-44s %s' % (lbl, stat(p)))

if os.path.exists(QP):
    wb = openpyxl.load_workbook(QP, read_only=True, data_only=True)
    print('\n   Quorum Processed.xlsx sheets: %s' % wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        n = 0
        for _ in ws.iter_rows(min_row=2, values_only=True):
            n += 1
        print('      %-12s data rows: %6d' % (sn, n))
    if 'Summary' in wb.sheetnames:
        ws = wb['Summary']
        hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        for r in ws.iter_rows(min_row=2, max_row=2, values_only=True):
            print('\n   process_quorum Summary sheet (its own persisted census):')
            for h, v in zip(hdr, r):
                print('      %-28s %s' % (h, v))
    # bucket census on the Members/Staffers/Excluded sheets
    for sn in ('Members', 'Staffers', 'Excluded'):
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        H = {h: i for i, h in enumerate(hdr)}
        buck = collections.Counter()
        rel = collections.Counter()
        st = collections.Counter()
        tot = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            tot += 1
            buck[str(r[H['bucket']]) if 'bucket' in H else '?'] += 1
            if 'cu_relevant' in H:
                rel[str(r[H['cu_relevant']])] += 1
            if 'stance' in H:
                st[str(r[H['stance']] or '')] += 1
        print('\n   sheet %-10s rows=%d' % (sn, tot))
        print('      bucket      : %s' % dict(buck.most_common()))
        print('      cu_relevant : %s' % dict(rel.most_common()))
        print('      stance      : %s' % dict(st.most_common()))
    wb.close()

qmt = rd(QMT)
print('\n   quorum_member_tweets.csv rows: %d' % len(qmt))
print("   RULE (process_quorum.py:150)  memtw=members[(members['cu_relevant']==1)]")
print("   RULE (process_quorum.py:119-122) TRUE duplicates = same Source + normalized Content + date;")
print("        df.drop_duplicates(subset='_k',keep='first')   -- retweets across sources survive")
if qmt:
    print('   quorum_member_tweets.csv columns: %s' % list(qmt[0].keys()))
    print('   platform (Document Type) split: %s'
          % dict(collections.Counter(r.get('Document Type') or '' for r in qmt).most_common()))
    print('   stance split: %s' % dict(collections.Counter(r.get('stance') or '' for r in qmt).most_common()))

print("\n   RULE (build_social_master.py:60)  bc=b[b['rel']==1][list(qc.columns)]")
print("        -> a direct BLUESKY post enters the Corpus ONLY if CU_TERMS matched its text.")
print("   RULE (build_social_master.py:53)  verified_handle=int(verified_bluesky_handle(...))")
print("        -> look-alike handles are INCLUDED in the Corpus but flagged 0 and dropped downstream.")

# ================= PART 2: OWN-SOURCE CHANNELS =================
print('\n' + '-' * 100)
print('PART 2  OWN-SOURCE CHANNELS (speech / hearing / interview) as build_media_blend reads them')
print('-' * 100)
roster = json.load(io.open(os.path.join(ROOT, 'data', 'engine', 'congestion.json'), encoding='utf-8'))
resolver = Resolver([x['member'] for x in roster], {x['member']: x.get('bioguide') for x in roster})
SIGN = {'support': 1.0, 'oppose': -1.0, 'neutral': 0.0}
TODAY = datetime.date.today()

CHANNELS = [
    ('speech', os.path.join(ROOT, 'speech_stance_final.csv'),
     [os.path.join(ROOT, 'speeches_cu.csv'), os.path.join(ROOT, 'data', 'stance', 'speech_gold_queue.csv')],
     'own_source_tier_weight', 1.0),
    ('hearing', os.path.join(ROOT, 'data', 'stance', 'hearing_stance_final.csv'),
     [os.path.join(ROOT, 'data', 'stance', 'hearings_cu.csv'),
      os.path.join(ROOT, 'data', 'stance', 'hearing_gold_queue.csv')],
     'hearing_tier_weight', 0.85),
    ('interview', os.path.join(ROOT, 'data', 'stance', 'interview_stance_final.csv'),
     [os.path.join(ROOT, 'data', 'stance', 'interviews_cu.csv'),
      os.path.join(ROOT, 'data', 'stance', 'interview_gold_queue.csv')],
     'own_source_tier_weight', 1.0),
]
for name, final, deps, weightfn, w in CHANNELS:
    print('\n== %s channel ==' % name.upper())
    print('   final:  %-52s %s' % (os.path.relpath(final, ROOT).replace('\\', '/'), stat(final)))
    for d in deps:
        print('   dep:    %-52s %s' % (os.path.relpath(d, ROOT).replace('\\', '/'), stat(d)))
    # staleness gate (_cu_stale, build_media_blend.py:130-137)
    stale = []
    if os.path.exists(final):
        fm = os.path.getmtime(final)
        stale = [os.path.basename(d) for d in deps if os.path.exists(d) and os.path.getmtime(d) > fm]
    print('   STALENESS GATE (_cu_stale, build_media_blend.py:130-137): %s'
          % ('STALE -> CHANNEL SKIPPED, newer inputs: %s' % stale if stale else 'fresh -> channel active'))
    rows = rd(final)
    print('   rows in: %d' % len(rows))
    if not rows:
        continue
    tiers = collections.Counter((r.get('tier') or '').strip().lower() for r in rows)
    stances = collections.Counter((r.get('stance') or '').strip().lower() for r in rows)
    print('   tier   : %s' % dict(tiers.most_common()))
    print('   stance : %s' % dict(stances.most_common()))
    if name == 'hearing':
        elig = [r for r in rows if MR.hearing_tier_weight(r.get('tier')) > 0.0]
        print("   RULE :192  _eligible=[r for r in _prows if MR.hearing_tier_weight(r.get('tier'))>0.0]")
        print('      gold-eligible: %d   non-gold EXCLUDED (_hreview): %d' % (len(elig), len(rows) - len(elig)))
        work = elig
    else:
        work = rows
    n_zero = n_unres = n_baddate = n_tier0 = n_nocat = n_signed = 0
    members = set()
    cats = collections.Counter()
    for r in work:
        sg = SIGN.get((r.get('stance') or '').strip().lower(), 0.0)
        if sg == 0.0:
            n_zero += 1
            continue
        m = resolver.resolve(r.get('member'), r.get('bioguide'))
        if not m:
            n_unres += 1
            continue
        try:
            datetime.date.fromisoformat((r.get('date') or '')[:10])
        except Exception:
            n_baddate += 1
            continue
        inten = (MR.hearing_tier_weight(r.get('tier')) if name == 'hearing'
                 else MR.own_source_tier_weight(r.get('tier')))
        if inten == 0.0:
            n_tier0 += 1
            continue
        cl = [c.strip() for c in (r.get('categories') or '').split(';') if c.strip()]
        if not cl:
            n_nocat += 1
            continue
        n_signed += 1
        members.add(m)
        for c in cl:
            cats[c] += 1
    print('   per-row disposition of the %d rows the channel works on:' % len(work))
    print("      RULE  if _sg==0.0: continue            neutral/unrelated  -> %6d  (unprinted)" % n_zero)
    print("      RULE  if not _m: continue              unresolved member  -> %6d  (unprinted)" % n_unres)
    print("      RULE  except Exception: continue       unparseable date   -> %6d  (unprinted)" % n_baddate)
    print("      RULE  if _inten==0.0: continue         unknown tier       -> %6d  (unprinted)" % n_tier0)
    print("      RULE  if not _cats: continue           no category        -> %6d  (unprinted)" % n_nocat)
    print('      SIGNED rows fed to the press side (own-source weight %.2f) -> %6d' % (w, n_signed))
    print('      members reached                                            -> %6d' % len(members))
    tot = n_zero + n_unres + n_baddate + n_tier0 + n_nocat + n_signed
    print('   CLOSES? %d = %d -> %s' % (len(work), tot, 'YES' if tot == len(work) else
                                        'NO  unaccounted=%d' % (len(work) - tot)))
    print('   per-event 1/n normalization applied (build_media_blend :167/:210/:252): one (member,date)'
          ' event contributes one unit regardless of passage count')
    print('   categories:')
    for k, v in sorted(cats.items(), key=lambda kv: -kv[1]):
        print('      %-38s %6d' % (k, v))

# ---- bluesky press-side channel: gated OFF ----------------------------------------------------
print('\n== BLUESKY press-side channel ==')
bcu = os.path.join(ROOT, 'data', 'stance', 'bluesky_cu.csv')
bsf = os.path.join(ROOT, 'data', 'stance', 'bluesky_stance_final.csv')
print('   data/stance/bluesky_cu.csv            %s  rows=%d' % (stat(bcu), len(rd(bcu))))
print('   data/stance/bluesky_stance_final.csv  %s  rows=%d' % (stat(bsf), len(rd(bsf))))
print("   RULE (build_media_blend.py:267-273) BLUESKY_ON default '0'; even at 1 the block prints a")
print('        refusal and injects NOTHING -- bluesky already rides the SOCIAL corpus, so a press-side')
print('        injection would double-count.  blc stays empty by construction.')
print('\n-- press_component_meta.json own_source_channels (engine self-report, for cross-check) --')
meta = json.load(io.open(os.path.join(ROOT, 'data', 'stance', 'press_component_meta.json'), encoding='utf-8'))
print(json.dumps(meta.get('own_source_channels'), indent=3))
print('\nDONE R5')
