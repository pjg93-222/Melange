# -*- coding: utf-8 -*-
"""LANE A / R2 -- routing & display-column census (read-only).

Answers step 3: does every media/press category the router can emit have (a) a workbook
column in 'Press Release - Raw' (the engine insertion point) and (b) a dashboard issue
column?  The July finding was that 'CU Tax Status' had no dashboard issue column.

Opens the live workbook read_only=True, data_only=True.  Writes nothing.
"""
import os, sys, io, json, collections, datetime
LIB = r'C:\Users\pgoff\AppData\Local\Temp\claude\C--Users-pgoff-OneDrive-Desktop-Melange-Systems--claude-worktrees-unfinished-work-handoff-a3bbad\5b7f7c28-909b-48f8-a495-2da7aebcc0fe\scratchpad\laneA\lib'
sys.path.insert(0, LIB)
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl
import media_rules as MR

ROOT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WBP = os.path.join(ROOT, 'Moneyball Main Version - Blended.xlsx')

print('=' * 100)
print('R2  ROUTING / DISPLAY-COLUMN CENSUS   run', datetime.datetime.now().isoformat(timespec='seconds'))
print('=' * 100)
print('workbook %s  mtime %s' % (os.path.getsize(WBP),
      datetime.datetime.fromtimestamp(os.path.getmtime(WBP)).isoformat(sep=' ')))

MEDIA_CATS = list(MR.RAW_PATTERNS.keys())     # == build_media_blend HDRMAP keys (the 20)
print('\nmedia router raw categories (media_rules.RAW_PATTERNS / build_media_blend HDRMAP): %d' % len(MEDIA_CATS))

wb = openpyxl.load_workbook(WBP, read_only=True, data_only=True)
print('sheets: %d' % len(wb.sheetnames))

# ---- A. 'Press Release - Raw' : the engine's write target -------------------------------------
rw = wb['Press Release - Raw']
rows = rw.iter_rows(min_row=1, max_row=1, values_only=True)
hdr_row = next(rows)
hdr = {}
for i, v in enumerate(hdr_row[:30], start=1):
    if v is not None:
        hdr[v] = i
print("\n-- A. sheet 'Press Release - Raw' (build_media_blend.py:279-282 write target) --")
print('header cells in cols 1..30: %d' % len(hdr))
for name, col in sorted(hdr.items(), key=lambda kv: kv[1]):
    print('   col %2d  %s' % (col, name))

print("\nRULE (build_media_blend.py:432-433)  col=hdr.get(hname);  if not col: print('  !! header not found:',hname); continue")
missing = [c for c in MEDIA_CATS if c not in hdr]
print('media categories WITHOUT a column in this sheet -> value never written: %d' % len(missing))
for c in missing:
    print('   MISSING COLUMN  %s' % c)
present = [c for c in MEDIA_CATS if c in hdr]
print('media categories WITH a column: %d' % len(present))
extra = [h for h in hdr if h not in MEDIA_CATS]
print('sheet headers that are not media categories (%d): %s' % (len(extra), extra))

# Z1 provenance note = which run last wrote the blend
z1 = None
for r in rw.iter_rows(min_row=1, max_row=1, min_col=26, max_col=26, values_only=True):
    z1 = r[0]
print('\nZ1 provenance note:')
print('   %s' % z1)

wb.close()

# ---- B. dashboard issue columns ---------------------------------------------------------------
print('\n-- B. dashboard issue columns --')
FL = json.load(io.open(os.path.join(ROOT, 'per_topic_alloc.json'), encoding='utf-8'))
print('per_topic_alloc.json keys = the desk/dashboard issue SOURCE list (make_member_desk.py:19,45): %d' % len(FL))
for k in FL:
    print('   %s' % k)

LBLCAT = {'Examination & Supervision': 'Examination & Supervision',
          'Credit Union Capital Requirements': 'CU Capital Requirements',
          'NCUA Budget & Share Insurance Fund': 'NCUA Budget & Insurance',
          'Interchange': 'Interchange', 'Payments': 'Payments',
          'Technology & Emerging Issues': 'Technology & Emerging',
          'Fraud Prevention': 'Fraud Prevention',
          'Data Privacy & Cybersecurity': 'Data Privacy & Cyber',
          'Charter Modernization & Governance': 'Charter Modernization',
          'Defense Credit Union Issues': 'Defense CU Issues',
          'Field of Membership': 'Field of Membership', 'CDFI Fund': 'CDFI Fund',
          'CFPB Restrain': 'CFPB', 'CFPB Reach (-)': 'CFPB',
          'Telephone Consumer Protection Act Issues': 'TCPA',
          'Credit Card & Consumer-Fee Mandates': 'Credit Card Mandates',
          'Member Business Lending': 'Member Business Lending',
          'Mortgage Lending & Housing': 'Mortgage & Housing',
          'Credit Union Tax Status': 'CU Tax Status'}
DROPPED = {'CFPB Reach (-)'}
source_issues = [l for l in FL if l not in DROPPED]
issue_cat = [LBLCAT.get(l, l) for l in source_issues]
print('\ndashboard issue contract: %d source issues -> %d canonical categories'
      % (len(source_issues), len(set(issue_cat))))
print('display-clone dropped by DROPPED_DISPLAY_CLONES: %s'
      % [l for l in FL if l in DROPPED])
print('canonical categories displayed as an issue column:')
for c in sorted(set(issue_cat)):
    print('   %s' % c)

# map every media category to its canonical name, then ask: is it a dashboard issue?
print('\n-- C. media category -> canonical -> dashboard issue? --')
print('%-38s %-28s %-9s %-9s' % ('media raw category', 'canonical (RAW_TO_STANCE)', 'wb col', 'dash col'))
disp_set = set(issue_cat)
no_dash = []
for c in MEDIA_CATS:
    canon = MR.RAW_TO_STANCE.get(c, c)
    incol = 'YES' if c in hdr else 'NO'
    indash = 'YES' if canon in disp_set else 'NO'
    if indash == 'NO':
        no_dash.append((c, canon))
    print('%-38s %-28s %-9s %-9s' % (c, canon, incol, indash))
print('\nmedia categories with NO dashboard issue column: %d' % len(no_dash))
for c, canon in no_dash:
    print('   %-38s -> %s' % (c, canon))

# ---- D. dashboard_data.json actual issue list -------------------------------------------------
print('\n-- D. dashboard_data.json shipped issue list --')
DDP = os.path.join(ROOT, 'dashboard_data.json')
print('dashboard_data.json mtime %s' % datetime.datetime.fromtimestamp(os.path.getmtime(DDP)).isoformat(sep=' '))
dd = json.load(io.open(DDP, encoding='utf-8'))
issues = dd.get('issues') or []
cats = (dd.get('engine') or {}).get('categories') or []
print('bundle issues (%d): %s' % (len(issues), issues))
print('bundle engine.categories (%d): %s' % (len(cats), cats))

# ---- E. category_stance.json matrix -----------------------------------------------------------
print('\n-- E. category_stance.json matrix categories (the per-category stance surface) --')
CSP = os.path.join(ROOT, 'category_stance.json')
cs = json.load(io.open(CSP, encoding='utf-8'))
mat = cs.get('matrix') or {}
mcats = collections.Counter()
for m, d in mat.items():
    for k in (d or {}):
        mcats[k] += 1
print('members in matrix: %d' % len(mat))
print('matrix categories (%d):' % len(mcats))
for k, v in sorted(mcats.items(), key=lambda kv: (-kv[1], kv[0])):
    print('   %-30s present on %d members' % (k, v))
print('\nDONE R2')
