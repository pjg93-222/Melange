# -*- coding: utf-8 -*-
"""N1a probe: what the desk sheet carries vs what the tone CSVs carry.  READ-ONLY.

Nothing under fec/ is executed: press_stream.py is a rules module (no top-level I/O,
verified line by line by the A5 lane) and is imported with MONEYBALL_ROOT pointed at the
LIVE tree so BASE_TONE_PATH/REFRESH_TONE_PATH resolve to the live inputs.
"""
import collections
import csv
import io
import os
import sys

LIVE = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\.claude\worktrees\unfinished-work-handoff-a3bbad'
os.environ['MONEYBALL_ROOT'] = LIVE
sys.path.insert(0, os.path.join(WT, 'fec'))
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

import openpyxl
import press_stream as PS

print('press_stream ROOT  :', PS.BASE_TONE_PATH)
print('press_stream REFR  :', PS.REFRESH_TONE_PATH)
print('exists base/refresh:', os.path.isfile(PS.BASE_TONE_PATH), os.path.isfile(PS.REFRESH_TONE_PATH))

WB = os.path.join(LIVE, 'Moneyball Main Version - Blended.xlsx')
wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
print('\nworkbook sheets containing "Press":', [s for s in wb.sheetnames if 'Press' in s])
pw = wb['Press Release Articles']
header = next(pw.iter_rows(min_row=1, max_row=1, values_only=True))
print('\n-- SHEET "Press Release Articles" header (%d columns) --' % len(header))
for i, h in enumerate(header):
    print('   %2d  %r' % (i, h))

ph = {h: i for i, h in enumerate(header)}
DESK_FIELDS = ['member', 'bioguide', 'url', 'stance_src', 'score', 'stance_sign',
               'score_graded', 'title', 'stance', 'stance_conf', 'date', 'categories',
               'relevance']
print('\n-- the 13 columns the desk actually reads (pv(r, k) call sites) --')
for k in DESK_FIELDS:
    print('   %-14s sheet index %s' % (k, ph.get(k, 'ABSENT')))

# type census of every desk-read column + a 3-row sample
types = {k: collections.Counter() for k in DESK_FIELDS}
samples = []
nrows = 0
stance_src_vals = collections.Counter()
date_shapes = collections.Counter()
billstatus = 0
for r in pw.iter_rows(min_row=2, values_only=True):
    nrows += 1
    for k in DESK_FIELDS:
        v = r[ph[k]] if k in ph else None
        types[k][type(v).__name__] += 1
    stance_src_vals[str(r[ph['stance_src']] if 'stance_src' in ph else None)] += 1
    d = r[ph['date']] if 'date' in ph else None
    s = str(d)
    date_shapes['%s|%s' % (type(d).__name__, 'digit4' if s[:4].isdigit() else 'NOT-digit4')] += 1
    if str(r[ph['url']] or '').startswith('billstatus:'):
        billstatus += 1
    if len(samples) < 3:
        samples.append({k: (r[ph[k]] if k in ph else None) for k in DESK_FIELDS})
wb.close()
print('\nsheet data rows: %d   billstatus: rows: %d' % (nrows, billstatus))
print('\n-- python type census per desk-read column --')
for k in DESK_FIELDS:
    print('   %-14s %s' % (k, dict(types[k])))
print('\n-- sheet date shape (what fdate() sees) --')
for k, v in date_shapes.most_common():
    print('   %-30s %6d' % (k, v))
print('\n-- sheet stance_src values (the `if src == \'bill\'` branch) --')
for k, v in stance_src_vals.most_common(12):
    print('   %-26s %6d' % (k, v))
print('\n-- 3 sheet rows verbatim --')
for s in samples:
    print('   ' + repr(s))


def rd(path):
    if not os.path.isfile(path):
        return []
    with io.open(path, encoding='utf-8-sig', newline='') as fh:
        return list(csv.DictReader(fh))


base, refr = rd(PS.BASE_TONE_PATH), rd(PS.REFRESH_TONE_PATH)
print('\n-- TONE CSV headers --')
print('base    (%d rows) %d cols: %s' % (len(base), len(base[0]), list(base[0])))
print('refresh (%d rows) %d cols: %s' % (len(refr), len(refr[0]), list(refr[0])))
print('\ndesk-read columns MISSING from base   :',
      [k for k in DESK_FIELDS if k not in base[0]])
print('desk-read columns MISSING from refresh:',
      [k for k in DESK_FIELDS if k not in refr[0]])
sheet_only_cols = [h for h in header if h and h not in base[0]]
print('\nsheet columns with NO base-CSV counterpart:', sheet_only_cols)
csv_only_cols = [c for c in base[0] if c not in ph]
print('base-CSV columns with NO sheet counterpart:', csv_only_cols)
print('refresh-CSV columns with NO sheet counterpart:', [c for c in refr[0] if c not in ph])

print('\n-- tone CSV date shapes (what fdate() would see UNCHANGED) --')
for tag, rows in (('base', base), ('refresh', refr)):
    shapes = collections.Counter()
    for r in rows:
        s = str(r.get('date'))
        shapes['digit4' if s[:4].isdigit() else 'NOT-digit4(%s)' % s[:10]] += 1
    print('   %-8s %s' % (tag, dict(shapes.most_common(6))))
print('\n-- tone CSV stance_src values --')
for tag, rows in (('base', base), ('refresh', refr)):
    print('   %-8s %s' % (tag, dict(collections.Counter(
        str(r.get('stance_src')) for r in rows).most_common(8))))
print('\n-- tone CSV relevance values --')
for tag, rows in (('base', base), ('refresh', refr)):
    print('   %-8s %s' % (tag, dict(collections.Counter(
        str(r.get('relevance')) for r in rows).most_common(8))))
print('\nDONE probe_map')
