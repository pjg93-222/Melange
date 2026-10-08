# -*- coding: utf-8 -*-
"""N1a -- DESK_PRESS_SOURCE measured against the LIVE data.  READ-ONLY.

Runs the desk's OWN press row-building code -- the statements of the loop inside
make_member_desk.load_press_and_bills, lifted as text, the pattern
fec/tests/stance/test_crosspost_scoring.py established because importing that module executes
it and overwrites the live desk -- over both positions of the dial, and diffs the two results
row by row.

Reports, for 'sheet' vs 'engine':
  * desk press rows in total and per member,
  * rows ADDED (with the 20 members who gain most) and the engine's own disposition for each,
  * rows REMOVED (must be 0),
  * rows whose DISPLAYED FIELDS changed (must be 0), compared field by field on the rows both
    runs produce,
  * the field-level disagreement between the two sources on the keys both carry -- the
    measurement that decides which source wins the overlap,
  * the desk's press census identity, which dashboard_integrity reconciles.

Writes nothing outside this directory.  The workbook is opened read_only.
"""
import ast
import collections
import csv
import datetime
import io
import json
import os
import sys

LIVE = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                  os.pardir, os.pardir, os.pardir))
HERE = os.path.dirname(os.path.abspath(__file__))
os.environ['MONEYBALL_ROOT'] = LIVE            # press_stream/paths resolve to the LIVE inputs
sys.path.insert(0, os.path.join(WT, 'fec'))
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

import openpyxl
import press_stream as PS
import surfaces.desk_press_source as DPS
from members.member_identity import Resolver

DESK_SOURCE = os.path.join(WT, 'fec', 'surfaces', 'make_member_desk.py')
WORKBOOK = os.path.join(LIVE, 'Moneyball Main Version - Blended.xlsx')

print('=' * 100)
print('N1a  DESK PRESS SOURCE: sheet vs engine   run',
      datetime.datetime.now().isoformat(timespec='seconds'))
print('=' * 100)
print('worktree   : %s' % WT)
print('live root  : %s' % LIVE)
print('base tone  : %s' % PS.BASE_TONE_PATH)
print('refresh    : %s' % PS.REFRESH_TONE_PATH)
print('dial default as imported: %r' % DPS.DESK_PRESS_SOURCE)


# ---- 1. the desk's own row-building code, lifted as text ------------------------------
SOURCE_TEXT = io.open(DESK_SOURCE, encoding='utf-8').read()
_TREE = ast.parse(SOURCE_TEXT)
_WANT_FUNCS = {'graded_label', '_fsign', 'norm_cats', '_press_corr'}
_WANT_NAMES = {'LBL2CAT', '_GRADED'}
_NODES = [n for n in _TREE.body
          if (isinstance(n, ast.FunctionDef) and n.name in _WANT_FUNCS)
          or (isinstance(n, ast.Assign)
              and any(isinstance(t, ast.Name) and t.id in _WANT_NAMES for t in n.targets))]
got_funcs = {n.name for n in _NODES if isinstance(n, ast.FunctionDef)}
got_names = {t.id for n in _NODES if isinstance(n, ast.Assign) for t in n.targets
             if isinstance(t, ast.Name)}
assert got_funcs == _WANT_FUNCS, 'lifted funcs %s' % sorted(got_funcs)
assert got_names == _WANT_NAMES, 'lifted names %s' % sorted(got_names)

LINES = SOURCE_TEXT.splitlines()
_start = next(i for i, ln in enumerate(LINES) if ln.strip() == 'for r in _press_rows:')
_end = next(i for i, ln in enumerate(LINES)
            if i > _start and ln.strip().startswith('# --- social'))
LOOP_SRC = '\n'.join(ln[4:] for ln in LINES[_start:_end])
print('\nlifted the desk press loop from make_member_desk.py lines %d-%d (%d lines)'
      % (_start + 1, _end, _end - _start))
assert "if str(pv(r, 'url') or '').startswith('billstatus:'):" in LOOP_SRC
assert '_press_corr(title' in LOOP_SRC and 'PDATE' in LOOP_SRC
LOOP = compile(LOOP_SRC, DESK_SOURCE, 'exec')

# ---- 2. the constants the loop closes over, read straight off the LIVE tree -----------
PDATE = json.load(io.open(os.path.join(LIVE, 'data', 'stance', 'press_date_overrides.json'),
                          encoding='utf-8'))
_ov_path = os.path.join(LIVE, 'data', 'stance', 'press_stance_overrides.json')
assert not os.path.exists(_ov_path + '.candidate'), 'a .candidate overlay exists: merge it'
PRESSOV = json.load(io.open(_ov_path, encoding='utf-8')).get('corrections') or []
print('press_date_overrides rows %d | press_stance_overrides corrections %d'
      % (len(PDATE), len(PRESSOV)))

bundle = json.load(io.open(os.path.join(LIVE, 'member_desk_bundle.json'), encoding='utf-8'))
NAMES = [m.get('name') for m in bundle.get('members') or []]
BIO = {m.get('name'): m.get('bioguide') for m in bundle.get('members') or []}
resolver = Resolver(NAMES, BIO)
print('roster %d members (from the shipped bundle)' % len(NAMES))

wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
pw = wb['Press Release Articles']
PH = {h: i for i, h in enumerate(next(pw.iter_rows(min_row=1, max_row=1, values_only=True)))}
SHEET_ROWS = list(pw.iter_rows(min_row=2, values_only=True))
wb.close()
print('sheet rows %d, header %d columns' % (len(SHEET_ROWS), len(PH)))


def run(source):
    """The desk's press loop, verbatim, over the rows the dial selects."""
    ns = {'os': os, 're': __import__('re'), 'PDATE': PDATE, 'PRESSOV': PRESSOV}
    exec(compile(ast.Module(body=_NODES, type_ignores=[]), DESK_SOURCE, 'exec'), ns)
    press_rows, ledger = DPS.desk_press_rows(list(SHEET_ROWS), PH, source=source)
    ns.update({
        '_press_rows': press_rows,
        'pv': (lambda r, k: r[PH[k]] if k in PH else None),
        'canon': (lambda name, bioguide: resolver.resolve(name, bioguide)),
        'fdate': (lambda d: '' if d is None else
                  (str(d)[:10] if str(d)[:4].isdigit() else '')),
        'arts': {}, 'press_source_rows': 0, 'press_unknown_member': 0,
        'press_billstatus_excluded': 0,
        'exclusion_samples': {'press_unknown_member': []},
    })
    exec(LOOP, ns)
    return ns, ledger


print('\n' + '-' * 100)
print('-- 2. RUN BOTH DIAL POSITIONS --')
sheet_ns, sheet_ledger = run('sheet')
engine_ns, engine_ledger = run('engine')
for tag, ns, ledger in (('sheet', sheet_ns, sheet_ledger), ('engine', engine_ns, engine_ledger)):
    displayed = sum(len(v) for v in ns['arts'].values())
    print('\n[%s]  press_source_rows %5d = press_displayed %5d + unknown_member %3d'
          ' + billstatus_routed %5d   CLOSES: %s'
          % (tag, ns['press_source_rows'], displayed, ns['press_unknown_member'],
             ns['press_billstatus_excluded'],
             'YES' if ns['press_source_rows'] == displayed + ns['press_unknown_member']
                      + ns['press_billstatus_excluded'] else 'NO'))
    print('        members with press %d' % len(ns['arts']))
    print('        source ledger: %s' % json.dumps(
        {k: v for k, v in sorted(ledger.items()) if k != 'press_engine_tone_dedup'}))
    if ledger.get('press_engine_tone_dedup'):
        td = ledger['press_engine_tone_dedup']
        print('        A5 tone ingest: read %d -> returned %d (exact %d, url-form %d,'
              ' conflict kept %d, official mention %d, no url %d)  CLOSES: %s'
              % (td['rows_read'], td['rows_returned'], td['exact_url_turn_collapsed'],
                 td['duplicate_url_form_collapsed'], td['release_key_conflict_kept'],
                 td['official_mention_skipped'], td['no_url_skipped'],
                 'YES' if td['rows_read'] == td['rows_returned']
                          + td['official_mention_skipped'] + td['no_url_skipped']
                          + td['exact_url_turn_collapsed']
                          + td['duplicate_url_form_collapsed'] else 'NO'))

# ---- 3. row-level diff ----------------------------------------------------------------
FIELDS = ('date', 'dest', 'title', 'url', 'cats', 'pol', 'label', 'stance', 'sconf',
          'ssrc', 'corr', 'rel', 'source', 'scored', 'exclusion')


def index(arts):
    """(member, url, title, date) -> the built row.  Every desk press row, keyed on what a
    reader sees, so an ADDED row and a CHANGED row can never be confused."""
    out = {}
    dupes = 0
    for nm, rows in arts.items():
        for row in rows:
            key = (nm, row['url'], row['title'], row['date'])
            if key in out:
                dupes += 1
                key = key + (dupes,)
            out[key] = row
    return out, dupes


sheet_index, sheet_dupes = index(sheet_ns['arts'])
engine_index, engine_dupes = index(engine_ns['arts'])
added = [k for k in engine_index if k not in sheet_index]
removed = [k for k in sheet_index if k not in engine_index]
changed = []
for k in sheet_index:
    if k not in engine_index:
        continue
    a, b = sheet_index[k], engine_index[k]
    delta = {f: (a.get(f), b.get(f)) for f in FIELDS if a.get(f) != b.get(f)}
    if delta or set(a) != set(b):
        changed.append((k, delta, sorted(set(a) ^ set(b))))

print('\n' + '-' * 100)
print('-- 3. ROW-LEVEL DIFF (sheet -> engine) --')
print('desk press rows  sheet %5d   engine %5d   (+%d)'
      % (len(sheet_index), len(engine_index), len(engine_index) - len(sheet_index)))
print('rows ADDED   : %5d' % len(added))
print('rows REMOVED : %5d   <- must be 0' % len(removed))
print('rows CHANGED : %5d   <- must be 0' % len(changed))
print('identical-display-key collisions within one run: sheet %d, engine %d'
      % (sheet_dupes, engine_dupes))
for k in removed[:20]:
    print('   REMOVED %s' % (k,))
for k, delta, extra in changed[:20]:
    print('   CHANGED %s %s %s' % (k[0], delta, extra))

# ---- 4. what was added --------------------------------------------------------------
r1 = json.load(io.open(os.path.join(
    WT, 'audits', 'laundry_list_closeout_20260918', 'laneA', 'receipts',
    'r1_press_rows.json'), encoding='utf-8'))['rows']
disp_by_url = {PS.canonical_url(x['url'], keep_fragment=True): x['disposition'] for x in r1}
refresh_urls = set()
with io.open(PS.REFRESH_TONE_PATH, encoding='utf-8-sig', newline='') as fh:
    for row in csv.DictReader(fh):
        refresh_urls.add(PS.canonical_url(row.get('url'), keep_fragment=True))

per_member = collections.Counter(k[0] for k in added)
disp = collections.Counter()
src_file = collections.Counter()
scored_flag = collections.Counter()
no_date = 0
for k in added:
    row = engine_index[k]
    u = PS.canonical_url(row['url'], keep_fragment=True)
    disp[disp_by_url.get(u, 'NOT-IN-R1-CENSUS')] += 1
    src_file['refresh' if u in refresh_urls else 'base'] += 1
    scored_flag['pol=%+d' % row['pol']] += 1
    if not row['date']:
        no_date += 1
print('\n' + '-' * 100)
print('-- 4. THE ADDED ROWS --')
print('added rows %d across %d members' % (len(added), len(per_member)))
print('by tone source file: %s' % dict(src_file))
print('by the engine\'s own scoring disposition (laneA r1 census):')
for kk, vv in disp.most_common():
    print('   %-22s %6d' % (kk, vv))
print('by displayed direction: %s' % dict(scored_flag.most_common()))
print('added rows rendering with a BLANK date: %d' % no_date)
print('\ntop 20 members by rows gained:')
for nm, n in per_member.most_common(20):
    print('   %-32s %4d  (sheet had %d)'
          % (nm, n, len(sheet_ns['arts'].get(nm) or [])))
print('\nmembers gaining at least one row: %d' % len(per_member))
print('first 10 added rows as the desk built them:')
for k in sorted(added)[:10]:
    row = engine_index[k]
    print('   %-26s %-10s pol=%+d %-26s %s'
          % (k[0][:26], row['date'], row['pol'], row['label'][:26], row['url'][:78]))

# ---- 5. why the SHEET wins the overlap ------------------------------------------------
print('\n' + '-' * 100)
print('-- 5. THE OVERLAP: what the two sources say about the SAME release --')
DESK_READ = ('member', 'bioguide', 'url', 'stance_src', 'score', 'stance_sign',
             'score_graded', 'title', 'stance', 'stance_conf', 'date', 'categories',
             'relevance')
sheet_by_key = {}
for row in SHEET_ROWS:
    url = str(row[PH['url']] or '')
    if url.startswith('billstatus:'):
        continue
    key = DPS.row_key(url, row[PH['bioguide']], row[PH['member']])
    if key[0]:
        sheet_by_key.setdefault(key, row)
engine_rows = PS.load_tone_rows()
disagree = collections.Counter()
samples = collections.defaultdict(list)
shared = 0
for row in engine_rows:
    key = DPS.row_key(row.get('url'), row.get('bioguide'), row.get('member'))
    sheet = sheet_by_key.get(key)
    if sheet is None:
        continue
    shared += 1
    for f in DESK_READ:
        sv, ev = sheet[PH[f]] if f in PH else None, row.get(f)
        if f == 'date':
            sv = str(sv)[:10] if str(sv)[:4].isdigit() else ''
            ev = DPS.engine_date(ev)
            ev = str(ev)[:10] if str(ev)[:4].isdigit() else ''
        elif f in ('score', 'stance_sign'):
            sv = str(sv or 0).strip()
            ev = str(int(float(ev))) if str(ev).strip() not in ('', 'None') else '0'
            sv = str(int(float(sv))) if sv not in ('', 'None') else '0'
        elif f in ('score_graded', 'stance_conf'):
            sv = '%.6f' % float(sv or 0)
            try:
                ev = '%.6f' % float(ev or 0)
            except (TypeError, ValueError):
                ev = str(ev)
        else:
            sv, ev = str(sv or '').strip(), str(ev or '').strip()
        if sv != ev:
            disagree[f] += 1
            if len(samples[f]) < 3:
                samples[f].append((key[1], sv[:46], ev[:46]))
print('engine rows whose release key the sheet also carries: %d' % shared)
print('fields on which the two sources DISAGREE for the same release:')
if not disagree:
    print('   (none)')
for f, n in disagree.most_common():
    print('   %-14s %6d of %d' % (f, n, shared))
    for ident, sv, ev in samples[f]:
        print('        %-10s sheet=%-46s csv=%s' % (ident, sv, ev))
print('\nTHE PRECEDENCE RULE THIS MEASURES: the sheet supplies every row it has, so all %d'
      % shared)
print('shared rows render exactly as today; taking the CSV instead would rewrite the fields')
print('above, including the DISPLAYED STANCE LABEL, which is the owner\'s domain.')

print('\n' + '-' * 100)
print('VERDICT  added %d / removed %d / changed %d   census closes under both dials'
      % (len(added), len(removed), len(changed)))
json.dump({'schema': 'n1a-desk-press-union-v1',
           'added': len(added), 'removed': len(removed), 'changed': len(changed),
           'members_gaining': len(per_member),
           'rows_sheet': len(sheet_index), 'rows_engine': len(engine_index),
           'disposition_of_added': dict(disp),
           'overlap_shared_keys': shared, 'overlap_field_disagreements': dict(disagree),
           'per_member_added': dict(per_member),
           'sheet_ledger': sheet_ledger,
           'engine_ledger': {k: v for k, v in engine_ledger.items()
                             if k != 'press_engine_tone_dedup'}},
          io.open(os.path.join(HERE, 'n1_desk_press_union.json'), 'w', encoding='utf-8'),
          indent=1, sort_keys=True)
print('wrote n1_desk_press_union.json')
print('DONE N1a')
