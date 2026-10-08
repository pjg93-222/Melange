# -*- coding: utf-8 -*-
"""N1a probe 2: sheet vs engine press rows keyed the way the engine keys them.  READ-ONLY.

Answers the three questions the union design turns on:
  1. how many engine rows the sheet already carries, how many it does not, and how many
     sheet rows the engine does not carry;
  2. for the rows BOTH carry, whether the 13 values the desk reads agree -- because the
     desk's row build is a pure function of those 13 values plus module constants that do
     not depend on the source, agreement there IS proof that the rendered row is identical;
  3. what the engine's own disposition says about the rows that would be added.
"""
import collections
import csv
import io
import json
import os
import sys

LIVE = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems'
WT = r'C:\Users\pgoff\OneDrive\Desktop\Melange Systems\.claude\worktrees\unfinished-work-handoff-a3bbad'
os.environ['MONEYBALL_ROOT'] = LIVE
sys.path.insert(0, os.path.join(WT, 'fec'))
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

import openpyxl
import press_stream as PS
from members.member_identity import Resolver

DESK_FIELDS = ('member', 'bioguide', 'url', 'stance_src', 'score', 'stance_sign',
               'score_graded', 'title', 'stance', 'stance_conf', 'date', 'categories',
               'relevance')

# --- the desk's own resolver, over the same roster the desk uses -----------------------
bundle = json.load(io.open(os.path.join(LIVE, 'member_desk_bundle.json'), encoding='utf-8'))
names = [m.get('name') for m in bundle.get('members') or []]
bio = {m.get('name'): m.get('bioguide') for m in bundle.get('members') or []}
resolver = Resolver(names, bio)
print('roster from the shipped bundle: %d members' % len(names))

# --- sheet side ------------------------------------------------------------------------
wb = openpyxl.load_workbook(os.path.join(LIVE, 'Moneyball Main Version - Blended.xlsx'),
                            read_only=True, data_only=True)
pw = wb['Press Release Articles']
ph = {h: i for i, h in enumerate(next(pw.iter_rows(min_row=1, max_row=1, values_only=True)))}
sheet_rows = []
for r in pw.iter_rows(min_row=2, values_only=True):
    sheet_rows.append({k: (r[ph[k]] if k in ph else None) for k in DESK_FIELDS})
wb.close()

sheet_press = [r for r in sheet_rows if not str(r['url'] or '').startswith('billstatus:')]
sheet_bill = len(sheet_rows) - len(sheet_press)
print('sheet rows %d = %d press + %d billstatus' % (len(sheet_rows), len(sheet_press), sheet_bill))

# --- engine side -----------------------------------------------------------------------
ledger = {}
engine_rows = PS.load_tone_rows(ledger=ledger)
print('engine rows from load_tone_rows(): %d   (ledger rows_read=%d, canon=%s)'
      % (len(engine_rows), ledger['rows_read'], ledger['canon']))

# --- one key for both sides: the engine's own turn key ---------------------------------
def key_of(row):
    return PS.tone_turn_key({'url': row.get('url'), 'bioguide': row.get('bioguide'),
                             'member': row.get('member')})

sheet_by_key = collections.defaultdict(list)
for r in sheet_press:
    sheet_by_key[key_of(r)].append(r)
engine_by_key = collections.defaultdict(list)
for r in engine_rows:
    engine_by_key[key_of(r)].append(r)

sheet_keys, engine_keys = set(sheet_by_key), set(engine_by_key)
overlap = sheet_keys & engine_keys
engine_only = engine_keys - sheet_keys
sheet_only = sheet_keys - engine_keys
print('\n-- keyed by (release_url_key, identity) --')
print('sheet press keys      : %6d  (rows %d)' % (len(sheet_keys), len(sheet_press)))
print('engine keys           : %6d  (rows %d)' % (len(engine_keys), len(engine_rows)))
print('BOTH   (overlap)      : %6d' % len(overlap))
print('ENGINE ONLY  -> added : %6d  rows %d'
      % (len(engine_only), sum(len(engine_by_key[k]) for k in engine_only)))
print('SHEET  ONLY  -> kept  : %6d  rows %d'
      % (len(sheet_only), sum(len(sheet_by_key[k]) for k in sheet_only)))
print('sheet keys carrying >1 row: %d' % sum(1 for k in sheet_by_key if len(sheet_by_key[k]) > 1))

# also the r9 key (canonical_url only, no identity) for continuity with the receipt
s_urls = set(PS.canonical_url(r['url'], keep_fragment=True) for r in sheet_press
             if resolver.resolve(r['member'], r['bioguide']))
e_urls = set(PS.canonical_url(r['url'], keep_fragment=True) for r in engine_rows)
print('\n-- r9 continuity (canonical_url only, member-resolvable sheet rows) --')
print('sheet urls %d  engine urls %d  engine-only %d  sheet-only %d'
      % (len(s_urls), len(e_urls), len(e_urls - s_urls), len(s_urls - e_urls)))

# --- 2. do the overlap rows agree field by field? -------------------------------------
def as_text(k, v):
    """The value as the desk's own code consumes it, so a type difference that the desk
    erases (int 0 vs '0') is not reported as a change."""
    if v is None:
        return ''
    if k == 'date':
        s = str(v)
        return s[:10] if s[:4].isdigit() else ''
    if k in ('score', 'stance_sign'):
        try:
            return '%d' % int(float(v))
        except (TypeError, ValueError):
            return str(v).strip()
    if k in ('score_graded', 'stance_conf'):
        try:
            return '%.6f' % float(v or 0)
        except (TypeError, ValueError):
            return str(v).strip()
    return str(v).strip()

diffs = collections.Counter()
examples = collections.defaultdict(list)
for k in sorted(overlap):
    s, e = sheet_by_key[k][0], engine_by_key[k][0]
    for f in DESK_FIELDS:
        sv, ev = as_text(f, s.get(f)), as_text(f, e.get(f))
        if sv != ev:
            diffs[f] += 1
            if len(examples[f]) < 4:
                examples[f].append((k[1], repr(s.get(f))[:60], repr(e.get(f))[:60]))
print('\n-- 2. OVERLAP rows: fields that DISAGREE between sheet and engine (%d keys) --'
      % len(overlap))
if not diffs:
    print('   (none -- every one of the 13 desk-read values agrees on every overlap key)')
for f, n in diffs.most_common():
    print('   %-14s %6d of %d keys' % (f, n, len(overlap)))
    for ident, sv, ev in examples[f]:
        print('        %-12s sheet=%-34s engine=%s' % (ident, sv, ev))

# --- 3. what would be added, and what the engine says about it ------------------------
r1 = json.load(io.open(os.path.join(
    WT, 'audits', 'laundry_list_closeout_20260918', 'laneA', 'receipts',
    'r1_press_rows.json'), encoding='utf-8'))['rows']
disp_by_url = {PS.canonical_url(x['url'], keep_fragment=True): x['disposition'] for x in r1}

added, unresolved = [], 0
disp = collections.Counter()
src_file = collections.Counter()
refresh_urls = set()
with io.open(PS.REFRESH_TONE_PATH, encoding='utf-8-sig', newline='') as fh:
    for r in csv.DictReader(fh):
        refresh_urls.add(PS.canonical_url(r.get('url'), keep_fragment=True))
for k in engine_only:
    for r in engine_by_key[k]:
        nm = resolver.resolve(r.get('member'), r.get('bioguide'))
        if not nm:
            unresolved += 1
            continue
        u = PS.canonical_url(r.get('url'), keep_fragment=True)
        added.append((nm, r))
        disp[disp_by_url.get(u, 'NOT-IN-R1-CENSUS')] += 1
        src_file['refresh' if u in refresh_urls else 'base'] += 1
print('\n-- 3. ADDED rows (engine-only, member resolvable by the desk resolver) --')
print('added rows             : %6d' % len(added))
print('engine-only but member UNRESOLVED by the desk (excluded, counted) : %d' % unresolved)
print('added rows by tone source file: %s' % dict(src_file))
print('\nadded rows by the ENGINE\'s own scoring disposition (laneA r1 census):')
for k, v in disp.most_common():
    print('   %-22s %6d' % (k, v))
per_member = collections.Counter(nm for nm, _ in added)
print('\ntop 20 members by rows gained:')
for nm, n in per_member.most_common(20):
    print('   %-30s %4d' % (nm, n))
print('members gaining at least one row: %d' % len(per_member))

# what fdate() would do to the added rows
blank = sum(1 for _, r in added if not as_text('date', r.get('date')))
print('\nadded rows whose date fdate() would BLANK (M/D/YYYY or NA): %d of %d'
      % (blank, len(added)))
print('added rows by stance_src: %s'
      % dict(collections.Counter(str(r.get('stance_src')) for _, r in added)))
print('\nDONE probe_overlap')
