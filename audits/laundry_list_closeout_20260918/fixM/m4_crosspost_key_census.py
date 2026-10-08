# -*- coding: utf-8 -*-
"""FIX N2 step 1 -- the cross-post key measured against the LIVE social corpus.

READ-ONLY.  No producer is executed.  Both shipped rules are lifted out of their producers
the way fec/tests/stance/test_crosspost_scoring.py lifts them (ast for the desk's
``social_scoring``, exec of the exact source lines for the engine's dedup block), so the
rules that run here are the shipped bytes, and the key comes from the shipped
``media_rules.crosspost_key``.

Reported:
  * duplicate groups and rows under the legacy 120-char key vs the full-text key
  * how many desk 'also posted on ...' notes change
  * scored-row counts at CROSSPOST_DEDUP=0 under both keys (must be IDENTICAL)
  * scored-row counts at CROSSPOST_DEDUP=1 under both keys (the rollback exposure)
  * the LiveBundle invariants of test_crosspost_scoring.py, re-checked under the new key
  * the engine site's ndupe / retained-row counts under both keys
"""
import ast
import collections
import io
import json
import os
import sys

LIVE = r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems"
DRAFT = r"C:\Users\pgoff\OneDrive\Desktop\Draft Systems"
HERE = os.path.dirname(os.path.abspath(__file__))
WT = os.path.abspath(os.path.join(HERE, os.pardir, os.pardir, os.pardir))
os.environ["MONEYBALL_ROOT"] = LIVE
os.environ["MONEYBALL_DRAFT"] = DRAFT
sys.path.insert(0, os.path.join(WT, "fec"))
import stance.media_rules as MR                                   # noqa: E402
from members.member_identity import Resolver                      # noqa: E402

DESK_SOURCE = os.path.join(WT, "fec", "surfaces", "make_member_desk.py")
BLEND_SOURCE = os.path.join(LIVE, "fec", "stance", "build_media_blend.py")

# ---- the desk rule, lifted without importing the desk's data producers ------------------
_TREE = ast.parse(io.open(DESK_SOURCE, encoding="utf-8").read())
_NODES = [n for n in _TREE.body
          if (isinstance(n, ast.FunctionDef) and n.name == "social_scoring")
          or (isinstance(n, ast.Assign)
              and any(isinstance(t, ast.Name) and t.id == "CROSSPOST_DEDUP"
                      for t in n.targets))]
NAMESPACE = {"os": os}
exec(compile(ast.Module(body=_NODES, type_ignores=[]), DESK_SOURCE, "exec"), NAMESPACE)
social_scoring = NAMESPACE["social_scoring"]

# ---- the engine rule, from the exact source lines of the LIVE producer -------------------
_lines = io.open(BLEND_SOURCE, encoding="utf-8").read().splitlines()
_start = next(i for i, l in enumerate(_lines) if l.strip().startswith("_dupe=_event in _seen"))
_end = next(i for i, l in enumerate(_lines) if l.strip().startswith("elif not _event[1]"))
BLEND_SRC = ("for _row_pass in (0,):\n"
             + "".join("    %s\n" % l[4:] for l in _lines[_start:_end + 1])
             + "    _kept=True\n")
BLEND_BLOCK = compile(BLEND_SRC, BLEND_SOURCE, "exec")


def blend_retained(events, dedup):
    state = {"_seen": set(), "ndupe": 0, "nlinkonly": 0, "DEDUP": dedup}
    kept = 0
    for event in events:
        state["_event"] = event
        state["_kept"] = False
        exec(BLEND_BLOCK, state)
        kept += bool(state["_kept"])
    return kept, state


print("=" * 108)
print("M4  CROSSPOST_KEY_CHARS measured on the LIVE social corpus")
print("=" * 108)
print("desk rule lifted from : %s" % DESK_SOURCE)
print("engine rule lifted from: %s" % BLEND_SOURCE)
print("shipped dial default   : CROSSPOST_KEY_CHARS=%r  (CROSSPOST_DEDUP=%r)"
      % (MR.CROSSPOST_KEY_CHARS, NAMESPACE["CROSSPOST_DEDUP"]))

# ---- corpus ------------------------------------------------------------------------------
import openpyxl                                                    # noqa: E402
book = openpyxl.load_workbook(os.path.join(LIVE, "Social Master.xlsx"),
                              read_only=True, data_only=True)
sheet = book["Corpus"]
head = {h: i for i, h in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))}


def cell(row, key):
    return row[head[key]] if key in head else None


stack = {}
stack_path = os.path.join(LIVE, "data", "stance", "tweet_stance_final.csv")
if os.path.exists(stack_path):
    import csv
    csv.field_size_limit(1 << 30)
    for sr in csv.DictReader(io.open(stack_path, encoding="utf-8-sig", newline="")):
        stack[(sr.get("member") or "", MR.normalized_text(sr.get("text")))] = sr

_roster = json.load(io.open(os.path.join(LIVE, "data", "engine", "congestion.json"),
                            encoding="utf-8"))
resolver = Resolver([x["member"] for x in _roster],
                    {x["member"]: x.get("bioguide") for x in _roster})
tierw = {"unanimous": 1.0, "majority": 0.7, "review": 0.35}
rows = []
skipped = collections.Counter()
for raw in sheet.iter_rows(min_row=2, values_only=True):
    skipped["source_rows"] += 1
    if str(cell(raw, "verified_handle") or "").strip() != "1":
        skipped["unverified_handle"] += 1
        continue
    member_raw = str(cell(raw, "member") or "")
    name = resolver.resolve(member_raw)
    if not name:
        skipped["unknown_member"] += 1
        continue
    text = str(cell(raw, "text") or "").strip()
    hit = stack.get((member_raw, MR.normalized_text(text)))
    stance = str((hit or {}).get("stack_stance") or cell(raw, "stance") or "neutral").lower()
    if stance == "unrelated":
        skipped["unrelated"] += 1
        continue
    if not MR.route_raw_categories(text):
        skipped["offtopic"] += 1
        continue
    rows.append({
        "name": name, "text": text,
        "platform": str(cell(raw, "platform") or "social"),
        "date": str(cell(raw, "date") or "")[:10],
        "sign": {"support": 1, "oppose": -1, "neutral": 0}.get(stance, 0),
        "weight": tierw.get(str((hit or {}).get("tier") or "").lower(), 0.5),
        "k120": MR.crosspost_key(text, chars=120),
        "kfull": MR.crosspost_key(text, chars=0),
    })
book.close()
print("\ncorpus rows %d -> desk-displayed social rows %d  (%s)"
      % (skipped["source_rows"], len(rows),
         ", ".join("%s=%d" % (k, v) for k, v in sorted(skipped.items()) if k != "source_rows")))


def replay(field):
    """The desk's own loop, in row order, keyed on `field`."""
    seen = {}
    out = []
    groups = collections.defaultdict(list)
    for row in rows:
        key = (row["name"], row[field])
        duplicate = key in seen
        if not duplicate:
            seen[key] = (row["platform"], row["date"])
        scored0, exclusion0, also0 = social_scoring(row["sign"], key[1], duplicate,
                                                    seen.get(key), dedup=False)
        scored1, exclusion1, _ = social_scoring(row["sign"], key[1], duplicate,
                                                seen.get(key), dedup=True)
        out.append({"row": row, "key": key, "duplicate": duplicate,
                    "scored0": scored0, "exclusion0": exclusion0, "also0": also0,
                    "scored1": scored1, "exclusion1": exclusion1})
        groups[key].append(out[-1])
    return out, groups


legacy, legacy_groups = replay("k120")
full, full_groups = replay("kfull")

print("\n" + "-" * 108)
print("A.  DUPLICATE GROUPS AND ROWS")
print("-" * 108)


def census(label, records, groups):
    multi = {k: v for k, v in groups.items() if len(v) > 1}
    repeats = sum(1 for r in records if r["duplicate"])
    print("  %-26s groups>1 = %5d   repeat rows = %5d   notes = %5d"
          % (label, len(multi), repeats, sum(1 for r in records if r["also0"])))
    return multi, repeats


legacy_multi, legacy_repeats = census("legacy key (120 chars)", legacy, legacy_groups)
full_multi, full_repeats = census("full key (whole text)", full, full_groups)
print("  false collisions dissolved: %d groups, %d rows"
      % (len(legacy_multi) - len(full_multi), legacy_repeats - full_repeats))
over120 = sum(1 for r in rows if len(r["kfull"]) > 120)
print("  rows longer than 120 normalised chars: %d of %d" % (over120, len(rows)))
# groups that the legacy key merged but the full key does NOT (the false collisions)
false_groups = 0
for key, members in legacy_multi.items():
    if len({m["row"]["kfull"] for m in members}) > 1:
        false_groups += 1
print("  legacy groups whose members do NOT share the full key: %d of %d"
      % (false_groups, len(legacy_multi)))

print("\n  A2. the SAME census at receipt r6's scope (stance non-empty AND verified_handle==1,")
print("      which is the 9,502-row population FINDINGS.md A3 counted 1,434/445/490 over):")
book2 = openpyxl.load_workbook(os.path.join(LIVE, "Social Master.xlsx"),
                               read_only=True, data_only=True)
sheet2 = book2["Corpus"]
head2 = {h: i for i, h in enumerate(next(sheet2.iter_rows(min_row=1, max_row=1,
                                                          values_only=True)))}
wide = []
for raw in sheet2.iter_rows(min_row=2, values_only=True):
    def c2(key):
        return raw[head2[key]] if key in head2 else None
    if str(c2("stance") or "").strip() == "" or str(c2("verified_handle") or "").strip() != "1":
        continue
    name = resolver.resolve(str(c2("member") or ""))
    if not name:
        continue
    text = str(c2("text") or "")
    wide.append((name, MR.crosspost_key(text, chars=120), MR.crosspost_key(text, chars=0)))
book2.close()
for label, index in (("legacy key (120 chars)", 1), ("full key (whole text)", 2)):
    buckets = collections.Counter((row[0], row[index]) for row in wide)
    multi = {k: n for k, n in buckets.items() if n > 1}
    print("      %-26s rows %5d   groups>1 = %5d   surplus rows = %5d"
          % (label, len(wide), len(multi), sum(n - 1 for n in multi.values())))
legacy_buckets = collections.defaultdict(set)
for name, k120, kfull in wide:
    legacy_buckets[(name, k120)].add(kfull)
false120 = {k: v for k, v in legacy_buckets.items() if len(v) > 1}
print("      legacy groups whose members do NOT share the full key (FALSE collisions): %d"
      % len(false120))

print("\n" + "-" * 108)
print("B.  DESK NOTES THAT CHANGE  ('also posted on <platform> <date>')")
print("-" * 108)
changed = lost = gained = same = 0
samples = []
for before, after in zip(legacy, full):
    assert before["row"] is after["row"]
    if before["also0"] == after["also0"]:
        same += 1
        continue
    changed += 1
    if before["also0"] and not after["also0"]:
        lost += 1
    elif after["also0"] and not before["also0"]:
        gained += 1
    if len(samples) < 12:
        samples.append({"member": before["row"]["name"],
                        "text": before["row"]["text"][:110],
                        "legacy_note": before["also0"], "full_note": after["also0"]})
print("  notes unchanged %d | changed %d  (of which %d removed, %d added, %d re-pointed)"
      % (same, changed, lost, gained, changed - lost - gained))
for s in samples:
    print("    %-22s legacy=%-34r full=%r" % (s["member"][:22], s["legacy_note"],
                                              s["full_note"]))
    print("       %s" % s["text"])

print("\n" + "-" * 108)
print("C.  SCORED ROWS -- must be IDENTICAL at CROSSPOST_DEDUP=0")
print("-" * 108)
for label, records in (("legacy key", legacy), ("full key", full)):
    print("  %-12s dial 0: %5d scored | dial 1: %5d scored | exclusions %s"
          % (label, sum(1 for r in records if r["scored0"]),
             sum(1 for r in records if r["scored1"]),
             dict(collections.Counter(r["exclusion0"] for r in records))))
identical = all(a["scored0"] == b["scored0"] and a["exclusion0"] == b["exclusion0"]
                for a, b in zip(legacy, full))
print("  every row's (scored, exclusion) identical at dial 0 under both keys: %s" % identical)
rollback = sum(1 for a, b in zip(legacy, full) if a["scored1"] != b["scored1"])
print("  rows whose dial-1 fate changes (would have been WRONGLY dropped on rollback): %d"
      % rollback)

print("\n" + "-" * 108)
print("D.  test_crosspost_scoring.py LiveBundle invariants, re-checked under the new key")
print("-" * 108)
bad = 0
for key, members in legacy_groups.items():
    noted = [m for m in full if m["row"] in [x["row"] for x in members] and m["also0"]]
    if len(noted) >= max(len(members), 2):
        bad += 1
print("  'no key has every row annotated' over legacy-key groups: violations = %d" % bad)
bad_first = sum(1 for k, members in full_groups.items() if members[0]["also0"])
print("  'the first post of a key never carries a note': violations = %d" % bad_first)
bad_noted = sum(1 for r in full if r["also0"] and
                not (r["scored0"] and r["exclusion0"] == "" and r["row"]["sign"] in (1, -1)
                     and r["key"][1]))
print("  'a noted row is scored, signed, texted and unexcluded': violations = %d" % bad_noted)
decompose = collections.Counter()
for r in full:
    if not r["duplicate"]:
        continue
    decompose["repeats"] += 1
    if not r["row"]["sign"]:
        decompose["neutral"] += 1
    elif not r["key"][1]:
        decompose["link_only"] += 1
    else:
        decompose["scored"] += 1
print("  repeats decompose exactly: %d = %d scored + %d neutral + %d link-only  -> %s"
      % (decompose["repeats"], decompose["scored"], decompose["neutral"],
         decompose["link_only"],
         decompose["repeats"] == decompose["scored"] + decompose["neutral"]
         + decompose["link_only"]))

print("\n" + "-" * 108)
print("E.  ENGINE SITE (build_media_blend.py's own dedup, LIVE bytes)")
print("-" * 108)
print("  lifted block:\n%s" % "".join("      %s\n" % l for l in BLEND_SRC.splitlines()))
for label, field in (("legacy key", "k120"), ("full key", "kfull")):
    events = [(r["name"], r[field]) for r in rows if r["sign"]]
    for dial in (False, True):
        kept, state = blend_retained(events, dial)
        print("  %-11s CROSSPOST_DEDUP=%d -> retained %5d | ndupe %5d | nlinkonly %4d"
              % (label, int(dial), kept, state["ndupe"], state["nlinkonly"]))

summary = {
    "desk_rows": len(rows),
    "legacy_groups": len(legacy_multi), "legacy_repeats": legacy_repeats,
    "full_groups": len(full_multi), "full_repeats": full_repeats,
    "false_collision_groups": false_groups,
    "rows_over_120_chars": over120,
    "notes_changed": changed, "notes_removed": lost, "notes_added": gained,
    "scored_identical_at_dial0": identical,
    "rollback_rows_rescued": rollback,
}
with io.open(os.path.join(HERE, "m4_crosspost_key_census.json"), "w", encoding="utf-8") as fh:
    json.dump(summary, fh, indent=1, ensure_ascii=False)
print("\nDONE M4")
