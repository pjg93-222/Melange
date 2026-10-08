# -*- coding: utf-8 -*-
"""FIX N2 step 2 -- the desk's OTHER reader of `seen`: the bluesky footprint.

make_member_desk.py ~2045 looks a bluesky post UP in `seen`, which is keyed by the
cross-post key.  That lookup had to move to the same key or it would silently stop
matching; this measures what moving it changes.  The routed-away posts are DISPLAY-ONLY
footprint records (scored=False, 'no CU relevance detected -- footprint only'), so nothing
here can move a score: the only effect is whether a bluesky post is shown as its own
footprint row or suppressed as already-counted social.

READ-ONLY.  social_media.db is opened in SQLite read-only (immutable) mode.
"""
import collections
import io
import json
import os
import sqlite3
import sys

LIVE = r"C:\Users\pgoff\OneDrive\Desktop\Melange Systems"
DRAFT = r"C:\Users\pgoff\OneDrive\Desktop\Draft Systems"
HERE = os.path.dirname(os.path.abspath(__file__))
WT = os.path.abspath(os.path.join(HERE, os.pardir, os.pardir, os.pardir))
os.environ["MONEYBALL_ROOT"] = LIVE
os.environ["MONEYBALL_DRAFT"] = DRAFT
sys.path.insert(0, os.path.join(WT, "fec"))
import stance.media_rules as MR                                    # noqa: E402
from members.member_identity import Resolver                       # noqa: E402
import openpyxl                                                    # noqa: E402

roster = json.load(io.open(os.path.join(LIVE, "data", "engine", "congestion.json"),
                           encoding="utf-8"))
resolver = Resolver([x["member"] for x in roster],
                    {x["member"]: x.get("bioguide") for x in roster})
bio2name = {m["bioguide"]: m["member"] for m in roster if m.get("bioguide")}

print("=" * 108)
print("M7  bluesky footprint: the desk's second reader of the cross-post key")
print("=" * 108)

# `seen` as the desk builds it: one entry per (member, key) over DISPLAYED social rows.
book = openpyxl.load_workbook(os.path.join(LIVE, "Social Master.xlsx"),
                              read_only=True, data_only=True)
sheet = book["Corpus"]
head = {h: i for i, h in enumerate(next(sheet.iter_rows(min_row=1, max_row=1,
                                                        values_only=True)))}
seen120, seenfull = set(), set()
for raw in sheet.iter_rows(min_row=2, values_only=True):
    def cell(key):
        return raw[head[key]] if key in head else None
    if str(cell("verified_handle") or "").strip() != "1":
        continue
    name = resolver.resolve(str(cell("member") or ""))
    if not name:
        continue
    text = str(cell("text") or "").strip()
    if not MR.route_raw_categories(text):
        continue
    seen120.add((name, MR.crosspost_key(text, chars=120)))
    seenfull.add((name, MR.crosspost_key(text, chars=0)))
book.close()
print("`seen` keys: legacy %d | full %d" % (len(seen120), len(seenfull)))

database = os.path.join(LIVE, "social_media.db")
uri = "file:%s?mode=ro&immutable=1" % database.replace("\\", "/")
connection = sqlite3.connect(uri, uri=True)
handles = {bg: (h or "") for bg, h in connection.execute(
    "SELECT bioguide, bluesky FROM members")}
counts = collections.Counter()
moved = []
for bg, _dt, text, _pid in connection.execute(
        "SELECT bioguide, created_at, text, post_id FROM posts WHERE platform='bluesky'"):
    counts["bluesky_post_rows"] += 1
    handle = handles.get(bg, "")
    if not handle or handle == "none-found" or "UNVERIFIED" in handle:
        counts["unverified_handle_excluded"] += 1
        continue
    name = resolver.resolve(bio2name.get(bg) or "", bg)
    if not name:
        counts["unknown_member_excluded"] += 1
        continue
    text = (text or "").strip()
    if not text:
        counts["empty_text_excluded"] += 1
        continue
    in120 = (name, MR.crosspost_key(text, chars=120)) in seen120
    infull = (name, MR.crosspost_key(text, chars=0)) in seenfull
    counts["routed_to_social_legacy"] += in120
    counts["routed_to_social_full"] += infull
    if in120 != infull:
        counts["changed"] += 1
        if len(moved) < 15:
            moved.append({"member": name, "legacy_routed": in120,
                          "full_routed": infull, "text": text[:120]})
connection.close()

for key in ("bluesky_post_rows", "unverified_handle_excluded", "unknown_member_excluded",
            "empty_text_excluded", "routed_to_social_legacy", "routed_to_social_full",
            "changed"):
    print("  %-30s %d" % (key, counts[key]))
print("\n  bluesky_scored_rows_routed_to_social: %d -> %d   (footprint rows shown: %+d)"
      % (counts["routed_to_social_legacy"], counts["routed_to_social_full"],
         counts["routed_to_social_legacy"] - counts["routed_to_social_full"]))
print("  every changed row is DISPLAY-ONLY (scored=False, 'no CU relevance detected "
      "-- footprint only'), so no score, tier, rank or allocation can move.")
for item in moved:
    print("    %-22s legacy_routed=%-5s full_routed=%-5s %s"
          % (item["member"][:22], item["legacy_routed"], item["full_routed"], item["text"]))

with io.open(os.path.join(HERE, "m7_bluesky_footprint_delta.json"), "w",
             encoding="utf-8") as fh:
    json.dump({"counts": dict(counts), "samples": moved}, fh, indent=1, ensure_ascii=False)
print("\nDONE M7")
