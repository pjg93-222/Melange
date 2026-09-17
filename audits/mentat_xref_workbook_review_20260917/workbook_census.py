"""Read-only census of every published Overview workbook: sheet list, and for the cycle sheets
plus a few key sheets, the header row number and header names (no data rows printed).

Usage: python workbook_census.py "<dir with *.xlsx>"
"""
import sys, re, os, glob
from openpyxl import load_workbook

folder = sys.argv[1]
cycle_re = re.compile(r"^\d{4}-\d{4}$")
KEY = ("Recipient Detail", "Contributor Detail", "Accounting Basis", "Recipient List", "Employee List", "Funding Cohorts", "Cycle Exceptions")
for path in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
    base = os.path.basename(path)
    if base.startswith("~$"):
        continue
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"\n### {base}: CANNOT OPEN {e!r}")
        continue
    names = wb.sheetnames
    print(f"\n### {base}  ({os.path.getsize(path)//1024} KB, mtime {__import__('datetime').datetime.fromtimestamp(os.path.getmtime(path)):%Y-%m-%d})")
    print("  sheets:", names)
    for name in names:
        if not (cycle_re.match(name) or name in KEY):
            continue
        ws = wb[name]
        header = None; hrow = 0
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
            vals = [str(v).strip() if v is not None else "" for v in row]
            nonempty = [v for v in vals if v]
            if len(nonempty) >= 4 and not any(len(v) > 60 for v in nonempty):
                header = nonempty; hrow = i; break
        dims = ws.calculate_dimension() if hasattr(ws, "calculate_dimension") else "?"
        print(f"  [{name}] header row {hrow} dims {dims}: {header}")
    wb.close()
