"""Read-only: sum the Banking Overview workbook's cycle sheets the way the console does.

Prints, per cycle sheet: header row found, column names, count of rows, gross sum of Amount,
net sum, sum where include_in_analysis == Yes (net and gross), count of Yes rows, and the
breakdown of the Yes rows by entity type / record type if such columns exist.
No writes.
"""
import sys, re
from decimal import Decimal
from collections import Counter, defaultdict
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
cycle_re = re.compile(r"^\d{4}-\d{4}$")
grand = defaultdict(Decimal)
for name in wb.sheetnames:
    if not cycle_re.match(name) and name != "Cycle Exceptions":
        continue
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    header = None
    header_row_no = 0
    for i, row in enumerate(rows, 1):
        vals = [str(v).strip() if v is not None else "" for v in row]
        low = [v.lower() for v in vals]
        if "amount" in low and any("include" in v for v in low):
            header = vals
            header_row_no = i
            break
        if i > 12:
            break
    if header is None:
        print(f"[{name}] header NOT FOUND in first 12 rows")
        continue
    idx = {h.lower(): k for k, h in enumerate(header) if h}
    print(f"\n[{name}] header row {header_row_no}: {header}")
    a = idx.get("amount")
    inc = next((k for h, k in idx.items() if "include" in h), None)
    ent = next((k for h, k in idx.items() if "entity" in h), None)
    rec = next((k for h, k in idx.items() if "record type" in h or h == "record_type"), None)
    memo = next((k for h, k in idx.items() if "memo" in h), None)
    status = next((k for h, k in idx.items() if "status" in h), None)
    n = 0; gross = Decimal(0); net = Decimal(0)
    yes_n = 0; yes_net = Decimal(0); yes_gross = Decimal(0)
    by_inc = Counter(); by_inc_amt = defaultdict(Decimal)
    by_ent = Counter(); by_ent_amt = defaultdict(Decimal)
    by_rec = Counter(); by_rec_amt = defaultdict(Decimal)
    by_memo = Counter(); by_memo_amt = defaultdict(Decimal)
    by_status = Counter(); by_status_amt = defaultdict(Decimal)
    for row in rows:
        if row is None:
            continue
        v = row[a] if a is not None and a < len(row) else None
        if v is None or v == "":
            continue
        try:
            d = Decimal(str(v))
        except Exception:
            continue
        n += 1; gross += abs(d); net += d
        flag = str(row[inc]).strip() if inc is not None and inc < len(row) and row[inc] is not None else ""
        by_inc[flag] += 1; by_inc_amt[flag] += d
        if flag.lower() == "yes":
            yes_n += 1; yes_net += d; yes_gross += abs(d)
            if ent is not None and ent < len(row):
                e = str(row[ent]); by_ent[e] += 1; by_ent_amt[e] += d
            if rec is not None and rec < len(row):
                r = str(row[rec]); by_rec[r] += 1; by_rec_amt[r] += d
            if memo is not None and memo < len(row):
                m = str(row[memo]); by_memo[m] += 1; by_memo_amt[m] += d
            if status is not None and status < len(row):
                s = str(row[status]); by_status[s] += 1; by_status_amt[s] += d
    print(f"  rows with amount: {n}  gross {gross:,.2f}  net {net:,.2f}")
    print(f"  include==Yes: {yes_n}  net {yes_net:,.2f}  gross {yes_gross:,.2f}")
    print("  by include flag:", {k: (c, f"{by_inc_amt[k]:,.2f}") for k, c in by_inc.most_common()})
    if by_ent:
        print("  Yes by entity:", {k: (c, f"{by_ent_amt[k]:,.2f}") for k, c in by_ent.most_common()})
    if by_rec:
        print("  Yes by record type:", {k: (c, f"{by_rec_amt[k]:,.2f}") for k, c in by_rec.most_common()})
    if by_memo:
        print("  Yes by memo:", {k: (c, f"{by_memo_amt[k]:,.2f}") for k, c in by_memo.most_common(6)})
    if by_status:
        print("  Yes by status:", {k: (c, f"{by_status_amt[k]:,.2f}") for k, c in by_status.most_common(8)})
    grand["yes_net"] += yes_net; grand["yes_gross"] += yes_gross; grand["all_net"] += net; grand["all_gross"] += gross
print("\nGRAND:", {k: f"{v:,.2f}" for k, v in grand.items()})
