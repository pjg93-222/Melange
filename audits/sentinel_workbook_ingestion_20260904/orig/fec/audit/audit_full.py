import os as _o,sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))));from paths import ROOT as MB_ROOT  # centralized root (arch review 2026-07-02)
"""FULL-BODY AUDIT — every module in one report.
Consolidates the integrity (22) + frontier (11) audits and adds correctness checks for the
new modules: money-response, electoral plays, co-sponsor routes, LT-priming, stance, determinism."""
import json,sys,os,subprocess,re,collections,importlib.util,csv,copy,hashlib
from pathlib import Path
import config_overlay as CO
import district.demographic_potential as _demographic_potential
from legislation.congress_bill_repository import validate_live_repository
from pipeline_lock import valid_inherited_lease
sys.stdout.reconfigure(encoding='utf-8',errors='replace')
B=MB_ROOT; FE=os.path.dirname(os.path.dirname(os.path.abspath(__file__))); ENG=FE+r'\engine'; DATA_FE=B+r'\fec'
def _pid_alive(pid):
    try:
        os.kill(int(pid),0); return True
    except (OSError,ValueError,TypeError):
        return False
def assert_no_foreign_pipeline_writer(lock_path=None,parent_pid=None):
    """A standalone audit must never read a mixed-generation artifact set."""
    lock_path=lock_path or os.path.join(B,'.pipeline_run.lock')
    parent_pid=os.getppid() if parent_pid is None else int(parent_pid)
    try:
        with open(lock_path,encoding='utf-8') as handle: owner=json.load(handle).get('pid')
    except (FileNotFoundError,OSError,ValueError,TypeError):
        return
    if (owner not in (os.getpid(),parent_pid) and _pid_alive(owner)
            and not valid_inherited_lease(lock_path)):
        raise RuntimeError('refusing audit while a foreign Moneyball pipeline owns the artifact set (pid %s)'%owner)
assert_no_foreign_pipeline_writer()
P,F=[],[]
def ck(name,cond,detail=''):
    (P if cond else F).append(name); print(('  OK   ' if cond else 'FAIL  ')+name+((' :: '+detail) if (detail and not cond) else ''))
def run(script):
    return subprocess.run([sys.executable,FE+'\\'+script],capture_output=True,text=True,encoding='utf-8',errors='replace').stdout
def _file_sha256(path):
    _h=hashlib.sha256()
    with open(path,'rb') as _handle:
        for _chunk in iter(lambda:_handle.read(1024*1024),b''): _h.update(_chunk)
    return _h.hexdigest()


print('========== FULL-BODY AUDIT ==========\n[1] Existing audits')
for s,nm,exp in [(r'audit\legacy_workbook_audit.py','integrity (22)',22),(r'audit\audit_frontier.py','frontier (11)',11)]:   # S3-03 audit run-2: assert COUNT too — a silently-skipped check must fail here, not vanish
    out=run(s); m=re.search(r'(\d+) passed, (\d+) FAILED',out)
    ck('%s sub-audit clean'%nm, bool(m) and m.group(2)=='0' and int(m.group(1))>=exp, out.strip().splitlines()[-1] if out else 'no output')
print('\n[2] Money-response (diminishing returns + capacity)')
spec=importlib.util.spec_from_file_location('mr',ENG+r'\money_response.py'); mr=importlib.util.module_from_spec(spec); spec.loader.exec_module(mr)
g='Jared Golden'
s1,s2,s4=mr.shift_lean(g,'cu',100000),mr.shift_lean(g,'cu',200000),mr.shift_lean(g,'cu',400000)
ck('money curve concave (doubling $ < doubling effect)', s2<2*s1 and s4<2*s2, '%.1f %.1f %.1f'%(s1,s2,s4))
big=mr.shift_lean('Jared Golden','cu',100000); small=mr.shift_lean('Steve Daines','cu',100000)
ck('capacity-aware (bigger-receipts member moves less)', abs(big)<abs(small))
# 2026-07-02: PT loss-frame asymmetry (Linde-Vis / K&T) is BY DESIGN: comp = -LOSSFRAME x cu, not the plain mirror.
ck('competitor money = -LOSSFRAME x cu (loss-frame asymmetry, PT)', abs(mr.shift_lean(g,'comp',100000)+mr.LOSSFRAME*s1)<1e-9)
_valid_ov=set(mr.ORDER); _deft={mr.defection(m,-30) for m in mr.C}-{None}
ck('defection yields ONLY valid overall tiers (no taxonomy leak)', _deft<=_valid_ov, 'invalid: %s'%(_deft-_valid_ov))
_probe=next(iter(mr.C)); _probe_old=mr.OV.get(_probe); mr.OV[_probe]='Opponent'
_opp_immune=mr.defection(_probe,-99) is None
if _probe_old is None: mr.OV.pop(_probe,None)
else: mr.OV[_probe]=_probe_old
ck('committed Opponents immune to money defection (structural probe)',_opp_immune)
import moneyball_core as _crown_core
import tier_util as _crown_tier_util
_crown_cs=json.load(open(B+r'\category_stance.json',encoding='utf-8'))
_crown_matrix=_crown_cs['matrix']; _crown_caps=(_crown_cs.get('gate') or {}).get('overall_capped') or {}
_crown_tier,_crown_trank=_crown_tier_util.legislative_press_tiers(list(_crown_matrix))
_crowned=[m for m in _crown_matrix if _crown_tier.get(m,0)==1]
_crown_floor_mismatches=[m for m in _crowned if mr.OV.get(m)!='Ally']
ck('Tier-1 display crowns are floored at Ally in money math (crowned remain movable)',
   bool(_crowned) and not _crown_floor_mismatches,
   '%d crowned members found; %d mismatched'%(len(_crowned),len(_crown_floor_mismatches)))
_champion_leaks=[m for m in mr.OV if mr.OV[m]=='Champion']
ck('Champion display label never leaks into money math',
   not _champion_leaks,'%d Champion values found'%len(_champion_leaks))
_crown_mismatches=[m for m,cm in _crown_matrix.items() if mr.OV.get(m)!=('Ally' if _crown_tier.get(m,0)==1 else _crown_core.overall_tier(cm,m,_crown_caps))]
ck('money engine uses specified crown-floor divergence from col G (not drift)',
   not _crown_mismatches,'%d mismatches'%len(_crown_mismatches))
# 2026-07-02: test the RUNG, not a named member — Durbin (the old exemplar) legitimately graduated to
# full Opponent once the tiered media weights priced his interchange opposition.
_hsk=[m for m in mr.C if mr.PROJ.get(mr.OV.get(m),mr.OV.get(m))=='Highly Skeptical']
ck('Highly Skeptical deep-but-movable (sub-threshold absorbed, big shock -> Opponent)',
   bool(_hsk) and all(mr.defection(m,-5) is None and mr.defection(m,-12)=='Opponent' for m in _hsk),
   '%d members at rung'%len(_hsk))
import moneyball_core as _band_core
_ob=dict((tier,cut) for cut,tier in _band_core.OVERALL_BANDS)
ck('overall bends are the strict fixed ladder, not roster-population quotas',
   _ob=={'Ally':3.0,'Highly Persuadable':2.0,'Persuadable':1.0,
         'Lean Persuadable':.4,'True Neutral':0.0,'Lean Skeptic':-.3,
         'Skeptic':-.75,'Highly Skeptical':-1.5} and
   _band_core.overall_tier_from_value(0)=='True Neutral' and
   _band_core.overall_tier_from_value(.01) not in ('Lean Skeptic','Skeptic','Highly Skeptical','Opponent') and
   _band_core.overall_tier_from_value(-.01) not in ('Highly Persuadable','Persuadable','Lean Persuadable','Ally'),
   str(_ob))
print('\n[3] Electoral plays (non-partisan, gated)')
EP=json.load(open(B+r'\data\money\electoral_plays.json'))
contest=[m for m,d in EP.items() if d.get('play')=='contest']; defend=[m for m,d in EP.items() if d.get('play')=='defend']
ck('contest plays gated to Opponents only', all(EP[m]['overall']=='Opponent' for m in contest))
ck('defend plays gated to allies only', all(EP[m]['overall'] in ('Champion','Defender','Ally') for m in defend))
pb=collections.Counter(EP[m]['party'] for m in contest)
_contest_n=sum(pb.values())
ck('contest plays not partisan-skewed (<70%% one party; enforced at n>=12)',
   _contest_n<12 or max(pb['D'],pb['R'])<=0.70*_contest_n,
   'D%d R%d (n=%d)'%(pb['D'],pb['R'],_contest_n))
print('\n[4] Stance matrix sanity')
CS=json.load(open(B+r'\category_stance.json')); MAT=CS['matrix']
from engine.pillar_policy import LABEL_EVIDENCE_WEIGHTS as _LABEL_EVIDENCE_WEIGHTS
ck('press-to-label evidence weight matches the diminished centralized policy',
   ((CS.get('thresholds') or {}).get('evidence_weights') or {})==_LABEL_EVIDENCE_WEIGHTS and
   (CS.get('thresholds') or {}).get('press_policy')=='active_diminished; may affect issue and overall labels',
   'artifact=%s policy=%s'%(((CS.get('thresholds') or {}).get('evidence_weights') or {}),
                            _LABEL_EVIDENCE_WEIGHTS))
_overall_bands=((CS.get('thresholds') or {}).get('overall') or {}).get('bands') or []
ck('overall band table matches kernel (all cutoffs + tiers, provenance)',
   len(_overall_bands)==len(mr.MC.OVERALL_BANDS) and
   all(abs((_overall_bands[_i].get('minimum') or 0)-_cut)<1e-12 and _overall_bands[_i].get('tier')==_tier
       for _i,(_cut,_tier) in enumerate(mr.MC.OVERALL_BANDS)))
ck('Durbin is Interchange Opponent', MAT['Richard Durbin']['Interchange']=='Opponent')
ck('Durbin anti on Credit Card (Opponent/Skeptic)', MAT['Richard Durbin']['Credit Card Mandates'] in ('Opponent','Skeptic'))
cnt=collections.Counter(MAT[m][c] for m in MAT for c in CS['cats'])
ck('champion rate disciplined (<2000 of ~9000 cells)', cnt['Champion']<2000, str(cnt['Champion']))
print('\n[5] Co-sponsor route (Pareto + MC reinforcement + determinism)')
def front(env):
    e=dict(os.environ,TOPIC='Mortgage',**env)
    out=subprocess.run([sys.executable,ENG+r'\cosponsor_route.py'],capture_output=True,text=True,encoding='utf-8',errors='replace',env=e).stdout
    d={}
    for ln in out.splitlines():
        m=re.search(r'([0-9,]+) +-> +\+([0-9.]+) +co-sponsors +risk +([0-9.]+)',ln)
        if m: d[int(m.group(1).replace(',',''))]=float(m.group(2))
    return d
tr=front({}); mc=front({'MCMODE':'1','MC':'30'}); mc2=front({'MCMODE':'1','MC':'30'})
shared=sorted(set(tr)&set(mc))
ck('route front non-empty', len(tr)>3, '%d pts'%len(tr))
ck('MC-exact >= tree-reach at shared costs (reinforcement)', all(mc[c]>=tr[c]-0.05 for c in shared) if shared else False)
ck('MC route deterministic (same seed -> identical)', mc==mc2)
print('\n[6] LT-priming validity')
pr=json.load(open(B+r'\data\network\lt_priming.json'))
ck('priming sorted cheapest-first', all(pr[i]['cost']<=pr[i+1]['cost'] for i in range(len(pr)-1)))
ck('every recipe closes the gap (pressure+money)', all(p['cost']>0 and p['gap']>0 for p in pr))
ck('no committed Opponents/Champions in primeable set', all(p['tier'] not in ('Champion','Defender','Opponent') for p in pr))
print('\n[7] Scenario runner (end-to-end)')
def scen():
    return subprocess.run([sys.executable,ENG+r'\scenario.py','--money','Jared Golden:comp:500000'],capture_output=True,text=True,encoding='utf-8',errors='replace').stdout
o1=scen(); o2=scen()
net=re.search(r'([+-][0-9.]+) co-sponsors\)',o1)   # the coalition/mobilization NET line
ck('scenario runs end-to-end (reports a net influence change)', bool(net), o1.strip().splitlines()[-1] if o1 else 'no output')
ck('scenario deterministic (same scenario -> identical)', o1==o2)
ck('competitor shock reduces CU influence (sign sane)', bool(net) and float(net.group(1))<=0, net.group(1) if net else '?')
# H2/H3 regression: competitor shock must reduce the FLIP channel too (caught the imputed-vs-raw ind footing bug)
flnet=re.search(r'winnability [0-9.]+ -> [0-9.]+  \(([+-][0-9.]+) expected flips\)',o1)
ck('competitor shock reduces winnability too (flip channel, consistent footing)', bool(flnet) and float(flnet.group(1))<=0, flnet.group(1) if flnet else '?')
# scenario CU spend now registers via flip channel (regression on the bug we just fixed)
cu_o=subprocess.run([sys.executable,ENG+r'\scenario.py','--money','Tim Scott:cu:300000'],capture_output=True,text=True,encoding='utf-8',errors='replace').stdout
fl=re.search(r'([+-][0-9.]+) expected flips',cu_o)
ck('CU spend registers gradual flip-value gain (not absorbed)', bool(fl) and float(fl.group(1))>0, fl.group(1) if fl else '?')
print('\n[8] Joint budget optimizer (fixed-state lazy greedy / modeled replanning)')
def opt(env): return subprocess.run([sys.executable,ENG+r'\optimize.py'],capture_output=True,text=True,encoding='utf-8',errors='replace',env=dict(os.environ,**env)).stdout
# OPT_OUT -> temp json so the test run doesn't clobber the production optimize_result.json (mirrors FRONTIER_OUT below).
# INC=5000 (divides PMAX=10000): optimize.py's cap check requires room for a FULL chunk, so INC>PMAX allocates nothing.
_OE={'BUDGET':'300000','INC':'5000','OPT_OUT':B+r'\data\lobbying\_audit_tmp_optimize.json'}
oo1=opt(_OE); oo2=opt(_OE)
gn=re.search(r'\(([+-][0-9.]+) expected flips\)',oo1)
ck('optimizer runs + reports a winnability frontier', bool(gn), oo1.strip().splitlines()[1] if oo1 else 'no output')
ck('optimizer improves winnability (monotone, sign sane)', bool(gn) and float(gn.group(1))>0, gn.group(1) if gn else '?')
ck('optimizer deterministic (same budget -> identical)', oo1==oo2)
# H1 regression: winning a member (conversion) must be a GAIN -> the spend->winnability path is monotone non-decreasing
_opt=json.load(open(B+r'\data\lobbying\_audit_tmp_optimize.json')); _pp=[v for _,v in _opt['path']]
_drops=sum(1 for i in range(1,len(_pp)) if _pp[i]<_pp[i-1]-1e-9)
ck('optimizer path monotone (conversion is a gain, not the old persuadability-loss)', _drops==0, '%d drops; base %.2f final %.2f'%(_drops,_pp[0],_pp[-1]))
ck('optimizer reports mode, observation status, and guarantee scope honestly',
   _opt.get('mode') in ('fixed_state_lazy_greedy','modeled_threshold_replanning') and
   isinstance(_opt.get('observations'),list) and isinstance(_opt.get('guarantee'),dict) and
   ((_opt.get('certificate_valid') and _opt['guarantee'].get('kind')=='exact_discrete_separable_concave')
    or (not _opt.get('certificate_valid') and _opt['guarantee'].get('kind')=='none')),
   str({k:_opt.get(k) for k in ('mode','observations','guarantee','certificate_valid')}))
ck('optimizer exact fixed-state certificate is internally consistent',
   (not _opt.get('certificate_valid')) or
   (_opt.get('fixed_state_concave') and not _opt.get('concavity_violations') and
    abs(_opt.get('opt_bound',0)-_opt.get('final',1))<=0.001 and
    _opt['guarantee'].get('certified_gain_ratio')==1.0),
   str(_opt.get('guarantee')))
# H3 regression: members with NO FEC data must NOT rank as top flip targets (ind=0 -> pflip~1 was misreading "missing" as "no opposition")
_s=importlib.util.spec_from_file_location('fmA',ENG+r'\flip_model.py'); _fm=importlib.util.module_from_spec(_s); _s.loader.exec_module(_fm)
_wn=collections.Counter()
for c in _fm.cats:
    cs,ct,nA,nB,mAt,mBt,tgt=_fm.topic_setup(c,_fm.CS0)
    for n in _fm.members: _wn[n]+=tgt[n]*_fm.pflip(_fm.cu0[n],_fm.ind0[n],_fm.R[n],nA[n],nB[n],mAt,mBt)
_nd=sum(1 for n,_ in _wn.most_common(20) if n in _fm.nodata)
ck('no zero-FEC-data member ranks top-20 winnability (missing != no-opposition)', _nd==0, '%d of top-20 are data-less'%_nd)
# HIGH-3 (corrected): impute ONLY no-FEC-record (receipts=0) blind members; trust FEC-verified-real ind=0 zeros
_blind_ok=all(_fm.ind0[n]>0 for n in _fm.ind_imputed)
_realzero=[n for n in _fm.members if (_fm.C[n].get('receipts') or 0)>0 and ((_fm.C[n]['banks'] or 0)+(_fm.C[n]['finserv'] or 0))==0]
_real_kept=all(_fm.ind0[n]==0 for n in _realzero)
ck('imputation limited to no-FEC-record members; verified-real ind=0 zeros trusted', _blind_ok and _real_kept, 'blind %d imputed, %d real-zeros trusted'%(len(_fm.ind_imputed),len(_realzero)))
# HIGH-1 regression: frontier cascade now uses edge-keyed CRN -> a single-topic run is deterministic across repeats
def bf(): return subprocess.run([sys.executable,ENG+r'\budget_frontier.py'],capture_output=True,text=True,encoding='utf-8',errors='replace',env=dict(os.environ,TOPIC='Interchange',FRONTIER_OUT=B+r'\data\money\_audit_tmp_frontier.json')).stdout   # temp out -> don't clobber the all-topics production json
bf1=bf(); bf2=bf()
ck('budget_frontier deterministic w/ edge-keyed CRN (single topic)', bf1==bf2 and 'Interchange' in bf1)

print('\n[NEW] Salience shock + press SNA layers (Noble-Turner / HCoI / Fink / Zhou-Elejalde / Joo)')
import json as _j
if os.path.exists(B+r'\salience_advisory.json'):
    _adv=_j.load(open(B+r'\salience_advisory.json'))['advisory']
    ck('salience advisory covers all issues', len(_adv)>=18, str(len(_adv)))
    _pr=[a['params'] for a in _adv]
    ck('salience regime probs valid (p+q+b~1, all in [0,1])',
       all(0<=x['p']<=1 and 0<=x['q']<=1 and 0<=x['b']<=1 and 0.94<=x['p']+x['q']+x['b']<=1.01 for x in _pr))
else: ck('salience advisory present (run salience_shock.py --advise)', False)
if os.path.exists(B+r'\dashboard_data.json'):
    _dd=_j.load(open(B+r'\dashboard_data.json',encoding='utf-8'))
    _sal=_dd.get('engine',{}).get('sal')
    ck('bundle exports Noble sal params (one per issue layer)', bool(_sal) and len(_sal)==len(_dd['issues']),
       '%s vs %s'%(len(_sal or []),len(_dd['issues'])))
if os.path.exists(B+r'\data\stance\press_hcoi.json'):
    _hc=_j.load(open(B+r'\data\stance\press_hcoi.json',encoding='utf-8'))
    _allm=[m2 for c in _hc['communities'] for m2 in c['members']]
    ck('HCoI communities non-trivial and disjoint (Louvain partition)',
       len(_hc['communities'])>=3 and len(_allm)==len(set(_allm)), '%d comms'%len(_hc['communities']))
if os.path.exists(B+r'\data\stance\press_engagement.json'):
    _pe=_j.load(open(B+r'\data\stance\press_engagement.json',encoding='utf-8'))
    ck('engagement layer: valid probabilities, no self-edges',
       all(0<e['p']<=1 and e['src']!=e['dst'] for e in _pe['directed']), '%d edges'%len(_pe['directed']))
if os.path.exists(B+r'\data\stance\stance_dynamics.json'):
    _sd=_j.load(open(B+r'\data\stance\stance_dynamics.json',encoding='utf-8'))
    ck('stance-dynamics imputation beats coin flip on masked cells (sign agreement > 0.55)',
       _sd['masking_validation']['sign_agreement']>0.55, str(_sd['masking_validation']))
    ck('imputation proposals carry provenance (user adjudicates; matrix untouched)',
       os.path.exists(B+r'\data\stance\stance_imputation_proposals.csv') and
       'provenance' in open(B+r'\data\stance\stance_imputation_proposals.csv',encoding='utf-8-sig').readline())
if os.path.exists(B+r'\data\stance\frame_alignment.json'):
    _fa=_j.load(open(B+r'\data\stance\frame_alignment.json',encoding='utf-8'))
    ck('frame-alignment rates in [0,1] across issues', all(0<=x['alignment']<=1 for x in _fa['alignment']))


print('\n[GOLD] Golden labels + domain classifier + collector DB')
if os.path.exists(B+'\\data\\stance\\cu_golden_labels.csv'):
    import csv as _csv
    _gl=list(_csv.DictReader(open(B+'\\data\\stance\\cu_golden_labels.csv',encoding='utf-8-sig')))
    _cls={'support','oppose','neutral','unrelated'}
    ck('golden label set parses, all labels in vocabulary', len(_gl)>=700 and all(r['label'] in _cls for r in _gl), '%d rows'%len(_gl))
    _sp=collections.Counter(r['split'] for r in _gl)
    ck("golden set preserves the user's train/test split and firewalls goldeval",
       _sp.get('train')==456 and _sp.get('test')==87 and _sp.get('goldeval',0)>0, str(dict(_sp)))
    _du=collections.Counter(r['url'] for r in _gl if r['url'])
    ck('golden set has no duplicate URLs', max(_du.values())==1 if _du else False)
    _cf=[r for r in _gl if r['conflict']]
    # two provenance formats are compliant: the arrow ('aug=X->queue=Y', unresolved flag) and the
    # user-adjudicated form 'user-resolved: <label> (was aug=X queue=Y)' (2026-07-03 overnight pass, 36 rows)
    ck('label conflicts are flagged, not silently resolved',
       all(('->' in r['conflict']) or r['conflict'].startswith('user-resolved:') for r in _cf), '%d flagged (user-resolved carries its trail inline)'%len(_cf))
else: ck('golden label set present (run build_golden_labels.py)', False)
if os.path.exists(B+'\\data\\stance\\stance_clf_composed.pkl'):
    try:
        import pickle as _pk
        _clf=_pk.load(open(B+'\\data\\stance\\stance_clf_composed.pkl','rb'))
        _pr=_clf.predict(['target: Credit Card Mandates. speaker: member. Senator introduces the Credit Card Competition Act to cap interchange.',
                          'target: Credit Unions (direct). speaker: member. Senator celebrates credit union community investment.'])
        ck('domain classifier loads and predicts sane labels', _pr[0]=='oppose' and _pr[1]=='support', str(list(_pr)))
    except Exception as e: ck('domain classifier loads', False, str(e)[:80])
if os.path.exists(B+'\\social_media.db'):
    import sqlite3 as _sq
    _c=_sq.connect(B+'\\social_media.db')
    _ok=_c.execute('PRAGMA integrity_check').fetchone()[0]
    _orph=_c.execute("SELECT COUNT(*) FROM posts WHERE platform IS NULL OR post_id IS NULL").fetchone()[0]
    ck('social_media.db integrity + no orphan posts', _ok=='ok' and _orph==0, '%s / %d orphans'%(_ok,_orph))
_tif=B+r'\data\stance\tweet_stance_final.csv'
if os.path.exists(_tif):
    with open(_tif,encoding='utf-8-sig',newline='') as _handle:
        _tir=csv.DictReader(_handle); _tifields=set(_tir.fieldnames or []); _tiout=list(_tir)
    ck('tweet stance output preserves image-reader provenance as a distinct signal',
       {'image_read','image_confidence','image_disposition'}.issubset(_tifields),
       'fields=%s'%sorted(_tifields))
    _barr_ids={'511658226','511657485','498264314','458365083'}
    _barr_out={row.get('quorum_id'):row.get('stack_stance') for row in _tiout
               if row.get('quorum_id') in _barr_ids}
    ck('Barr credit-card late-fee tweet gold is applied by exact Quorum ID',
       _barr_out=={qid:'support' for qid in _barr_ids},str(_barr_out))
    _followup_ids={'559107834','559097784','552715581','552700010','545211528','433093684',
                   '591392565','591305285','539036369','536549427','536545303','445704177'}
    _followup_out={row.get('quorum_id'):row.get('stack_stance') for row in _tiout
                   if row.get('quorum_id') in _followup_ids}
    ck('Britt/Lummis user-ratified social evidence is support by exact Quorum ID',
       _followup_out=={qid:'support' for qid in _followup_ids},str(_followup_out))
    _cbdc_out={row.get('quorum_id'):row.get('stack_stance') for row in _tiout
               if row.get('quorum_id') in {'620647802','615403327'}}
    _registry=CO.merged_rows('data/stance/cu_valence_registry.json','entries',('key',),data_root=B)
    _cbdc=next((row for row in _registry if row.get('key')=='cbdc'),{})
    ck('CBDC polarity is anti-CU and anti-CBDC Davidson posts compose to support',
       _cbdc.get('valence')==-1 and _cbdc_out=={'620647802':'support','615403327':'support'},
       'registry=%s tweets=%s'%(_cbdc,_cbdc_out))
_tirows=CO.merged_rows('data/stance/tweet_image_annotations.json','annotations',('quorum_id',),data_root=B)
_tibad=sum(1 for row in _tirows if not row.get('quorum_id') or
           row.get('review_status')!='approved' or
           row.get('cu_alignment') not in ('support','oppose','neutral','unrelated') or
           not 0<=float(row.get('confidence',-1))<=1)
ck('approved tweet-image annotations carry exact ID, stance, confidence, and review status',
   _tibad==0,'%d invalid annotations'%_tibad)
_sgr=B+r'\data\stance\speech_gold_retired.json'
if os.path.exists(_sgr):
    _sgdoc=json.load(open(_sgr)); _sgrows=_sgdoc.get('rows') or []
    ck('retired speech gold is limited to explicit neutral precision negatives',
       _sgdoc.get('schema')=='speech-gold-retired-v1' and
       all(row.get('gold_stance')=='neutral' and row.get('reason') for row in _sgrows),
       '%d retired rows'%len(_sgrows))

print('\n========================================')

# RECEPTIVITY LAYER (2026-07-02): coverage + documented caps.
if os.path.exists(B+r'\member_receptivity.json'):
    _rj=json.load(open(B+r'\member_receptivity.json'))['members']
    ck('receptivity layer covers all members', len(_rj)==533, str(len(_rj)))
    ck('receptivity multipliers within documented caps',
       all(0.30<=v['seed_prob']<=0.90 and 0.50<=v['pressure_mult']<=1.60 and 0.80<=v['edge_mult']<=1.20 for v in _rj.values()))

# W1 guard 2026-07-02: the press-pillar SUMPRODUCT weight row once held stray LET-formula remnants that
# evaluated to 0 on the three TOP-tier categories — every weight must be a positive literal number.
import openpyxl as _px
_wb=_px.load_workbook(B+r'\Moneyball Main Version - Blended.xlsx',data_only=True)
_w535=[_wb['Press Release - Normalized'].cell(535,c).value for c in range(5,24)]
ck('press-pillar weight row 535: 19 numeric positive weights (no formula remnants)',
   all(isinstance(v,(int,float)) and v>0 for v in _w535),str(_w535))
# 2026-07-04 integrity fix guard: the press component must be DYNAMIC (built from the current tone
# csv), never the frozen 07-02 archive (which silently excluded every post-archive article: 197
# members frozen forever). Checks the blend's meta artifact is fresh vs the csv and coverage grew.
_pm_p=B+r'\data\stance\press_component_meta.json'
if os.path.exists(_pm_p):
    _pm=json.load(open(_pm_p))
    _pm_hashes=_pm.get('inputs_sha256') or {}
    _pm_hashes_current=bool(_pm_hashes) and all(
        os.path.isfile(os.path.join(B,*_rel.split('/'))) and
        _file_sha256(os.path.join(B,*_rel.split('/')))==_sha
        for _rel,_sha in _pm_hashes.items())
    _tone_p=os.path.join(os.path.dirname(B),'Draft Systems','cu_press_stance_tone.csv')
    _fresh=(os.path.getmtime(_tone_p)-_pm.get('built_from_csv_mtime',0))<8*86400 if os.path.exists(_tone_p) else True   # csv may be up to one weekly cycle newer; beyond that the blend stopped ingesting
    ck('media-blend press component is DYNAMIC and fresh (not the frozen archive)',
       _fresh and _pm_hashes_current and _pm.get('members_nonzero',0)>=197,
       'members=%s fresh=%s hashes=%s'%(_pm.get('members_nonzero'),_fresh,_pm_hashes_current))
    # 2026-07-04: intensity must be NON-DEGENERATE. The first dynamic build clipped |press_value|
    # to a 0.25 floor that 100% of articles saturated -> the component silently became a pure
    # article-count measure (flipped Warren +310 ranks). Guard: recorded pv stats must show real
    # spread (many distinct values, median strictly inside (0, max)).
    _pvs=_pm.get('pv_stats') or {}
    ck('press component intensity non-degenerate (no clip saturation / constant weights)',
       _pvs.get('n_distinct',0)>=20 and 0<_pvs.get('median',0)<_pvs.get('max',0),
       'pv_stats=%s'%_pvs)
    # 2026-07-04 stack wiring guard: when the press_stance_final.csv sidecar exists, the stack must
    # actually drive the signs (>=90% stack vs DEBATE-fallback) and the unrelated->0 mention-noise
    # kill must be active. A silent fall-back to DEBATE (0.571 adversarial-frame) is a regression.
    _sk=_pm.get('stack') or {}
    if _sk.get('sidecar'):
        _tot=_sk.get('stack_signed',0)+_sk.get('debate_fallback',0)
        ck('press stack drives the dynamic component (not silently on DEBATE fallback)',
           _tot>0 and _sk.get('stack_signed',0)/_tot>=0.9 and _sk.get('unrelated_zeroed',0)>0,
           'stack=%s'%_sk)
        ck('user-ratified press gold overrides reach the scoring component',
           _sk.get('gold_overrides',0)>0,'stack=%s'%_sk)
    # 2026-07-16 own-source channel gate: speech, hearing, and interview sidecars were
    # integrity-checked before publication.  Keep those gates observable and fail loud if a
    # later build silently disables a channel, admits learned hearing rows, or double-counts
    # Bluesky (which is already represented in the social component).
    _osc=_pm.get('own_source_channels') or {}
    _validated=('speech','hearing','interview')
    ck('validated own-source press channels are enabled and loaded',
       all((_osc.get(ch) or {}).get('enabled') is True and
           (_osc.get(ch) or {}).get('loaded') is True and
           (_osc.get(ch) or {}).get('signed_gold_rows',
                                    (_osc.get(ch) or {}).get('signed_rows',0))>0
           for ch in _validated),
       'own_source_channels=%s'%_osc)
    _hm=_osc.get('hearing') or {}
    ck('hearing channel remains gold-only',
       _hm.get('tier_policy')=='gold only' and _hm.get('signed_gold_rows',0)>0 and
       _hm.get('non_gold_rows_excluded',0)>0,
       'hearing=%s'%_hm)
    _bm=_osc.get('bluesky') or {}
    ck('Bluesky remains single-counted',
       _bm.get('enabled') is False and _bm.get('loaded') is False and
       'no press-side duplicate' in _bm.get('tier_policy',''),
       'bluesky=%s'%_bm)

    _sidecar_specs=(
        ('speech','speeches_cu.csv','speech_stance_final.csv','data/stance/speech_gold_queue.csv'),
        ('hearing','data/stance/hearings_cu.csv','data/stance/hearing_stance_final.csv','data/stance/hearing_gold_queue.csv'),
        ('interview','data/stance/interviews_cu.csv','data/stance/interview_stance_final.csv','data/stance/interview_gold_queue.csv'),
    )
    _valid_stances={'support','oppose','neutral'}
    _valid_tiers={'gold','unanimous','majority','review'}
    _sidecars_ok=True; _sidecar_detail=[]
    for _ch,_src_name,_final_name,_gold_name in _sidecar_specs:
        _src_path=os.path.join(B,_src_name); _final_path=os.path.join(B,_final_name)
        _gold_path=os.path.join(B,_gold_name)
        _files_present=all(os.path.exists(p) for p in (_src_path,_final_path,_gold_path))
        _src_rows=[]; _final_rows=[]
        if _files_present:
            with open(_src_path,encoding='utf-8-sig',newline='') as _fh:
                _src_rows=list(csv.DictReader(_fh))
            with open(_final_path,encoding='utf-8-sig',newline='') as _fh:
                _final_rows=list(csv.DictReader(_fh))
        _row_parity=_files_present and len(_src_rows)>0 and len(_src_rows)==len(_final_rows)
        _identity_ok=_row_parity and all(
            all((src.get(field) or '')==(final.get(field) or '')
                for field in ('bioguide','member','date','categories'))
            for src,final in zip(_src_rows,_final_rows)
        )
        _values_ok=_row_parity and all(
            row.get('stance') in _valid_stances and row.get('tier') in _valid_tiers
            for row in _final_rows
        )
        _fresh=_files_present and os.path.getmtime(_final_path)>=max(
            os.path.getmtime(_src_path),os.path.getmtime(_gold_path))
        _channel_ok=_files_present and _row_parity and _identity_ok and _values_ok and _fresh
        _sidecars_ok=_sidecars_ok and _channel_ok
        _sidecar_detail.append('%s:%d/%d identity=%s values=%s fresh=%s'%(
            _ch,len(_src_rows),len(_final_rows),_identity_ok,_values_ok,_fresh))
    ck('own-source channel sidecars preserve source row identity',
       _sidecars_ok,'; '.join(_sidecar_detail))
    _ixr_p=B+r'\data\stance\interviews_cu.receipt.json'
    _isr_p=B+r'\data\stance\interview_stance_final.receipt.json'
    _ixr=json.load(open(_ixr_p,encoding='utf-8')) if os.path.exists(_ixr_p) else {}
    _isr=json.load(open(_isr_p,encoding='utf-8')) if os.path.exists(_isr_p) else {}
    _ixdb=_ixr.get('database') or {}; _ixout=_ixr.get('extract') or {}
    _dbp=B+r'\social_media.db'; _ixcsv=B+r'\data\stance\interviews_cu.csv'
    _iscsv=B+r'\data\stance\interview_stance_final.csv'
    _ix_projection=hashlib.sha256(); _ix_rows=0; _ix_max_rowid=0; _ix_max_date=''
    if os.path.exists(_dbp):
        import sqlite3 as _ix_sqlite
        _ix_con=_ix_sqlite.connect('file:'+_dbp.replace('\\','/')+'?mode=ro',uri=True)
        for _ix_row in _ix_con.execute(
                'select rowid,member,date,chamber,package,nwords,text from speeches '
                'where source=? order by rowid',('INTERVIEW',)):
            _ix_rows+=1; _ix_max_rowid=max(_ix_max_rowid,int(_ix_row[0] or 0))
            _ix_max_date=max(_ix_max_date,str(_ix_row[2] or ''))
            _ix_projection.update(json.dumps(_ix_row,ensure_ascii=False,separators=(',',':'),default=str)
                                  .encode('utf-8','replace'))
            _ix_projection.update(b'\n')
        _ix_con.close()
    _ix_current=(os.path.exists(_dbp) and os.path.exists(_ixcsv) and
        _ixdb.get('size')==os.path.getsize(_dbp) and _ixdb.get('mtime_ns')==os.stat(_dbp).st_mtime_ns and
        _ixdb.get('interview_rows')==_ix_rows and _ixdb.get('max_rowid')==_ix_max_rowid and
        _ixdb.get('max_date')==_ix_max_date and
        _ixdb.get('projection_sha256')==_ix_projection.hexdigest() and
        _ixout.get('output_sha256')==_file_sha256(_ixcsv))
    _is_inputs=_isr.get('inputs_sha256') or {}
    _is_inputs_current=(set(_is_inputs)=={
        'data/stance/interviews_cu.csv','data/stance/interview_gold_queue.csv'} and all(
        os.path.isfile(os.path.join(B,*_rel.split('/'))) and
        _file_sha256(os.path.join(B,*_rel.split('/')))==_sha
        for _rel,_sha in _is_inputs.items()))
    from stance.own_source_stance_model import input_fingerprint as _own_source_input_fingerprint
    _is_current=(os.path.exists(_iscsv) and _isr.get('channel')=='interview' and
        _isr.get('source_receipt_sha256')==(_file_sha256(_ixr_p) if os.path.exists(_ixr_p) else None) and
        _is_inputs.get('data/stance/interviews_cu.csv')==(_file_sha256(_ixcsv) if os.path.exists(_ixcsv) else None) and
        _is_inputs_current and _isr.get('model_fingerprint')==_own_source_input_fingerprint() and
        _isr.get('output_sha256')==_file_sha256(_iscsv))
    ck('interview DB projection -> extract -> stance chain is receipt/hash bound',
       _ixr.get('schema')=='interview-cu-extract-receipt-v1' and _ix_current and
       _isr.get('schema')=='own-source-channel-score-receipt-v1' and _is_current,
       'extract_current=%s stance_current=%s'%(_ix_current,_is_current))
    # 2026-07-19: cu_passages merges overlapping ~500-char windows, so a committee chair's final
    # speaking turn can absorb the transcript APPENDIX (prepared statements, written questions for
    # the record) into one blob that is then stance-graded as if the member had said it. Measured
    # today: 27 of 775 hearing passages (3.5%), but only 5 of the 263 GOLD rows -- and gold is the
    # only tier build_media_blend.py feeds into the press pillar. 4 of those 5 change no outcome
    # (the member's clean gold rows already agree); the 5th is a tie-break on a 1/1/1 split. So the
    # contamination is real but immaterial TODAY, and the gold labels are the user's own calls --
    # nothing is relabelled here. This is a tripwire so it cannot grow silently as the corpus does.
    _apx_re=re.compile(r'\[whereupon|prepared statements, additional material|'
                       r'supplied for the record follow|\[additional material|'
                       r'responses to written questions|\[the following|insert(ed)? for the record',re.I)
    _qfr_re=re.compile(r'\bQ\.\d+\.|\bA\.\d+\.')
    def _apx_contaminated(text):
        return bool(_apx_re.search(text)) or len(_qfr_re.findall(text))>=2 or len(text)>2000
    _hs_src=os.path.join(B,'data','stance','hearings_cu.csv'); _hs_fin=os.path.join(B,'data','stance','hearing_stance_final.csv')
    if os.path.exists(_hs_src) and os.path.exists(_hs_fin):
        with open(_hs_src,encoding='utf-8-sig',newline='') as _fh: _hsrc=list(csv.DictReader(_fh))
        with open(_hs_fin,encoding='utf-8-sig',newline='') as _fh: _hfin=list(csv.DictReader(_fh))
        if len(_hsrc)==len(_hfin):
            _apx_gold=sum(1 for s,f in zip(_hsrc,_hfin)
                          if f.get('tier')=='gold' and _apx_contaminated(s.get('text') or ''))
            _apx_all=sum(1 for s in _hsrc if _apx_contaminated(s.get('text') or ''))
            ck('transcript-appendix passages stay a negligible share of pillar-feeding gold hearings',
               _apx_gold<=10,
               '%d gold rows contaminated (bound 10); %d/%d overall'%(_apx_gold,_apx_all,len(_hsrc)))
        else:
            ck('transcript-appendix passages stay a negligible share of pillar-feeding gold hearings',
               False,'hearings_cu/hearing_stance_final row counts differ')
_pra_p=B+r'\data\stance\press_domain_routing_audit.json'
if os.path.exists(_pra_p):
    _pra=json.load(open(_pra_p)); _prrows=_pra.get('rows') or []
    from press_stream import load_tone_rows as _load_press_tone_rows
    _expected_press_rows=len(_load_press_tone_rows())
    _prbad=sum(1 for row in _prrows if not row.get('routing_basis') or
               not isinstance(row.get('effective_categories'),list) or
               ('General Credit Union' in row.get('effective_categories',[]) and
                len(row.get('effective_categories',[]))>1))
    ck('press domain ledger covers every input row with one auditable disposition',
       _pra.get('schema')=='press-domain-routing-audit-v1' and
       len(_prrows)==_expected_press_rows and _prbad==0,
       'rows=%d expected=%d bad=%d'%(len(_prrows),_expected_press_rows,_prbad))
    ck('press routing rejects unsupported declared tags instead of trusting them blindly',
       sum(bool(row.get('discarded_declared')) for row in _prrows)>0,
       'no discarded declared tags recorded')
    _barr_press=[row for row in _prrows if row.get('member_source')=='Andy Barr' and
                 row.get('title')=='House, Senate resolutions would void CFPB credit card late fees rule']
    ck('Barr late-fee CRA press evidence is scored as user-ratified support',
       bool(_barr_press) and all(row.get('gold_override')=='support' for row in _barr_press),
       str(_barr_press[:2]))
else:
    ck('press domain routing audit present',False,'press_domain_routing_audit.json missing')
_lra_p=B+r'\legislation_routing_audit.json'
if os.path.exists(_lra_p):
    _lra=json.load(open(_lra_p)); _lrrows=_lra.get('rows') or []
    _target={('H.R.7733','117'):'CDFI Fund',('S.869','118'):'CDFI Fund',
             ('H.R.7036','118'):'Data Privacy & Cyber',('H.R.7022','117'):'Data Privacy & Cyber',
             ('S.4698','117'):'Data Privacy & Cyber'}
    _lrbad=0
    for (_bill,_congress),_cat in _target.items():
        _seen={row.get('effective_category') for row in _lrrows
               if row.get('bill')==_bill and str(row.get('congress')) in (_congress,_congress+'.0')
               and row.get('effective_category')}
        _lrbad += int(bool(_seen) and _seen!={_cat})
    _ldup=sum((_lra.get('summary') or {}).get('duplicates',{}).values())
    ck('legislation ledger covers every source row and scores one strongest-role event',
       _lra.get('schema')=='legislation-routing-audit-v1' and
       len(_lrrows)==_lra.get('source_rows') and _ldup>0,
       'rows=%d source=%s duplicates=%d'%(len(_lrrows),_lra.get('source_rows'),_ldup))
    ck('pinned CDFI/privacy bills have exactly one effective policy column',
       _lrbad==0 and (_lra.get('summary') or {}).get('effective_categories',{}).get('CDFI Fund',0)>0,
       '%d targeted bill routing errors'%_lrbad)
    # 2026-07-19 coverage gaps found by fec/legislation/bill_coverage_audit.py. Row conservation above proves
    # nothing VANISHED, but nothing asserted that every row reached a SLOT, that all 18 slots
    # survived (the exact 2026-07-02 dropped-18th-category failure mode), that the unmatched-member
    # set stayed put across a roster swap, or that two bill ids never collapse to one bill_key
    # (bill_key discards the R/J/Res segment -- bill_rules.py:129-132).
    import bill_rules as _BR
    _lrsum=_lra.get('summary') or {}; _lrdisp=_lrsum.get('dispositions') or {}
    ck('no bill orphaned without a policy slot',
       _lrdisp.get('unknown-section',0)==0,
       'unknown-section=%d'%_lrdisp.get('unknown-section',0))
    _lrcats=set(_lrsum.get('effective_categories') or {})
    ck('all 18 canonical categories survive routing',
       _lrcats==set(_BR.CATEGORIES),
       'missing=%s stray=%s'%(sorted(set(_BR.CATEGORIES)-_lrcats),sorted(_lrcats-set(_BR.CATEGORIES))))
    # WHO MAY LEGITIMATELY BE UNMATCHED (corrected 2026-07-29). The archive is
    # historical, the roster is the sitting 119th, so a sponsor can be absent
    # from the roster for TWO honest reasons -- this allowlist previously
    # encoded only the first, so the check failed on the second and reported a
    # data error that was not one. Each name is annotated; verified absent from
    # the 533 roster under every spelling before being listed here.
    _lrdelegates={'Stacey Plaskett','James Moylan'}      # non-voting delegates
    _lrdeparted={
        'Ed Perlmutter',        # retired, end of the 117th
        'Michael C. Burgess',   # retired, end of the 118th
        'Alexander X. Mooney',  # left the House after the 118th
        'Marco Rubio',          # resigned Jan 2025 (Secretary of State)
    }
    _lrknown=_lrdelegates|_lrdeparted
    _lrunmap=_lrsum.get('unmatched_members') or {}
    _lrun=set(_lrunmap)
    ck('unmatched members stay the known delegates and departed members',
       _lrun<=_lrknown,'unexpected=%s'%sorted(_lrun-_lrknown))
    # The invariant the allowlist alone cannot express, and the actual hazard
    # this family of checks exists for: a SITTING member must never fall out of
    # name resolution. If archive spelling drifts from roster spelling (middle
    # initials, suffixes, the 120th-Congress roster swap), the member silently
    # stops matching and their legislative record vanishes from the model.
    ck('no sitting roster member is unmatched by name resolution',
       not (_lrun & set(MAT)),'roster members unmatched=%s'%sorted(_lrun & set(MAT)))
    # Bound the blast radius so a bad archive refresh cannot quietly multiply
    # unmatched rows while every name still sits on the allowlist.
    _lrunrows=sum(int(v or 0) for v in _lrunmap.values())
    ck('unmatched sponsor rows stay a negligible share of the archive',
       _lrunrows<=50 and _lrunrows<=0.01*max(1,len(_lrrows)),
       '%d unmatched rows of %d'%(_lrunrows,len(_lrrows)))
    _lrkey={}
    for _r in _lrrows:
        if _r.get('disposition')!='selected': continue
        _lrkey.setdefault((_BR.bill_key(_r.get('bill')),str(_r.get('congress'))),set()).add(str(_r.get('bill')))
    _lrcol={k:sorted(v) for k,v in _lrkey.items() if k[0] and len(v)>1}
    ck('no two scored bill ids collapse to one bill_key',
       not _lrcol,'%d colliding keys, e.g. %s'%(len(_lrcol),sorted(_lrcol.values())[:2]))
    # The 17 uncovered members were individually adjudicated 2026-07-19: 16 absent from the archive
    # entirely, and Clyburn's rows all user-excluded. The floor exists so a name-resolution regression
    # or bad archive refresh cannot silently zero out coverage; a legitimate roster change at the
    # 120th-Congress rollover re-bases this number (see moneyball-congress-rollover-campaign).
    _lrarch_p=B+r'\archive_leg.json'
    if os.path.exists(_lrarch_p):
        with open(_lrarch_p,encoding='utf-8') as _fh: _lrarch=json.load(_fh)
        ck('legislative evidence coverage holds its floor',
           len(_lrarch)>=516,
           '%d members carry >=1 archive bill (floor 516 of 533)'%len(_lrarch))
    else:
        ck('legislative evidence coverage holds its floor',False,'archive_leg.json missing')
else:
    ck('legislation routing audit present',False,'legislation_routing_audit.json missing')
# 2026-07-04 nomological fix guard: receptivity must run on the BEHAVIORAL basis (the lexical Big5
# was INDICTED by the pre-registered label-free validation — big5_nomological_prereg.md). A silent
# fall-back to lexical (missing dispositions artifact) is a regression to an indicted instrument.
_rc_p=B+r'\member_receptivity.json'
if os.path.exists(_rc_p):
    _rc=json.load(open(_rc_p))
    _bd_p=B+r'\behavioral_dispositions.json'
    _bd_ok=os.path.exists(_bd_p) and len((json.load(open(_bd_p)) or {}).get('members',{}))>=450
    ck('receptivity runs on the behavioral basis (lexical Big5 indicted 2026-07-04)',
       _rc.get('basis')=='behavioral' and _bd_ok,
       'basis=%s dispositions_ok=%s'%(_rc.get('basis'),_bd_ok))
    # ROUND-7/7b N-slot guard (2026-07-11 promotion, user-ratified, pre-registered gate 5/5): the
    # susceptibility N-slot must remain affective-content N. The 2026-07-12 criterion audit clarifies
    # that E/party residualization is a confound control rather than independent validity evidence; it
    # does not silently overturn the locked promotion. Coverage floor 150.
    _nsl=_rc.get('n_slot'); _nsrc=sum(1 for v in _rc.get('members',{}).values() if v.get('n_source')=='affective')
    ck('receptivity N-slot = affective-content N (round-7/7b; no silent behavioral/text revert)',
       _nsl=='affective' and _nsrc>=150,'n_slot=%s affective_members=%d'%(_nsl,_nsrc))
    # COOP-SLOT guard (audit 2026-07-10, mirrors the N-slot guard): the cooperation-role slot must be the promoted
    # share-based cooperation_v2 (@94f09cf); a silent RECEPT_COOP=v1 revert or a coverage collapse
    # (coop_v2_candidate staleness / c2-floor starvation) would sail through the basis gate above.
    _csl=_rc.get('coop_slot'); _csrc=sum(1 for v in _rc.get('members',{}).values() if v.get('coop_source')=='coop_v2')
    _cv=[v.get('cooperation_v2_z') for v in _rc.get('members',{}).values() if v.get('cooperation_v2_z') is not None]
    ck('receptivity cooperation-role slot = cooperation_v2 (promotion @94f09cf; coverage>=250, |z|<=5, no silent v1 revert)',
       _csl=='v2' and _csrc>=250 and len(_cv)>=250 and all(abs(x)<=5.0 for x in _cv),
       'coop_slot=%s coop_v2_members=%d'%(_csl,_csrc))
    # 2026-07-13 construct-contamination audit (user observed cooperation/A friction): cooperation_v2
    # is a cross-party cosponsorship SHARE; qualified Big-Five A is a package-disjoint low-antagonism
    # interview facet. They must remain separately sourced and must not collapse into one score. The
    # |r|<.30 diagnostic is a contamination alarm (<9% shared variance), not a validity claim.
    _cvp=B+r'\coop_v2_candidate.json'; _avp=B+r'\data\stance\interview_antagonism.json'
    if os.path.exists(_cvp) and os.path.exists(_avp):
        import numpy as _np, math as _math, csv as _csv
        _cvm=json.load(open(_cvp,encoding='utf-8')).get('members',{}); _avm=json.load(open(_avp,encoding='utf-8')).get('members',{})
        _co_bad=sum(1 for m,d in _cvm.items() if d.get('cooperation_v2_z') is not None and
                    ((_rc.get('members',{}).get(m) or {}).get('coop_source')!='coop_v2' or
                     abs(((_rc.get('members',{}).get(m) or {}).get('cooperation_v2_z') or 0)-d['cooperation_v2_z'])>.002))
        ck('cooperation and low-antagonism A remain separately sourced (zero substitutions)',
           _co_bad==0, '%d cooperation_v2 values substituted/mismatched'%_co_bad)
        _con={d['member']:d for d in json.load(open(B+r'\data\engine\congestion.json',encoding='utf-8'))}
        _g={d.get('full_name'):d.get('gender') for d in _csv.DictReader(open(DATA_FE+r'\legislators.csv',encoding='utf-8-sig'))}
        _rows=[]
        for _m in set(_cvm)&set(_avm)&set(_con):
            _x=(_cvm.get(_m) or {}).get('cooperation_v2_z')
            if _x is None: continue
            _c=_con[_m]; _p=str(_c.get('party') or '').lower()
            _rows.append([_x,_avm[_m]['agreeableness_z'],1.0 if _p.startswith('r') else 0.0,
                          1.0 if _g.get(_m)=='F' else 0.0,1.0 if _c.get('chamber')=='Senate' else 0.0,
                          _math.log1p(_avm[_m].get('words') or 0)])
        _z=_np.asarray(_rows,float); _raw=float(_np.corrcoef(_z[:,0],_z[:,1])[0,1]) if len(_z)>5 else 9
        _X=_np.c_[_np.ones(len(_z)),_z[:,2:]]
        _rx=_z[:,0]-_X@_np.linalg.lstsq(_X,_z[:,0],rcond=None)[0]; _ry=_z[:,1]-_X@_np.linalg.lstsq(_X,_z[:,1],rcond=None)[0]
        _part=float(_np.corrcoef(_rx,_ry)[0,1]) if len(_z)>5 else 9
        ck('cooperation vs low-antagonism A cross-contamination below alarm threshold',
           len(_z)>=180 and abs(_raw)<.30 and abs(_part)<.30,
           'n=%d raw_r=%+.3f partial_r(party,gender,chamber,interview-volume)=%+.3f'%(len(_z),_raw,_part))
    if _bd_ok:
        _rv=[d['reactivity_z'] for d in json.load(open(_bd_p))['members'].values() if d.get('reactivity_z') is not None]
        ck('reactivity_z disposition sane (coverage >=150, |z|<=5)',
           len(_rv)>=150 and all(abs(v)<=5.0 for v in _rv),'%d covered'%len(_rv))
else:
    # (2026-07-09 fix: this branch previously fired a MISLABELED press-component check name)
    ck('receptivity artifact present (member_receptivity.json)',False,'missing — chain step 1 (receptivity_layer.py) did not run')
# S2-04 guard (audit 2026-07-02): every workbook reader hardcodes range(3,536) for the 533-member roster.
# If a member row is ever added/removed, those readers silently truncate or NUL the tail. Assert on
# R&A col A — TYPED names, immune to the phase-2 openpyxl-save cache death (Blended col A is
# formula-backed and legitimately reads empty between Excel opens).
_roster=len(json.load(open(B+r'\data\engine\congestion.json')))
_ra_ws=_wb['Rank & Allocation']; _wbrows=sum(1 for r in range(3,536) if _ra_ws.cell(r,1).value)
ck('workbook member rows == roster size (no silent row drift)', _wbrows==_roster, 'workbook %d vs roster %d'%(_wbrows,_roster))
# Roadmap-#9 phase 1 (2026-07-02): the whole raw->normalized->pillar->score->rank chain is reproduced
# in Python (compute_pillars.py) and must match the sheet EXACTLY — catches any silent formula edit.
os.environ['PILLAR_SNAPSHOT_FREEZE']='1'   # the audit must NEVER advance the pillar snapshot vintage (2026-07-17 guard)
try:
    _pp=run('engine/compute_pillars.py')
finally:
    os.environ.pop('PILLAR_SNAPSHOT_FREEZE',None)
ck('python pillar shadow-compute parity (rank math reproducible, 533/533)','PARITY: CLEAN' in _pp,
   _pp.strip().splitlines()[-1] if _pp else 'no output')
# rank-integrity guards (2026-07-17): (a) the pillar snapshot must not silently serve a stale
# leg/inf/demo basis (the 2026-07-16 522-rank incident class); (b) the own-source channel inputs
# must not silently freeze while their channels are ON (extractors are not chain-wired yet —
# collection runs on Paul's 3-week cadence, so the ceiling is cadence+grace = 28 days).
_snp=B+r'\data\engine\pillar_inputs_snapshot.json'
_sn_stamp=(json.load(open(_snp,encoding='utf-8')).get('stamped') if os.path.exists(_snp) else None)
if _sn_stamp:
    import datetime as _dt2
    _sn_age=(_dt2.datetime.now().astimezone()-_dt2.datetime.fromisoformat(_sn_stamp)).days
    ck('pillar snapshot vintage <= 14 days (stale-basis tripwire)', _sn_age<=14, 'stamped %s (%dd)'%(_sn_stamp,_sn_age))
else:
    ck('pillar snapshot vintage <= 14 days (stale-basis tripwire)', True, 'unstamped pre-guard snapshot (stamps on next compute run)')
_dbm=os.path.getmtime(B+r'\social_media.db') if os.path.exists(B+r'\social_media.db') else None
_chstale=[]
for _cf in ('speeches_cu.csv','data/stance/hearings_cu.csv','data/stance/interviews_cu.csv','data/stance/bluesky_cu.csv'):
    _p2=B+'\\'+_cf
    if _dbm and os.path.exists(_p2) and (_dbm-os.path.getmtime(_p2))>28*86400:
        _chstale.append('%s (%.0fd behind corpus)'%(_cf,(_dbm-os.path.getmtime(_p2))/86400))
ck('own-source channel inputs within 28 days of the corpus (channels are default-ON)', not _chstale, str(_chstale))

if os.path.exists(B+r'\data\stance\stance_training.jsonl'):
    import hashlib as _hl, json as _js
    _tr=set(); _te=set()
    for _l in open(B+r'\data\stance\stance_training.jsonl',encoding='utf-8'):
        _r=_js.loads(_l)
        _vs=[_r.get('text',''),_r.get('title',''),_r.get('stance_span','')]
        _hs={_hl.md5(v.strip().lower()[:400].encode()).hexdigest() for v in _vs if v and v.strip()}
        if _r['split']=='train': _tr|=_hs
        elif _r['split']=='test': _te|=_hs
    _ov=len(_tr&_te)
    ck('train/test split-leak (text-hash, incl. augment views) within known bound', _ov<=4,
       '%d collisions (4 known: same-ACU-article rows 75/214; fix = same-text same-split on next label refresh)'%_ov)

print('\n[IND] District industry')
_ind_path=B+r'\data\lobbying\district_industry.json'
if not os.path.exists(_ind_path):
    print('  SKIP  district_industry.json not present (demographics campaign data not deployed)')
else:
    try:
        _ind=json.load(open(_ind_path,encoding='utf-8'))
        _ind_geos=_ind.get('geos') or {}
        _ind_rows=list(_ind_geos.values())
        ck('ind-schema district-industry-v1',
           _ind.get('schema')=='district-industry-v1', str(_ind.get('schema')))
        _ind_districts=sum(1 for _geo in _ind_rows if _geo.get('kind')=='district')
        ck('ind-geos 486 total and 435 districts',
           len(_ind_geos)==486 and _ind_districts==435,
           '%d geos, %d districts'%(len(_ind_geos),_ind_districts))
        _ind_sectors={'CNS%02d'%_i for _i in range(1,21)}
        ck('ind-sector_meta exactly CNS01..CNS20',
           set(_ind.get('sector_meta') or {})==_ind_sectors,
           str(sorted((_ind.get('sector_meta') or {}).keys())))
        try:
            _ind_recon=[float(_geo['shift_share_total'].get('reconciliation_abs_err'))
                        for _geo in _ind_rows if _geo.get('shift_share_total') is not None]
            _ind_max_recon=max(_ind_recon,default=float('inf'))
        except (AttributeError,TypeError,ValueError):
            _ind_max_recon=float('inf')
        ck('ind-shift-share max reconciliation_abs_err <= 0.5',
           _ind_max_recon<=0.5, 'max %s'%_ind_max_recon)
        _ind_bad_makeup=[]
        for _ind_key,_geo in _ind_geos.items():
            if (_geo.get('employment') or {}).get('C000',0)>0:
                try:
                    _ind_makeup_total=sum(float(_v) for _v in (_geo.get('makeup_pct') or {}).values()
                                          if _v is not None)
                except (AttributeError,TypeError,ValueError):
                    _ind_makeup_total=float('inf')
                if not 0.985<=_ind_makeup_total<=1.015:
                    _ind_bad_makeup.append(_ind_key)
        ck('ind-positive-employment makeup_pct sums within [0.985, 1.015]',
           not _ind_bad_makeup,
           '%d geos outside tolerance: %s'%(len(_ind_bad_makeup),_ind_bad_makeup[:5]))
        import datetime as _ind_dt
        _ind_age=None
        try:
            _ind_generated=_ind_dt.datetime.fromisoformat(
                str(_ind.get('generated_at') or '').replace('Z','+00:00'))
            if _ind_generated.tzinfo is None:
                _ind_generated=_ind_generated.replace(tzinfo=_ind_dt.timezone.utc)
            _ind_age=(_ind_dt.datetime.now(_ind_dt.timezone.utc)-
                      _ind_generated.astimezone(_ind_dt.timezone.utc)).total_seconds()/86400
        except (TypeError,ValueError):
            pass
        ck('ind-generated_at parses ISO and artifact age between -1 and 450 days',
           _ind_age is not None and -1<=_ind_age<=450,
           'ANNUAL REFRESH DUE: rerun fetch_cd_geography/fetch_lodes_wac/build_district_industry (+fetch_cbp_4digit, validate_district_industry)')
        _ind_bundle_path=B+r'\member_desk_bundle.json'
        if os.path.exists(_ind_bundle_path):
            _ind_bundle=json.load(open(_ind_bundle_path,encoding='utf-8'))
            _ind_member_count=sum(1 for _member in _ind_bundle.get('members',[])
                                  if _member.get('industry') is not None)
            ck('ind-member desk non-null industry coverage >= 500',
               _ind_member_count>=500, '%d members'%_ind_member_count)
            _ind4_path=B+r'\data\lobbying\district_industry_4digit.json'
            if os.path.exists(_ind4_path):
                _ind4=json.load(open(_ind4_path,encoding='utf-8'))
                ck('ind4-schema district-industry-4digit-v1',
                   _ind4.get('schema')=='district-industry-4digit-v1',
                   str(_ind4.get('schema')))
                ck('ind4-district coverage exactly 435',
                   len(_ind4.get('districts') or {})==435,
                   '%d districts'%len(_ind4.get('districts') or {}))
                _potential_expected_members=copy.deepcopy(
                    _ind_bundle.get('members',[]))
                for _member in _potential_expected_members:
                    _packed=_member.get('industry') or {}
                    _geo=_packed.get('geo')
                    _source=_ind_geos.get(_geo)
                    _expected=_demographic_potential.evaluate(
                        _geo,_source,_ind4,_ind.get('sector_meta') or {})
                    _member['demographic_potential']=_expected
                _potential_stats=_demographic_potential.score_members(
                    _potential_expected_members,_ind4)
                _potential_bad=[]
                for _actual_member,_expected_member in zip(
                        _ind_bundle.get('members',[]),
                        _potential_expected_members):
                    _actual=_actual_member.get('demographic_potential') or {}
                    _expected=(
                        _expected_member.get('demographic_potential') or {})
                    if _actual!=_expected:
                        _potential_bad.append(_actual_member.get('name'))
                _potential_count=_potential_stats['qualified_members']
                _potential_meta=_ind_bundle.get('demographic_potential_meta') or {}
                ck('demographic-potential qualification, score, rank, and tier recompute exactly',
                   not _potential_bad,
                   '%d mismatches: %s'%(len(_potential_bad),_potential_bad[:5]))
                ck('demographic-potential metadata count matches qualifying members',
                   _potential_count>0 and
                   _potential_count==_potential_meta.get('qualified_members'),
                   'members %d, metadata %s'%(
                       _potential_count,_potential_meta.get('qualified_members')))
                ck('demographic-potential metadata count matches qualifying markets',
                   _potential_stats['qualified_markets']>0 and
                   _potential_stats['qualified_markets']==
                   _potential_meta.get('qualified_markets'),
                   'markets %d, metadata %s'%(
                       _potential_stats['qualified_markets'],
                       _potential_meta.get('qualified_markets')))
                ck('demographic-potential score weights are pinned and sum to 100%',
                   _potential_meta.get('score_weights')==
                   _demographic_potential.SCORE_WEIGHTS and
                   abs(sum((_potential_meta.get('score_weights') or {}).values())
                       -1.0)<1e-9,
                   str(_potential_meta.get('score_weights')))
                ck('demographic-potential weights are explicitly judgmental, not fitted',
                   _potential_meta.get('weight_basis')==
                   _demographic_potential.WEIGHT_BASIS and
                   _potential_meta.get('not_empirically_optimized') is True and
                   bool(_potential_meta.get('optimization_requirement')),
                   str(_potential_meta.get('weight_basis')))
                _potential_defs=_potential_meta.get(
                    'component_definitions') or []
                ck('demographic-potential formula and official-source catalog is complete',
                   [row.get('key') for row in _potential_defs
                    if isinstance(row,dict)]==
                   list(_demographic_potential.SCORE_WEIGHTS) and
                   _potential_meta.get('source_catalog')==
                   _demographic_potential.SOURCE_CATALOG,
                   '%d definitions / %d sources'%(
                       len(_potential_defs),
                       len(_potential_meta.get('source_catalog') or {})))
                _sensitivity=_potential_meta.get('sensitivity_audit') or {}
                ck('demographic-potential sensitivity audit covers equal and leave-one-out weights',
                   _sensitivity.get('version')==
                   _demographic_potential.SENSITIVITY_VERSION and
                   _sensitivity.get('scenario_count')==9 and
                   len(_sensitivity.get('leave_one_out') or [])==8,
                   str(_sensitivity))
                _structural=_potential_meta.get(
                    'structural_adjustment') or {}
                ck('demographic-potential structural companion recomputes exactly',
                   _structural==
                   _potential_stats.get('structural_adjustment') and
                   _structural.get('version')==
                   _demographic_potential.STRUCTURAL_ADJUSTMENT_VERSION and
                   _structural.get('adjusted_markets')>0,
                   str(_structural.get('comparison_with_baseline')))
                _community_path=B+r'\data\district\district_community.json'
                _community=(
                    json.load(open(_community_path,encoding='utf-8'))
                    if os.path.exists(_community_path) else {})
                _community_coverage=_community.get('coverage') or {}
                _ak_community=(
                    (_community.get('districts') or {}).get('AK-00') or {})
                ck('ACS community artifact is complete 2024 119th-Congress context',
                   _community.get('schema')=='district-community-acs-v1' and
                   _community.get('vintage')==2024 and
                   _community.get('congress')==119 and
                   _community_coverage.get('complete_districts',0)>=435 and
                   _community_coverage.get('complete_states',0)>=50 and
                   (_ak_community.get('estimates') or {}).get(
                       'population')==740133 and
                   _ak_community.get('geo_id')=='5001900US0200',
                   str(_community_coverage))
                ck('ACS community variable/profile/source catalogs are complete',
                   set(_community.get('profile_families') or {})==
                   {'DP03','DP04','DP05'} and
                   set(_community.get('variable_catalog') or {})>={
                       'population','median_age',
                       'median_household_income','unemployed',
                       'below_poverty','median_home_value',
                       'median_gross_rent'} and
                   len(_community.get('sources') or {})==11,
                   '%d variables / %d sources'%(
                       len(_community.get('variable_catalog') or {}),
                       len(_community.get('sources') or {})))
                _incomplete_large_bad=[
                    _member.get('name')
                    for _member in _ind_bundle.get('members',[])
                    if ((_member.get('demographic_potential') or {}).get(
                            'qualified') and
                        (((_member.get('employers') or {}).get('large') or {})
                         .get('complete') is False) and
                        ((_member.get('demographic_potential') or {}).get(
                            'available') or {}).get(
                            'large_employer_reach') is not False)
                ]
                ck('incomplete CBP large-employer cells are missing/neutral in score',
                   not _incomplete_large_bad,
                   '%d violations: %s'%(
                       len(_incomplete_large_bad),
                       _incomplete_large_bad[:5]))
            else:
                ck('district_industry_4digit.json present for potential cross-check',
                   False,'four-digit artifact missing')
            _employer_member_count=sum(
                1 for _member in _ind_bundle.get('members',[])
                if _member.get('employers') is not None)
            _employer_bad=sum(
                1 for _member in _ind_bundle.get('members',[])
                for _row in ((_member.get('employers') or {}).get('top_federal') or [])
                if (not str(_row.get('name') or '').strip() or
                    str(_row.get('name') or '').strip().casefold()=='multiple recipients' or
                    _row.get('amount') is None)
            )
            ck('district employer candidates attached to >= 500 members',
               _employer_member_count>=500,'%d members'%_employer_member_count)
            ck('district employer candidate rows are named, valued, and non-aggregate',
               _employer_bad==0,'%d malformed rows'%_employer_bad)
            _desk_html_path=B+r'\cu-member-desk.html'
            if os.path.exists(_desk_html_path):
                _desk_html=open(_desk_html_path,encoding='utf-8').read()
                ck('desk renders Demographic Potential sticker and filter',
                   'Demographic Potential only' in _desk_html and
                   'class=potential-tag' in _desk_html)
                ck('desk renders Demographic Potential score, rank, and sorting',
                   'Sort: District opportunity' in _desk_html and
                   'Sort: Structurally adjusted potential' in _desk_html and
                   'Potential score / 100' in _desk_html and
                   'potential.evidence_coverage' in _desk_html and
                   "rstate.sort==='potential'" in _desk_html)
                ck('desk renders weight provenance, formulas, sources, and sensitivity',
                   'Weight provenance:' in _desk_html and
                   'Weights are policy assumptions, not empirical coefficients.' in
                   _desk_html and
                   'potentialMeta.component_definitions' in _desk_html and
                   'potentialMeta.source_catalog' in _desk_html and
                   'Sensitivity audit:' in _desk_html and
                   'Rank range across sensitivity scenarios' in _desk_html)
                ck('desk renders ACS community controls and structural companion',
                   'Structural community context' in _desk_html and
                   '2024 ACS 1-year' in _desk_html and
                   'Structurally adjusted / 100' in _desk_html and
                   'Structural companion score:' in _desk_html and
                   'potential.structural_adjustment' in _desk_html)
                ck('desk renders demographic industry pie chart',
                   "opts.label||'Main industries in the district'" in _desk_html and
                   'pieSVG(shares,meta,92' in _desk_html)
                ck('desk renders employer candidate lists with method caveat',
                   'Largest employer candidates' in _desk_html and
                   'employer candidates, not verified employee-count rankings' in _desk_html)
            else:
                ck('cu-member-desk.html present for demographic display audit',
                   False,'dashboard missing')
            _potential_method_path=os.path.join(
                B,'DEMOGRAPHIC_POTENTIAL_METHOD.md')
            _potential_method=(
                open(_potential_method_path,encoding='utf-8').read()
                if os.path.exists(_potential_method_path) else '')
            ck('demographic-potential formula is preserved in tracked methodology',
               'Demographic Potential Score = 100 * sum' in
               _potential_method and
               'Missing evidence' in _potential_method and
               'judgmental policy weights' in _potential_method and
               'Exact formulas' in _potential_method and
               'Source lineage' in _potential_method and
               'Structural controls and companion score' in
               _potential_method and
               'not a prediction' in _potential_method,
               _potential_method_path)
        else:
            print('  SKIP  ind-desk coverage (member_desk_bundle.json not present)')
    except Exception as exc:
        ck('ind-artifact well-formed', False, repr(exc))

# run-5b guards (audit 2026-07-06): the psychology/media-layer artifacts had ZERO audit coverage —
# a truncated big5_llm_opus.json or malformed reference/co-appearance edges would sail through.
# Schema + sanity floors only (freshness is meaningless mid-collector-crawl); each skips if absent.
if os.path.exists(B+r'\data\members\big5_llm_opus.json'):
    _b5=json.load(open(B+r'\data\members\big5_llm_opus.json',encoding='utf-8')).get('traits',{})
    _bad=sum(1 for d in _b5.values() if not all(k in d for k in ('O','C','E','A','ES')) or not d.get('n_chunks'))
    ck('big5_llm artifact schema (5 traits + n_chunks per member, coverage >=400)',
       len(_b5)>=400 and _bad==0, '%d members, %d malformed'%(len(_b5),_bad))
if os.path.exists(B+r'\member_references.json'):
    _re_=json.load(open(B+r'\member_references.json',encoding='utf-8')).get('edges',[])
    _bad=sum(1 for e in _re_ if not all(k in e for k in ('a','b','n_raw','n_scored','pos','neg','net')) or e['n_scored']<(e['pos']+e['neg']))
    ck('reference-layer edge schema (keys present, scored>=pos+neg, >=1000 edges)',
       len(_re_)>=1000 and _bad==0, '%d edges, %d malformed'%(len(_re_),_bad))
# run-5h guard (user 2026-07-06 'make sure the displays are entirely accurate'): the member-desk
# pillar bars + weights MUST match the rank authority (ranks_python.json) — they silently drifted
# (Crapo influence 1.000 in the authority displayed as 0.222 from a stale workbook column; weights
# 0.40/0.20/0.35/0.05 vs the then-live 0.25/0.30/0.40/0.05), making the #1 rank look unjustified on the
# desk. This asserts the DISPLAYED decomposition reconstructs the DISPLAYED score for every member.
if os.path.exists(B+r'\member_desk_bundle.json') and os.path.exists(B+r'\ranks_python.json'):
    _bun=json.load(open(B+r'\member_desk_bundle.json',encoding='utf-8')); _ra=json.load(open(B+r'\ranks_python.json',encoding='utf-8'))
    _W=_bun.get('weights',{}); _pcts=_ra.get('pcts',{})
    from engine.pillar_policy import PILLAR_WEIGHTS as _PW, scoring_policy as _scoring_policy
    _expectedW={'leg':_PW[2],'press':_PW[3],'comm':_PW[4],'demo':_PW[5]}
    ck('press sentiment is active at a diminished weight and may affect labels',
       all(_W.get(_key)==_value for _key,_value in _expectedW.items()) and
       _ra.get('scoring_policy')==_scoring_policy(),
       'bundle=%s policy=%s'%(_W,_ra.get('scoring_policy')))
    _bad=[]; _nomatch=0
    for _m in _bun.get('members',[]):
        _nm=_m.get('name'); _p=_pcts.get(_nm)
        if not _p: continue
        # displayed pillars == authority pcts
        if any(abs((_m.get(_dk) or 0)-_p[_ak])>1e-4 for _dk,_ak in (('press_n','press'),('leg_n','leg'),('comm_n','inf'),('demo_n','dem'))):
            _nomatch+=1
        # displayed weights*displayed pillars reconstruct displayed score
        _rc=_W.get('leg',0)*(_m.get('leg_n') or 0)+_W.get('press',0)*(_m.get('press_n') or 0)+_W.get('comm',0)*(_m.get('comm_n') or 0)+_W.get('demo',0)*(_m.get('demo_n') or 0)
        if abs(_rc-(_m.get('score') or 0))>0.003: _bad.append(_nm)
    ck('desk pillar bars match rank authority (no stale-workbook drift)', _nomatch==0, '%d members mismatch authority pcts'%_nomatch)
    ck('desk displayed pillars*weights reconstruct displayed score', len(_bad)==0, '%d members fail reconstruction: %s'%(len(_bad),_bad[:3]))
    # Paul-directed fill contract (2026-07-17): every member renders a Big-Five panel — gold-anchored
    # values primary, labeled estimate layer where gold is absent; estimate slots must carry a
    # provenance label. Guard-only display returns via DESK_B5_FALLBACK=0 (then this check adapts).
    _fill_on=os.environ.get('DESK_B5_FALLBACK','1')!='0'
    if _fill_on:
        _b5empty=[]; _fb_unlabeled=0
        for _m in _bun.get('members',[]):
            _p=_m.get('psy') or {}
            _g=(_p.get('big5_pct') or [None]*5); _f=(_p.get('big5_fb_pct') or [None]*5); _s=(_p.get('big5_fb_src') or [None]*5)
            _mrg=[(_g[_i] if _i<len(_g) and _g[_i] is not None else (_f[_i] if _i<len(_f) else None)) for _i in range(5)]
            if not any(_v is not None for _v in _mrg): _b5empty.append(_m.get('name'))
            for _i in range(5):
                if (_i>=len(_g) or _g[_i] is None) and _i<len(_f) and _f[_i] is not None and not (_i<len(_s) and _s[_i]):
                    _fb_unlabeled+=1
        ck('every member renders a Big-Five panel (gold or labeled estimate; Paul fill 2026-07-17)',
           not _b5empty, str(_b5empty[:6]))
        ck('every estimate slot carries a provenance label', _fb_unlabeled==0, '%d unlabeled'%_fb_unlabeled)
    _fr=json.load(open(B+r'\data\members\formal_visibility_roles.json',encoding='utf-8'))
    from engine import influence_roles as _ir
    _frm=_fr.get('members') or {}; _rb=_ra.get('formal_role_influence_bonus') or {}
    _rd=_ra.get('formal_role_display_weight') or {}
    _rn=_ra.get('formal_role_influence_nominal') or {}; _rr=_ra.get('formal_role_influence_eligibility') or {}
    _rbase=_ra.get('influence_raw_base') or {}; _radj=_ra.get('influence_raw_adjusted') or {}
    _cseats={r[0]:str(r[4] or '') for r in _wb['Committee Data'].iter_rows(min_row=2,values_only=True) if r[0]}
    _desk_by={m.get('name'):m for m in _bun.get('members',[])}
    _rolebad=0
    for _nm in _ra.get('rank',{}):
        _entry=_frm.get(_nm) or {}; _nom=float(_entry.get('influence_bonus') or 0.0)
        _want,_reason,_overlap=_ir.eligible_bonus(_entry,_cseats.get(_nm,''))
        _rolebad += int(abs(float(_rn.get(_nm,999))-_nom)>1e-9 or _rr.get(_nm)!=_reason)
        _rolebad += int(abs(float(_rb.get(_nm,999)))>1e-9)
        _rolebad += int(abs(float(_rd.get(_nm,999))-_want)>1e-9)
        _rolebad += int(abs(float(_radj.get(_nm,999))-float(_rbase.get(_nm,0)))>1e-9)
        _dm=_desk_by.get(_nm) or {}
        _rolebad += int(abs(float(_dm.get('formal_role_influence_bonus',999)))>1e-9)
        _rolebad += int(abs(float(_dm.get('formal_role_display_weight',999))-_want)>1e-9)
        _rolebad += int(abs(float(_dm.get('comm_raw',999))-float(_radj.get(_nm,0)))>1e-9)
    _rmeta=_ra.get('formal_role_influence') or {}
    ck('formal office remains titled display context and never changes Influence or rank',
       _fr.get('schema')=='party-elite-influence-roles-v3' and
       _fr.get('valid_until','')>__import__('datetime').date.today().isoformat() and
       len(_frm)>=20 and 'Chuck Grassley' not in _frm and _rolebad==0 and
        _rb.get('Mike Johnson')==0.0 and _rb.get('Tim Scott')==0.0 and _rb.get('John Barrasso')==0.0 and
        _rd.get('Mike Johnson')==3.0 and _rmeta.get('applied_party_elites')==0 and
        _rmeta.get('listed_party_elites')==len(_frm) and _rmeta.get('display_context_party_elites',0)>0,
        '%d listed party elites, %d rank-applied, %d display contexts, %d decomposition errors'%(
        len(_frm),_rmeta.get('applied_party_elites',0),_rmeta.get('display_context_party_elites',0),_rolebad))
    # Blue/Pink remain independently constructed. The equal-weight E remains a frozen comparator;
    # the active OOF-gated E is audited below with dominance and role-adjusted visibility included.
    _psy=[m.get('psy') or {} for m in _bun.get('members',[])]
    _bp=[p for p in _psy if len(p.get('big5_pct') or [])==5 and len(p.get('llm_pct') or [])==5]
    _same_n=sum(1 for p in _bp if p['big5_pct'][4]==p['llm_pct'][4])
    ck('desk broad-gold and raw Pink Neuroticism reads do not collapse to one value',
       len(_bp)>=400 and _same_n<len(_bp), '%d paired members, %d identical N percentiles'%(len(_bp),_same_n))
    _blue=next((_p for _p in (B+r'\data\members\big5_blue_validated.json',B+r'\data\members\big5_blue_emb.json',B+r'\data\members\big5_v2.json') if os.path.exists(_p)),B+r'\data\members\big5_v2.json')
    _pink=B+r'\data\members\big5_llm_opus.json'
    if os.path.exists(_blue) and os.path.exists(_pink) and os.path.getmtime(_pink)<os.path.getmtime(_blue):
        _unflag=sum(1 for p in _bp if not ((p.get('big5_vintage') or {}).get('pink_older_days')>0))
        ck('desk preserves mixed Blue/Pink corpus-vintage metadata for archived diagnostics',
           _unflag==0, '%d/%d paired members lack the stale-vintage disclosure metadata'%(_unflag,len(_bp)))
    _egp=B+r'\data\stance\big5_e_gold_anchored.json'; _evp=B+r'\data\stance\big5_external_gold_validation.json'
    if os.path.exists(_egp) and os.path.exists(_evp):
        _eg=json.load(open(_egp,encoding='utf-8')); _ev=json.load(open(_evp,encoding='utf-8'))
        _egm=_eg.get('members') or {}; _epr=_ev.get('panel_results') or {}
        _ebad=sum(1 for d in _egm.values() if d.get('E') is None or d.get('percentile') is None or
                  d.get('channel_weights')!={'blue_visibility':.5,'pink_observer':.5})
        ck('external Senate gold preserves the discriminant equal-weight Extraversion comparator',
           _eg.get('schema')=='big5-e-gold-anchored-v1' and _eg.get('status')=='ACTIVE_EXTERNAL_GOLD_ANCHORED' and
           _ev.get('verdict')=='PASS' and len(_egm)>=450 and _ebad==0 and
           all((d.get('e_r') or 0)>=.30 and d.get('diagonal_dominant') and (d.get('e_permutation_p_one_sided') or 1)<.05
               for d in _epr.values()),
           '%d E members, %d malformed; panel r=%s'%(
           len(_egm),_ebad,{k:round(v.get('e_r',0),3) for k,v in _epr.items()}))
    _mgp=B+r'\data\stance\big5_gold_anchored.json'
    if os.path.exists(_mgp):
        _mg=json.load(open(_mgp,encoding='utf-8')); _mgt=_mg.get('traits') or {}; _mge=_mgt.get('E') or {}
        _mgem=_mge.get('members') or {}; _mgew=_mge.get('selected_weights') or {}
        _mgebad=sum(1 for d in _mgem.values() if d.get('E') is None or d.get('percentile') is None)
        _mgpan=(_mge.get('validation') or {}).get('panels') or {}
        ck('multitrait OOF Senate gold activates E with dominance and role-adjusted visibility',
           _mg.get('schema')=='big5-gold-anchored-v2' and set(_mgt)=={'O','C','E','A','N'} and
           _mge.get('status')=='ACTIVE_EXTERNAL_GOLD_ANCHORED' and len(_mgem)>=450 and _mgebad==0 and
           _mgew.get('interaction_dominance',0)>=.1 and _mgew.get('role_adjusted_visibility',0)>=.1 and
           all(d.get('diagonal_dominant') and (d.get('r') or 0)>=.2 and
               (d.get('permutation_p_one_sided') or 1)<.1 for d in _mgpan.values()) and
           all((d.get('validation') or {}).get('predictions_are_out_of_fold') for d in _mgt.values()),
           '%d E members, %d malformed; weights=%s; panel r=%s'%(
           len(_mgem),_mgebad,_mgew,{k:round(v.get('r',0),3) for k,v in _mgpan.items()}))
        _expected_status={'O':'ACTIVE_ONE_PANEL_ASSOCIATION_SUPPORTED',
                          'C':'ACTIVE_ONE_PANEL_ASSOCIATION_SUPPORTED',
                          'E':'ACTIVE_EXTERNAL_GOLD_ANCHORED',
                          'A':'ACTIVE_EXTERNAL_GOLD_ANCHORED',
                          'N':'ACTIVE_ONE_PANEL_ASSOCIATION_SUPPORTED'}
        ck('all broad Big Five traits activate with explicit panel-strength labels',
           {k:v.get('status') for k,v in _mgt.items()}==_expected_status and
           all(any((v.get('association_gates') or {}).values()) for v in _mgt.values()) and
           all(len(v.get('members') or {})>=300 for v in _mgt.values()),
           'statuses=%s; association gates=%s'%(
           {k:v.get('status') for k,v in _mgt.items()},
           {k:v.get('association_gates') for k,v in _mgt.items()}))
    # Source-primary Blue is the completed Round-4 promotion, not the pooled lexical v2 draft.
    _bep=B+r'\data\members\big5_blue_emb.json'
    if os.path.exists(_bep):
        _be=json.load(open(_bep,encoding='utf-8')); _bem=_be.get('traits',{}); _bev=_be.get('validation') or {}
        _bars=_bev.get('bars') or {}; _bad=sum(1 for d in _bem.values() if
             not all(k in d for k in ('O','C','E','A','ES','N','tokens','n_source')) or
             abs(float(d['ES'])+float(d['N']))>.002 or d.get('n_source')!='affective_n')
        ck('source-primary Blue promotion preserves frozen MTMM/criterion gate and affective-N slot',
           _bev.get('verdict')=='PASS' and all(_bars.values()) and
           _bev.get('mtmm_same_trait_cross_genre',0)>1.15*_bev.get('mtmm_cross_trait_within_genre',1) and
           set(_be.get('source_genres') or [])=={'CREC','INTERVIEW','social'} and
           len(_bem)>=450 and _bad==0,
           '%d members, %d malformed, verdict=%s, MTMM=%s/%s'%
           (len(_bem),_bad,_bev.get('verdict'),_bev.get('mtmm_same_trait_cross_genre'),
            _bev.get('mtmm_cross_trait_within_genre')))
    _bvp=B+r'\data\members\big5_blue_validated.json'
    if os.path.exists(_bvp):
        _bv=json.load(open(_bvp,encoding='utf-8')); _bvm=_bv.get('traits',{}); _vv=_bv.get('validation') or {}; _rr=_bv.get('reliability') or {}
        _diag=_vv.get('diagonal') or {}; _known=_vv.get('known_groups') or {}; _rt=_rr.get('traits') or {}
        _bad=sum(1 for d in _bvm.values() if not all(k in d for k in ('O','C','E','A','N','ES','tokens')) or
                 (d.get('N') is not None and abs(float(d['ES'])+float(d['N']))>.002))
        ck('validated multimethod Big Five preserves all 5 reliability/known-group/discriminant gates',
           len(_bvm)>=450 and _bad==0 and _vv.get('diagonal_pass_n')==5 and _vv.get('known_groups_pass_n')==5 and
           len(_diag)==5 and all(d.get('pass') for d in _diag.values()) and
           len(_known)==5 and all(d.get('pass') for d in _known.values()) and
           _rr.get('pass_n')==5 and _rr.get('package_overlap')==0 and len(_rt)==5 and all(d.get('pass') for d in _rt.values()),
           '%d members, %d malformed, reliability=%s, known=%s, diagonal=%s, package overlap=%s'%
           (len(_bvm),_bad,_rr.get('pass_n'),_vv.get('known_groups_pass_n'),_vv.get('diagonal_pass_n'),_rr.get('package_overlap')))
        _ocebad=sum(1 for d in _bvm.values() if any(d.get(k) is None for k in ('O','C','E')))
        _ncovered=sum(1 for d in _bvm.values() if d.get('N') is not None)
        _coverage=_vv.get('coverage') or {}
        ck('independent O/C/E coverage is not truncated by the narrower affective-N stratum',
           _ocebad==0 and len(_bvm)>_ncovered and
           all(_coverage.get(k)==len(_bvm) for k in ('O','C','E')) and _coverage.get('N')==_ncovered,
           '%d text-basis members, %d O/C/E gaps, %d N-covered; recorded coverage=%s'%
           (len(_bvm),_ocebad,_ncovered,_coverage))
    # Round-9 qualified A facet remains separate from the cooperation engine role.
    _avp=B+r'\data\stance\interview_antagonism.json'
    if os.path.exists(_avp):
        _av=json.load(open(_avp,encoding='utf-8')); _avm=_av.get('members',{})
        _mal=sum(1 for d in _avm.values() if not all(k in d for k in
                 ('antagonism_z','agreeableness_z','raw_antagonism','half_a','half_b','words'))
                 or abs(d['agreeableness_z']+d['antagonism_z'])>.002)
        ck('validated A-antagonism facet preserves qualified schema and inverse pole',
           _av.get('n')==len(_avm) and len(_avm)>=300 and _av.get('reliability_sb',0)>=.65 and _mal==0
           and (_av.get('source') or {}).get('sha256')=='15FFB00BA199708C17733260205E0580CE766F3B37E0C69A74965B35B4AB0575'
           and 'residualization-only' in ((_av.get('controls') or {}).get('role') or ''),
           '%d members, SB=%s, %d malformed/non-inverse'%
           (len(_avm),_av.get('reliability_sb'),_mal))
        _desk_av={m.get('name'):(m.get('psy') or {}).get('a_validated_pct')
                  for m in _bun.get('members',[]) if (m.get('psy') or {}).get('a_validated_pct') is not None}
        ck('validated A facet propagates to desk without replacing engine/text A',
           len(_desk_av)==len(_avm) and all(0<=v<=1 for v in _desk_av.values())
           and all((m.get('psy') or {}).get('a_validated_scope')=='interview/media-covered stratum only'
                   for m in _bun.get('members',[]) if (m.get('psy') or {}).get('a_validated_pct') is not None),
           '%d artifact members, %d desk values'%(len(_avm),len(_desk_av)))
        _scope=_av.get('scope') or {}; _sel=_av.get('selection_audit') or {}
        ck('validated A claim is sampling-frame scoped (no full-roster percentile/imputation)',
           _scope.get('covered_n')==322 and _scope.get('roster_n')==533
           and _scope.get('imputation')=='forbidden'
           and _sel.get('verdict')=='DEMOGRAPHICALLY SELECTED MEDIA-COVERED STRATUM',
           'scope=%s selection=%s'%(_scope,_sel.get('verdict')))
        # Every rendered paragraph must be byte-identical to a correctly labelled, basis-bound
        # brief. API-produced prose additionally requires the full-corpus semantic receipt; a
        # mechanical-only artifact is not eligible to render as "reviewed".
        _narp=B+r'\member_narratives_v2.json'; _narbp=B+r'\data\members\profile_basis_v2.json'
        _narrp=B+r'\data\members\member_narratives_v2_semantic_receipt.json'
        _narevp=B+r'\data\members\member_narrative_evidence_v4.json'
        _nardoc=json.load(open(_narp,encoding='utf-8')) if os.path.exists(_narp) else {}
        _narlab=(_nardoc.get('schema'),_nardoc.get('status')) in {
                 ('member-narratives-v2-local','ACTIVE_LOCAL_ZERO_API'),
                 ('member-narratives-v2-api','ACTIVE_API_GENERATED_USER_RATIFIED')}
        _narbound=bool(_nardoc) and os.path.exists(_narbp) and _nardoc.get('basis_sha256')==_file_sha256(_narbp)
        _narreceipt=json.load(open(_narrp,encoding='utf-8')) if os.path.exists(_narrp) else {}
        _narapi=_nardoc.get('schema')=='member-narratives-v2-api'
        _narreviewed=(not _narapi) or (
            _narreceipt.get('schema')=='member-narratives-v2-semantic-receipt-v1'
            and _narreceipt.get('status')=='PASS'
            and _narreceipt.get('basis_sha256')==_nardoc.get('basis_sha256')
            and os.path.exists(_narevp)
            and _narreceipt.get('evidence_sha256')==_file_sha256(_narevp)
            and (_narreceipt.get('counts') or {}).get('semantic_pass')==533
            and _narreceipt.get('active_artifact_sha256')==_file_sha256(_narp))
        _narsrc=(_nardoc.get('members') or {}) if (_narlab and _narbound and _narreviewed) else {}
        _nar=[(m.get('psy') or {}).get('narrative') or {} for m in _bun.get('members',[])]
        _nar_present=[n for n in _nar if n]
        _withheld=sum(1 for m in _bun.get('members',[]) if
                      (m.get('psy') or {}).get('narrative_status')=='withheld_pending_guarded_local_rebuild_and_user_approval')
        _ungrounded=0
        for _dm in _bun.get('members',[]):
            _dn=(_dm.get('psy') or {}).get('narrative') or {}
            if not _dn: continue
            _src=_narsrc.get(_dm.get('name')) or {}
            if _src.get('ok') is not True or any(not _src.get(_p) or _dn.get(_p)!=_src.get(_p)
                                                 for _p in ('p1','p2','p3')): _ungrounded+=1
        _safe=(not _nar_present and _withheld==533) or (
              bool(_narsrc) and len(_nar_present)+_withheld==533 and _ungrounded==0)
        ck('current narratives cannot silently fall back to legacy E/audio prose', _safe,
           '%d desk narratives, %d explicitly withheld, %d paragraphs not source-identical '
           '(labelled=%s basis-bound=%s semantic-receipt=%s, %d briefs)'%
           (len(_nar_present),_withheld,_ungrounded,_narlab,_narbound,_narreviewed,len(_narsrc)))
# 2026-07-13 personal-texture catch-net: this display-only artifact must cover every member for whom
# explicit autobiographical evidence was found, retain exact provenance, and propagate byte-for-field
# into the member desk. Coverage must retain explicit evidence for >=90% of candidate-bearing members;
# this ratio gate rewards precision without forcing false positives merely to hit a roster count.
# No low-confidence or source-less inference may enter a profile.
_ptp=B+r'\data\members\member_personal_texture.json'
if os.path.exists(_ptp):
    _pt=json.load(open(_ptp,encoding='utf-8')); _ptm=_pt.get('members',{}); _pp=_pt.get('provenance') or {}
    _bad=0; _ni=0; _local_narrative_leak=0; _bad_display_gate=0
    for _m,_d in _ptm.items():
        _it=_d.get('items') or []; _ni+=len(_it)
        if _d.get('n_items')!=len(_it): _bad+=1
        _bad+=sum(1 for x in _it if not all(x.get(k) not in (None,'') for k in
                 ('category','claim','confidence','evidence','source','classifier')) or x.get('confidence')=='low'
                 or not isinstance(x.get('narrative_eligible'),bool)
                 or not isinstance(x.get('display_eligible'),bool))
        _local_narrative_leak+=sum(1 for x in _it if str(x.get('classifier','')).startswith('local_explicit_')
                                    and x.get('narrative_eligible') is not False)
        _bad_display_gate+=sum(1 for x in _it if x.get('display_eligible') is True and (
            (str(x.get('classifier','')).startswith('local_explicit_') and
             (x.get('classifier')!='local_explicit_v6' or x.get('source')!='INTERVIEW')) or
            (not str(x.get('classifier','')).startswith('local_explicit_') and
             x.get('classifier')!='fable')))
    ck('personal-texture artifact precision/schema/provenance (no volume reward)',
        len(_ptm)>=.90*_pp.get('candidate_member_count',9999) and _ni>0 and _bad==0 and
        _local_narrative_leak==0 and _bad_display_gate==0 and
        _pp.get('candidate_count',0)>=_pp.get('semantic_gate_count',1) and
        _pp.get('classified_count',0)>=_pp.get('semantic_gate_count',1),
        '%d members, %d details, %d malformed, %d local-to-narrative leaks, %d invalid display admissions, candidates=%s classified=%s gate=%s'%
        (len(_ptm),_ni,_bad,_local_narrative_leak,_bad_display_gate,
         _pp.get('candidate_count'),_pp.get('classified_count'),_pp.get('semantic_gate_count')))
    if os.path.exists(B+r'\member_desk_bundle.json'):
        _dm={m.get('name'):m.get('personal') for m in json.load(open(B+r'\member_desk_bundle.json',encoding='utf-8')).get('members',[])}
        _mis=sum(1 for m,d in _ptm.items() if _dm.get(m)!=d)
        ck('personal-texture layer propagates exactly to member profiles (display-only)',
           _mis==0 and sum(1 for d in _dm.values() if d)==len(_ptm),
           '%d payload mismatches; artifact=%d desk=%d'%(_mis,len(_ptm),sum(1 for d in _dm.values() if d)))
else:
    ck('personal-texture artifact present (member_personal_texture.json)',False,
       'missing — run personal_texture_layer.py --score --go or --local, then rebuild the desk')
# Wikipedia biographies are revision-bound, extractive display context.  Full-roster accounting is
# required: a page may be explicitly missing, but it may not silently disappear between the roster,
# the summary artifact, and the desk payload.
_wkp=B+r'\data\members\member_wikipedia_summaries.json'
if os.path.exists(_wkp):
    _wk=json.load(open(_wkp,encoding='utf-8')); _wkm=_wk.get('members') or {}; _wkc=_wk.get('coverage') or {}
    _wk_roster=json.load(open(B+r'\data\engine\congestion.json',encoding='utf-8'))
    _wk_roster_n=len(_wk_roster)
    _wk_bad=sum(1 for _name,_row in _wkm.items() if (
        _row.get('status') not in ('ok','missing') or not _row.get('bioguide') or
        (_row.get('status')=='ok' and (
            not _row.get('url') or not _row.get('revision_id') or
            not (_row.get('summary') or {}).get('overview')))))
    ck('Wikipedia biography summaries account for the full current roster',
       _wk.get('schema')=='member-wikipedia-biographies-v1' and len(_wkm)==_wk_roster_n and
       _wkc.get('rows')==_wk_roster_n and not _wkc.get('missing') and _wk_bad==0,
       '%d rows, %d summarized, %d missing, %d malformed'%(
           len(_wkm),_wkc.get('ok',0),len(_wkc.get('missing') or []),_wk_bad))
    if os.path.exists(B+r'\member_desk_bundle.json'):
        _desk_bio={m.get('name'):(m.get('bio') or {}).get('wikipedia')
                   for m in json.load(open(B+r'\member_desk_bundle.json',encoding='utf-8')).get('members',[])}
        _wk_mis=sum(1 for name,row in _wkm.items() if row.get('status')=='ok' and
                    ((_desk_bio.get(name) or {}).get('revision_id')!=row.get('revision_id')))
        ck('Wikipedia summaries propagate revision-identically to member profiles',
           _wk_mis==0,'%d revision mismatches'%_wk_mis)
else:
    ck('Wikipedia biography summary artifact present',False,
       'missing — run fec/members/wikipedia_biographies.py, then rebuild the desk')
# run-5j guard (2026-07-06): exact-duplicate speech rows (same member+package+text) inflate word
# counts + big5 chunk sampling. The CREC backfill introduced 922 (no pre-insert dedup); collector
# legs now carry stored-package guards. Assert the corpus stays exact-dup-clean (threshold covers
# the small user-gated legacy ABC/CNN set, currently 0 post-dedup).
if os.path.exists(B+r'\social_media.db'):
    import sqlite3 as _sq
    _c=_sq.connect(B+r'\social_media.db',timeout=60); _c.execute('PRAGMA query_only=1')
    _ex=_c.execute("SELECT COALESCE(SUM(c-1),0) FROM (SELECT COUNT(*) c FROM speeches GROUP BY member,package,text HAVING c>1)").fetchone()[0]
    _c.close()
    ck('speech corpus exact-duplicate free (member+package+text)', _ex<=50, '%d exact-dup surplus rows'%_ex)
if os.path.exists(B+r'\data\network\co_appearance.json'):
    _co=json.load(open(B+r'\data\network\co_appearance.json',encoding='utf-8')).get('edges',[])
    _bad=sum(1 for e in _co if not all(k in e for k in ('a','b','n','last','bipartisan')) or e['n']<1)
    ck('co-appearance edge schema (keys present, n>=1, non-empty)', len(_co)>=100 and _bad==0,
       '%d pairs, %d malformed'%(len(_co),_bad))
# 2026-07-13 production audio guard: zero verified members is a valid fail-closed state. Every
# accepted clip must have two distinct, dated, method-identified target-only reviews, and no display
# producer may fall back to the retired model-only verifier or legacy 88-feature cache.
from audio_rules import MIN_MEMBER_TARGET_SEC as _MIN_AUDIO_MEMBER_SEC
_intake_path=B+r'\data\engine\collection_intake_receipt.json'
_audio_candidate_path=B+r'\data\members\b5_audio_reviews_v2.json.candidate'
_press_feed_path=B+r'\data\stance\press_refresh_feed_receipt.json'
_intake=json.load(open(_intake_path,encoding='utf-8')) if os.path.exists(_intake_path) else {}
_audio_candidate=json.load(open(_audio_candidate_path,encoding='utf-8')) if os.path.exists(_audio_candidate_path) else {}
_press_feed=json.load(open(_press_feed_path,encoding='utf-8')) if os.path.exists(_press_feed_path) else {}
_candidate_partition=_audio_candidate.get('partition') or {}; _candidate_source=_audio_candidate.get('source_manifest') or {}
_expected_audio_rows=int(((_intake.get('house') or {}).get('all_clip_rows')) or -1)
_manifest_path=_candidate_source.get('path') or ''
_candidate_hash_ok=bool(_manifest_path and os.path.isfile(_manifest_path) and
    _intake.get('schema')=='collection-intake-receipt-v1' and _intake.get('status')=='ACCEPTED' and
    _candidate_source.get('sha256')==_file_sha256(_manifest_path) and
    (_intake.get('inputs') or {}).get('clip_manifest_sha256')==_candidate_source.get('sha256'))
ck('accepted audio manifest is fully partitioned into a candidate-only review intake',
   _audio_candidate.get('schema')=='b5-audio-review-v2-candidate' and
   _audio_candidate.get('status')=='CANDIDATE_ONLY_CONSUMED_BY_NOTHING_AWAITING_DOUBLE_REVIEW' and
   _candidate_hash_ok and _candidate_partition.get('total')==_expected_audio_rows and
   _candidate_partition.get('queued')+_candidate_partition.get('excluded')==_candidate_partition.get('total') and
   (_audio_candidate.get('contract') or {}).get('active_artifacts_written') is False,
   'total=%s queued=%s excluded=%s hash=%s'%(_candidate_partition.get('total'),
     _candidate_partition.get('queued'),_candidate_partition.get('excluded'),_candidate_hash_ok))
_press_output_hashes=_press_feed.get('outputs_sha256') or {}
_press_outputs_ok=(set(_press_output_hashes)=={
    'data/stance/press_refresh_tone.csv','data/stance/press_refresh_fulltext.jsonl',
    'data/press/official_mentions.jsonl'})
for _rel,_sha in _press_output_hashes.items():
    _pp=os.path.join(B,*_rel.split('/'))
    if not os.path.isfile(_pp) or _file_sha256(_pp)!=_sha: _press_outputs_ok=False
_press_source_root=_press_feed.get('source_run') or ''
_press_source_hashes=_press_feed.get('source_sha256') or {}
_press_sources_ok=(bool(_press_source_root and os.path.isdir(_press_source_root)) and
    set(_press_source_hashes)=={'press_delta_scored.csv','press_delta_host_articles.csv',
        'press_delta_article_mentions.csv','press_network_retry_receipt.json'})
for _name,_sha in _press_source_hashes.items():
    _pp=os.path.join(_press_source_root,_name)
    if not os.path.isfile(_pp) or _file_sha256(_pp)!=_sha: _press_sources_ok=False
_press_intake_binding=_press_feed.get('collection_intake') or {}
_press_intake_source_hashes=(_intake.get('inputs') or {}).get('press_source_sha256') or {}
_press_intake_ok=(os.path.isfile(_intake_path) and
    _intake.get('schema')=='collection-intake-receipt-v1' and _intake.get('status')=='ACCEPTED' and
    _press_intake_binding.get('sha256')==_file_sha256(_intake_path) and
    os.path.normcase(os.path.abspath((_intake.get('inputs') or {}).get('press_delta_root') or ''))==
        os.path.normcase(os.path.abspath(_press_source_root)) and
    (_intake.get('inputs') or {}).get('press_retry_receipt_sha256')==
        _press_source_hashes.get('press_network_retry_receipt.json') and
    set(_press_intake_source_hashes)=={
        'press_delta_scored.csv','press_delta_host_articles.csv',
        'press_delta_article_mentions.csv'} and
    all(_press_intake_source_hashes.get(_name)==_press_source_hashes.get(_name)
        for _name in _press_intake_source_hashes))
_press_counts=_press_feed.get('counts') or {}; _press_contract=_press_feed.get('contract') or {}
_press_tone_path=B+r'\data\stance\press_refresh_tone.csv'
_press_mentions_path=B+r'\data\press\official_mentions.jsonl'
with open(_press_tone_path,encoding='utf-8-sig',newline='') as _pfh:
    _press_tone_actual=sum(1 for _row in csv.DictReader(_pfh))
with open(_press_mentions_path,encoding='utf-8') as _pfh:
    _press_mentions_actual=sum(1 for _line in _pfh if _line.strip())
ck('press refresh feed is hash-bound and keeps official mentions out of stance/speech',
   _press_feed.get('schema')=='press-refresh-engine-feed-v1' and
   _press_feed.get('status')=='ACTIVE_TYPED_SUPPLEMENT' and _press_outputs_ok and
   _press_sources_ok and _press_intake_ok and
   _press_counts.get('tone_rows')==_press_tone_actual and _press_tone_actual>0 and
   _press_counts.get('official_mentions')==_press_mentions_actual and _press_mentions_actual>0 and
   _press_contract.get('scored_host_articles_enter_press_stance') is True and
   _press_contract.get('official_mentions_enter_press_stance') is False and
   _press_contract.get('official_mentions_enter_speech_turns') is False,
   'tone=%s mentions=%s outputs=%s sources=%s intake=%s'%(_press_counts.get('tone_rows'),
     _press_counts.get('official_mentions'),_press_outputs_ok,_press_sources_ok,_press_intake_ok))
_email_manifest_path=B+r'\outputs\email_package\EMAIL_PACKAGE_MANIFEST.json'
try:
    _email_manifest=json.load(open(_email_manifest_path,encoding='utf-8'))
except Exception:
    _email_manifest={}
_email_inputs_ok=True
for _item in (_email_manifest.get('inputs') or [])+(_email_manifest.get('monolith_inputs') or []):
    _ep=os.path.join(B,_item.get('name') or '')
    if not os.path.isfile(_ep) or _file_sha256(_ep)!=_item.get('sha256'):
        _email_inputs_ok=False
_email_packages_ok=bool(_email_manifest.get('packages'))
for _rel,_item in (_email_manifest.get('packages') or {}).items():
    _ep=os.path.join(B,*_rel.split('/'))
    if (not os.path.isfile(_ep) or _file_sha256(_ep)!=_item.get('sha256') or
            os.path.getsize(_ep)>14*1024*1024):
        _email_packages_ok=False
ck('email delivery packages are current, hash-bound, and within the attachment budget',
   _email_manifest.get('schema')=='email-package-v2' and
   _email_manifest.get('status')=='CURRENT_HASH_BOUND' and _email_inputs_ok and
   _email_packages_ok and len(_email_manifest.get('inputs') or [])==8 and
   len(_email_manifest.get('monolith_inputs') or [])==1,
   'inputs=%s packages=%s'%(_email_inputs_ok,_email_packages_ok))
try:
    _bill_live=validate_live_repository(
        Path(B)/'member_desk_bundle.json',
        Path(B)/'bill_repository_live.json',
        Path(B)/'data'/'engine'/'congress_bill_refresh_receipt.json',
    )
    _bill_live_error=''
except Exception as _bill_exc:
    _bill_live={}; _bill_live_error='%s: %s'%(type(_bill_exc).__name__,_bill_exc)
ck('Congress.gov bill repository is weekly-current, complete, and hash-bound',
   bool((_bill_live.get('bills') or {})) and not (_bill_live.get('failures') or {}),
   _bill_live_error or 'bills=%d failures=%d'%(
       len(_bill_live.get('bills') or {}),len(_bill_live.get('failures') or {})))
_ar=json.load(open(B+r'\data\members\b5_audio_reviews_v2.json',encoding='utf-8'))
_au=json.load(open(B+r'\data\members\b5_audio_descriptors_v2.json',encoding='utf-8'))
_abad=0
for _rel,_clip in (_au.get('members') or {}).items():
    if (_clip.get('target_elapsed_sec',0)<_MIN_AUDIO_MEMBER_SEC or _clip.get('source_n',0)<2 or
        _clip.get('program_n',0)<2 or not _clip.get('cross_program')): _abad+=1
    for _row in _clip.get('clips') or []:
        if (_row.get('venue')!='on_air_interview' or not _row.get('outlet') or
            not _row.get('program') or not _row.get('aired_at') or not _row.get('recording_context') or
            _row.get('target_words',0)<50 or _row.get('pitch_deviation_from_speaker_baseline_z') is None or
            not (_row.get('topic') or {}).get('id')): _abad+=1
for _rel,_clip in (_au.get('rejected') or {}).items():
    if not isinstance(_clip,list) or not _clip: _abad+=1
for _rel,_clip in ((_ar.get('clips') or {}).items()):
    _revs=_clip.get('reviews') or []; _ids=[r.get('reviewer') for r in _revs]
    if ((_clip.get('venue')!='on_air_interview' or not _clip.get('outlet') or not _clip.get('program') or
         not _clip.get('aired_at') or not _clip.get('recording_context') or not _clip.get('speech_id') or
         not (_clip.get('topic') or {}).get('id')) or
        len(_revs)!=2 or len(set(_ids))!=2 or any(r.get('identity')!='TARGET' or
        r.get('speaker')!='TARGET_ONLY' or not r.get('reviewed_at') or not r.get('review_method') or
        r.get('transcript_alignment_confirmed') is not True or r.get('interviewer_excluded') is not True
        or r.get('speech_unit_confirmed') is not True or r.get('topic_from_aligned_transcript_confirmed') is not True
        for r in _revs)): _abad+=1
_acon=_ar.get('contract') or {}; _mcon=_au.get('measurement_contract') or {}; _dcon=_au.get('deployment_contract') or {}
ck('audio v2 fail-closed on-air, aligned, program-conditioned contract active',
   _ar.get('schema')=='b5-audio-review-v2' and _au.get('schema')=='b5-audio-descriptors-v2' and
   _au.get('status')=='ACTIVE_FAIL_CLOSED' and _dcon.get('unverified_audio_enters_profiles') is False and
   _dcon.get('on_air_interviews_only') is True and _dcon.get('program_conditioned_percentiles') is True and
   _dcon.get('within_speaker_pitch_deviation') is True and _dcon.get('transcript_identified_topics') is True and
    _dcon.get('audio_may_modify_big_five') is False and _acon.get('transcript_alignment_and_interviewer_exclusion_double_confirmed') is True and
    _acon.get('speech_unit_and_transcript_topic_double_confirmed') is True and
    _acon.get('minimum_confirmed_target_seconds_per_member')==_MIN_AUDIO_MEMBER_SEC and
    _acon.get('minimum_independent_programs_per_member')==2 and
    _acon.get('minimum_speech_units_per_member')==5 and _acon.get('minimum_distinct_topics_per_member')==2 and
    _mcon.get('minimum_target_seconds_per_member')==_MIN_AUDIO_MEMBER_SEC and
    _mcon.get('minimum_speech_units_per_member')==5 and _mcon.get('minimum_distinct_topics_per_member')==2 and
    _mcon.get('minimum_aligned_target_words_per_clip')==50 and
   _mcon.get('percentile_normalization')=='within outlet+program; never raw cross-speaker pitch/energy' and _abad==0,
   '%d accepted clips, %d verified members, %d contract errors'%(_au.get('accepted_clip_n',0),len(_au.get('members') or {}),_abad))
_ncode=open(FE+r'\members\b5_narratives.py',encoding='utf-8').read()
_dcode=open(FE+r'\surfaces\make_member_desk.py',encoding='utf-8').read()
_wcode=open(FE+r'\weekly_refresh.py',encoding='utf-8').read()
ck('legacy audio and paid narrative paths retired from unattended/display consumption',
   'b5_audio_cache.json' not in _ncode and 'b5_audio_verify.py' not in _wcode and
   'b5_audio_build.py' not in _wcode and 'b5_narratives.py' not in _wcode and
   '--allow-paid-api' in _ncode and 'member_narratives.json' not in _dcode,
   'weekly/display code must consume guarded v2 artifacts only')
_ci=json.load(open(B+r'\data\members\corpus_integrity_report.json',encoding='utf-8'))
_cidb=_ci.get('database') or {}; _cic=_ci.get('accepted_model') or {}; _ciq=_ci.get('quarantined_chrg') or {}
ck('corpus integrity gate protects accepted scores and verifies repaired CHRG attribution',
   _ci.get('schema')=='corpus-integrity-v1' and _ci.get('status')=='PASS' and
   not _ci.get('hard_errors') and _cic.get('cache_exact_match') is True and _cic.get('cache_content_match') is True and
   _cidb.get('exact_duplicate_surplus')==0 and _ciq.get('enters_default_big5') is False and
   (_ciq.get('service_chamber_mismatch_rows',0)+_ciq.get('no_member_service_rows',0))==0 and
   _ciq.get('repair_status')=='source-rebuilt and verified',
   'status=%s hard=%s cache=%s invalid_chrg=%d'%
   (_ci.get('status'),_ci.get('hard_errors'),_cic.get('cache_exact_match'),
    _ciq.get('service_chamber_mismatch_rows',0)+_ciq.get('no_member_service_rows',0)))
_bpanel=json.load(open(B+r'\behavior_panel.json',encoding='utf-8'))
_bdisp=json.load(open(B+r'\behavioral_dispositions.json',encoding='utf-8'))
ck('corrupt legacy CHRG hearing counts cannot re-enter behavioral diligence',
   _bdisp.get('schema')=='behavioral-dispositions-v2' and _bdisp.get('status')=='PASS' and
   'chrg_hearings' not in (_bpanel.get('outcomes') or []) and
   all('chrg_hearings' not in (row or {}) for row in (_bpanel.get('members') or {}).values()) and
   'diligence=roll-call participation+LEP follow-through' in (_bdisp.get('method') or '') and
   'diligence=roll-call participation+LEP follow-through+CHRG' not in (_bdisp.get('method') or '') and
   ((_bdisp.get('integrity') or {}).get('diligence') or {}).get('legacy_chrg_enters_diligence') is False and
   (_bdisp.get('coverage') or {}).get('diligence_z',0)>=525 and
   len(((_bdisp.get('inputs_sha256') or {}).get('data/members/lep_criterion.json') or ''))==64,
   'behavior panel/disposition method must keep quarantined CHRG out of C')
# voice-ES display artifact RETIRED 2026-07-12 (user): the +0.208 look was gender-confounded (locked retest
# partial r -0.03). The desk shows behavioral reactivity on the role radar and independently ranked
# Blue/Pink N=-ES on the text radar; no voice-ES schema check remains. Guard against a silent
# regression to the retired artifact:
ck('voice-ES retired from display (b5_es_voice not re-wired into the desk)',
   'b5_es_voice' not in open(FE+r'\surfaces\make_member_desk.py',encoding='utf-8').read().split('RETIRED')[-1],
   'desk must not consume the gender-confounded voice-ES artifact')
_desk_tpl=open(FE+r'\surfaces\templates\make_member_desk_00.tpl',encoding='utf-8').read()
ck('desk presents the accepted five-axis display under canonical Big-Five names',
   "const B5AX=['Sociability','Openness','Diligence','Cooperativeness','Neuroticism'];" in _desk_tpl and
   'Gold-validated: gold-anchored axes use the OOF Senate-gold composites' in _desk_tpl and
   all(len((m.get('psy') or {}).get('big5_pct') or [])==5 for m in _bun.get('members',[])
       if (m.get('psy') or {}).get('big5_pct')),
   'accepted display must carry five canonically named axes and disclose the E scope')
_comm=json.load(open(B+r'\communication_style.json',encoding='utf-8'))
_csrc=_comm.get('members') or {}
_gdoc=json.load(open(B+r'\data\stance\big5_gold_anchored.json',encoding='utf-8')); _gtraits=_gdoc.get('traits') or {}
_egsrc=((_gtraits.get('E') or {}).get('members') or {})
_gorder=('E','O','C','A','N')
_cbad=0; _egbad=0; _gbad=0
for _row in _bun.get('members',[]):
    _name=_row.get('name'); _psy=_row.get('psy') or {}; _src=_csrc.get(_name) or {}
    if _src:
        _got=_psy.get('communication_style') or {}
        _adj=(_src.get('public_exposure') or {}).get('role_adjusted_percentile')
        _act=((_src.get('facets') or {}).get('activity') or {}).get('percentile')
        _cbad += int((_got.get('public_exposure') or {}).get('role_adjusted_percentile') != _adj)
        _cbad += int(_psy.get('e_public_proxy_pct') != _adj)
        # Public exposure propagates independently.  The Big-Five E slot is the OOF-gated
        # four-channel composite where both primary channels exist and remains missing otherwise.
        if _psy.get('big5_pct'):
            _egbad += int(_psy['big5_pct'][0] != (_egsrc.get(_name) or {}).get('percentile'))
            _gbad += sum(int(_psy['big5_pct'][i] != (((_gtraits.get(trait) or {}).get('members') or {}).get(_name) or {}).get('percentile'))
                         for i,trait in enumerate(_gorder))
        _cbad += int(((_psy.get('behav_pct') or [None,None,None,None])[3]) != _act)
_mike=_csrc.get('Mike Johnson') or {}; _mexp=_mike.get('public_exposure') or {}
ck('observable communication facets and role-adjusted exposure propagate as separate context',
   _comm.get('schema')=='communication-style-v1' and len(_csrc)==533 and _cbad==0 and _egbad==0 and _gbad==0 and
   len(_egsrc)>=450 and
   _mexp.get('formal_role')=='Speaker of the House' and
   _mexp.get('role_adjusted_percentile',1)<_mexp.get('observed_visibility_percentile',0) and
    all(s in _desk_tpl for s in ('Observed communication style','Public exposure context','mean words/turn','Warmth cues',
                                 'positive-energy/excitement cues',
                                 )),
   '%d style members, %d style errors, %d E errors, %d broad-trait errors; Mike raw %.3f adjusted %.3f'%
   (len(_csrc),_cbad,_egbad,_gbad,_mexp.get('observed_visibility_percentile',-1),_mexp.get('role_adjusted_percentile',-1)))
_goldn=((_gtraits.get('N') or {}).get('members') or {})
_golddesk={m.get('name'):((m.get('psy') or {}).get('big5_pct') or [None]*5)[4]
          for m in _bun.get('members',[]) if ((m.get('psy') or {}).get('big5_pct') or [None]*5)[4] is not None}
_profm=(json.load(open(B+r'\data\members\profile_basis_v2.json',encoding='utf-8')).get('members') or {})
_narrow_n=sum(1 for d in _profm.values() if ((d.get('personality_facets') or {}).get('Neuroticism_affective_content') or {}).get('z') is not None)
ck('broad one-panel Neuroticism occupies the canonical slot while affective N remains separate',
   len(_golddesk)==len(_goldn)>=450 and all(_golddesk.get(m)==d.get('percentile') for m,d in _goldn.items()) and
   _narrow_n>=450 and 'Neuroticism' in _desk_tpl,
   '%d broad-N source members, %d desk N values, %d retained affective-N facets'%(len(_goldn),len(_golddesk),_narrow_n))
_narr2=os.path.exists(B+r'\member_narratives_v2.json')
_nbad=sum(1 for _row in _bun.get('members',[]) if ((_row.get('psy') or {}).get('narrative') is not None or
          (_row.get('psy') or {}).get('narrative_status')!='withheld_pending_guarded_local_rebuild_and_user_approval')) if not _narr2 else 0
ck('legacy E/audio narratives fail closed pending guarded local rebuild',
   (_narr2 or _nbad==0) and 'member_narratives_v2.json' in _dcode and
   'withheld_pending_guarded_local_rebuild_and_user_approval' in _dcode,
   ('guarded v2 narrative artifact present' if _narr2 else '%d desk rows leaked legacy prose or missed withholding status'%_nbad))
ck('desk suppresses synthetic/default-only dominance values and keeps behavior display N-free',
    all(s in _desk_tpl for s in ('Conversational dominance',
                                 'insufficient observed interview coverage')) and
   'Affective N' not in _desk_tpl and
   "const BLAB=['Cross-party cooperation','Follow-through','Conversational dominance','Media appearances','Issue breadth'];" in _desk_tpl and
   "if _r2.get('dominance_v2') is not None:" in
       open(FE+r'\surfaces\make_member_desk.py',encoding='utf-8').read(),
   'behavior display must expose coverage, plot five behavioral axes, and omit Affective N')
if os.path.exists(B+r'\data\members\behavioral_roles_v2_candidate.json') and os.path.exists(B+r'\member_desk_bundle.json'):
    _r2=json.load(open(B+r'\data\members\behavioral_roles_v2_candidate.json',encoding='utf-8')).get('members',{})
    _dpsy={m.get('name'):(m.get('psy') or {}) for m in _bun.get('members',[])}
    _r2bad=0
    for _name,_d in _dpsy.items():
        _src=_r2.get(_name)
        if _src:
            _r2bad += int(abs(float(_d.get('dominance_v2_z',999))-float(_src['dominance_v2']))>.000001)
            _r2bad += int(abs(float(_d.get('expressiveness_v2_z',999))-float(_src['expressiveness_v2']))>.000001)
            _r2bad += int((_d.get('behavior_scope') or {}).get('interaction_turns')!=int(_src['qualifying_turns']))
            _r2bad += int(_d.get('expressiveness_v2_pct') is None or (_d.get('behav_pct') or [None,None,None])[2] is None)
        else:
            _r2bad += int(any(k in _d for k in ('dominance_v2_z','expressiveness_v2_z','expressiveness_v2_pct')))
    _producer=open(FE+r'\surfaces\make_member_desk.py',encoding='utf-8').read()
    _engine_refs=[]
    # The locked study producer, audio validator, and OOF-gold display/profile estimator are research
    # provenance, not influence/rank engines. The legislation shadow wrappers
    # copy this artifact only as an immutable input to their isolated
    # dashboard/member-desk branch; they do not expose it to rank computation.
    # Any additional Python consumer is a wiring regression.
    _r2_allowed={'make_member_desk.py','audit_full.py','dashboard_integrity.py','b5_narratives.py','communication_style.py',
                 'b5_profile_basis.py','b5_gold_multitrait_validate.py',
                 'behavioral_roles_v2_study.py','interview_audio_e_validate.py',
                 'run_legislation_dashboard_member_desk.py',
                 'run_legislation_shadow_recompute.py'}
    for _scan_dir in (FE,os.path.join(FE,'members'),os.path.join(FE,'legislation')):
        for _fn in os.listdir(_scan_dir):
            if not _fn.endswith('.py') or _fn in _r2_allowed: continue
            if 'behavioral_roles_v2_candidate.json' in open(os.path.join(_scan_dir,_fn),encoding='utf-8',errors='ignore').read(): _engine_refs.append(_fn)
    ck('qualified conversation roles v2 propagate exactly to the desk and remain engine-isolated',
       len(_r2)>=250 and _r2bad==0 and not _engine_refs and
       all(s in _desk_tpl for s in ('Conversation dynamics','Verbal expressiveness')) and
       '_ROLES2_ART' in _producer,
       '%d source members, %d propagation errors, engine refs=%s'%(len(_r2),_r2bad,_engine_refs))
ck('desk preserves missing low-antagonism behavior and its selected-stratum scope',
   all(s in _desk_tpl for s in ('Low-antagonism A facet',
                                'insufficient interview coverage')),
   'low antagonism must remain missing outside the package-qualified stratum, never imputed as A')
_qproc=open(FE+r'\network\process_quorum.py',encoding='utf-8').read()
ck('Quorum member resolver refuses surname-only identity guesses',
   'Never infer identity from surname alone' in _qproc and 'if len(cands)==1: return cands[0]' not in _qproc,
   'same-surname outsiders must be segregated, not inserted into a legislator corpus')
# V-1 (audit 2026-07-08): vote_category_lean.json is a manually-maintained FROZEN input -- its producer
# vote_by_category.py is in no orchestrator and its cleaned-roll-call CSV is an external file with no in-repo
# producer. Surface the frozen-ness: warn if the vote leans fall >30d behind the newest bills archive, so a
# silently-stale Path V / per_category_stance signal is visible instead of aging unnoticed.
_vcl=B+r'\vote_category_lean.json'
if os.path.exists(_vcl):
    import glob as _g2
    _arch=sorted(_g2.glob(B+r'\congress_bills_archive_*.xlsx'))
    _stale=bool(_arch) and os.path.getmtime(_vcl) < os.path.getmtime(_arch[-1])-30*86400
    ck('vote_category_lean.json fresh vs bills archive (manual frozen input; see V-1)', not _stale,
       'vote leans %s'%(('STALE >30d behind '+os.path.basename(_arch[-1])) if _stale else 'current'))
# say-do corroboration gate coherence (2026-07-07): every capped cell must DISPLAY Ally with its raw
# tier preserved as Champion/Defender; every proven cell must BE Champion/Defender in the matrix.
_g=(CS.get('gate') or {})
if _g:
    _cap=_g.get('capped') or {}; _prov=_g.get('proven') or {}
    _b1=sum(1 for m,cc in _cap.items() for c,raw in cc.items()
            if MAT.get(m,{}).get(c)!='Ally' or raw not in ('Champion','Defender'))
    _b2=sum(1 for m,cs_ in _prov.items() for c in cs_ if MAT.get(m,{}).get(c) not in ('Champion','Defender'))
    ck('say-do gate coherent (capped->Ally w/ raw C/D preserved; proven are C/D cells)', _b1==0 and _b2==0,
       '%d capped cells / %d proven cells; %d+%d incoherent'%(sum(len(v) for v in _cap.values()),
       sum(len(v) for v in _prov.values()),_b1,_b2))
    # Only authored/original-cosponsored legislation and aligned recorded votes qualify as action.
    _cor=_g.get('corrob') or {}
    _b3=sum(1 for m,cc in _cor.items() for c,pth in cc.items()
            if MAT.get(m,{}).get(c) not in ('Champion','Defender') or c in (_cap.get(m) or {})
            or pth not in ('A','V'))
    _cbp=collections.Counter(pth for cc in _cor.values() for pth in cc.values())
    ck('say-do corroboration paths coherent (only A/V cells survive as C/D)', _b3==0 and _cbp['J']==0,
       'corroborated A%d V%d; %d incoherent'%(_cbp['A'],_cbp['V'],_b3))
    _oc=_g.get('overall_capped') or {}; _thr=float(_g.get('overall_delivery_threshold') or 2/3)
    _b4=sum(1 for m,d in _oc.items() if d.get('raw')!='Champion' or d.get('final')!='Ally'
            or not d.get('evidence_available') or not set(d.get('cap_reasons') or []).issubset(
                {'insufficient_top_tier_breadth','top_tier_action_coverage','high_volume_low_followthrough'})
            or not d.get('cap_reasons'))
    ck('overall say-do caps have an auditable fixed-rule reason',_b4==0,
       '%d overall caps, %d incoherent'%(len(_oc),_b4))
    _leg_ready_path=B+r'\legislation_shadow_readiness.json'
    _leg_ready=(json.load(open(_leg_ready_path,encoding='utf-8'))
                if os.path.exists(_leg_ready_path) else {})
    _cfpb_routes=_leg_ready.get('cfpb_route_checks') or {}
    ck('CFPB direction contract is unchanged (restrain positive; reach negative)',
       bool(_cfpb_routes) and all(_cfpb_routes.values()),
       _cfpb_routes or 'legislation_shadow_readiness.json missing')
    import moneyball_core as _mc
    from tier_util import (legislative_press_tiers as _lpt,
                           load_tier_weights as _ltw,
                           load_tier_positions as _ltp,
                           load_tier_overrides as _lto,
                           load_primacy_factors as _lpf,
                           load_primacy_cu_credit as _lpc,
                           PRIMACY_BOARD as _tu_board,
                           PRIMACY_MIN_CUPRIM as _tu_floor)
    # Board basis floor. 'factor' is the production default (2026-07-29
    # directive); 'legxfactor' is the revert switch. Anything else means the
    # PRIMACY_BOARD dial was typo'd, which would otherwise reorder the whole
    # board silently.
    ck('primacy board basis is a recognized mode',
       _tu_board in ('factor','legxfactor'),_tu_board)
    # Evidence floor ("purity ranks, evidence qualifies", user-ratified
    # 2026-07-29): production default 0.75 cu_primary credit to hold a tier
    # seat; 0 disables. Import already rejects negatives -- this pins the
    # ratified production value so a drifted env cannot silently reshape
    # Tier 2/3.
    ck('evidence floor holds the ratified production value',
       _tu_floor==0.75 or (_tu_board!='factor' and _tu_floor>=0),
       'PRIMACY_MIN_CUPRIM=%r'%_tu_floor)
    _tier_weights=_ltw()
    ck('legacy display-tier policy pair remains 84%% Legislative / 16%% Press',
       all(abs(a-b)<1e-12 for a,b in zip(_tier_weights,(0.84,0.16))),str(_tier_weights))
    _tier_positions=_ltp()
    _adjudicated_members={
        'Brad Sherman','Bill Huizenga','Young Kim',
        'Catherine Cortez Masto','Tim Scott'}
    ck('primacy board has no fixed display positions',
       _tier_positions=={},_tier_positions)
    _tier,_tier_rank,_tier_detail=_lpt(list(MAT),with_detail=True)
    ck('primacy board has no fixed tier placements',
       _lto()=={},_lto())
    # Enforcement, not just configuration: under the factor board no member
    # below the evidence floor may hold a tier seat, and all 85 seats must
    # still be filled by qualified members.
    _cuc=_lpc(B)
    _subfloor=sorted(m for m,t in _tier.items()
                     if t>0 and _tu_board=='factor' and _tu_floor>0
                     and float(_cuc.get(m,0.0))<_tu_floor)
    ck('no tiered member sits below the evidence floor',
       not _subfloor and sum(1 for t in _tier.values() if t>0)==85,
       'sub-floor tiered=%s seats=%d'%(_subfloor[:6],
                                       sum(1 for t in _tier.values() if t>0)))
    _engine_tier,_engine_tier_rank,_engine_priority=_mc.primacy_engine_state(list(MAT))
    ck('allocation engine consumes primacy tiers rather than rank scores',
       _engine_tier==_tier
       and _engine_tier_rank==_tier_rank
       and all(abs(_engine_priority[m]-_mc.primacy_engine_value(_tier[m]))<1e-12
               for m in MAT),
       str(_mc.PRIMACY_ENGINE_LEVELS))
    _primacy=_lpf(B)
    ck('primacy factors cover the full 533-member rank roster',
       len(_primacy)==len(MAT)==533,
       '%d factors vs %d stance members'%(len(_primacy),len(MAT)))
    from engine import primacy_factors as _pf
    _adjudicated=_pf.load_adjudicated_credits(
        os.path.join(B,_pf.ADJUDICATION_LEDGER))
    _primacy_doc=json.load(open(
        os.path.join(B,'data', 'engine', 'primacy_factors.json'),encoding='utf-8'))
    ck('five uniform sponsor-equivalent CU adjudication credits enter factor math',
       set(_adjudicated)==_adjudicated_members
       and all(row.get('denominator_delta')==1.0
               and row.get('numerator_delta')==1.0
               and len(row.get('events') or [])==1
               for row in _adjudicated.values())
       and (_primacy_doc.get('_constants') or {}).get(
           'active_adjudication_events')==5
       and all((_primacy_doc.get('members') or {}).get(member,{}).get(
                   'adjudicated_credit')==1.0
               for member in _adjudicated_members),
       str({member:(_primacy_doc.get('members') or {}).get(member,{})
            .get('adjudicated_events')
            for member in sorted(_adjudicated_members)}))
    _math_review_path=os.path.join(
        B,'audits','legfix5','production_promotion',
        'greenlight-20260728a-primacy-v1',
        'primacy_math_opus5_review.json')
    _math_review=(json.load(open(_math_review_path,encoding='utf-8'))
                  if os.path.exists(_math_review_path) else {})
    _math_gates=_math_review.get('gates') or {}
    ck('local Opus-5-method primacy math receipt closes every gate',
       _math_review.get('status')=='PASS_SCORE_BASED_ADJUDICATION'
       and bool(_math_gates) and all(_math_gates.values())
       and (_math_review.get('execution') or {}).get(
           'provider_call_performed') is False,
       str([name for name,passed in _math_gates.items() if not passed])
       if _math_gates else 'receipt missing')
    _culling_path=B+r'\data\legislation\legislation_culling_status.json'
    _culling=(json.load(open(_culling_path,encoding='utf-8'))
              if os.path.exists(_culling_path) else {})
    _cc=_culling.get('counts') or {}
    ck('culling surface reports the hash-bound production promotion state',
       _culling.get('status')=='PASS_PRODUCTION_CULLING_INTEGRATED'
       and _culling.get('phase')=='production'
       and _culling.get('promotion_performed') is True
       and _culling.get('production_workbook_written') is True
       and _culling.get('rank_written') is True
       and (_culling.get('promotion_receipt') or {}).get('status')
           =='PASS_PRODUCTION_INPUTS_PROMOTED',
       _culling.get('status') or 'legislation_culling_status.json missing')
    ck('culling ledger closes all scoring, held, and excluded decisions',
       _cc.get('scoring_ready_candidates')==60
       and _cc.get('scoring_ready_category_edges')==74
       and _cc.get('held_decision_units')==1184
       and (_cc.get('integration_state_counts') or {}).get(
           'excluded_false_positive')==2908,
       _cc or 'legislation_culling_status.json missing')
    _tag_path=B+r'\bill_beneficiary_tags.ndjson'
    _tag_rows=([json.loads(line) for line in
                open(_tag_path,encoding='utf-8') if line.strip()]
               if os.path.exists(_tag_path) else [])
    _tag_count=len(_tag_rows)
    _tag_keys={'%s|%s'%(str(row.get('bill') or '').strip(),
                        str(row.get('congress') or '').strip())
               for row in _tag_rows}
    # 2026-08-10 bank-anchor expansion 633->640 -> 644 crypto/fintech association sweep
    ck('merged beneficiary universe contains exactly 644 classified acts',
       _tag_count==644,'%d tags'%_tag_count)
    _desk_path=B+r'\member_desk_bundle.json'
    _desk_for_closure=(json.load(open(_desk_path,encoding='utf-8'))
                       if os.path.exists(_desk_path) else {})
    _desk_acts={'%s|%s'%(str(bill.get('bill') or '').strip(),
                         str(bill.get('congress') or '').strip())
                for member in _desk_for_closure.get('members',[])
                for bill in (member.get('bills') or [])}
    ck('member desk and beneficiary universe close on the same 644 acts',
       len(_desk_acts)==644 and _desk_acts==_tag_keys,
       '%d desk acts vs %d tags'%(len(_desk_acts),len(_tag_keys)))
    _membership_path=B+r'\data\members\primacy_act_memberships.ndjson'
    _membership_rows=(
        [json.loads(line) for line in
         open(_membership_path,encoding='utf-8') if line.strip()]
        if os.path.exists(_membership_path) else [])
    _membership_keys={
        '%s|%s'%(str(row.get('bill') or '').strip(),
                 str(row.get('congress') or '').strip())
        for row in _membership_rows
    }
    ck('primacy amendment sidecar is 117 unique sponsor-only memberships',
       len(_membership_rows)==len(_membership_keys)==117
       and all('Amdt.' in str(row.get('bill') or '')
               and row.get('role')=='sponsor'
               for row in _membership_rows),
       '%d rows / %d unique acts'%(
           len(_membership_rows),len(_membership_keys)))
    _registry_receipt_path=B+r'\data\legislation\registry_anchor_reconciliation.json'
    _registry_receipt=(
        json.load(open(_registry_receipt_path,encoding='utf-8'))
        if os.path.exists(_registry_receipt_path) else {})
    _registry_counts=_registry_receipt.get('counts') or {}
    _registry_gates=_registry_receipt.get('gates') or {}
    _registry_membership_path=B+r'\data\members\registry_anchor_memberships.ndjson'
    _registry_memberships=(
        [json.loads(line) for line in
         open(_registry_membership_path,encoding='utf-8') if line.strip()]
        if os.path.exists(_registry_membership_path) else [])
    _registry_keys={
        '%s|%s'%(str(row.get('bill') or '').strip(),
                 str(row.get('congress') or '').strip())
        for row in _registry_memberships
    }
    ck('all 55 registry anchors are classified after six official-source repairs',
       _registry_counts.get('declared_registry_anchors')==55
       and _registry_counts.get('missing_before')==6
       and _registry_counts.get('missing_after')==0
       and _registry_counts.get('recovered_anchors')==6
       and _registry_counts.get('classified_act_universe')==644
       and _registry_receipt.get('missing_after')==[],
       _registry_counts or 'registry reconciliation missing')
    ck('registry sidecar resolves 73 official member links across six acts',
       len(_registry_memberships)==73 and len(_registry_keys)==6,
       '%d rows / %d acts'%(
           len(_registry_memberships),len(_registry_keys)))
    ck('every registry reconciliation gate is closed',
       bool(_registry_gates) and all(_registry_gates.values()),
       _registry_gates or 'registry gates missing')
    _barr=_tier_detail.get('Andy Barr') or {}
    _hill=_tier_detail.get('French Hill') or {}
    ck('net conduct ledger reproduces Barr and Hill with both sides visible',
       abs(_barr.get('gross_severity',-1)-7.25)<1e-12
       and abs(_barr.get('delivered_credit',-1)-0.0)<1e-12
       and abs(_barr.get('net_severity',-1)-7.25)<1e-12
       and _barr.get('tier_ceiling')==3
       and abs(_hill.get('gross_severity',-1)-8.0)<1e-12
       and abs(_hill.get('delivered_credit',-1)-3.0)<1e-12
       and abs(_hill.get('net_severity',-1)-5.0)<1e-12
       and _hill.get('tier_ceiling')==2,
       'Barr %s; Hill %s'%(_barr,_hill))
    _grievance_detail=(
        (_barr.get('grievances') or [])
        + (_hill.get('grievances') or []))
    ck('grievance evidence carries a human reason and cure for every event',
       len(_grievance_detail)==4
       and all(str(row.get('reason') or '').strip()
               and str(row.get('cure') or '').strip()
               for row in _grievance_detail),
       '%d detailed grievance events'%len(_grievance_detail))
    _hill_deliveries=_hill.get('deliveries') or []
    ck('delivery evidence names the bill, instrument, and positive credit',
       {row.get('bill') for row in _hill_deliveries}
           >= {'H.R.975','H.R.4437','H.R.4544'}
       and all(str(row.get('instrument') or '').strip()
               and float(row.get('credit') or 0)>0
               for row in _hill_deliveries),
       '%d detailed delivery events'%len(_hill_deliveries))
    ck('conduct ceilings keep constrained members out of Tier 1',
       _tier.get('Andy Barr') not in (1,2)
       and _tier.get('French Hill')!=1,
       'Barr tier=%s rank=%s; Hill tier=%s rank=%s'%
       (_tier.get('Andy Barr'),_tier_rank.get('Andy Barr'),
        _tier.get('French Hill'),_tier_rank.get('French Hill')))
    _display={m:_mc.overall_display_tier(MAT.get(m,{}),_tier.get(m,0),m,_oc) for m in MAT}
    _champ={m for m,v in _display.items() if v=='Champion'}
    _tier1={m for m,t in _tier.items() if t==1}
    ck('overall Champion crown exactly equals primacy Tier 1',
       _champ==_tier1 and len(_champ)==13,
       '%d crowns vs %d Tier-1; missing %s; extra %s'%(
           len(_champ),len(_tier1),sorted(_tier1-_champ),sorted(_champ-_tier1)))
    # RE-PINNED 2026-07-29 (second revision, uncertainty penalty K=0.25):
    # the factor-only directive first re-seated Huizenga and Kim in Tier 2;
    # the same-day "certainty outranks estimation" penalty then re-seated Kim
    # in Tier 1 at slot 13 on her full-evidence measurement (no ledger
    # re-size) and left Huizenga at 16 -- his gap is structural (1.0 credits
    # of bank/crypto ballast), not measurement noise. Seats are pinned
    # EXACTLY rather than relaxed, so any further drift still fails. Under
    # PRIMACY_BOARD=legxfactor the original all-five-Tier-1 contract holds.
    # THIRD revision 2026-07-29: amendments now carry measure weight 0.5
    # ("0.5 applies broad"), which re-seats Kim to Tier 2 at 14 and lifts
    # Crapo into Tier 1 at 13. Kim's Tier 1 seat under the uncertainty
    # penalty was a 0.002 margin; halving 173 amendments moved enough of the
    # roster to close it. Accepted with the directive.
    # 2026-08-10 re-ratified (Paul: independence ruling - other industries
    # are not held to the CU gold standard; "if it checks out it checks
    # out"): tier-membership guards under the 640-act universe (valence +
    # CU axis + bank-anchor expansion). Exact-rank pins retired with the
    # 633-act era; limits are now the tier outer boundaries (T1<=13,
    # T2<=37), so lawful cross-industry data changes no longer fire this
    # check while tier eviction still does.
    # 2026-08-12 re-pinned (Paul: "Huizenga dropping out is fine"): the five
    # ACU-SUPPORT crypto acts are now priced by adjudicated CU tensor cells
    # rather than the blanket acu_position floor, which retired the 0.5
    # shelter those acts gave his portfolio. Tier 2 is the ratified seat; the
    # T2 outer boundary applies.
    _seating_expected=({'Brad Sherman':1,'Catherine Cortez Masto':2,
                        'Tim Scott':1,'Bill Huizenga':2,'Young Kim':2}
                       if _tu_board=='factor' else
                       {m:1 for m in _adjudicated_members})
    _seating_limit=({'Brad Sherman':13,'Catherine Cortez Masto':37,
                     'Tim Scott':13,'Bill Huizenga':37,'Young Kim':37}
                    if _tu_board=='factor' else
                    {m:(12 if m=='Tim Scott' else 13)
                     for m in _adjudicated_members})
    ck('adjudicated five hold their ratified seats under the active board basis',
       all(_tier.get(member)==_seating_expected[member]
           and _tier_rank.get(member)<=_seating_limit[member]
           for member in _adjudicated_members),
       'basis=%s %s'%(_tu_board,
       str({member:(_tier.get(member),_tier_rank.get(member)) for member in
            sorted(_adjudicated_members)})))
    # Ratified constraint recorded with K=0.25 ("Vargas needs to be in"): the
    # penalty strength was chosen over 0.5-1.0 specifically so Juan Vargas --
    # thin in volume (6.25 credits), not in CU purity (4.0 cu_primary) --
    # keeps a Tier 1 seat. Guards against future K drift evicting him.
    ck('Vargas holds Tier 1 under the uncertainty-penalized factor board',
       _tu_board!='factor'
       or (_tier.get('Juan Vargas')==1 and _tier_rank.get('Juan Vargas')<=13),
       '(%s, %s)'%(_tier.get('Juan Vargas'),_tier_rank.get('Juan Vargas')))
    # The artifact must carry the ratified penalty constant: a stale
    # pre-penalty primacy_factors.json silently degrades the board to the
    # K=0 ranking, so pin the constant here to fail loud instead.
    _pf_doc_k=json.load(open(os.path.join(B,'data', 'engine', 'primacy_factors.json'),
                             encoding='utf-8'))
    _pf_k=(_pf_doc_k.get('_constants') or {}).get('uncertainty_penalty_k')
    ck('primacy artifact carries the ratified uncertainty penalty K=0.25',
       _tu_board!='factor' or _pf_k==0.25,
       'artifact K=%r'%(_pf_k,))
    # Amendments carry measure weight 0.5 on the tier board (user-ratified
    # 2026-07-29). Pin the constant so a drifted dial cannot silently restore
    # full bill weight to 173 acts.
    _pf_amw=(_pf_doc_k.get('_constants') or {}).get('amendment_measure_weight')
    ck('primacy artifact carries the ratified amendment weight 0.5',
       _pf_amw==0.5,'artifact amendment weight=%r'%(_pf_amw,))
    # The detector fails OPEN to bill weight, so a regex miss would silently
    # score an amendment at 1.0. Assert that every act whose id even LOOKS
    # like an amendment is actually recognized as one.
    import legislation_measure_key as _MK
    _desk_doc=json.load(open(os.path.join(B,'member_desk_bundle.json'),
                             encoding='utf-8'))
    _champions_boards_path=os.path.join(
        os.environ['LOCALAPPDATA'],'Melange','lda_rebuild','derived',
        'champions_v2','boards.json')
    _champions_boards_missing=not os.path.exists(_champions_boards_path)
    _champions_boards=({} if _champions_boards_missing else
                       json.load(open(_champions_boards_path,encoding='utf-8')))
    _champion_threat_rows=[(sector,row) for sector,boards in
                           (_champions_boards.get('sectors') or {}).items()
                           for row in (boards.get('threats') or [])]
    _tier1_threat_rows=[(sector,row) for sector,row in _champion_threat_rows
                        if int(row.get('tier') or 0)==1]
    ck('champions boards obey champion-cannot-be-threat law',
       not _tier1_threat_rows and not _champions_boards_missing,
       'boards.json missing -- run champions_v2' if _champions_boards_missing
       else ', '.join('%s: %s (%s)'%(sector,row.get('name'),row.get('bioguide'))
                      for sector,row in _tier1_threat_rows))
    _live_crowns={str((row.get('bio') or {}).get('bioguide') or '').strip():
                  row.get('name') for row in (_desk_doc.get('members') or [])
                  if int(row.get('tier') or 0)==1
                  and str((row.get('bio') or {}).get('bioguide') or '').strip()}
    _stale_crown_threats=[(sector,row) for sector,row in _champion_threat_rows
                          if row.get('bioguide') in _live_crowns]
    ck('champions threat boards fresh vs live crown set',
       not _stale_crown_threats and not _champions_boards_missing,
       'boards.json missing -- run champions_v2' if _champions_boards_missing
       else '%s; champions_v2 runs before make_member_desk -- rerun chain'%
       ', '.join('%s: %s (%s)'%(
           sector,_live_crowns.get(row.get('bioguide')),row.get('bioguide'))
           for sector,row in _stale_crown_threats))
    _champion_seat_rows=[(sector,row) for sector,boards in
                         (_champions_boards.get('sectors') or {}).items()
                         for row in (boards.get('champions') or [])]
    _threat_seat_bios={row.get('bioguide') for _s,row in _champion_threat_rows}
    _cross_violations=[(sector,row) for sector,row in _champion_seat_rows
                       if row.get('bioguide') in _threat_seat_bios]
    ck('champions boards obey cross-sector exclusivity (champion seat yields)',
       not _cross_violations and not _champions_boards_missing,
       'boards.json missing -- run champions_v2' if _champions_boards_missing
       else ', '.join('%s: %s (%s)'%(sector,row.get('name'),row.get('bioguide'))
                      for sector,row in _cross_violations))
    _amd_ids={str(_x.get('bill') or '')
              for _rec in (_desk_doc.get('members') or [])
              for _x in (_rec.get('bills') or [])
              if 'amdt' in str(_x.get('bill') or '').lower()}
    _amd_missed=sorted(i for i in _amd_ids if not _MK.is_display_amendment(i))
    ck('every amendment-shaped act id is recognized by the measure-weight bridge',
       not _amd_missed,'unrecognized=%s'%_amd_missed[:8])
    if os.path.exists(B+r'\dashboard_data.json'):
        _dash=json.load(open(B+r'\dashboard_data.json',encoding='utf-8'))
        _tc={'Champion':'C','Defender':'F','Ally':'A','Highly Persuadable':'H','Persuadable':'P',
             'Lean Persuadable':'G','True Neutral':'T','Lean Skeptic':'L','Neutral':'N',
             'Skeptic':'S','Highly Skeptical':'K','Opponent':'O','—':'-'}
        _expected_gen=''.join(_tc.get(_display.get(m,'—'),'-') for m in _dash.get('name',[]))
        ck('dashboard overall labels exactly match shared crown law',_dash.get('gen')==_expected_gen,
           'dashboard and kernel label strings differ')
        _dash_detail=dict(zip(_dash.get('name',[]),_dash.get('tierdetail',[])))
        ck('dashboard carries aligned grievance and delivery tier evidence',
           _dash_detail.get('Andy Barr')==_barr
           and _dash_detail.get('French Hill')==_hill,
           'tierdetail missing or not member-aligned')
        _dash_engine=_dash.get('engine') or {}
        _dash_const=_dash_engine.get('const') or {}
        _dash_members=_dash.get('name') or []
        ck('dashboard scenario engine carries primacy tier priority only',
           _dash_engine.get('tierp')==
               [round(_engine_priority[m],6) for m in _dash_members]
           and 'sco' not in _dash_engine
           and 'smax' not in _dash_const,
           'tierp=%d sco=%s smax=%s'%(
               len(_dash_engine.get('tierp') or []),
               'sco' in _dash_engine,
               'smax' in _dash_const))
    if os.path.exists(B+r'\member_desk_bundle.json'):
        _desk=json.load(open(B+r'\member_desk_bundle.json',encoding='utf-8'))
        _desk_detail={row.get('name'):row.get('tier_detail')
                      for row in _desk.get('members',[])}
        ck('member desk carries aligned grievance and delivery tier evidence',
           _desk_detail.get('Andy Barr')==_barr
           and _desk_detail.get('French Hill')==_hill,
           'tier_detail missing or not member-aligned')
    _dashboard_html_path=B+r'\Moneyball Dashboard.html'
    _desk_html_path=B+r'\cu-member-desk.html'
    _dashboard_html=(open(_dashboard_html_path,encoding='utf-8').read()
                     if os.path.exists(_dashboard_html_path) else '')
    _desk_html=(open(_desk_html_path,encoding='utf-8').read()
                if os.path.exists(_desk_html_path) else '')
    ck('dashboard HTML renders conduct reason and cure',
       'Cure:' in _dashboard_html
       and 'Refused credit-union asks' in _dashboard_html,
       'Moneyball Dashboard.html')
    ck('member desk HTML renders cure plus all three CU delivery anchors',
       'Cure:' in _desk_html
       and all(bill in _desk_html
               for bill in ('H.R.975','H.R.4437','H.R.4544')),
       'cu-member-desk.html')
    _wb_display={_ra_ws.cell(r,1).value:_ra_ws.cell(r,7).value for r in range(3,536) if _ra_ws.cell(r,1).value}
    ck('workbook stance column exactly matches shared crown law',
       all(_wb_display.get(m)==v for m,v in _display.items()) and len(_wb_display)==len(_display),
       'workbook %d vs matrix %d rows'%(len(_wb_display),len(_display)))
else:
    ck('say-do gate block present in category_stance.json', False, 'gate key missing — producer predates 2026-07-07 landing')

# ---- Speaker-identity layer candidates (audio-identity design 2026-07-16, user-directed build) ----
# Research-only .json.candidate artifacts consumed by NOTHING until user ratification. Every check
# below SKIPs (informational line, no P/F entry) while its artifact/source is absent, so the live
# audit is unaffected until the layer actually produces candidates.
print('\n[AUDIO-ID] Speaker-identity layer candidates (research-only)')
import glob as _gid
_idb=B+r'\b5_audio_enroll_bank.json.candidate'
_idt=B+r'\b5_audio_identity_thresholds.json.candidate'
_ida=sorted(_gid.glob(B+r'\b5_audio_attrib_*.json.candidate'))
# Main per-corpus verdict artifacts = everything that is not a rejects/mismatch ledger (suffix
# rule, not a corpus whitelist, so a new corpus can never silently escape the pairing check).
_ida_main=[f for f in _ida if not os.path.basename(f).endswith('_ledger.json.candidate')]
if os.path.exists(_idb) or _ida:
    _idbad=[]
    if os.path.exists(_idb):
        _idd=json.load(open(_idb,encoding='utf-8'))
        _idbank_ok=(_idd.get('schema')=='b5-audio-enroll-bank-v1' and
                    _idd.get('status')=='RESEARCH_CANDIDATE_NOT_CONSUMED' and
                    isinstance(_idd.get('theta_enroll'),(int,float)) and
                    len(str(_idd.get('thresholds_sha256') or ''))==64)
        if os.path.exists(_idt):
            import hashlib as _idh
            _hh=_idh.sha256(open(_idt,'rb').read()).hexdigest()
            _idbank_ok=(_idbank_ok and _idd.get('thresholds_sha256')==_hh)
        if not _idbank_ok:
            _idbad.append(os.path.basename(_idb))
    for _f in _ida:
        _idd=json.load(open(_f,encoding='utf-8'))
        _base=os.path.basename(_f)[:-len('.json.candidate')]
        _want=('b5-audio-attrib-rejects-v1' if _base.endswith('rejects_ledger') else
               'b5-audio-attrib-mismatch-v1' if _base.endswith('mismatch_ledger') else 'b5-audio-attrib-v1')
        if _idd.get('schema')!=_want or _idd.get('status')!='RESEARCH_CANDIDATE_NOT_CONSUMED':
            _idbad.append(os.path.basename(_f))
    ck('identity candidates schema-tagged + status RESEARCH_CANDIDATE_NOT_CONSUMED', not _idbad,
       'bad: %s'%_idbad)
    # No trait/narrative/refresh/receptivity consumer may reference the candidate artifacts.
    _idrefs=[]
    for _scan_dir in (FE,os.path.join(FE,'members')):
        for _fn in os.listdir(_scan_dir):
            if (_fn.endswith('.py') and (_fn.startswith(('b5_traits','big5_traits')) or
                _fn in ('b5_narratives.py','weekly_refresh.py','receptivity_layer.py'))
                and any(_s in open(os.path.join(_scan_dir,_fn),encoding='utf-8',errors='ignore').read()
                        for _s in ('b5_audio_enroll_bank','b5_audio_attrib_'))):
                _idrefs.append(_fn)
    ck('no engine/display consumer references identity candidates (traits/narratives/weekly/receptivity)',
       not _idrefs, str(_idrefs))
else:
    print('  SKIP  identity enroll-bank/attribution candidates absent (layer not yet run)')
if os.path.exists(_idt):
    _idd=json.load(open(_idt,encoding='utf-8'))
    _th=_idd.get('theta_accept')
    _te=_idd.get('theta_eer'); _tr=_idd.get('eer_rate')
    ck('identity thresholds calibrated (passing status, score/rate distinction, stop bars)',
       _idd.get('schema')=='b5-audio-identity-thresholds-v1' and
       _idd.get('status')=='CALIBRATION_PASS' and
       isinstance(_th,(int,float)) and isinstance(_te,(int,float)) and
       -1<=_te<=_th<=1 and isinstance(_idd.get('theta_enroll'),(int,float)) and
       -1<=_idd.get('theta_enroll')<=1 and isinstance(_tr,(int,float)) and 0<=_tr<=.05 and
       _idd.get('caption_mismatch_n',0)>=148 and
       _idd.get('caption_mismatch_rejection_rate',0)>=.99 and
       (_idd.get('sex_parity') or {}).get('threshold_was_retuned_for_parity') is False,
       'theta_accept=%s theta_eer=%s eer_rate=%s caption=%s/%s'%(
       _th,_te,_tr,_idd.get('caption_mismatch_n'),_idd.get('caption_mismatch_rejection_rate')))
else:
    print('  SKIP  identity thresholds candidate absent (calibration not yet run)')
if _ida_main:
    _idmiss=[os.path.basename(_f) for _f in _ida_main
             if not os.path.exists(os.path.join(os.path.dirname(_f),
                os.path.basename(_f)[:-len('.json.candidate')].rsplit('_',1)[0]+'_rejects_ledger.json.candidate'))]
    ck('every attribution candidate ships its rejects ledger beside it', not _idmiss,
       'missing ledger for: %s'%_idmiss)
else:
    print('  SKIP  no per-corpus attribution candidates (rejects-ledger pairing not applicable)')
_idsrc=FE+r'\members\b5_audio_attrib.py'
if os.path.exists(_idsrc):
    _badw=[_ln for _ln,_line in enumerate(open(_idsrc,encoding='utf-8',errors='ignore').read().splitlines(),1)
           if ('b5_audio_reviews_v2.json' in _line or 'b5_audio_descriptors_v2.json' in _line)
           and ('dump_atomic' in _line or 'open(' in _line)]
    ck('attribution module has no write reference to live review/descriptor artifacts', not _badw,
       'lines %s'%_badw)
else:
    print('  SKIP  b5_audio_attrib.py absent (identity layer not yet merged)')


# ============================ NETWORK CONFIGURATIONS ============================
# Five networks ship (cu | holistic | holistic_norm | bank_reg | crypto) and until now the audit
# tree contained no reference to NET_CONFIG at all -- they had kernel tests but no audit. These
# checks read ARTIFACTS ONLY; none runs a producer. Deliberately NOT asserted: "every config
# loads", because proving loadability means an openpyxl read of a 63,707-pair workbook per config
# for near-zero marginal catch over presence + selection agreement below.
print('\n[NET] Network configurations (registry / artifact / bundle tri-state)')
try:
    import glob as _g
    _prev_nc = os.environ.pop('NET_CONFIG', None)
    try:
        import moneyball_core as _MCA
        _NET = _MCA.NETWORKS
        _dpath = B + r'\dashboard_data.json'
        _D = json.load(open(_dpath, encoding='utf-8')) if os.path.exists(_dpath) else None
        ck('dashboard_data.json present for config audit', _D is not None, _dpath)

        if _D is not None:
            _pf = lambda w: w / (w + 6.0)

            # --- RANK 1: the PUBLISHED network must not drift when a sector is rebuilt -----------
            # Compared at the bundle's OWN precision: build_dashboard_bundle.py:532 rounds to 3.
            # Comparing at 4 reports ~87% spurious mismatches -- measured, and it looks exactly
            # like catastrophic drift while nothing is wrong.
            _idx = {n: i for i, n in enumerate(_D['name'])}
            _bun = {(min(a, b), max(a, b)): float(w) for a, b, w in _D['engine']['adj']}
            _adj = _MCA.load_network()
            _liv = {}
            for _u, _nb in _adj.items():
                if _u not in _idx: continue
                for _v, _w in _nb.items():
                    if _v not in _idx or _u == _v: continue
                    _k = (min(_idx[_u], _idx[_v]), max(_idx[_u], _idx[_v]))
                    _liv.setdefault(_k, float(_w))
            _mis = [k for k in _bun if k not in _liv or abs(_bun[k] - round(_liv[k], 3)) > 1e-9]
            ck('published CU adjacency in the bundle matches the CU artifact (no drift on rebuild)',
               set(_bun) == set(_liv) and not _mis,
               '%d bundle vs %d kernel dyads, %d weight mismatches at the bundle precision (r3)'
               % (len(_bun), len(_liv), len(_mis)))
            ck('published CU network records its provenance in the bundle',
               bool((_D.get('cu') or {}).get('graph')),
               'cu.graph absent -- the graph every published figure rests on is untraceable')

            # --- RANK 2: registry / artifact / bundle agree, or are HONESTLY absent -------------
            _FAM = {'holistic': 'holistic_network_*.xlsx',
                    'holistic_norm': 'holistic_network_N*.xlsx'}
            for _k in sorted(_NET):
                if _k == 'cu': continue
                _spec = _NET[_k]
                _pat = ('sector_network_%s_*.xlsx' % _spec['sector']) if _spec.get('family') == 'sector' \
                    else _FAM.get(_k, '')
                _on_disk = bool(_g.glob(os.path.join(B, _pat))) if _pat else False
                _in_bundle = _k in _D
                ck('NET_CONFIG=%s: artifact on disk and bundle presence agree' % _k,
                   _on_disk == _in_bundle,
                   'disk=%s (%s) bundle=%s -- a config the bundle advertises but cannot source, or '
                   'an artifact silently dropped from the surface' % (_on_disk, _pat, _in_bundle))

            # --- RANK 3: sector floors recompute; CU comparators are LIVE, not pinned -----------
            _cuw = sorted(float(w) for _, _, w in _D['engine']['adj'])
            _cu_live = {'cu_dyads': len(_cuw), 'cu_p_max': round(_pf(_cuw[-1]), 4),
                        'cu_dyads_p_ge_010': sum(1 for w in _cuw if _pf(w) >= 0.10)}
            for _k, _v in sorted(_NET.items()):
                if _v.get('family') != 'sector' or _k not in _D: continue
                _f = (_D[_k] or {}).get('engine_floor') or {}
                ck('sector %s quotes the LIVE CU baseline, not a frozen one' % _k,
                   all(_f.get(_c) == _cu_live[_c] for _c in _cu_live),
                   'pinned %s vs live %s -- the banner would keep quoting a stale CU network'
                   % ({c: _f.get(c) for c in _cu_live}, _cu_live))
                _aw = sorted(float(w) for _, _, w in (_D['engine'].get('adj_' + _k) or []))
                if _aw:
                    ck('sector %s engine_floor recomputes from its own adjacency' % _k,
                       _f.get('p_max') == round(_pf(_aw[-1]), 4)
                       and _f.get('dyads_p_ge_010') == sum(1 for w in _aw if _pf(w) >= 0.10),
                       'stated %s / %s vs recomputed %.4f / %d'
                       % (_f.get('p_max'), _f.get('dyads_p_ge_010'), _pf(_aw[-1]),
                          sum(1 for w in _aw if _pf(w) >= 0.10)))

            # --- RANK 4: family separation, measured on disk ------------------------------------
            # 24 lines across 17 files glob congress_cosponsor_network_*; three take max-by-mtime
            # with NO AUG and NO family filter, so ONE misnamed file silently re-points the press,
            # HCoI and stance layers. This is the check that catches that before they do.
            _cu_glob = set(map(os.path.basename, _g.glob(os.path.join(B, 'congress_cosponsor_network_*.xlsx'))))
            _other = set(map(os.path.basename, _g.glob(os.path.join(B, 'holistic_network_*.xlsx')))) | \
                     set(map(os.path.basename, _g.glob(os.path.join(B, 'sector_network_*.xlsx'))))
            ck('network artifact families are disjoint at the live root',
               _cu_glob.isdisjoint(_other), 'collisions: %s' % sorted(_cu_glob & _other))
            if _cu_glob:
                # The three unfiltered max-by-mtime consumers (press_engagement_layer, press_hcoi,
                # stance_dynamics) take the NEWEST file in this glob with no AUG and no family
                # filter. The hazard is therefore a stray file DISPLACING the published network by
                # mtime -- which is not hypothetical: an accidental import wrote a duplicate
                # holistic AUG artifact on 2026-08-19 and it immediately became the selected graph.
                # So the check that matters is "would they pick the graph the bundle published?".
                # (The earlier form asserted the picked name startswith the glob's own prefix, which
                # is tautologically true for anything the glob returned and could never fail -- a
                # negative control caught it.)
                _picked = os.path.basename(max(_g.glob(os.path.join(B, 'congress_cosponsor_network_*.xlsx')),
                                               key=os.path.getmtime))
                _published = (_D.get('cu') or {}).get('graph')
                ck('the CU file unfiltered max-by-mtime consumers would pick IS the published graph',
                   bool(_published) and _picked == _published,
                   'they would load %r but the bundle published %r -- a stray file in the CU glob '
                   'silently re-points the press, HCoI and stance layers' % (_picked, _published))

            # --- RANK 5: key-quad completeness, and the converse -------------------------------
            for _k in sorted(_NET):
                if _k == 'cu' or _k not in _D: continue
                _a = _D['engine'].get('adj_' + _k); _e = _D.get('edges_' + _k); _l = _D.get('elay_' + _k)
                ck('NET_CONFIG=%s ships a complete key quad (adj/edges/elay)' % _k,
                   bool(_a) and bool(_e) and _l is not None,
                   'adj=%s edges=%s elay=%s' % (bool(_a), bool(_e), _l is not None))
                if _e is not None and _l is not None:
                    ck('NET_CONFIG=%s elay is one provenance code per drawn edge' % _k,
                       len(_l) == len(_e), '%d codes vs %d edges' % (len(_l), len(_e)))
                if _D.get('x_' + _k) is not None:
                    ck('NET_CONFIG=%s layout arrays are roster-length and paired' % _k,
                       len(_D['x_' + _k]) == len(_D['name']) and len(_D.get('y_' + _k) or []) == len(_D['name']),
                       'x=%d y=%d roster=%d' % (len(_D['x_' + _k]),
                                                len(_D.get('y_' + _k) or []), len(_D['name'])))

            # --- RANK 6: JS tokens equal the registry, in template AND deployed html -------------
            for _lbl, _f2 in (('template', FE + r'\surfaces\templates\make_dashboard_html_00.tpl'),
                              ('deployed', B + r'\Moneyball Dashboard.html')):
                if not os.path.exists(_f2): continue
                _src = open(_f2, encoding='utf-8', errors='replace').read()
                _btn = set(re.findall(r'data-net=([A-Za-z_]+)', _src))
                ck('%s network toggle offers exactly the registry configs' % _lbl,
                   _btn == set(_NET), 'buttons=%s registry=%s' % (sorted(_btn), sorted(_NET)))

            # --- RANK 8: the sector honesty block, scoped to sectors (never blanket) ------------
            # holistic/holistic_norm legitimately carry NO engine_floor, so a blanket assertion
            # would go red on correct artifacts.
            for _k, _v in sorted(_NET.items()):
                if _v.get('family') != 'sector' or _k not in _D: continue
                _m = _D[_k] or {}
                _missing = [f for f in ('engine_floor', 'concentration', 'scope', 'members_isolated')
                            if _m.get(f) is None]
                ck('promoted sector %s ships its full honesty block' % _k,
                   not _missing, 'missing: %s' % _missing)
                if _m.get('members_tied') is not None and _m.get('members_isolated') is not None:
                    ck('promoted sector %s accounts for every roster member' % _k,
                       _m['members_tied'] + _m['members_isolated'] == len(_D['name']),
                       'tied %s + isolated %s != roster %d'
                       % (_m['members_tied'], _m['members_isolated'], len(_D['name'])))
    finally:
        if _prev_nc is None: os.environ.pop('NET_CONFIG', None)
        else: os.environ['NET_CONFIG'] = _prev_nc
except Exception as _e:
    # An exception here is a FAILURE, never a silent skip: a section that cannot run is a section
    # that proves nothing, and this audit's own contract is that unreadable output is failure.
    ck('network configuration section ran to completion', False,
       '%s: %s' % (type(_e).__name__, str(_e)[:160]))

print('\n[SYS] System manifest vs disk (containerization phase 0)')
# The expectation is data/system/system_manifest.json AS COMMITTED — produced by a PRIOR run of
# fec/audit/build_system_manifest.py. The observation is live os.stat/os.scandir/hashlib. This
# section deliberately does NOT import the producer: the checker asserts that build_system_manifest
# is absent from sys.modules, so a future "helpful" regeneration here turns the section red instead
# of turning it vacuous. Skips (absent off-repo roots on a fresh clone) are printed with their
# reason and are not verdicts; a manifest with an empty store list is a FAILURE, not a free pass.
try:
    # Loaded by path, the way this file already loads money_response — fec/ is on sys.path but
    # fec/audit/ is not, and this must not depend on fec/audit becoming a package.
    _smc_spec = importlib.util.spec_from_file_location(
        'system_manifest_check',
        os.path.join(os.path.dirname(os.path.abspath(__file__)), 'system_manifest_check.py'))
    _SMC = importlib.util.module_from_spec(_smc_spec); _smc_spec.loader.exec_module(_SMC)
    _sys_manifest = _SMC.default_manifest_path(B)
    print('  manifest      : %s' % _sys_manifest)

    def _sys_notice(_text):
        print('  NOTICE %s' % _text)

    _sys_tally = _SMC.check_manifest(_sys_manifest, ck, notice=_sys_notice, git_root=B)
    for _nm, _why in _sys_tally.get('skips', []):
        print('  SKIP   %s :: %s' % (_nm, _why))
    print('  manifest agreement: %d checks, %d failed, %d skipped, %d notices'
          % (_sys_tally['checks'], _sys_tally['failed'], _sys_tally['skipped'],
             _sys_tally['notices']))
    # A section that emitted nothing proves nothing. This repo has shipped two checks that could
    # not fail; an empty tally is the shape both of them had.
    ck('system manifest section emitted checks', _sys_tally['checks'] > 0,
       'the manifest-vs-disk section ran but asserted nothing')
except Exception as _e:
    # Same contract as the [NET] section above: an exception here is a FAILURE, never a skip.
    ck('system manifest section ran to completion', False,
       '%s: %s' % (type(_e).__name__, str(_e)[:200]))

print('FULL-BODY AUDIT — %d passed, %d FAILED'%(len(P),len(F)))
if F:
    print('  FAILURES:',F)
    sys.exit(1)
else: print('  ALL MODULES CLEAN.')
